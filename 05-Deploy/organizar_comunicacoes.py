#!/usr/bin/env python3
"""Organiza ofícios e TSIDs recebidos em arquivos ZIP na estrutura do projeto.

Materializa para cada um dos 113 auditados os documentos que lhe foram
enviados. Anexos comuns são replicados na pasta de cada destinatário. Os ZIPs
de entrada são preservados como fontes originais depois da validação.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import shutil
import sys
import unicodedata
import zipfile
from pathlib import Path, PurePosixPath

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_RESULTADO = ROOT / "02-Execucao/03-Execucao_Procedimentos/02-Resultados_Auditoria/03-pos-comentarios-gestor/resultado_auditoria.json"
DEFAULT_AUDITADOS = ROOT / "02-Execucao/03-Execucao_Procedimentos/01-Insumos/bd_auditados.xlsx"
DEFAULT_OFICIOS_ZIP = ROOT / "oficios_apresentacao.zip"
DEFAULT_TSID01_ZIP = ROOT / "TSID01.zip"
DEFAULT_TSID23_ZIP = ROOT / "TSID2 e 3.zip"
DESTINO_OFICIOS = ROOT / "99-Gestao/01-Oficios_Apresentacao"
DESTINO_TSIDS = ROOT / "99-Gestao/02-TSIDs"
DESTINO_ORIGINAIS = ROOT / "99-Gestao/03-Fontes_Originais_Comunicacoes"

# Nomes dos expedientes que diferem das siglas vigentes no cadastro.
ALIASES_EXPLICITOS = {
    "CIDADES": "SECID",
    "ESPORTE": "SEEL",
    "IORJ": "IOERJ",
    "REPREBSB": "SERGB",
    "SEC EST CULTURA E ECONOMIA CR 370 2026": "SECEC",
    "TRABALHO": "SETRAB",
}


def normalizar(texto: object) -> str:
    valor = unicodedata.normalize("NFKD", str(texto or ""))
    valor = "".join(c for c in valor if not unicodedata.combining(c))
    return " ".join(re.sub(r"[^A-Z0-9]+", " ", valor.upper()).split())


def nome_zip(info: zipfile.ZipInfo) -> str:
    if info.flag_bits & 0x800:
        return info.filename
    try:
        return info.filename.encode("cp437").decode("cp850")
    except UnicodeError:
        return info.filename


def caminho_relativo_seguro(nome: str) -> Path:
    partes = [p for p in PurePosixPath(nome.replace("\\", "/")).parts if p not in ("", ".")]
    if not partes or any(p == ".." for p in partes):
        raise ValueError(f"Caminho inseguro no ZIP: {nome!r}")
    return Path(*partes)


def escrever_se_igual_ou_ausente(destino: Path, dados: bytes) -> None:
    if destino.is_file():
        if destino.read_bytes() != dados:
            raise FileExistsError(f"Arquivo divergente já existe: {destino}")
        return
    destino.parent.mkdir(parents=True, exist_ok=True)
    destino.write_bytes(dados)


def carregar_auditados(resultado: Path, cadastro: Path) -> list[dict[str, str]]:
    dados = json.loads(resultado.read_text(encoding="utf-8"))
    selecionados = {
        str(item.get("sigla") or sigla).strip()
        for sigla, item in dados.items()
        if isinstance(item, dict) and item.get("foi_auditado") and item.get("respondeu_questionario")
    }
    if len(selecionados) != 113:
        raise ValueError(f"Esperados 113 auditados avaliados; encontrados {len(selecionados)}.")

    wb = load_workbook(cadastro, read_only=True, data_only=True)
    linhas = wb.active.iter_rows(values_only=True)
    cabecalho = [str(v or "").strip() for v in next(linhas)]
    auditados: dict[str, dict[str, str]] = {}
    for valores in linhas:
        item = dict(zip(cabecalho, valores))
        sigla = str(item.get("sigla") or "").strip()
        if sigla in selecionados:
            auditados[sigla] = {
                "sigla": sigla,
                "orgao": str(item.get("orgao") or "").strip(),
                "esfera": str(item.get("esfera") or "").strip().upper(),
            }
    ausentes = sorted(selecionados - set(auditados))
    if ausentes:
        raise ValueError(f"Auditados ausentes do cadastro: {', '.join(ausentes)}")
    return [auditados[sigla] for sigla in dados if sigla in auditados]


def criar_aliases(auditados: list[dict[str, str]]) -> dict[str, str]:
    aliases: dict[str, str] = {}
    siglas = {item["sigla"] for item in auditados}
    for item in auditados:
        sigla = item["sigla"]
        aliases[normalizar(sigla)] = sigla
        if item["esfera"] == "M":
            municipio = re.sub(r"^Prefeitura Municipal d[aeo]s?\s+", "", item["orgao"], flags=re.I)
            aliases[normalizar(municipio)] = sigla
    aliases.update({normalizar(alias): sigla for alias, sigla in ALIASES_EXPLICITOS.items() if sigla in siglas})
    return aliases


def identificador_destinatario(nome: str) -> str:
    stem = PurePosixPath(nome).stem
    for padrao in (
        r"^TSID0?[23]\s*-\s*",
        r"^Comentarios do Gestor\s*-\s*",
        r"^Relatório Individual Preliminar\s*-\s*",
    ):
        stem = re.sub(padrao, "", stem, flags=re.I)
    return normalizar(stem)


def ler_arquivos_zip(path: Path) -> dict[str, bytes]:
    resultado: dict[str, bytes] = {}
    with zipfile.ZipFile(path) as zf:
        for info in zf.infolist():
            nome = nome_zip(info)
            if info.is_dir() or PurePosixPath(nome).name.startswith("~$"):
                continue
            caminho_relativo_seguro(nome)
            if nome in resultado:
                raise ValueError(f"Nome duplicado no ZIP {path.name}: {nome}")
            resultado[nome] = zf.read(info)
    return resultado


def localizar_unico(arquivos: dict[str, bytes], nome_final: str) -> tuple[str, bytes]:
    achados = [(PurePosixPath(n).name, d) for n, d in arquivos.items() if PurePosixPath(n).name == nome_final]
    if len(achados) != 1:
        raise ValueError(f"Esperado um arquivo {nome_final!r}; encontrados {len(achados)}.")
    return achados[0]


def indexar_por_destinatario(
    arquivos: dict[str, bytes], prefixo: str, aliases: dict[str, str]
) -> tuple[dict[str, tuple[str, bytes]], list[str]]:
    indice: dict[str, tuple[str, bytes]] = {}
    ignorados: list[str] = []
    for nome, dados in arquivos.items():
        if not nome.startswith(prefixo + "/"):
            continue
        sigla = aliases.get(identificador_destinatario(nome))
        if not sigla:
            ignorados.append(nome)
            continue
        if sigla in indice:
            raise ValueError(f"Mais de um arquivo em {prefixo} corresponde a {sigla}.")
        indice[sigla] = (PurePosixPath(nome).name, dados)
    return indice, ignorados


def organizar(args: argparse.Namespace) -> dict:
    entradas = [args.oficios_zip, args.tsid01_zip, args.tsid23_zip]
    for entrada in entradas:
        if not entrada.is_file():
            raise FileNotFoundError(f"Arquivo de entrada não localizado: {entrada}")

    auditados = carregar_auditados(args.resultado_auditoria, args.auditados)
    aliases = criar_aliases(auditados)
    oficios = ler_arquivos_zip(args.oficios_zip)
    tsid01 = ler_arquivos_zip(args.tsid01_zip)
    tsid23 = ler_arquivos_zip(args.tsid23_zip)

    _, estado = localizar_unico(tsid01, "TSID01 - Estado.docx")
    _, prefeituras_docx = localizar_unico(tsid01, "TSID01 - Prefeituras.docx")
    _, prefeituras_pdf = localizar_unico(tsid01, "TSID01 - Prefeituras.pdf")
    comuns_tsid02 = [
        localizar_unico(tsid23, "ANEXO I - QUESTIONARIO.pdf"),
        localizar_unico(tsid23, "ANEXO II - GLOSSARIO.xlsx"),
    ]
    tsid02, ignorados02 = indexar_por_destinatario(tsid23, "02-TSID02", aliases)
    tsid03, ignorados03 = indexar_por_destinatario(tsid23, "04-TSID03/TSIDs", aliases)
    questionarios03, ignorados_q = indexar_por_destinatario(tsid23, "04-TSID03/Questionários", aliases)
    relatorios03, ignorados_r = indexar_por_destinatario(tsid23, "04-TSID03/Relatórios", aliases)

    indices = {
        "TSID02": tsid02,
        "TSID03": tsid03,
        "questionário enviado com TSID03": questionarios03,
        "relatório preliminar enviado com TSID03": relatorios03,
    }
    for descricao, indice in indices.items():
        ausentes = [a["sigla"] for a in auditados if a["sigla"] not in indice]
        if ausentes:
            raise ValueError(f"{descricao}: faltam destinatários: {', '.join(ausentes)}")

    registros: list[dict[str, str]] = []
    for nome, dados in oficios.items():
        relativo = caminho_relativo_seguro(nome)
        if relativo.parts[0].lower() == "assinados":
            relativo = Path(*relativo.parts[1:])
        destino = DESTINO_OFICIOS / relativo
        escrever_se_igual_ou_ausente(destino, dados)
        registros.append(registro("oficio_apresentacao", "", destino, dados))

    for item in auditados:
        sigla = item["sigla"]
        pasta = DESTINO_TSIDS / sigla
        fontes01 = [("TSID01 - Estado.docx", estado)] if item["esfera"] == "E" else [
            ("TSID01 - Prefeituras.docx", prefeituras_docx),
            ("TSID01 - Prefeituras.pdf", prefeituras_pdf),
        ]
        conjuntos = {
            "TSID01": fontes01,
            "TSID02": [tsid02[sigla], *comuns_tsid02],
            "TSID03": [tsid03[sigla], questionarios03[sigla], relatorios03[sigla]],
        }
        for tipo, fontes in conjuntos.items():
            for nome, dados in fontes:
                destino = pasta / tipo / nome
                escrever_se_igual_ou_ausente(destino, dados)
                registros.append(registro(tipo, sigla, destino, dados))

    DESTINO_ORIGINAIS.mkdir(parents=True, exist_ok=True)
    originais = []
    for entrada in entradas:
        destino = DESTINO_ORIGINAIS / entrada.name
        if destino.is_file() and destino.read_bytes() != entrada.read_bytes():
            raise FileExistsError(f"Fonte original divergente já existe: {destino}")
        if not destino.exists():
            shutil.move(str(entrada), destino)
        elif entrada.resolve() != destino.resolve():
            entrada.unlink()
        originais.append({"arquivo": str(destino.relative_to(ROOT)), "sha256": hashlib.sha256(destino.read_bytes()).hexdigest()})

    ignorados = sorted(set(ignorados02 + ignorados03 + ignorados_q + ignorados_r))
    manifesto = {
        "auditados_organizados": len(auditados),
        "oficios_organizados": len(oficios),
        "arquivos_por_auditado": {
            "estadual": {"TSID01": 1, "TSID02": 3, "TSID03": 3},
            "municipal": {"TSID01": 2, "TSID02": 3, "TSID03": 3},
        },
        "fontes_originais": originais,
        "arquivos_nao_associados_aos_113_auditados": ignorados,
        "registros": registros,
    }
    DESTINO_ORIGINAIS.joinpath("manifesto-importacao.json").write_text(
        json.dumps(manifesto, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    with DESTINO_ORIGINAIS.joinpath("manifesto-importacao.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=["tipo", "auditado", "arquivo", "sha256"], delimiter=";")
        writer.writeheader()
        writer.writerows(registros)
    return manifesto


def registro(tipo: str, auditado: str, destino: Path, dados: bytes) -> dict[str, str]:
    return {
        "tipo": tipo,
        "auditado": auditado,
        "arquivo": str(destino.relative_to(ROOT)),
        "sha256": hashlib.sha256(dados).hexdigest(),
    }


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--oficios-zip", type=Path, default=DEFAULT_OFICIOS_ZIP)
    parser.add_argument("--tsid01-zip", type=Path, default=DEFAULT_TSID01_ZIP)
    parser.add_argument("--tsid23-zip", type=Path, default=DEFAULT_TSID23_ZIP)
    parser.add_argument("--resultado-auditoria", type=Path, default=DEFAULT_RESULTADO)
    parser.add_argument("--auditados", type=Path, default=DEFAULT_AUDITADOS)
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    try:
        manifesto = organizar(parse_args(argv))
    except Exception as exc:
        print(f"ERRO: {exc}", file=sys.stderr)
        return 1
    print(f"Ofícios organizados: {manifesto['oficios_organizados']}")
    print(f"Auditados com TSIDs organizados: {manifesto['auditados_organizados']}")
    print(f"Arquivos de outros destinatários preservados no manifesto: {len(manifesto['arquivos_nao_associados_aos_113_auditados'])}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
