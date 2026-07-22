#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Consolida coleta, avaliação e impactos dos comentários do gestor."""

from __future__ import annotations

import argparse
import html
import json
import re
import textwrap
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable
from xml.etree import ElementTree

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_LSS = ROOT / "02-Execucao/05-Comentarios_Gestor/questionario_comentarios_gestor.lss"
DEFAULT_RESULTADO = ROOT / (
    "02-Execucao/03-Execucao_Procedimentos/02-Resultados_Auditoria/resultado_auditoria.json"
)
DEFAULT_AJUSTES = ROOT / (
    "02-Execucao/01-Questionario/02-Ajustes_Respostas/"
    "ajustes_respostas_questionario_pos_avaliacao_evidencias.xlsx"
)
DEFAULT_AVALIACAO_FINAL = ROOT / (
    "02-Execucao/05-Comentarios_Gestor/03-Produtos_Pos_Comentarios/"
    "avaliacao_comentarios_gestor.xlsx"
)
DEFAULT_AJUSTES_COMENTARIOS = ROOT / (
    "02-Execucao/05-Comentarios_Gestor/02-Avaliacao_Comentarios_Gestor/"
    "ajustes_respostas_questionario_pos_comentarios_gestor.xlsx"
)
DEFAULT_IMPACTOS = ROOT / (
    "02-Execucao/05-Comentarios_Gestor/03-Produtos_Pos_Comentarios/"
    "impactos-comentarios-gestor.xlsx"
)
DEFAULT_OUTPUT = Path(__file__).resolve().parent

RESPOSTAS = [
    "Concorda e ja atendeu",
    "Concorda e esta atendendo",
    "Concorda, sem medida adotada",
    "Discorda",
]
CORES = {
    "Concorda e ja atendeu": "#228B22",
    "Concorda e esta atendendo": "#9ACD32",
    "Concorda, sem medida adotada": "#FFA500",
    "Discorda": "#F1613F",
    "Situacao encontrada inexistente": "#D2D3CF",
}
CORES_DECISOES = {
    "Acolhida": "#70AD47",
    "Parcialmente acolhida": "#FFC000",
    "Não acolhida": "#D9534F",
}
ROTULOS_CATEGORIAS = {
    "Concorda e ja atendeu": "Concorda e j\u00e1 atendeu \u00e0s propostas de encaminhamento",
    "Concorda e esta atendendo": "Concorda e j\u00e1 est\u00e1 atendendo \u00e0s propostas de encaminhamento",
    "Concorda, sem medida adotada": (
        "Concorda, mas ainda n\u00e3o adotou nenhuma medida para atender \u00e0s propostas de encaminhamento"
    ),
    "Discorda": "Discorda da sinaliza\u00e7\u00e3o de inadequa\u00e7\u00e3o",
    "Situacao encontrada inexistente": "Situa\u00e7\u00e3o encontrada inexistente",
}
TITULOS_ACHADOS = {
    1: "Estrutura de TIC insuficiente para coordenar, gerir e sustentar a tecnologia da informacao",
    2: "Governanca de TIC insuficiente para avaliar, dirigir e monitorar a tecnologia da informacao",
    3: "Planejamento de TIC insuficiente para orientar a gestao, o orcamento e as contratacoes de TIC",
    4: "Capacidade institucional insuficiente para sustentar a gestao de TIC e seguranca da informacao",
    5: "Gestao de servicos de TIC insuficiente para assegurar controle sobre servicos, ativos e incidentes",
    6: "Fragilidades na governanca tecnica da fase preparatoria das contratacoes de TIC",
}
TEMAS = {
    "Formalizacao, normas e governanca": (
        r"formaliz|normativ|regimento|estatuto|portaria|governan|comite|comit[eê]|pap[eé]is|responsabil"
    ),
    "Planejamento e planos de acao": (
        r"planej|plano de a[cç][aã]o|pdti|pdtic|peti|cronograma|em elabora[cç][aã]o|em andamento"
    ),
    "Forca de trabalho e competencias": (
        r"servidor|funcion[aá]ri|equipe|pessoal|quadro|for[cç]a de trabalho|capacita|compet[eê]ncia|concurso"
    ),
    "Orcamento, recursos e contratacao": (
        r"or[cç]ament|recurso|aquisi[cç][aã]o|contrata[cç][aã]o|licita[cç][aã]o|etp|termo de refer[eê]ncia|\btr\b"
    ),
    "Ferramentas, ativos e processos operacionais": (
        r"sistema|ferrament|glpi|invent[aá]rio|cat[aá]logo|incidente|configura[cç][aã]o|sla|ans|ativo"
    ),
    "Dependencia de terceiros ou estrutura compartilhada": (
        r"terceir|prestador|fornecedor|proderj|prefeitura|compartilhad|outsourc|empresa contratada"
    ),
    "Documentos e evidencias adicionais": (
        r"anex|evid[eê]ncia|document|comprova|processo sei|\bsei[- ]"
    ),
    "Pedido de orientacao ou esclarecimento": r"orienta[cç][aã]o|esclarec|d[uú]vida|clareza|como atender",
}


@dataclass(frozen=True)
class Situacao:
    codigo: str
    achado: int
    ordem: int
    texto: str


def texto(valor: Any) -> str:
    if valor is None:
        return ""
    try:
        if pd.isna(valor):
            return ""
    except (TypeError, ValueError):
        pass
    return " ".join(str(valor).replace("\xa0", " ").split()).strip()


def chave(valor: Any) -> str:
    base = unicodedata.normalize("NFKD", texto(valor)).encode("ascii", "ignore").decode("ascii")
    return base.casefold()


def limpar_html(valor: Any) -> str:
    valor = str(valor or "")
    if "?" in valor:
        try:
            valor = valor.encode("latin-1").decode("utf-8")
        except (UnicodeEncodeError, UnicodeDecodeError):
            pass
    return texto(re.sub(r"<[^>]+>", " ", html.unescape(valor)))


def inteiro(valor: Any) -> int:
    try:
        return int(float(valor or 0))
    except (TypeError, ValueError):
        return 0


def percentual(parte: int, total: int) -> str:
    return f"{(100 * parte / total):.1f}%".replace(".", ",") if total else "0,0%"


def classificar_resposta(valor: Any) -> str:
    normalizado = chave(valor)
    if normalizado.startswith("discorda"):
        return "Discorda"
    if "ja atendeu" in normalizado:
        return "Concorda e ja atendeu"
    if "ja esta atendendo" in normalizado:
        return "Concorda e esta atendendo"
    if normalizado.startswith("concorda"):
        return "Concorda, sem medida adotada"
    return "Nao classificada"


def ler_xlsx(path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    linhas = ws.iter_rows(values_only=True)
    cabecalho = [texto(v) for v in next(linhas)]
    return [dict(zip(cabecalho, linha)) for linha in linhas if any(v is not None for v in linha)]


def consolidar_submissoes(linhas: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    concluidas = [linha for linha in linhas if linha.get("submitdate")]
    validas = []
    excluidas = []
    for linha in concluidas:
        organizacao = texto(linha.get("firstname"))
        token = texto(linha.get("token"))
        if chave(organizacao) == "teste zip" or token == "123":
            excluidas.append({**linha, "motivo_exclusao": "Registro de teste"})
        else:
            validas.append(linha)

    por_token: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for linha in validas:
        por_token[texto(linha.get("token")) or chave(linha.get("firstname"))].append(linha)

    finais = []
    for grupo in por_token.values():
        ordenado = sorted(grupo, key=lambda item: item.get("submitdate") or datetime.min)
        finais.append(ordenado[-1])
        for antiga in ordenado[:-1]:
            excluidas.append({**antiga, "motivo_exclusao": "Substituida por reenvio posterior do mesmo token"})
    return sorted(finais, key=lambda item: chave(item.get("firstname"))), excluidas


def metadados_lss(path: Path) -> tuple[dict[str, Situacao], dict[str, str]]:
    root = ElementTree.parse(path).getroot()
    questoes = {}
    for row in root.findall("./questions/rows/row"):
        questoes[texto(row.findtext("title"))] = row

    situacoes = {}
    for codigo, row in questoes.items():
        match = re.fullmatch(r"A(\d+)G(\d+)Conc", codigo)
        if not match:
            continue
        pergunta = limpar_html(row.findtext("question"))
        trecho = re.search(r'situa[cç][aã]o\s+"(.+?)"\s+apontada', pergunta, flags=re.I)
        situacoes[codigo] = Situacao(
            codigo=codigo,
            achado=int(match.group(1)),
            ordem=int(match.group(2)),
            texto=texto(trecho.group(1) if trecho else pergunta).rstrip("."),
        )

    questoes_reavaliacao = {}
    for row in root.findall("./groups/rows/row"):
        nome = texto(row.findtext("group_name"))
        match = re.search(r"Reavalia[cç][aã]o de evid[eê]ncias\s*-\s*(q\d+)", nome, flags=re.I)
        if not match:
            continue
        codigo = match.group(1).lower()
        descricao = limpar_html(row.findtext("description"))
        trecho = re.search(rf"Quest[aã]o-base\s+{codigo}:\s*(.+?)(?:A avalia[cç][aã]o|$)", descricao, flags=re.I)
        questoes_reavaliacao[codigo] = texto(trecho.group(1) if trecho else descricao)
    return situacoes, questoes_reavaliacao


def resultado_final_ajuste(row: dict[str, Any]) -> str:
    revisor = texto(row.get("Avaliacao do auditor revisor") or row.get("Avaliação do auditor revisor"))
    if revisor:
        return revisor
    return texto(row.get("Resultado da avaliacao do juiz") or row.get("Resultado da avaliação do juiz"))


def elegibilidade_reavaliacao(path: Path) -> dict[str, dict[str, list[str]]]:
    elegiveis: dict[str, dict[str, list[str]]] = defaultdict(lambda: defaultdict(list))
    for row in pd.read_excel(path).to_dict(orient="records"):
        if "nao conforme" not in chave(resultado_final_ajuste(row)):
            continue
        auditado = texto(row.get("Auditado")).upper()
        codigo_item = texto(row.get("Codigo do item avaliado") or row.get("Código do item avaliado"))
        match = re.match(r"^(q\d{4})", codigo_item, flags=re.I)
        if auditado and match:
            elegiveis[auditado][match.group(1).lower()].append(codigo_item)
    return {org: dict(questoes) for org, questoes in elegiveis.items()}


def extrair_secao_achados(
    respostas: list[dict[str, Any]], situacoes: dict[str, Situacao]
) -> list[dict[str, Any]]:
    saida = []
    for row in respostas:
        organizacao = texto(row.get("firstname"))
        for codigo, situacao in situacoes.items():
            valor = texto(row.get(codigo))
            if not valor:
                continue
            prefixo = codigo[:-4]
            saida.append(
                {
                    "organizacao": organizacao,
                    "id_submissao": row.get("id"),
                    "achado": situacao.achado,
                    "codigo_situacao": prefixo,
                    "situacao": situacao.texto,
                    "resposta_original": valor,
                    "categoria": classificar_resposta(valor),
                    "comentario": texto(row.get(prefixo + "Com")),
                    "justificativa": texto(row.get(prefixo + "Jus")),
                    "arquivos": inteiro(row.get(prefixo + "Evi[filecount]")),
                }
            )
    return saida


def extrair_reavaliacao(
    respostas: list[dict[str, Any]], elegiveis: dict[str, dict[str, list[str]]], questoes: dict[str, str]
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    detalhes = []
    cobertura = []
    por_org = {texto(row.get("firstname")).upper(): row for row in respostas}
    for organizacao, bases in elegiveis.items():
        row = por_org.get(organizacao)
        for base, itens in bases.items():
            comentario = texto(row.get(f"REV{base.upper()}Com")) if row else ""
            arquivos = inteiro(row.get(f"REV{base.upper()}Evi[filecount]")) if row else 0
            registro = {
                "organizacao": organizacao,
                "questao_base": base,
                "texto_questao": questoes.get(base, ""),
                "itens_nao_conformes": "; ".join(itens),
                "quantidade_itens": len(itens),
                "respondeu_etapa": bool(row),
                "comentario": comentario,
                "arquivos": arquivos,
                "submeteu_reavaliacao": bool(comentario or arquivos),
            }
            cobertura.append(registro)
            if registro["submeteu_reavaliacao"]:
                detalhes.append(registro)
    return detalhes, cobertura


def extrair_nao_respondentes(respostas: list[dict[str, Any]], siglas: set[str]) -> list[dict[str, Any]]:
    saida = []
    por_org = {texto(row.get("firstname")).upper(): row for row in respostas}
    for sigla in sorted(siglas):
        row = por_org.get(sigla)
        saida.append(
            {
                "organizacao": sigla,
                "manifestou": bool(row and row.get("NRStatus")),
                "status": texto(row.get("NRStatus")) if row else "",
                "justificativa": texto(row.get("NRTexto")) if row else "",
                "arquivos": inteiro(row.get("NREvi[filecount]")) if row else 0,
            }
        )
    return saida


def limpar_registros_dataframe(df: pd.DataFrame) -> list[dict[str, Any]]:
    return df.astype(object).where(pd.notna(df), None).to_dict(orient="records")


def normalizar_decisao(valor: Any) -> str:
    normalizado = chave(valor)
    if "parcialmente acolhida" in normalizado:
        return "Parcialmente acolhida"
    if "nao acolhida" in normalizado:
        return "Não acolhida"
    if "inconclusiva" in normalizado:
        return "Não acolhida"
    if "acolhida" in normalizado:
        return "Acolhida"
    return texto(valor) or "Sem decisão"


def carregar_avaliacao_final(
    path: Path,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    secoes = pd.read_excel(path, sheet_name=None)
    if "Seção 1 - situações" not in secoes or "Seção 2 - itens" not in secoes:
        raise ValueError("A avaliação final deve conter as abas 'Seção 1 - situações' e 'Seção 2 - itens'.")

    secao_1 = limpar_registros_dataframe(secoes["Seção 1 - situações"])
    secao_2 = limpar_registros_dataframe(secoes["Seção 2 - itens"])
    manifestacoes = []
    for row in secao_1:
        codigo = texto(row.get("codigo") or row.get("Código"))
        match = re.match(r"^A(\d+)G\d+$", codigo, flags=re.I)
        if not match:
            continue
        decisao = texto(row.get("Decisão revisada")) or texto(row.get("decisao"))
        manifestacao_equipe = (
            texto(row.get("Manifestação revisada da equipe"))
            or texto(row.get("manifestacao_equipe"))
        )
        manifestacoes.append(
            {
                "organizacao": texto(row.get("Auditado")),
                "achado": int(match.group(1)),
                "codigo_situacao": codigo,
                "situacao": texto(row.get("situacao")),
                "resposta_original": texto(row.get("Alternativa selecionada pelo gestor")),
                "categoria": classificar_resposta(row.get("Alternativa selecionada pelo gestor")),
                "decisao": normalizar_decisao(decisao),
                "situacao_atual": texto(row.get("situacao_atual")),
                "manifestacao_equipe": manifestacao_equipe,
                "case_id": texto(row.get("case_id")),
            }
        )

    for row in secao_2:
        decisao = texto(row.get("Decisão revisada")) or texto(row.get("decisao"))
        row["decisao_final"] = normalizar_decisao(decisao)
        row["manifestacao_equipe_final"] = (
            texto(row.get("Manifestação revisada da equipe"))
            or texto(row.get("manifestacao_equipe"))
        )
    return manifestacoes, secao_1, secao_2


def carregar_ajustes_comentarios(path: Path) -> dict[str, list[dict[str, Any]]]:
    return {
        nome: limpar_registros_dataframe(df)
        for nome, df in pd.read_excel(path, sheet_name=None).items()
    }


def carregar_impactos(path: Path) -> list[dict[str, Any]]:
    return limpar_registros_dataframe(pd.read_excel(path, sheet_name="Impactos"))


def resumo_impactos(registros: list[dict[str, Any]]) -> dict[str, Any]:
    df = pd.DataFrame(registros)
    variacao = pd.to_numeric(df["variacao_igovti"], errors="coerce")
    anterior = pd.to_numeric(df["igovti_anterior"], errors="coerce")
    atual = pd.to_numeric(df["igovti_atual"], errors="coerce")
    situacoes_antes = pd.to_numeric(df["situacoes_antes"], errors="coerce").fillna(0)
    situacoes_atuais = pd.to_numeric(df["situacoes_atuais"], errors="coerce").fillna(0)
    situacoes_removidas = pd.to_numeric(df["situacoes_removidas"], errors="coerce").fillna(0)
    achados_antes = pd.to_numeric(df["achados_antes"], errors="coerce").fillna(0)
    achados_atuais = pd.to_numeric(df["achados_atuais"], errors="coerce").fillna(0)
    achados_removidos = pd.to_numeric(df["achados_removidos"], errors="coerce").fillna(0)
    algum_impacto = (situacoes_removidas > 0) | (achados_removidos > 0) | (variacao.abs() > 1e-12)
    alterados = variacao[variacao.abs() > 1e-12]
    return {
        "universo": len(df),
        "situacoes_antes": int(situacoes_antes.sum()),
        "situacoes_atuais": int(situacoes_atuais.sum()),
        "situacoes_removidas": int(situacoes_removidas.sum()),
        "organizacoes_com_situacoes_removidas": int((situacoes_removidas > 0).sum()),
        "achados_antes": int(achados_antes.sum()),
        "achados_atuais": int(achados_atuais.sum()),
        "achados_removidos": int(achados_removidos.sum()),
        "organizacoes_com_achados_removidos": int((achados_removidos > 0).sum()),
        "organizacoes_com_igovti_alterado": int((variacao.abs() > 1e-12).sum()),
        "organizacoes_com_igovti_aumentado": int((variacao > 1e-12).sum()),
        "organizacoes_com_igovti_reduzido": int((variacao < -1e-12).sum()),
        "igovti_medio_anterior": float(anterior.mean()),
        "igovti_medio_atual": float(atual.mean()),
        "variacao_media_igovti": float(variacao.mean()),
        "variacao_media_entre_alterados": float(alterados.mean()) if len(alterados) else 0.0,
        "variacao_maxima_igovti": float(variacao.max()),
        "organizacoes_com_algum_impacto": int(algum_impacto.sum()),
    }


def textos_significativos(registros: Iterable[dict[str, Any]], campos: tuple[str, ...]) -> list[tuple[str, str]]:
    ignorar = {
        "", "de acordo", "idem", "idem as anteriores", "justificativa em respostas anteriores",
        "nao ha comentarios", "sem comentarios", "nao se aplica",
    }
    saida = []
    for item in registros:
        for campo in campos:
            valor = texto(item.get(campo))
            reduzido = chave(valor).strip(" .;:")
            if reduzido not in ignorar and len(reduzido) >= 12:
                saida.append((texto(item.get("organizacao")), valor))
    return saida


def analisar_temas(textos: list[tuple[str, str]]) -> list[dict[str, Any]]:
    saida = []
    for tema, padrao in TEMAS.items():
        ocorrencias = [(org, valor) for org, valor in textos if re.search(padrao, chave(valor), flags=re.I)]
        saida.append(
            {
                "tema": tema,
                "mencoes": len(ocorrencias),
                "organizacoes": len({org for org, _ in ocorrencias}),
            }
        )
    return sorted(saida, key=lambda item: (-item["organizacoes"], -item["mencoes"], item["tema"]))


def configurar_grafico() -> None:
    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 9,
            "axes.titlesize": 11,
            "axes.titleweight": "bold",
            "axes.edgecolor": "#B7B7B7",
            "axes.spines.top": False,
            "axes.spines.right": False,
            "figure.facecolor": "white",
            "axes.facecolor": "white",
        }
    )


def salvar_figura(fig: plt.Figure, path: Path) -> None:
    fig.savefig(path, dpi=180, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def grafico_participacao(path: Path, totais: list[tuple[str, int, int]]) -> None:
    labels = [acentuar_rotulo(item[0]) for item in totais][::-1]
    valores = [100 * item[1] / item[2] if item[2] else 0 for item in totais][::-1]
    fig, ax = plt.subplots(figsize=(9, 3.8))
    bars = ax.barh(labels, valores, color="#4472C4", height=0.55)
    ax.set_xlim(0, 100)
    ax.set_xlabel(acentuar_rotulo("Percentual de organizacoes"))
    ax.grid(axis="x", color="#E7E6E6", linewidth=0.8)
    for bar, (_, numerador, denominador) in zip(bars, totais[::-1]):
        ax.text(bar.get_width() + 1, bar.get_y() + bar.get_height() / 2,
                f"{numerador}/{denominador} ({percentual(numerador, denominador)})", va="center")
    ax.set_title(acentuar_rotulo("Participacao na etapa de comentarios do gestor"))
    salvar_figura(fig, path)


def grafico_panorama_manifestacoes(path: Path, contagem: Counter) -> None:
    categorias = RESPOSTAS
    valores = [contagem[categoria] for categoria in categorias]
    total = sum(valores)
    labels = [textwrap.fill(ROTULOS_CATEGORIAS[categoria], width=34) for categoria in categorias]
    fig, ax = plt.subplots(figsize=(12, 4.1))
    bars = ax.bar(labels, valores, color=[CORES[categoria] for categoria in categorias], width=0.52)
    maior = max(valores or [1])
    ax.set_ylim(0, maior * 1.18)
    ax.grid(axis="y", color="#D9D9D9", linestyle="--", linewidth=0.7, alpha=0.75)
    ax.set_axisbelow(True)
    ax.spines["left"].set_visible(False)
    ax.spines["bottom"].set_color("#D9D9D9")
    ax.tick_params(axis="x", length=0, labelsize=8.5)
    ax.tick_params(axis="y", length=0)
    for bar, valor in zip(bars, valores):
        ax.text(
            bar.get_x() + bar.get_width() / 2,
            valor + maior * 0.025,
            f"{valor} ({percentual(valor, total)})",
            ha="center",
            va="bottom",
            fontsize=9,
            color="#333333",
        )
    fig.subplots_adjust(bottom=0.3)
    salvar_figura(fig, path)


def grafico_situacoes_achado(
    path: Path,
    numero_achado: int,
    situacoes: dict[str, Situacao],
    manifestacoes: list[dict[str, Any]],
    total_respondentes: int,
) -> None:
    categorias = [*RESPOSTAS, "Situacao encontrada inexistente"]
    metadados = sorted(
        (item for item in situacoes.values() if item.achado == numero_achado),
        key=lambda item: item.ordem,
    )
    por_codigo: dict[str, Counter] = defaultdict(Counter)
    for item in manifestacoes:
        if item["achado"] == numero_achado:
            por_codigo[item["codigo_situacao"]][item["categoria"]] += 1

    linhas = []
    for situacao in metadados:
        codigo = situacao.codigo[:-4]
        contagem = por_codigo[codigo]
        aplicaveis = sum(contagem[categoria] for categoria in RESPOSTAS)
        contagem["Situacao encontrada inexistente"] = max(0, total_respondentes - aplicaveis)
        linhas.append((textwrap.fill(situacao.texto, width=48), contagem))

    linhas = linhas[::-1]
    labels = [label for label, _ in linhas]
    altura = max(3.5, 1.05 * len(linhas) + 2.1)
    fig, ax = plt.subplots(figsize=(12.8, altura))
    esquerda = np.zeros(len(linhas))
    for categoria in categorias:
        valores = np.array([contagem[categoria] for _, contagem in linhas], dtype=float)
        ax.barh(
            labels,
            valores,
            left=esquerda,
            color=CORES[categoria],
            label=ROTULOS_CATEGORIAS[categoria],
            height=0.58,
        )
        for indice, valor in enumerate(valores):
            if not valor:
                continue
            pct = 100 * valor / total_respondentes if total_respondentes else 0
            rotulo = f"{int(valor)}\n({str(f'{pct:.2f}').replace('.', ',')}%)" if pct >= 5 else str(int(valor))
            ax.text(
                esquerda[indice] + valor / 2,
                indice,
                rotulo,
                ha="center",
                va="center",
                fontsize=7.2,
                color="#202020",
                linespacing=0.9,
            )
        esquerda += valores

    ax.set_xlim(0, max(total_respondentes, 1))
    ax.grid(axis="x", color="#D9D9D9", linestyle="--", linewidth=0.7, alpha=0.75)
    ax.set_axisbelow(True)
    ax.spines["left"].set_visible(False)
    ax.spines["bottom"].set_color("#D9D9D9")
    ax.tick_params(axis="y", length=0, labelsize=8)
    ax.tick_params(axis="x", length=0)
    ax.legend(
        loc="lower center",
        bbox_to_anchor=(0.5, 1.01),
        ncol=3,
        frameon=False,
        fontsize=7.5,
        handlelength=1.8,
        columnspacing=1.2,
    )
    fig.subplots_adjust(left=0.32, top=0.8)
    salvar_figura(fig, path)


def grafico_empilhado(
    path: Path, linhas: list[tuple[str, Counter]], titulo: str, figsize: tuple[float, float]
) -> None:
    labels = [linha[0] for linha in linhas][::-1]
    totais = [sum(linha[1].values()) for linha in linhas][::-1]
    fig, ax = plt.subplots(figsize=figsize)
    esquerda = np.zeros(len(linhas))
    for categoria in RESPOSTAS:
        valores = np.array(
            [100 * linha[1][categoria] / sum(linha[1].values()) if sum(linha[1].values()) else 0
             for linha in linhas[::-1]]
        )
        ax.barh(labels, valores, left=esquerda, color=CORES[categoria],
                label=acentuar_rotulo(categoria), height=0.62)
        for i, valor in enumerate(valores):
            if valor >= 6:
                ax.text(esquerda[i] + valor / 2, i, f"{valor:.1f}%".replace(".", ","),
                        ha="center", va="center", fontsize=8, color="white" if categoria == "Discorda" else "#1F1F1F")
        esquerda += valores
    for i, total in enumerate(totais):
        ax.text(101, i, f"n={total}", va="center", fontsize=8)
    ax.set_xlim(0, 108)
    ax.set_xlabel(acentuar_rotulo("Percentual das manifestacoes"))
    ax.set_title(acentuar_rotulo(titulo))
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.18), ncol=2, frameon=False)
    ax.grid(axis="x", color="#E7E6E6", linewidth=0.8)
    salvar_figura(fig, path)


def grafico_reavaliacao(path: Path, cobertura: list[dict[str, Any]]) -> None:
    por_base: dict[str, Counter] = defaultdict(Counter)
    for item in cobertura:
        por_base[item["questao_base"]]["elegiveis"] += 1
        por_base[item["questao_base"]]["submissoes"] += int(item["submeteu_reavaliacao"])
    ranking = sorted(por_base.items(), key=lambda par: (-par[1]["submissoes"], par[0]))[:15][::-1]
    labels = [base.upper() for base, _ in ranking]
    valores = [c["submissoes"] for _, c in ranking]
    totais = [c["elegiveis"] for _, c in ranking]
    fig, ax = plt.subplots(figsize=(9, 6))
    bars = ax.barh(labels, valores, color="#5B9BD5", height=0.62)
    ax.set_xlabel(acentuar_rotulo("Quantidade de organizacoes que solicitaram reavaliacao"))
    ax.set_title(acentuar_rotulo("Questoes-base com maior numero de pedidos de reavaliacao"))
    ax.grid(axis="x", color="#E7E6E6", linewidth=0.8)
    for bar, valor, total in zip(bars, valores, totais):
        ax.text(valor + 0.3, bar.get_y() + bar.get_height() / 2,
                f"{valor} de {total} elegiveis", va="center", fontsize=8)
    salvar_figura(fig, path)


def grafico_resultados_avaliacao(
    path: Path, decisoes_secao_1: Counter, decisoes_secao_2: Counter
) -> None:
    categorias = list(CORES_DECISOES)
    linhas = [
        ("Situações encontradas", decisoes_secao_1),
        ("Reavaliação de respostas e evidências", decisoes_secao_2),
    ]
    labels = [f"{nome}\n(n={sum(contagem.values())})" for nome, contagem in linhas][::-1]
    fig, ax = plt.subplots(figsize=(11.5, 3.8))
    esquerda = np.zeros(len(linhas))
    for categoria in categorias:
        valores = np.array(
            [
                100 * contagem[categoria] / sum(contagem.values())
                if sum(contagem.values()) else 0
                for _, contagem in linhas[::-1]
            ]
        )
        ax.barh(
            labels,
            valores,
            left=esquerda,
            color=CORES_DECISOES[categoria],
            label=categoria,
            height=0.52,
        )
        for indice, valor in enumerate(valores):
            quantidade = linhas[::-1][indice][1][categoria]
            if valor >= 4:
                ax.text(
                    esquerda[indice] + valor / 2,
                    indice,
                    f"{quantidade}\n({str(f'{valor:.1f}').replace('.', ',')}%)",
                    ha="center",
                    va="center",
                    fontsize=8,
                    color="white" if categoria == "Não acolhida" else "#202020",
                )
        esquerda += valores
    ax.set_xlim(0, 105)
    ax.set_xlabel("Percentual dos casos avaliados")
    ax.set_title("Resultado consolidado da avaliação dos comentários do gestor")
    ax.grid(axis="x", color="#E7E6E6", linewidth=0.8)
    ax.set_axisbelow(True)
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.22), ncol=3, frameon=False)
    salvar_figura(fig, path)


def grafico_impactos_organizacoes(path: Path, resumo: dict[str, Any]) -> None:
    universo = resumo["universo"]
    linhas = [
        ("Algum impacto final", resumo["organizacoes_com_algum_impacto"]),
        ("Situações removidas", resumo["organizacoes_com_situacoes_removidas"]),
        ("Aumento do iGovTI", resumo["organizacoes_com_igovti_aumentado"]),
        ("Achados afastados", resumo["organizacoes_com_achados_removidos"]),
    ][::-1]
    labels = [item[0] for item in linhas]
    valores = [item[1] for item in linhas]
    fig, ax = plt.subplots(figsize=(9.5, 4.2))
    bars = ax.barh(labels, valores, color=["#A5A5A5", "#5B9BD5", "#70AD47", "#4472C4"])
    ax.set_xlim(0, max(valores) * 1.35)
    ax.set_xlabel("Quantidade de organizações")
    ax.set_title("Organizações alcançadas pelos impactos dos comentários do gestor")
    ax.grid(axis="x", color="#E7E6E6", linewidth=0.8)
    ax.set_axisbelow(True)
    for bar, valor in zip(bars, valores):
        ax.text(
            valor + 0.5,
            bar.get_y() + bar.get_height() / 2,
            f"{valor} de {universo} ({percentual(valor, universo)})",
            va="center",
            fontsize=8.5,
        )
    salvar_figura(fig, path)


def grafico_estoques_antes_depois(path: Path, resumo: dict[str, Any]) -> None:
    fig, axes = plt.subplots(1, 2, figsize=(10.8, 4.2))
    paineis = [
        (
            axes[0],
            "Situações encontradas",
            [resumo["situacoes_antes"], resumo["situacoes_atuais"]],
            resumo["situacoes_antes"] - resumo["situacoes_atuais"],
        ),
        (
            axes[1],
            "Achados",
            [resumo["achados_antes"], resumo["achados_atuais"]],
            resumo["achados_antes"] - resumo["achados_atuais"],
        ),
    ]
    for ax, titulo, valores, reducao in paineis:
        bars = ax.bar(["Antes", "Após comentários"], valores, color=["#A5A5A5", "#4472C4"], width=0.55)
        ax.set_ylim(0, max(valores) * 1.18)
        ax.set_title(titulo)
        ax.grid(axis="y", color="#E7E6E6", linewidth=0.8)
        ax.set_axisbelow(True)
        for bar, valor in zip(bars, valores):
            ax.text(
                bar.get_x() + bar.get_width() / 2,
                valor + max(valores) * 0.025,
                f"{valor:,}".replace(",", "."),
                ha="center",
                fontsize=9,
            )
        ax.text(
            0.5,
            max(valores) * 1.10,
            f"Redução líquida: {reducao}",
            ha="center",
            fontsize=8.5,
            color="#404040",
        )
    fig.suptitle("Situações e achados antes e após os comentários do gestor", fontweight="bold")
    fig.tight_layout(rect=(0, 0, 1, 0.92))
    salvar_figura(fig, path)


def grafico_evolucao_igovti(path: Path, resumo: dict[str, Any]) -> None:
    valores = [100 * resumo["igovti_medio_anterior"], 100 * resumo["igovti_medio_atual"]]
    fig, ax = plt.subplots(figsize=(7.5, 4.5))
    bars = ax.bar(["Antes", "Após comentários"], valores, color=["#A5A5A5", "#70AD47"], width=0.5)
    ax.set_ylim(0, max(valores) * 1.35)
    ax.set_ylabel("iGovTI médio (%)")
    ax.set_title("Evolução da média do iGovTI após os comentários do gestor")
    ax.grid(axis="y", color="#E7E6E6", linewidth=0.8)
    ax.set_axisbelow(True)
    for bar, valor in zip(bars, valores):
        ax.text(
            bar.get_x() + bar.get_width() / 2,
            valor + max(valores) * 0.035,
            f"{valor:.2f}%".replace(".", ","),
            ha="center",
            fontsize=10,
        )
    variacao = 100 * resumo["variacao_media_igovti"]
    ax.text(
        0.5,
        max(valores) * 1.20,
        f"Variação média: +{str(f'{variacao:.2f}').replace('.', ',')} ponto percentual",
        ha="center",
        fontsize=9,
        color="#404040",
    )
    salvar_figura(fig, path)


def ajustar_planilha(ws) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
    for column in ws.columns:
        values = [len(texto(cell.value)) for cell in list(column)[:200]]
        ws.column_dimensions[column[0].column_letter].width = min(max(values or [10]) + 2, 60)


def gravar_memoria(path: Path, abas: dict[str, list[dict[str, Any]]]) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for nome, linhas in abas.items():
        ws = wb.create_sheet(nome[:31])
        if not linhas:
            ws.append(["Sem registros"])
            continue
        colunas = list(linhas[0])
        ws.append(colunas)
        for linha in linhas:
            ws.append([linha.get(coluna) for coluna in colunas])
        ajustar_planilha(ws)
    wb.save(path)


def acentuar_rotulo(valor: str) -> str:
    substituicoes = {
        "Participacao": "Participa\u00e7\u00e3o",
        "participacao": "participa\u00e7\u00e3o",
        "comentarios": "coment\u00e1rios",
        "organizacoes": "organiza\u00e7\u00f5es",
        "manifestacoes": "manifesta\u00e7\u00f5es",
        "reavaliacao": "reavalia\u00e7\u00e3o",
        "Reavaliacao": "Reavalia\u00e7\u00e3o",
        "Questoes": "Quest\u00f5es",
        "questao": "quest\u00e3o",
        "ja": "j\u00e1",
        "esta atendendo": "est\u00e1 atendendo",
    }
    resultado = valor
    for origem, destino in substituicoes.items():
        resultado = re.sub(rf"\b{re.escape(origem)}\b", destino, resultado)
    return resultado


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--respostas", type=Path, required=True, help="Exportacao XLSX do LimeSurvey.")
    parser.add_argument("--lss", type=Path, default=DEFAULT_LSS)
    parser.add_argument("--resultado-auditoria", type=Path, default=DEFAULT_RESULTADO)
    parser.add_argument("--ajustes-evidencias", type=Path, default=DEFAULT_AJUSTES)
    parser.add_argument(
        "--avaliacao-final",
        type=Path,
        default=DEFAULT_AVALIACAO_FINAL,
        help="Avaliação consolidada das seções 1 e 2.",
    )
    parser.add_argument(
        "--ajustes-comentarios",
        type=Path,
        default=DEFAULT_AJUSTES_COMENTARIOS,
        help="Ajustes consolidados decorrentes dos comentários do gestor.",
    )
    parser.add_argument(
        "--impactos",
        type=Path,
        default=DEFAULT_IMPACTOS,
        help="Comparação dos resultados anterior e posterior aos comentários.",
    )
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    for path in (
        args.respostas,
        args.lss,
        args.resultado_auditoria,
        args.ajustes_evidencias,
        args.avaliacao_final,
        args.ajustes_comentarios,
        args.impactos,
    ):
        if not path.exists():
            parser.error(f"Arquivo nao encontrado: {path}")

    args.output_dir.mkdir(parents=True, exist_ok=True)
    img_dir = args.output_dir / "img"
    dados_dir = args.output_dir / "dados"
    img_dir.mkdir(exist_ok=True)
    dados_dir.mkdir(exist_ok=True)

    brutas = ler_xlsx(args.respostas)
    respostas, excluidas = consolidar_submissoes(brutas)
    situacoes, questoes_reavaliacao = metadados_lss(args.lss)
    manifestacoes_coletadas = extrair_secao_achados(respostas, situacoes)
    elegiveis = elegibilidade_reavaliacao(args.ajustes_evidencias)
    reavaliacoes, cobertura = extrair_reavaliacao(respostas, elegiveis, questoes_reavaliacao)
    manifestacoes_finais, avaliacao_secao_1, avaliacao_secao_2 = carregar_avaliacao_final(
        args.avaliacao_final
    )
    ajustes_comentarios = carregar_ajustes_comentarios(args.ajustes_comentarios)
    impactos = carregar_impactos(args.impactos)
    impactos_resumo = resumo_impactos(impactos)

    if len(manifestacoes_coletadas) != len(manifestacoes_finais):
        raise ValueError(
            "Divergência entre as manifestações coletadas e avaliadas: "
            f"{len(manifestacoes_coletadas)} na coleta e {len(manifestacoes_finais)} na avaliação final."
        )

    resultado = json.loads(args.resultado_auditoria.read_text(encoding="utf-8"))
    nao_responderam = {sigla for sigla, item in resultado.items() if not item.get("respondeu_questionario")}
    nao_respondentes = extrair_nao_respondentes(respostas, nao_responderam)
    respondentes_igovti = sum(bool(item.get("respondeu_questionario")) for item in resultado.values())

    temas_achados = analisar_temas(
        textos_significativos(manifestacoes_coletadas, ("comentario", "justificativa"))
    )
    temas_reavaliacao = analisar_temas(textos_significativos(reavaliacoes, ("comentario",)))

    configurar_grafico()
    regulares = len({item["organizacao"] for item in manifestacoes_finais})
    elegiveis_que_responderam = len({item["organizacao"] for item in cobertura if item["respondeu_etapa"]})
    orgs_reavaliacao = len({item["organizacao"] for item in reavaliacoes})
    grafico_participacao(
        img_dir / "01-participacao.png",
        [
            ("Todas as organizações", len(respostas), len(resultado)),
            ("Respondentes do iGovTI", regulares, respondentes_igovti),
            ("Elegíveis à reavaliação: responderam", elegiveis_que_responderam, len(elegiveis)),
            ("Elegíveis à reavaliação: apresentaram pedido", orgs_reavaliacao, len(elegiveis)),
            ("Sem resposta válida ao iGovTI", sum(item["manifestou"] for item in nao_respondentes), len(nao_respondentes)),
        ],
    )
    geral = Counter(item["categoria"] for item in manifestacoes_finais)
    grafico_panorama_manifestacoes(img_dir / "02-panorama-geral.png", geral)
    por_achado = defaultdict(Counter)
    for item in manifestacoes_finais:
        por_achado[item["achado"]][item["categoria"]] += 1
    grafico_empilhado(
        img_dir / "03-manifestacoes-por-achado.png",
        [(f"Achado {numero}", por_achado[numero]) for numero in range(1, 7)],
        "Manifestacoes dos gestores por achado", (9, 5.2),
    )
    for numero in range(1, 7):
        grafico_situacoes_achado(
            img_dir / f"achado-{numero}-situacoes.png",
            numero,
            situacoes,
            manifestacoes_finais,
            regulares,
        )
    grafico_reavaliacao(img_dir / "04-reavaliacoes-por-questao.png", cobertura)
    decisoes_secao_1 = Counter(item["decisao"] for item in manifestacoes_finais)
    decisoes_secao_2 = Counter(item["decisao_final"] for item in avaliacao_secao_2)
    grafico_resultados_avaliacao(
        img_dir / "05-resultados-avaliacao.png",
        decisoes_secao_1,
        decisoes_secao_2,
    )
    grafico_impactos_organizacoes(
        img_dir / "06-impactos-organizacoes.png",
        impactos_resumo,
    )
    grafico_estoques_antes_depois(
        img_dir / "07-saldo-situacoes-achados.png",
        impactos_resumo,
    )
    grafico_evolucao_igovti(
        img_dir / "08-evolucao-igovti.png",
        impactos_resumo,
    )

    respostas_resumo = [
        {
            "organizacao": texto(row.get("firstname")),
            "id_submissao": row.get("id"),
            "data_submissao": row.get("submitdate"),
            "secao": "Ausencia de resposta" if texto(row.get("NRStatus")) else "Achados e/ou reavaliacao",
        }
        for row in respostas
    ]
    exclusoes_resumo = [
        {
            "organizacao": texto(row.get("firstname")),
            "id_submissao": row.get("id"),
            "data_submissao": row.get("submitdate"),
            "motivo_exclusao": row.get("motivo_exclusao"),
        }
        for row in excluidas
    ]
    gravar_memoria(
        dados_dir / "memoria-calculo-comentarios-gestor.xlsx",
        {
            "Submissoes validas": respostas_resumo,
            "Exclusoes": exclusoes_resumo,
            "Manifestacoes coletadas": manifestacoes_coletadas,
            "Reavaliacoes submetidas": reavaliacoes,
            "Cobertura reavaliacao": cobertura,
            "Nao respondentes": nao_respondentes,
            "Temas achados": temas_achados,
            "Temas reavaliacao": temas_reavaliacao,
            "Avaliacao final secao 1": avaliacao_secao_1,
            "Avaliacao final secao 2": avaliacao_secao_2,
            "Ajustes pos comentarios": ajustes_comentarios.get("Ajustes", []),
            "Pendencias ajustes": ajustes_comentarios.get("Pendências", []),
            "Impactos finais": impactos,
        },
    )
    ajustes_aplicados = ajustes_comentarios.get("Ajustes", [])
    pendencias_ajustes = ajustes_comentarios.get("Pendências", [])
    (dados_dir / "resumo-execucao.json").write_text(
        json.dumps(
            {
                "fonte": str(args.respostas.resolve()),
                "submissoes_brutas": len(brutas),
                "submissoes_validas": len(respostas),
                "submissoes_excluidas": len(excluidas),
                "manifestacoes_situacoes": len(manifestacoes_finais),
                "discordancias": sum(
                    item["categoria"] == "Discorda" for item in manifestacoes_finais
                ),
                "arquivos_achados": sum(item["arquivos"] for item in manifestacoes_coletadas),
                "reavaliacoes_submetidas": len(reavaliacoes),
                "arquivos_reavaliacao": sum(item["arquivos"] for item in reavaliacoes),
                "organizacoes_nao_respondentes_que_manifestaram": sum(item["manifestou"] for item in nao_respondentes),
                "avaliacao_secao_1": dict(decisoes_secao_1),
                "avaliacao_secao_2": dict(decisoes_secao_2),
                "ajustes_aplicados": len(ajustes_aplicados),
                "organizacoes_com_ajustes": len(
                    {texto(item.get("Auditado")) for item in ajustes_aplicados}
                ),
                "itens_ajustados_distintos": len(
                    {texto(item.get("Código do item avaliado")) for item in ajustes_aplicados}
                ),
                "pendencias_ajustes": len(pendencias_ajustes),
                "impactos": impactos_resumo,
            },
            ensure_ascii=False,
            indent=2,
            default=str,
        ),
        encoding="utf-8",
    )

    print(f"Memoria: {dados_dir / 'memoria-calculo-comentarios-gestor.xlsx'}")
    print(f"Graficos: {img_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
