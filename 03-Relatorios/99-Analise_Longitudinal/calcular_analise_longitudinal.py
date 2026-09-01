#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Calcula testes pareados do iGovTI 2023–2026 em dois cenários de 2026.

O teste principal é o Wilcoxon dos postos sinalizados, bilateral, com tratamento
de zeros de Pratt e p-valor por permutação. Os p-valores dos indicadores são
ajustados em conjunto pelo método sequencial de Holm. O teste t pareado é
registrado como análise de sensibilidade, não como critério principal. O cenário
base contém a autodeclaração saneada; o cenário final incorpora a avaliação das
evidências e dos comentários dos gestores.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.lines import Line2D
from matplotlib.colors import ListedColormap
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from scipy.stats import PermutationMethod, rankdata, shapiro, ttest_rel, wilcoxon


ROOT = Path(__file__).resolve().parents[2]
SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_SOURCE_BASE = ROOT / (
    "02-Execucao/01-Questionario/04-Resultados_iGovTI/"
    "01-pos-ajuste-inicial/comparacao-2023-2026/"
    "20260621-comparacao-iGovTI-2023-2026.xlsx"
)
DEFAULT_SOURCE_FINAL = ROOT / (
    "02-Execucao/01-Questionario/04-Resultados_iGovTI/"
    "03-pos-comentarios-gestor/comparacao-2023-2026/"
    "20260716-comparacao-iGovTI-2023-2026.xlsx"
)
DEFAULT_OUTPUT = SCRIPT_DIR

SHEET_PARES = "Resultados pareados"
SHEET_DESCRICOES = "Estatísticas agregados"
ALPHA = 0.05
TOLERANCIA_ZERO = 1e-12
DECIMAIS_DIFERENCAS = 12
SEMENTE_PADRAO = 20260811
PERMUTACOES_PADRAO = 99_999
BOOTSTRAPS_PADRAO = 50_000

CORES = {
    "aumento": "#2F7D32",
    "reducao": "#B33A3A",
    "sem_evidencia": "#68737D",
    "estavel": "#C7CDD1",
    "ano_2023": "#D59A2F",
    "cenario_base": "#78B159",
    "cenario_final": "#1B5E3C",
    "grade": "#D9DEE2",
    "texto": "#263238",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--comparacao-base", type=Path, default=DEFAULT_SOURCE_BASE)
    parser.add_argument("--comparacao-final", type=Path, default=DEFAULT_SOURCE_FINAL)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--alpha", type=float, default=ALPHA)
    parser.add_argument("--permutacoes", type=int, default=PERMUTACOES_PADRAO)
    parser.add_argument("--bootstraps", type=int, default=BOOTSTRAPS_PADRAO)
    parser.add_argument("--semente", type=int, default=SEMENTE_PADRAO)
    return parser.parse_args()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for bloco in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(bloco)
    return digest.hexdigest()


def caminho_relativo(path: Path) -> str:
    try:
        return str(path.resolve().relative_to(ROOT.resolve()))
    except ValueError:
        return str(path.resolve())


def carregar_base(path: Path) -> tuple[pd.DataFrame, dict[str, str], list[str]]:
    if not path.exists():
        raise FileNotFoundError(f"Base de comparação não encontrada: {path}")
    xls = pd.ExcelFile(path)
    faltantes = {SHEET_PARES, SHEET_DESCRICOES} - set(xls.sheet_names)
    if faltantes:
        raise ValueError(f"Abas obrigatórias ausentes: {sorted(faltantes)}")

    pares = pd.read_excel(path, sheet_name=SHEET_PARES)
    descricoes_df = pd.read_excel(path, sheet_name=SHEET_DESCRICOES)
    if descricoes_df[["indicador", "descricao"]].isna().any().any():
        raise ValueError("A aba de estatísticas contém indicador ou descrição em branco.")
    descricoes = dict(zip(descricoes_df["indicador"], descricoes_df["descricao"]))
    indicadores = list(descricoes)

    colunas_identificacao = {
        "grupo", "sigla_2023", "sigla_2026", "tipo_pareamento",
        "nivel_maturidade_2023", "nivel_maturidade_2026",
    }
    faltantes = colunas_identificacao - set(pares.columns)
    for indicador in indicadores:
        faltantes.update(
            coluna for coluna in (f"{indicador}_2023", f"{indicador}_2026")
            if coluna not in pares.columns
        )
    if faltantes:
        raise ValueError(f"Colunas obrigatórias ausentes: {sorted(faltantes)}")
    if pares.empty:
        raise ValueError("A base pareada está vazia.")
    if pares["sigla_2023"].isna().any() or pares["sigla_2026"].isna().any():
        raise ValueError("Há identificadores de organização em branco na base pareada.")
    if pares["sigla_2023"].duplicated().any() or pares["sigla_2026"].duplicated().any():
        raise ValueError("Cada organização deve aparecer uma única vez em cada ano.")

    for indicador in indicadores:
        for ano in (2023, 2026):
            coluna = f"{indicador}_{ano}"
            pares[coluna] = pd.to_numeric(pares[coluna], errors="raise")
            if pares[coluna].isna().any():
                raise ValueError(f"Há valores ausentes em {coluna}.")
            fora = ~pares[coluna].between(-TOLERANCIA_ZERO, 1 + TOLERANCIA_ZERO)
            if fora.any():
                raise ValueError(f"Há pontuações fora do intervalo [0, 1] em {coluna}.")
    return pares, descricoes, indicadores


def holm(p_valores: np.ndarray, alpha: float) -> tuple[np.ndarray, np.ndarray]:
    """Ajusta p-valores pelo procedimento sequencial de Holm."""
    p = np.asarray(p_valores, dtype=float)
    if p.ndim != 1 or np.isnan(p).any() or ((p < 0) | (p > 1)).any():
        raise ValueError("P-valores inválidos para o ajuste de Holm.")
    m = len(p)
    ordem = np.argsort(p, kind="stable")
    ordenados = p[ordem]
    ajustados_ordenados = np.maximum.accumulate(
        np.array([(m - posicao) * valor for posicao, valor in enumerate(ordenados)])
    )
    ajustados_ordenados = np.minimum(ajustados_ordenados, 1.0)
    ajustados = np.empty(m, dtype=float)
    ajustados[ordem] = ajustados_ordenados
    return ajustados, ajustados <= alpha


def intervalo_bootstrap_media(
    diferencas: np.ndarray, n_bootstraps: int, rng: np.random.Generator
) -> tuple[float, float]:
    medias = np.empty(n_bootstraps, dtype=float)
    tamanho_lote = 5_000
    n = len(diferencas)
    for inicio in range(0, n_bootstraps, tamanho_lote):
        fim = min(inicio + tamanho_lote, n_bootstraps)
        indices = rng.integers(0, n, size=(fim - inicio, n))
        medias[inicio:fim] = diferencas[indices].mean(axis=1)
    inferior, superior = np.quantile(medias, [0.025, 0.975])
    return float(inferior), float(superior)


def correlacao_bisserial_postos(diferencas: np.ndarray) -> float:
    nao_zero = diferencas[~np.isclose(diferencas, 0.0, atol=TOLERANCIA_ZERO)]
    if len(nao_zero) == 0:
        return 0.0
    postos = rankdata(np.abs(nao_zero), method="average")
    positivos = float(postos[nao_zero > 0].sum())
    negativos = float(postos[nao_zero < 0].sum())
    return (positivos - negativos) / (positivos + negativos)


def cohen_dz(diferencas: np.ndarray) -> float:
    desvio = float(np.std(diferencas, ddof=1))
    if math.isclose(desvio, 0.0, abs_tol=TOLERANCIA_ZERO):
        return 0.0
    return float(np.mean(diferencas) / desvio)


def analisar(
    pares: pd.DataFrame,
    descricoes: dict[str, str],
    indicadores: list[str],
    alpha: float,
    n_permutacoes: int,
    n_bootstraps: int,
    semente: int,
) -> pd.DataFrame:
    resultados: list[dict[str, Any]] = []
    for posicao, indicador in enumerate(indicadores):
        anterior = pares[f"{indicador}_2023"].to_numpy(dtype=float)
        atual = pares[f"{indicador}_2026"].to_numpy(dtype=float)
        diferencas = np.round(atual - anterior, DECIMAIS_DIFERENCAS)
        rng_permutacao = np.random.default_rng(np.random.SeedSequence([semente, posicao, 1]))
        rng_bootstrap = np.random.default_rng(np.random.SeedSequence([semente, posicao, 2]))

        if np.all(np.isclose(diferencas, 0.0, atol=TOLERANCIA_ZERO)):
            estatistica_w, p_wilcoxon = 0.0, 1.0
        else:
            teste_w = wilcoxon(
                diferencas,
                zero_method="pratt",
                alternative="two-sided",
                method=PermutationMethod(
                    n_resamples=n_permutacoes,
                    batch=min(5_000, n_permutacoes),
                    rng=rng_permutacao,
                ),
            )
            estatistica_w = float(teste_w.statistic)
            p_wilcoxon = float(teste_w.pvalue)

        teste_t = ttest_rel(atual, anterior, alternative="two-sided")
        ic_t = teste_t.confidence_interval(confidence_level=0.95)
        ic_boot_inf, ic_boot_sup = intervalo_bootstrap_media(
            diferencas, n_bootstraps, rng_bootstrap
        )
        p_shapiro = float(shapiro(diferencas).pvalue) if len(diferencas) >= 3 else math.nan
        media_diferenca = float(np.mean(diferencas))
        mediana_diferenca = float(np.median(diferencas))

        resultados.append(
            {
                "ordem": posicao + 1,
                "indicador": indicador,
                "descricao": descricoes[indicador],
                "n_pares": len(diferencas),
                "media_2023": float(np.mean(anterior)),
                "mediana_2023": float(np.median(anterior)),
                "media_2026": float(np.mean(atual)),
                "mediana_2026": float(np.median(atual)),
                "diferenca_media": media_diferenca,
                "diferenca_mediana": mediana_diferenca,
                "ic95_bootstrap_media_inferior": ic_boot_inf,
                "ic95_bootstrap_media_superior": ic_boot_sup,
                "avancos": int(np.sum(diferencas > TOLERANCIA_ZERO)),
                "regressoes": int(np.sum(diferencas < -TOLERANCIA_ZERO)),
                "estabilidades": int(np.sum(np.abs(diferencas) <= TOLERANCIA_ZERO)),
                "wilcoxon_w": estatistica_w,
                "wilcoxon_p_bruto": p_wilcoxon,
                "r_bisserial_postos": correlacao_bisserial_postos(diferencas),
                "teste_t_estatistica": float(teste_t.statistic),
                "teste_t_gl": int(teste_t.df),
                "teste_t_p": float(teste_t.pvalue),
                "teste_t_ic95_inferior": float(ic_t.low),
                "teste_t_ic95_superior": float(ic_t.high),
                "cohen_dz": cohen_dz(diferencas),
                "shapiro_p_diferencas": p_shapiro,
            }
        )

    df = pd.DataFrame(resultados)
    ajustados, rejeitar = holm(df["wilcoxon_p_bruto"].to_numpy(), alpha)
    df["wilcoxon_p_holm"] = ajustados
    df["significativo_holm"] = rejeitar
    t_ajustados, t_rejeitar = holm(df["teste_t_p"].to_numpy(), alpha)
    df["teste_t_p_holm"] = t_ajustados
    df["teste_t_significativo_holm"] = t_rejeitar
    df["conclusao"] = np.where(
        rejeitar & (df["diferenca_media"] > TOLERANCIA_ZERO),
        "Aumento estatisticamente significativo",
        np.where(
            rejeitar & (df["diferenca_media"] < -TOLERANCIA_ZERO),
            "Redução estatisticamente significativa",
            "Sem evidência estatística de mudança",
        ),
    )
    return df


def validar_consistencia(
    pares: pd.DataFrame, resultados: pd.DataFrame, indicadores: list[str]
) -> list[dict[str, Any]]:
    verificacoes: list[dict[str, Any]] = []

    def adicionar(nome: str, condicao: bool, detalhe: str) -> None:
        verificacoes.append(
            {"verificacao": nome, "resultado": "OK" if condicao else "FALHA", "detalhe": detalhe}
        )
        if not condicao:
            raise AssertionError(f"{nome}: {detalhe}")

    adicionar("Quantidade de pares", len(pares) == 68, f"{len(pares)} pares encontrados")
    adicionar("Quantidade de indicadores", len(indicadores) == 16, f"{len(indicadores)} indicadores")
    adicionar("Identificadores únicos em 2023", pares["sigla_2023"].is_unique, "sem duplicidades")
    adicionar("Identificadores únicos em 2026", pares["sigla_2026"].is_unique, "sem duplicidades")
    adicionar(
        "Contagem de direções",
        bool(((resultados["avancos"] + resultados["regressoes"] + resultados["estabilidades"]) == len(pares)).all()),
        "avanços + regressões + estabilidades = n em todos os indicadores",
    )
    adicionar(
        "Diferenças reproduzidas",
        all(
            np.allclose(
                np.round(pares[f"{i}_2026"] - pares[f"{i}_2023"], DECIMAIS_DIFERENCAS).mean(),
                resultados.loc[resultados["indicador"].eq(i), "diferenca_media"].iloc[0],
                atol=1e-14,
            )
            for i in indicadores
        ),
        "médias das diferenças conferidas contra a base",
    )
    adicionar(
        "P-valores ajustados válidos",
        bool(
            resultados["wilcoxon_p_holm"].between(0, 1).all()
            and resultados["teste_t_p_holm"].between(0, 1).all()
        ),
        "Wilcoxon e teste t: todos no intervalo [0, 1]",
    )
    return verificacoes


def validar_cenarios(
    pares_base: pd.DataFrame,
    pares_final: pd.DataFrame,
    descricoes_base: dict[str, str],
    descricoes_final: dict[str, str],
    indicadores_base: list[str],
    indicadores_final: list[str],
) -> list[dict[str, Any]]:
    verificacoes: list[dict[str, Any]] = []

    def adicionar(nome: str, condicao: bool, detalhe: str) -> None:
        verificacoes.append(
            {"verificacao": nome, "resultado": "OK" if condicao else "FALHA", "detalhe": detalhe}
        )
        if not condicao:
            raise AssertionError(f"{nome}: {detalhe}")

    chaves = ["grupo", "sigla_2023", "sigla_2026", "tipo_pareamento"]
    adicionar(
        "Mesma coorte nos cenários",
        pares_base[chaves].reset_index(drop=True).equals(pares_final[chaves].reset_index(drop=True)),
        "identificadores e pareamentos conferidos",
    )
    adicionar(
        "Mesmo conjunto de indicadores",
        indicadores_base == indicadores_final and descricoes_base == descricoes_final,
        f"{len(indicadores_base)} indicadores com descrições coincidentes",
    )
    adicionar(
        "Resultados de 2023 idênticos",
        all(
            np.allclose(
                pares_base[f"{indicador}_2023"],
                pares_final[f"{indicador}_2023"],
                atol=1e-14,
            )
            for indicador in indicadores_base
        ),
        "a referência de 2023 é a mesma nos dois cenários",
    )
    return verificacoes


def comparar_cenarios(resultados_base: pd.DataFrame, resultados_final: pd.DataFrame) -> pd.DataFrame:
    colunas = [
        "ordem", "indicador", "descricao", "media_2023", "media_2026",
        "diferenca_media", "diferenca_mediana", "ic95_bootstrap_media_inferior",
        "ic95_bootstrap_media_superior", "avancos", "regressoes", "estabilidades",
        "wilcoxon_p_bruto", "wilcoxon_p_holm", "significativo_holm",
        "r_bisserial_postos", "teste_t_p_holm", "teste_t_significativo_holm",
        "cohen_dz", "conclusao",
    ]
    base = resultados_base[colunas].add_suffix("_base")
    final = resultados_final[colunas].add_suffix("_final")
    comparacao = base.merge(
        final,
        left_on="indicador_base",
        right_on="indicador_final",
        validate="one_to_one",
    )
    comparacao.insert(0, "ordem", comparacao.pop("ordem_base"))
    comparacao.insert(1, "indicador", comparacao.pop("indicador_base"))
    comparacao.insert(2, "descricao", comparacao.pop("descricao_base"))
    comparacao = comparacao.drop(columns=["ordem_final", "indicador_final", "descricao_final"])
    comparacao["diferenca_final_menos_base"] = (
        comparacao["media_2026_final"] - comparacao["media_2026_base"]
    )

    def sintetizar(linha: pd.Series) -> str:
        base_conclusao = linha["conclusao_base"]
        final_conclusao = linha["conclusao_final"]
        if base_conclusao == final_conclusao:
            return f"Conclusão concordante: {base_conclusao.lower()}"
        if "Aumento" in base_conclusao and "Sem evidência" in final_conclusao:
            return "Aumento apenas no cenário base"
        if "Sem evidência" in base_conclusao and "Redução" in final_conclusao:
            return "Redução apenas no cenário final"
        return f"Conclusões distintas: base={base_conclusao}; final={final_conclusao}"

    comparacao["sintese_cenarios"] = comparacao.apply(sintetizar, axis=1)
    return comparacao.sort_values("ordem").reset_index(drop=True)


def base_organizacoes(pares: pd.DataFrame, indicadores: list[str]) -> pd.DataFrame:
    colunas = [
        "grupo", "sigla_2023", "sigla_2026", "tipo_pareamento",
        "nivel_maturidade_2023", "nivel_maturidade_2026",
    ]
    saida = pares[colunas].copy()
    for indicador in indicadores:
        anterior = pares[f"{indicador}_2023"].astype(float)
        atual = pares[f"{indicador}_2026"].astype(float)
        saida[f"{indicador}_2023"] = anterior
        saida[f"{indicador}_2026"] = atual
        saida[f"delta_{indicador}"] = np.round(atual - anterior, DECIMAIS_DIFERENCAS)
    saida["direcao_iGovTI"] = np.select(
        [saida["delta_iGovTI"] > TOLERANCIA_ZERO, saida["delta_iGovTI"] < -TOLERANCIA_ZERO],
        ["Avanço", "Regressão"],
        default="Estabilidade",
    )
    return saida


def analisar_por_grupo(
    pares_por_cenario: dict[str, pd.DataFrame],
    descricoes: dict[str, str],
    indicadores: list[str],
) -> pd.DataFrame:
    """Produz leitura descritiva por grupo, sem inferência para subgrupos pequenos."""
    registros: list[dict[str, Any]] = []
    for cenario, pares in pares_por_cenario.items():
        for grupo, subconjunto in pares.groupby("grupo", sort=False, dropna=False):
            for ordem, indicador in enumerate(indicadores, start=1):
                anterior = subconjunto[f"{indicador}_2023"].to_numpy(dtype=float)
                atual = subconjunto[f"{indicador}_2026"].to_numpy(dtype=float)
                diferencas = np.round(atual - anterior, DECIMAIS_DIFERENCAS)
                registros.append(
                    {
                        "cenario": cenario,
                        "grupo": str(grupo),
                        "n_organizacoes": len(subconjunto),
                        "ordem": ordem,
                        "indicador": indicador,
                        "descricao": descricoes[indicador],
                        "media_2023": float(np.mean(anterior)),
                        "media_2026": float(np.mean(atual)),
                        "diferenca_media": float(np.mean(diferencas)),
                        "diferenca_mediana": float(np.median(diferencas)),
                        "avancos": int(np.sum(diferencas > TOLERANCIA_ZERO)),
                        "regressoes": int(np.sum(diferencas < -TOLERANCIA_ZERO)),
                        "estabilidades": int(np.sum(np.abs(diferencas) <= TOLERANCIA_ZERO)),
                    }
                )
    return pd.DataFrame(registros)


def estilo_graficos() -> None:
    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 9.5,
            "axes.labelcolor": CORES["texto"],
            "xtick.color": CORES["texto"],
            "ytick.color": CORES["texto"],
            "axes.edgecolor": CORES["grade"],
            "figure.facecolor": "white",
            "axes.facecolor": "white",
        }
    )


def limpar_eixo(ax: plt.Axes, eixo_grade: str = "x") -> None:
    ax.grid(axis=eixo_grade, color=CORES["grade"], linewidth=0.7, alpha=0.75)
    ax.set_axisbelow(True)
    ax.spines[["top", "right", "left"]].set_visible(False)
    ax.tick_params(axis="y", length=0)


def salvar(fig: plt.Figure, path: Path) -> None:
    fig.tight_layout()
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def grafico_medias_cenarios(comparacao: pd.DataFrame, path: Path) -> None:
    dados = comparacao.iloc[::-1].reset_index(drop=True)
    y = np.arange(len(dados))
    fig, ax = plt.subplots(figsize=(10.8, 8.2))
    for posicao, linha in dados.iterrows():
        ax.plot(
            [linha["media_2023_base"], linha["media_2026_base"], linha["media_2026_final"]],
            [posicao, posicao, posicao],
            color="#AAB2B8", linewidth=1.8, zorder=1,
        )
    ax.scatter(
        dados["media_2023_base"], y, color=CORES["ano_2023"], marker="o",
        s=48, label="2023", zorder=3,
    )
    ax.scatter(
        dados["media_2026_base"], y, color=CORES["cenario_base"], marker="^", s=52,
        label="2026 — cenário base", zorder=3,
    )
    ax.scatter(
        dados["media_2026_final"], y, color=CORES["cenario_final"], marker="s", s=48,
        label="2026 — cenário final", zorder=3,
    )
    ax.set_yticks(y, dados["descricao"])
    colunas = ["media_2023_base", "media_2026_base", "media_2026_final"]
    ax.set_xlim(0, max(0.52, float(dados[colunas].max().max()) + 0.06))
    ax.set_xlabel("Média do indicador (escala de 0 a 1)")
    ax.legend(
        frameon=False,
        ncol=3,
        loc="lower center",
        bbox_to_anchor=(0.5, 1.01),
        borderaxespad=0,
    )
    limpar_eixo(ax)
    salvar(fig, path)


def grafico_variacoes_ic(resultados: pd.DataFrame, path: Path, rotulo_cenario: str) -> None:
    dados = resultados.sort_values("diferenca_media").reset_index(drop=True)
    y = np.arange(len(dados))
    cores = np.where(
        dados["significativo_holm"],
        np.where(dados["diferenca_media"] > 0, CORES["aumento"], CORES["reducao"]),
        CORES["sem_evidencia"],
    )
    erros = np.vstack(
        [
            dados["diferenca_media"] - dados["ic95_bootstrap_media_inferior"],
            dados["ic95_bootstrap_media_superior"] - dados["diferenca_media"],
        ]
    )
    fig, ax = plt.subplots(figsize=(10.8, 8.2))
    ax.axvline(0, color="#20262B", linewidth=1.0)
    for posicao in range(len(dados)):
        ax.errorbar(
            dados.loc[posicao, "diferenca_media"], posicao,
            xerr=erros[:, posicao].reshape(2, 1), fmt="o",
            color=cores[posicao], ecolor=cores[posicao], capsize=3.5,
            markersize=6.5, linewidth=1.4,
        )
    ax.set_yticks(y, dados["descricao"])
    ax.set_xlabel(f"Variação média (2026 – 2023), {rotulo_cenario}, com IC 95% por bootstrap")
    legenda = [
        Line2D([0], [0], marker="o", color="none", markerfacecolor=CORES["aumento"],
               markeredgecolor=CORES["aumento"], label="Aumento significativo"),
        Line2D([0], [0], marker="o", color="none", markerfacecolor=CORES["reducao"],
               markeredgecolor=CORES["reducao"], label="Redução significativa"),
        Line2D([0], [0], marker="o", color="none", markerfacecolor=CORES["sem_evidencia"],
               markeredgecolor=CORES["sem_evidencia"], label="Sem evidência após Holm"),
    ]
    ax.legend(handles=legenda, frameon=False, loc="lower right")
    limpar_eixo(ax)
    salvar(fig, path)


def grafico_conclusoes_cenarios(comparacao: pd.DataFrame, path: Path) -> None:
    dados = comparacao.iloc[::-1].reset_index(drop=True)
    matriz = np.zeros((len(dados), 2), dtype=int)
    textos: list[list[str]] = []
    for posicao, linha in dados.iterrows():
        linha_textos = []
        for coluna, sufixo in enumerate(("base", "final")):
            conclusao = str(linha[f"conclusao_{sufixo}"])
            if conclusao.startswith("Aumento"):
                valor, simbolo = 1, "↑"
            elif conclusao.startswith("Redução"):
                valor, simbolo = -1, "↓"
            else:
                valor, simbolo = 0, "—"
            matriz[posicao, coluna] = valor
            p = float(linha[f"wilcoxon_p_holm_{sufixo}"])
            p_texto = f"{p:.5f}" if p < 0.001 else f"{p:.3f}"
            linha_textos.append(f"{simbolo}\np={p_texto}")
        textos.append(linha_textos)

    cmap = ListedColormap([CORES["reducao"], CORES["sem_evidencia"], CORES["aumento"]])
    fig, ax = plt.subplots(figsize=(8.8, 8.2))
    ax.imshow(matriz, cmap=cmap, vmin=-1, vmax=1, aspect="auto")
    for linha in range(len(dados)):
        for coluna in range(2):
            ax.text(
                coluna, linha, textos[linha][coluna], ha="center", va="center",
                color="white", fontweight="bold", fontsize=8.5,
            )
    ax.set_yticks(np.arange(len(dados)), dados["descricao"])
    ax.set_xticks([0, 1], ["2026 — cenário base", "2026 — cenário final"])
    ax.tick_params(axis="x", top=True, labeltop=True, bottom=False, labelbottom=False, length=0)
    ax.tick_params(axis="y", length=0)
    for spine in ax.spines.values():
        spine.set_visible(False)
    ax.set_xticks(np.arange(-0.5, 2, 1), minor=True)
    ax.set_yticks(np.arange(-0.5, len(dados), 1), minor=True)
    ax.grid(which="minor", color="white", linewidth=2)
    ax.tick_params(which="minor", bottom=False, left=False)
    salvar(fig, path)


def formatar_planilha(path: Path) -> None:
    wb = load_workbook(path)
    preenchimento = PatternFill("solid", fgColor="1F4E78")
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = preenchimento
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for coluna in ws.columns:
            letra = coluna[0].column_letter
            maior = max(len(str(c.value or "")) for c in coluna[: min(len(coluna), 250)])
            ws.column_dimensions[letra].width = min(max(maior + 2, 11), 48)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.000000"
    wb.save(path)


def json_seguro(valor: Any) -> Any:
    if isinstance(valor, dict):
        return {str(k): json_seguro(v) for k, v in valor.items()}
    if isinstance(valor, list):
        return [json_seguro(v) for v in valor]
    if isinstance(valor, (np.integer,)):
        return int(valor)
    if isinstance(valor, (np.floating, float)):
        return None if not math.isfinite(float(valor)) else float(valor)
    if isinstance(valor, (np.bool_,)):
        return bool(valor)
    return valor


def escrever_produtos(
    output_dir: Path,
    sources: dict[str, Path],
    pares: dict[str, pd.DataFrame],
    resultados: dict[str, pd.DataFrame],
    comparacao: pd.DataFrame,
    indicadores: list[str],
    verificacoes: list[dict[str, Any]],
    args: argparse.Namespace,
) -> tuple[Path, Path, list[Path]]:
    dados_dir = output_dir / "dados"
    img_dir = output_dir / "img"
    dados_dir.mkdir(parents=True, exist_ok=True)
    img_dir.mkdir(parents=True, exist_ok=True)
    xlsx_path = dados_dir / "resultados-testes-pareados-igovti-2023-2026.xlsx"
    json_path = dados_dir / "resumo-testes-pareados-igovti-2023-2026.json"

    metodologia = pd.DataFrame(
        [
            ("Base do cenário inicial", caminho_relativo(sources["base"])),
            ("SHA-256 do cenário inicial", sha256(sources["base"])),
            ("Base do cenário final", caminho_relativo(sources["final"])),
            ("SHA-256 do cenário final", sha256(sources["final"])),
            ("Cenário inicial", "Respostas autodeclaradas após ajustes iniciais de saneamento"),
            ("Cenário final", "Respostas após avaliação de evidências e comentários dos gestores"),
            ("Unidade de análise", "Organização presente em 2023 e 2026"),
            ("Número de pares", len(pares["base"])),
            ("Indicadores testados", len(indicadores)),
            ("Diferença", "Pontuação de 2026 menos pontuação de 2023"),
            ("Teste principal", "Wilcoxon dos postos sinalizados, bilateral"),
            ("Tratamento de zeros", "Pratt"),
            ("P-valor", f"Permutação Monte Carlo; {args.permutacoes:,} reamostragens"),
            ("Multiplicidade", "Holm para controle do erro familiar nos 16 testes"),
            ("Nível de significância", args.alpha),
            ("Intervalo da variação média", f"Bootstrap pareado percentil; {args.bootstraps:,} reamostragens"),
            ("Tamanho de efeito não paramétrico", "Correlação bisserial de postos nas diferenças não nulas"),
            ("Análise de sensibilidade", "Teste t pareado bilateral e Cohen dz"),
            ("Semente", args.semente),
            ("Arredondamento pré-Wilcoxon", f"{DECIMAIS_DIFERENCAS} casas decimais"),
        ],
        columns=["parametro", "valor"],
    )
    organizacoes_base = base_organizacoes(pares["base"], indicadores)
    organizacoes_final = base_organizacoes(pares["final"], indicadores)
    descricoes = dict(zip(resultados["base"]["indicador"], resultados["base"]["descricao"]))
    resultados_por_grupo = analisar_por_grupo(pares, descricoes, indicadores)
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        comparacao.to_excel(writer, sheet_name="Comparação cenários", index=False)
        resultados["base"].to_excel(writer, sheet_name="Testes cenário base", index=False)
        resultados["final"].to_excel(writer, sheet_name="Testes cenário final", index=False)
        resultados_por_grupo.to_excel(writer, sheet_name="Análise por grupo", index=False)
        organizacoes_base.to_excel(writer, sheet_name="Pares cenário base", index=False)
        organizacoes_final.to_excel(writer, sheet_name="Pares cenário final", index=False)
        metodologia.to_excel(writer, sheet_name="Metodologia", index=False)
        pd.DataFrame(verificacoes).to_excel(writer, sheet_name="Validações", index=False)
    formatar_planilha(xlsx_path)

    metadata = {
        "gerado_em_utc": datetime.now(timezone.utc).isoformat(),
        "fontes": {
            cenario: {"path": caminho_relativo(path), "sha256": sha256(path)}
            for cenario, path in sources.items()
        },
        "n_pares": len(pares["base"]),
        "n_indicadores": len(indicadores),
        "alpha": args.alpha,
        "teste_principal": "Wilcoxon bilateral; zero_method=pratt; p-valor por permutação",
        "permutacoes": args.permutacoes,
        "correcao_multiplicidade": "Holm",
        "bootstraps": args.bootstraps,
        "semente": args.semente,
        "resultados": {
            cenario: tabela.to_dict(orient="records")
            for cenario, tabela in resultados.items()
        },
        "comparacao_cenarios": comparacao.to_dict(orient="records"),
        "resultados_por_grupo": resultados_por_grupo.to_dict(orient="records"),
        "validacoes": verificacoes,
    }
    json_path.write_text(
        json.dumps(json_seguro(metadata), ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    estilo_graficos()
    graficos = [
        img_dir / "01-medias-tres-cenarios.png",
        img_dir / "02-variacoes-cenario-base-ic95.png",
        img_dir / "03-variacoes-cenario-final-ic95.png",
        img_dir / "04-conclusoes-cenarios.png",
    ]
    grafico_medias_cenarios(comparacao, graficos[0])
    grafico_variacoes_ic(resultados["base"], graficos[1], "cenário base")
    grafico_variacoes_ic(resultados["final"], graficos[2], "cenário final")
    grafico_conclusoes_cenarios(comparacao, graficos[3])
    for obsoleto in [
        "01-medias-2023-2026.png",
        "02-variacoes-medias-ic95.png",
        "03-direcao-variacoes-organizacoes.png",
    ]:
        (img_dir / obsoleto).unlink(missing_ok=True)
    return xlsx_path, json_path, graficos


def main() -> None:
    args = parse_args()
    if not 0 < args.alpha < 1:
        raise ValueError("--alpha deve estar entre 0 e 1.")
    if args.permutacoes < 999:
        raise ValueError("--permutacoes deve ser ao menos 999.")
    if args.bootstraps < 1_000:
        raise ValueError("--bootstraps deve ser ao menos 1.000.")
    pares_base, descricoes_base, indicadores_base = carregar_base(args.comparacao_base)
    pares_final, descricoes_final, indicadores_final = carregar_base(args.comparacao_final)
    resultados_base = analisar(
        pares_base, descricoes_base, indicadores_base, args.alpha,
        args.permutacoes, args.bootstraps, args.semente,
    )
    resultados_final = analisar(
        pares_final, descricoes_final, indicadores_final, args.alpha,
        args.permutacoes, args.bootstraps, args.semente,
    )
    comparacao = comparar_cenarios(resultados_base, resultados_final)
    verificacoes = []
    for cenario, verificacoes_cenario in [
        ("Cenário base", validar_consistencia(pares_base, resultados_base, indicadores_base)),
        ("Cenário final", validar_consistencia(pares_final, resultados_final, indicadores_final)),
        (
            "Entre cenários",
            validar_cenarios(
                pares_base, pares_final, descricoes_base, descricoes_final,
                indicadores_base, indicadores_final,
            ),
        ),
    ]:
        verificacoes.extend({"cenario": cenario, **item} for item in verificacoes_cenario)
    xlsx_path, json_path, graficos = escrever_produtos(
        args.output_dir,
        {"base": args.comparacao_base, "final": args.comparacao_final},
        {"base": pares_base, "final": pares_final},
        {"base": resultados_base, "final": resultados_final},
        comparacao,
        indicadores_base,
        verificacoes,
        args,
    )
    print(f"Cenário base: {caminho_relativo(args.comparacao_base)}")
    print(f"Cenário final: {caminho_relativo(args.comparacao_final)}")
    print(f"Organizações pareadas: {len(pares_base)}")
    print(f"Indicadores testados por cenário: {len(indicadores_base)}")
    print(f"Significativos no cenário base: {int(resultados_base['significativo_holm'].sum())}")
    print(f"Significativos no cenário final: {int(resultados_final['significativo_holm'].sum())}")
    print(f"Planilha: {caminho_relativo(xlsx_path)}")
    print(f"JSON: {caminho_relativo(json_path)}")
    for grafico in graficos:
        print(f"Gráfico: {caminho_relativo(grafico)}")


if __name__ == "__main__":
    main()
