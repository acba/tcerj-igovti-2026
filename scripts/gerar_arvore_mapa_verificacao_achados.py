#!/usr/bin/env python3
"""Gera uma árvore SVG/HTML do mapa de verificação de achados."""

from __future__ import annotations

import argparse
import re
import textwrap
from collections import OrderedDict
from html import escape
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_MAPA = ROOT / "02-Execucao/03-Execucao_Procedimentos/01-Insumos/mapa-verificacao-achados.xlsx"
DEFAULT_QUESTIONARIO = ROOT / "01-Planejamento/02-Metodologia_iGovTI/igovti_2026.md"
DEFAULT_OUTPUT = Path("/tmp/arvore-mapa-verificacao-achados.html")

ACHADO_COLORS = ["#1f6f8b", "#8a5a00", "#5b6e2d", "#8c3f58", "#476a92", "#7b5e7a"]
SOURCE_STYLES = {
    "questionario": {"color": "#2563eb", "label": "Questionário"},
    "avaliacao_evidencias_ajustes": {"color": "#dc2626", "label": "Avaliação de evidências"},
}


def normalizar_texto(valor) -> str:
    if valor is None:
        return ""
    return re.sub(r"\s+", " ", str(valor)).strip()


def sheet_rows(workbook, name: str) -> list[dict[str, object]]:
    ws = workbook[name]
    headers = [ws.cell(3, col).value for col in range(1, ws.max_column + 1)]
    rows: list[dict[str, object]] = []
    for row in range(4, ws.max_row + 1):
        item = {headers[col - 1]: ws.cell(row, col).value for col in range(1, ws.max_column + 1)}
        if any(value not in (None, "") for value in item.values()):
            rows.append(item)
    return rows


def parse_questionario(path: Path) -> dict[str, dict[str, object]]:
    text = path.read_text(encoding="utf-8")
    sections: dict[str, str] = {}
    pattern = r"^### (q\d+)(?:\s+\[[^\]]+\])?\s*$"
    for match in re.finditer(pattern, text, flags=re.M):
        code = match.group(1)
        start = match.end()
        next_match = re.search(pattern, text[start:], flags=re.M)
        end = start + next_match.start() if next_match else len(text)
        sections[code] = text[start:end]

    result: dict[str, dict[str, object]] = {}
    for code, section in sections.items():
        info = {"question": "", "options": {}, "details": {}, "subquestions": {}}
        question_match = re.search(r"^question:\s*\*\*(.*?)\*\*\s*$", section, flags=re.M)
        if question_match:
            info["question"] = normalizar_texto(question_match.group(1))

        for key, label in re.findall(r"^-\s*([A-Z])\s*\|\s*(.+)$", section, flags=re.M):
            pos = section.find(f"- {key} | {label}")
            prefix = section[:pos]
            last_options = prefix.rfind("options:")
            last_details = prefix.rfind("detail_options:")
            last_subquestions = prefix.rfind("subquestions:")
            if last_options > last_details and last_options > last_subquestions:
                target = "options"
            elif last_subquestions > last_options and last_subquestions > last_details:
                target = "subquestions"
            else:
                target = "details"
            info[target][key] = normalizar_texto(label)  # type: ignore[index]
        result[code] = info
    return result


def base_code(item: str) -> str:
    match = re.match(r"(q\d+)", item or "")
    return match.group(1) if match else item


def item_text(item: str, questionario: dict[str, dict[str, object]], variaveis: dict[str, str]) -> str:
    item = normalizar_texto(item)
    if not item:
        return ""
    if item in variaveis:
        return variaveis[item]
    info = questionario.get(base_code(item))
    if not info:
        return item

    suffix_match = re.search(r"\[([A-Z])\]", item)
    if suffix_match:
        suffix = suffix_match.group(1)
        for bucket in ("details", "subquestions", "options"):
            values = info.get(bucket, {})
            if isinstance(values, dict) and suffix in values:
                return str(values[suffix])
    return str(info.get("question") or item)


def source_color(source: str) -> str:
    return SOURCE_STYLES.get(source, {"color": "#6b7280"})["color"]


def wrap_text(text: str, width: int) -> list[str]:
    text = normalizar_texto(text)
    if not text:
        return [""]
    return textwrap.wrap(text, width=width, break_long_words=False, break_on_hyphens=False) or [""]


def svg_text(
    text: str,
    x: int,
    y: int,
    width_chars: int,
    *,
    size: float = 13,
    weight: str = "400",
    fill: str = "#182026",
    max_lines: int = 4,
    line_h: int = 16,
) -> str:
    raw = "" if text is None else str(text)
    lines = wrap_text(raw, width_chars)
    clipped = False
    if len(lines) > max_lines:
        lines = lines[:max_lines]
        lines[-1] = lines[-1].rstrip(" .,:;") + "..."
        clipped = True
    spans = [
        f'<tspan x="{x}" dy="{0 if index == 0 else line_h}">{escape(line)}</tspan>'
        for index, line in enumerate(lines)
    ]
    title = f"<title>{escape(raw)}</title>" if clipped or len(raw) > 100 else ""
    return (
        f'<text x="{x}" y="{y}" font-size="{size}" font-weight="{weight}" '
        f'fill="{fill}">{title}{"".join(spans)}</text>'
    )


def svg_rect(x: int, y: int, width: int, height: int, fill: str, stroke: str, *, sw: float = 1) -> str:
    return (
        f'<rect x="{x}" y="{y}" width="{width}" height="{height}" rx="7" ry="7" '
        f'fill="{fill}" stroke="{stroke}" stroke-width="{sw}" />'
    )


def svg_path(x1: int, y1: float, x2: int, y2: float, stroke: str, cls: str) -> str:
    mid = (x1 + x2) / 2
    data = f"M{x1},{y1} C{mid},{y1} {mid},{y2} {x2},{y2}"
    return f'<path class="{cls}" d="{data}" fill="none" stroke="{stroke}" stroke-width="1.2" />'


def montar_grupos(procedimentos: list[dict[str, object]], acoes: list[dict[str, object]]):
    action_by_id = {str(action["id"]): action for action in acoes if action.get("id")}
    groups = []
    for procedimento in procedimentos:
        ids = re.findall(r"AV\d+", str(procedimento.get("logica_achado") or ""))
        seen = []
        for action_id in ids:
            if action_id in action_by_id and action_id not in seen:
                seen.append(action_id)
        by_situation: OrderedDict[str, list[dict[str, object]]] = OrderedDict()
        for action_id in seen:
            action = action_by_id[action_id]
            description = normalizar_texto(action.get("descricao_situacao_inconforme")) or "(sem descrição)"
            by_situation.setdefault(description, []).append(action)
        groups.append({"proc": procedimento, "situations": by_situation})
    return groups


def gerar_html(mapa: Path, questionario_path: Path) -> str:
    workbook = load_workbook(mapa, data_only=True, read_only=True)
    procedimentos = sheet_rows(workbook, "Procedimentos de Auditoria")
    acoes = sheet_rows(workbook, "Ações de Verificação")
    variaveis_rows = sheet_rows(workbook, "Variáveis Temporárias")

    questionario = parse_questionario(questionario_path)
    variaveis = {
        normalizar_texto(row.get("nome")): normalizar_texto(row.get("descricao"))
        for row in variaveis_rows
        if normalizar_texto(row.get("nome"))
    }
    groups = montar_grupos(procedimentos, acoes)

    width = 3220
    header_h = 174
    row_h = 98
    gap_sit = 20
    gap_ach = 40
    x_ach, w_ach = 36, 430
    x_sit, w_sit = 540, 690
    x_av, w_av = 1310, 430
    x_item, w_item = 1810, 1310

    rows = []
    y = header_h
    for group_index, group in enumerate(groups):
        for situation, actions in group["situations"].items():
            for action in actions:
                rows.append((group_index, group["proc"], situation, action, y))
                y += row_h
            y += gap_sit
        y += gap_ach
    height = y + 80

    ach_bounds: dict[int, list[int]] = {}
    sit_bounds: dict[tuple[int, str], list[int]] = {}
    for group_index, _proc, situation, _action, row_y in rows:
        ach_bounds.setdefault(group_index, [row_y, row_y])
        ach_bounds[group_index][1] = row_y
        sit_bounds.setdefault((group_index, situation), [row_y, row_y])
        sit_bounds[(group_index, situation)][1] = row_y

    svg = [
        f'<svg id="treeSvg" xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" '
        f'viewBox="0 0 {width} {height}" role="img" aria-labelledby="svgTitle svgDesc">',
        '<title id="svgTitle">Árvore do mapa de verificação de achados</title>',
        '<desc id="svgDesc">Diagrama em colunas conectando achados, situações inconformes, ações de verificação e itens avaliados.</desc>',
        svg_rect(24, 18, width - 48, 116, "#ffffff", "#cbd2d9"),
        svg_text("Mapa de Verificação de Achados - Visão em Árvore", 46, 50, 90, size=24, weight="700", max_lines=1),
        svg_text(
            f"Fonte: {mapa} | {len(procedimentos)} achados/procedimentos, "
            f"{len(sit_bounds)} situações inconformes, {len(acoes)} ações de verificação. "
            f"Itens do questionário incorporados a partir de {questionario_path}.",
            46,
            82,
            205,
            size=13,
            fill="#4b5560",
            max_lines=2,
        ),
    ]
    for label, x, box_width in [
        ("Achado", x_ach, w_ach),
        ("Situação inconforme", x_sit, w_sit),
        ("Ação de verificação", x_av, w_av),
        ("Item e condição verificada", x_item, w_item),
    ]:
        svg.append(svg_rect(x, 146, box_width, 30, "#eef2f5", "#d5dbe1"))
        svg.append(svg_text(label, x + 12, 166, 46, size=13, weight="700", fill="#35424d", max_lines=1))

    for (group_index, situation), (start_y, end_y) in sit_bounds.items():
        ach_mid = (ach_bounds[group_index][0] + ach_bounds[group_index][1] + 76) / 2
        sit_mid = (start_y + end_y + 76) / 2
        svg.append(svg_path(x_ach + w_ach, ach_mid, x_sit, sit_mid, "#bac2ca", f"edge ach-{group_index + 1}"))
    for group_index, _proc, situation, action, row_y in rows:
        sit_mid = (sit_bounds[(group_index, situation)][0] + sit_bounds[(group_index, situation)][1] + 76) / 2
        av_mid = row_y + 38
        source = normalizar_texto(action.get("id_fonte_informacao"))
        svg.append(svg_path(x_sit + w_sit, sit_mid, x_av, av_mid, "#c4cbd2", f"edge ach-{group_index + 1}"))
        svg.append(svg_path(x_av + w_av, av_mid, x_item, av_mid, source_color(source), f"edge ach-{group_index + 1}"))

    for group_index, group in enumerate(groups):
        proc = group["proc"]
        start_y, end_y = ach_bounds.get(group_index, [header_h, header_h])
        box_h = max(110, end_y - start_y + 76)
        color = ACHADO_COLORS[group_index % len(ACHADO_COLORS)]
        svg.append(f'<g class="node achado ach-{group_index + 1}" data-achado="{group_index + 1}">')
        svg.append(svg_rect(x_ach, start_y, w_ach, box_h, "#ffffff", color, sw=1.4))
        svg.append(f'<rect x="{x_ach}" y="{start_y}" width="8" height="{box_h}" rx="4" fill="{color}" />')
        svg.append(
            svg_text(
                f"Achado {proc.get('numero_achado')}: {proc.get('nome_achado')}",
                x_ach + 18,
                start_y + 26,
                50,
                size=14,
                weight="700",
                max_lines=5,
                line_h=17,
            )
        )
        svg.append(
            svg_text(
                f"{proc.get('id')} | Lógica: {proc.get('logica_achado')}",
                x_ach + 18,
                start_y + box_h - 30,
                52,
                size=11,
                fill="#54616d",
                max_lines=2,
                line_h=13,
            )
        )
        svg.append("</g>")

    for (group_index, situation), (start_y, end_y) in sit_bounds.items():
        box_h = max(84, end_y - start_y + 76)
        color = ACHADO_COLORS[group_index % len(ACHADO_COLORS)]
        av_count = sum(1 for row in rows if row[0] == group_index and row[2] == situation)
        svg.append(f'<g class="node situacao ach-{group_index + 1}" data-achado="{group_index + 1}">')
        svg.append(svg_rect(x_sit, start_y, w_sit, box_h, "#ffffff", "#d8dde3"))
        svg.append(f'<circle cx="{x_sit + 17}" cy="{start_y + 21}" r="6" fill="{color}" />')
        svg.append(svg_text(situation, x_sit + 34, start_y + 25, 76, size=13, weight="700", max_lines=4))
        svg.append(svg_text(f"{av_count} ações de verificação", x_sit + 34, start_y + box_h - 17, 70, size=11, fill="#60707c", max_lines=1))
        svg.append("</g>")

    for group_index, _proc, _situation, action, row_y in rows:
        ach_color = ACHADO_COLORS[group_index % len(ACHADO_COLORS)]
        source = normalizar_texto(action.get("id_fonte_informacao"))
        source_border = source_color(source)
        condition = normalizar_texto(action.get("situacao_inconforme"))
        item = normalizar_texto(action.get("informacao_requerida"))
        criteria = normalizar_texto(action.get("criterio"))
        item_label = item_text(item, questionario, variaveis)

        svg.append(f'<g class="node acao ach-{group_index + 1}" data-achado="{group_index + 1}">')
        svg.append(svg_rect(x_av, row_y, w_av, 76, "#ffffff", source_border, sw=2))
        svg.append(svg_text(normalizar_texto(action.get("id")), x_av + 12, row_y + 22, 10, size=13, weight="700", fill=ach_color, max_lines=1))
        svg.append(svg_text(source, x_av + 76, row_y + 22, 34, size=12, fill="#34404a", max_lines=1))
        svg.append(svg_text(normalizar_texto(action.get("tipo_encaminhamento")), x_av + 76, row_y + 42, 34, size=11, fill="#6b7280", max_lines=1))
        svg.append(svg_text("Fonte", x_av + 12, row_y + 62, 15, size=10.5, fill=source_border, max_lines=1))
        svg.append("</g>")

        svg.append(f'<g class="node item ach-{group_index + 1}" data-achado="{group_index + 1}">')
        svg.append(svg_rect(x_item, row_y, w_item, 76, "#ffffff", source_border, sw=2))
        svg.append(svg_text(item_label, x_item + 12, row_y + 20, 150, size=12, weight="700", max_lines=2, line_h=14))
        svg.append(svg_text(f"{item} => {condition}", x_item + 12, row_y + 50, 150, size=11.5, weight="700", fill=source_border, max_lines=1))
        if criteria:
            svg.append(svg_text(f"Critério: {criteria}", x_item + 12, row_y + 68, 155, size=10, fill="#697783", max_lines=1))
        svg.append("</g>")

    svg.append("</svg>")

    achado_legend = "".join(
        f'<button class="chip" data-achado="{index + 1}" style="--c:{ACHADO_COLORS[index % len(ACHADO_COLORS)]}">'
        f'Achado {group["proc"].get("numero_achado")}</button>'
        for index, group in enumerate(groups)
    )
    source_legend = "".join(
        f'<span class="source-key" style="--c:{style["color"]}">{escape(style["label"])} '
        f'<code>{escape(source)}</code></span>'
        for source, style in SOURCE_STYLES.items()
    )

    return f"""<!doctype html>
<html lang="pt-BR">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Árvore do Mapa de Verificação de Achados</title>
<style>
  :root {{ color-scheme: light; --bg:#f5f6f8; --ink:#182026; --muted:#5f6b76; --line:#d7dde3; }}
  * {{ box-sizing: border-box; }}
  body {{ margin:0; font-family: Arial, Helvetica, sans-serif; background:var(--bg); color:var(--ink); }}
  header {{ padding:18px 24px 14px; background:#fff; border-bottom:1px solid var(--line); position:sticky; top:0; z-index:5; }}
  h1 {{ margin:0 0 8px; font-size:20px; line-height:1.25; }}
  .meta {{ color:var(--muted); font-size:13px; line-height:1.45; }}
  .toolbar {{ display:flex; align-items:center; gap:10px; flex-wrap:wrap; margin-top:12px; }}
  .chip {{ border:1px solid var(--line); background:#fff; color:#26323b; border-radius:6px; padding:7px 10px; cursor:pointer; font-size:13px; border-left:6px solid var(--c); }}
  .chip.active {{ background:#eef4f7; outline:2px solid #b9d3df; }}
  .btn {{ border:1px solid var(--line); background:#26323b; color:#fff; border-radius:6px; padding:7px 11px; cursor:pointer; font-size:13px; }}
  .source-legend {{ display:flex; gap:12px; flex-wrap:wrap; margin-top:10px; color:#4b5560; font-size:12px; }}
  .source-key {{ display:inline-flex; align-items:center; gap:6px; padding:5px 8px; border:1px solid var(--line); border-left:6px solid var(--c); border-radius:6px; background:#fff; }}
  main {{ padding:14px 18px 24px; }}
  .canvas {{ overflow:auto; background:#fff; border:1px solid var(--line); border-radius:8px; box-shadow:0 1px 2px rgba(0,0,0,.04); }}
  svg {{ display:block; min-width:2800px; }}
  text {{ font-family: Arial, Helvetica, sans-serif; }}
  .node rect, .edge {{ vector-effect: non-scaling-stroke; }}
  .dim {{ opacity:.13; }}
  .note {{ margin:12px 4px 0; color:var(--muted); font-size:12px; }}
</style>
</head>
<body>
<header>
  <h1>Árvore do Mapa de Verificação de Achados</h1>
  <div class="meta">Visualização derivada de <code>{escape(str(mapa))}</code>. Cada linha conecta uma ação de verificação ao item/condição que ela avalia. O texto do item do questionário aparece acima do código e da condição.</div>
  <div class="toolbar">
    <button class="btn" id="showAll">Mostrar todos</button>
    {achado_legend}
  </div>
  <div class="source-legend">{source_legend}</div>
</header>
<main>
  <div class="canvas">{"".join(svg)}</div>
  <p class="note">Leitura recomendada: da esquerda para a direita. As bordas das ações e dos itens indicam a fonte de informação. Passe o cursor sobre textos truncados para ver o conteúdo completo quando disponível pelo navegador.</p>
</main>
<script>
const chips = document.querySelectorAll('.chip');
const showAll = document.getElementById('showAll');
function setFilter(id) {{
  document.querySelectorAll('.node,.edge').forEach(el => {{
    if (!id) {{ el.classList.remove('dim'); return; }}
    el.classList.toggle('dim', !el.classList.contains('ach-' + id));
  }});
  chips.forEach(c => c.classList.toggle('active', c.dataset.achado === id));
}}
chips.forEach(c => c.addEventListener('click', () => setFilter(c.dataset.achado)));
showAll.addEventListener('click', () => setFilter(null));
</script>
</body>
</html>
"""


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--mapa", type=Path, default=DEFAULT_MAPA, help="Planilha mapa-verificacao-achados.xlsx.")
    parser.add_argument("--questionario", type=Path, default=DEFAULT_QUESTIONARIO, help="Questionário Markdown usado para obter textos dos itens.")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="Arquivo HTML de saída.")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    html = gerar_html(args.mapa, args.questionario)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(html, encoding="utf-8")
    print(args.output)


if __name__ == "__main__":
    main()
