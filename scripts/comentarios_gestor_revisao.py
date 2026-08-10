"""Classificação e gate de revisão humana dos comentários do gestor."""
from __future__ import annotations

import hashlib
import json
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


STATUS_APROVADO = "Aprovado"
STATUS_REPROVADO = "Reprovado"
STATUS_PENDENTE = "Pendente"


def _ler_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.is_file():
        return []
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def _texto(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _hash_resultado(registro: dict[str, Any]) -> str:
    bruto = json.dumps(registro.get("result") or {}, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(bruto.encode("utf-8")).hexdigest()


def _dimensoes(registro: dict[str, Any], secao: str) -> dict[str, str]:
    dimensoes: dict[str, str] = {}
    for conclusao in (registro.get("result") or {}).get("conclusoes") or []:
        codigo = _texto(conclusao.get("item_codigo")) or _texto(registro.get("codigo"))
        if secao == "2":
            estado = _texto(conclusao.get("estado"))
            if codigo and estado:
                dimensoes[f"item:{codigo}"] = estado
            continue
        temporal = _texto(conclusao.get("estado_temporal"))
        if temporal:
            dimensoes[f"situacao:{codigo}"] = temporal
        for motivo in conclusao.get("conclusoes_motivos") or []:
            if not isinstance(motivo, dict):
                continue
            motivo_id = _texto(motivo.get("id_motivo"))
            estado = _texto(motivo.get("estado_motivo"))
            if motivo_id and estado:
                dimensoes[f"motivo:{motivo_id}"] = estado
    return dimensoes


def _opinioes_vigentes(paths: list[Path]) -> dict[str, dict[str, dict[str, Any]]]:
    grupos: dict[str, dict[str, dict[str, Any]]] = defaultdict(dict)
    for path in paths:
        for registro in _ler_jsonl(path):
            if registro.get("status") != "completed" or not registro.get("case_id"):
                continue
            if not (registro.get("result") or {}).get("conclusoes"):
                continue
            model = _texto(registro.get("model_key")) or _texto(registro.get("model"))
            anterior = grupos[registro["case_id"]].get(model)
            if anterior is None or _texto(registro.get("finished_at")) >= _texto(anterior.get("finished_at")):
                grupos[registro["case_id"]][model] = registro
    return grupos


def detectar_prioridades(
    *,
    secoes: dict[str, tuple[list[Path], Path]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    """Retorna casos prioritários, divergências e universo consolidado."""
    prioridades: list[dict[str, Any]] = []
    divergencias: list[dict[str, Any]] = []
    universo: list[dict[str, Any]] = []
    for secao, (analyses, consolidado_path) in secoes.items():
        opinioes = _opinioes_vigentes(analyses)
        for parecer in _ler_jsonl(consolidado_path):
            if parecer.get("status") != "completed":
                continue
            cid = _texto(parecer.get("case_id"))
            identity = _texto(parecer.get("identity"))
            hash_resultado = _hash_resultado(parecer)
            origem = _texto(parecer.get("origem_decisao")) or "juiz_ia"
            motivos_prioridade: set[str] = set()
            votos_dimensao: dict[str, list[tuple[str, str]]] = defaultdict(list)
            for model, opiniao in opinioes.get(cid, {}).items():
                for dimensao, estado in _dimensoes(opiniao, secao).items():
                    votos_dimensao[dimensao].append((model, estado))
            juiz = _dimensoes(parecer, secao)
            for dimensao, votos in sorted(votos_dimensao.items()):
                contagem = Counter(estado for _, estado in votos)
                if len(contagem) < 2:
                    continue
                maximo = max(contagem.values())
                lideres = sorted(estado for estado, total in contagem.items() if total == maximo)
                tipo = ""
                if len(lideres) > 1:
                    tipo = "empate"
                elif maximo > len(votos) / 2 and origem == "juiz_ia" and juiz.get(dimensao) != lideres[0]:
                    tipo = "juiz_contra_maioria"
                if not tipo:
                    continue
                motivos_prioridade.add(tipo)
                divergencias.append({
                    "Seção": secao,
                    "Auditado": parecer.get("auditado", ""),
                    "Código": parecer.get("codigo", ""),
                    "Dimensão": dimensao,
                    "Tipo da prioridade": tipo,
                    "Votos": "; ".join(f"{model}={estado}" for model, estado in sorted(votos)),
                    "Contagem": "; ".join(f"{estado}={total}" for estado, total in sorted(contagem.items())),
                    "Maioria": lideres[0] if len(lideres) == 1 and maximo > len(votos) / 2 else "",
                    "Decisão do juiz": juiz.get(dimensao, ""),
                    "Case ID": cid,
                    "Identidade do parecer": identity,
                    "Hash do parecer": hash_resultado,
                })
            universo.append({
                "Seção": secao,
                "Auditado": parecer.get("auditado", ""),
                "Código": parecer.get("codigo", ""),
                "Case ID": cid,
                "Identidade do parecer": identity,
                "Hash do parecer": hash_resultado,
                "Origem da decisão": origem,
                "Prioritário": "Sim" if motivos_prioridade else "Não",
                "Motivos da prioridade": "; ".join(sorted(motivos_prioridade)),
            })
            if motivos_prioridade:
                revisao = parecer.get("revisao_humana") or {}
                prioridades.append({
                    "Seção": secao,
                    "Auditado": parecer.get("auditado", ""),
                    "Código": parecer.get("codigo", ""),
                    "Motivos da prioridade": "; ".join(sorted(motivos_prioridade)),
                    "Case ID": cid,
                    "Identidade do parecer": identity,
                    "Hash do parecer": hash_resultado,
                    "Origem da decisão": origem,
                    "Status da revisão": STATUS_APROVADO if origem == "revisao_humana" else STATUS_PENDENTE,
                    "Revisor": _texto(revisao.get("revisor")),
                    "Data da revisão": _texto(revisao.get("data_revisao")),
                    "Fundamento da revisão": _texto(revisao.get("fundamento")),
                })
    prioridades.sort(key=lambda item: (item["Seção"], item["Auditado"], item["Código"]))
    divergencias.sort(key=lambda item: (item["Seção"], item["Auditado"], item["Código"], item["Dimensão"]))
    universo.sort(key=lambda item: (item["Seção"], item["Auditado"], item["Código"]))
    return prioridades, divergencias, universo


def _aprovacoes_existentes(path: Path) -> dict[tuple[str, str, str], dict[str, Any]]:
    if not path.is_file():
        return {}
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Casos prioritários" not in wb.sheetnames:
            return {}
        ws = wb["Casos prioritários"]
        rows = ws.iter_rows(values_only=True)
        headers = [_texto(value) for value in next(rows)]
        existentes: dict[tuple[str, str, str], dict[str, Any]] = {}
        for values in rows:
            row = dict(zip(headers, values))
            key = (
                _texto(row.get("Case ID")),
                _texto(row.get("Identidade do parecer")),
                _texto(row.get("Hash do parecer")),
            )
            if all(key):
                existentes[key] = row
        return existentes
    finally:
        wb.close()


def _formatar(ws: Any) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column in ws.columns:
        letter = column[0].column_letter
        ws.column_dimensions[letter].width = min(70, max(12, max(len(_texto(cell.value)) for cell in column) + 2))


def sincronizar_planilha_revisao(
    *,
    path: Path,
    prioridades: list[dict[str, Any]],
    divergencias: list[dict[str, Any]],
    universo: list[dict[str, Any]],
) -> Path:
    existentes = _aprovacoes_existentes(path)
    for item in prioridades:
        key = (item["Case ID"], item["Identidade do parecer"], item["Hash do parecer"])
        anterior = existentes.get(key)
        if anterior:
            for campo in ("Status da revisão", "Revisor", "Data da revisão", "Fundamento da revisão"):
                item[campo] = _texto(anterior.get(campo)) or item[campo]
    wb = Workbook()
    wb.remove(wb.active)
    conjuntos = (
        ("Casos prioritários", prioridades),
        ("Divergências", divergencias),
        ("Universo consolidado", universo),
    )
    for nome, rows in conjuntos:
        ws = wb.create_sheet(nome)
        headers = list(rows[0]) if rows else ["Seção", "Auditado", "Código", "Case ID"]
        ws.append(headers)
        for row in rows:
            ws.append([row.get(header, "") for header in headers])
        _formatar(ws)
        if nome == "Casos prioritários" and "Status da revisão" in headers and ws.max_row >= 2:
            column = headers.index("Status da revisão") + 1
            validation = DataValidation(
                type="list",
                formula1=f'"{STATUS_PENDENTE},{STATUS_APROVADO},{STATUS_REPROVADO}"',
                allow_blank=False,
            )
            ws.add_data_validation(validation)
            validation.add(f"{ws.cell(1, column).column_letter}2:{ws.cell(1, column).column_letter}{ws.max_row}")
    controle = wb.create_sheet("Controle")
    pendentes = sum(item.get("Status da revisão") != STATUS_APROVADO for item in prioridades)
    controle.append(["Indicador", "Valor"])
    controle.append(["Pareceres consolidados", len(universo)])
    controle.append(["Casos prioritários", len(prioridades)])
    controle.append(["Prioridades aprovadas", len(prioridades) - pendentes])
    controle.append(["Prioridades pendentes", pendentes])
    _formatar(controle)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def validar_aprovacoes(path: Path, prioridades: list[dict[str, Any]]) -> dict[str, Any]:
    existentes = _aprovacoes_existentes(path)
    pendencias: list[dict[str, str]] = []
    for item in prioridades:
        key = (item["Case ID"], item["Identidade do parecer"], item["Hash do parecer"])
        row = existentes.get(key)
        motivo = ""
        if row is None:
            motivo = "aprovação ausente ou obsoleta"
        elif _texto(row.get("Status da revisão")) != STATUS_APROVADO:
            motivo = f"status={_texto(row.get('Status da revisão')) or STATUS_PENDENTE}"
        elif not _texto(row.get("Revisor")) or not _texto(row.get("Data da revisão")):
            motivo = "aprovação sem revisor ou data"
        if motivo:
            pendencias.append({
                "auditado": _texto(item.get("Auditado")),
                "secao": _texto(item.get("Seção")),
                "codigo": _texto(item.get("Código")),
                "case_id": item["Case ID"],
                "motivo": motivo,
            })
    return {
        "status": "approved" if not pendencias else "awaiting_review",
        "prioridades": len(prioridades),
        "aprovadas": len(prioridades) - len(pendencias),
        "pendentes": pendencias,
        "planilha": str(path),
    }
