#!/usr/bin/env python3
"""Renderiza relatorio Markdown/Jinja e converte o resultado para DOCX.

O fluxo reproduz o utilizado pelo webapp-streamlit-argos:

1. renderizacao Jinja2 com StrictUndefined e suporte a includes;
2. numeracao de figuras e tabelas e conversao de referencias cruzadas;
3. conversao de quebras de pagina e sublinhados para sintaxe Pandoc;
4. geracao DOCX com pypandoc e documento de referencia opcional;
5. ajuste de fonte, tamanho e alinhamento das tabelas com python-docx.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import sys
import tempfile
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any

try:
    import pypandoc
    from docx import Document
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
    from docx.shared import Pt
    from jinja2 import Environment, FileSystemLoader, StrictUndefined
except ImportError as exc:  # pragma: no cover - mensagem operacional
    raise SystemExit(
        "Dependencias ausentes. Instale-as com: "
        "python -m pip install -r scripts/requirements-relatorio.txt"
    ) from exc


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE = (
    ROOT
    / "03-Relatorios"
    / "02-Relatorios_Individuais_Preliminares"
    / "relatorio-individual-preliminar-template.md"
)
ARGOS_REFERENCE_DOC = Path(
    "/home/acba/workspace/webapp-streamlit-argos/docs/template-base-estilos-sigiloso.docx"
)


@dataclass
class Achado:
    numero: int | str
    nome: str
    situacoes_encontradas: list[str] = field(default_factory=list)
    evidencias: list[str] = field(default_factory=list)
    encaminhamentos: list[dict[str, str]] = field(default_factory=list)


@dataclass
class Auditado:
    nome: str
    sigla: str
    achados: list[Achado] = field(default_factory=list)

    @property
    def tem_achados(self) -> bool:
        return bool(self.achados)

    def get_achado_por_nome(self, nome_achado: str) -> Achado | None:
        return next((achado for achado in self.achados if achado.nome == nome_achado), None)

    def get_plano_acao(self) -> list[dict[str, str | int]]:
        itens: list[dict[str, str | int]] = []
        vistos: set[tuple[str, str, str]] = set()
        for achado in self.achados:
            for encaminhamento in achado.encaminhamentos:
                tipo = str(encaminhamento.get("tipo", "Recomendação"))
                texto = str(encaminhamento["encaminhamento"])
                chave = (str(achado.numero), tipo, texto)
                if chave in vistos:
                    continue
                vistos.add(chave)
                itens.append(
                    {
                        "achado_num": achado.numero,
                        "tipo": tipo,
                        "encaminhamento": texto,
                    }
                )
        return itens


def data_hoje_abnt() -> str:
    meses = {
        1: "jan.",
        2: "fev.",
        3: "mar.",
        4: "abr.",
        5: "maio",
        6: "jun.",
        7: "jul.",
        8: "ago.",
        9: "set.",
        10: "out.",
        11: "nov.",
        12: "dez.",
    }
    hoje = date.today()
    return f"{hoje.day} {meses[hoje.month]} {hoje.year}"


def data_hoje() -> str:
    return date.today().strftime("%d/%m/%Y")


def cross_ref_figuras(texto: str) -> str:
    mapa: dict[str, int] = {}
    for match in re.finditer(r"(?:\{#fig:([^#]+)#\}|\[@fig:([^\]]+)\])", texto):
        identificador = match.group(1) or match.group(2)
        if identificador not in mapa:
            mapa[identificador] = len(mapa) + 1

    def substituir_referencia(match: re.Match[str]) -> str:
        identificador = match.group(1)
        return f"Figura {mapa[identificador]}"

    texto = re.sub(r"\[@fig:([^\]]+)\]", substituir_referencia, texto)

    def numerar_figura(match: re.Match[str]) -> str:
        legenda, caminho, atributos, identificador = match.groups()
        atributos = atributos or ""
        return f"![Figura {mapa[identificador]} - {legenda}]{caminho}{atributos}"

    return re.sub(
        r"!\[([^\]]*)\](\([^)]*\))\s*(\{[^}]*\})?\s*\{#fig:([^#]+)#\}",
        numerar_figura,
        texto,
    )


def cross_ref_tabelas(texto: str) -> str:
    mapa: dict[str, int] = {}
    for match in re.finditer(r"(?:\{#tbl:([^#]+)#\}|\[@tbl:([^\]]+)\])", texto):
        identificador = match.group(1) or match.group(2)
        if identificador not in mapa:
            mapa[identificador] = len(mapa) + 1

    def substituir_referencia(match: re.Match[str]) -> str:
        identificador = match.group(1)
        return f"Tabela {mapa[identificador]}"

    texto = re.sub(r"\[@tbl:([^\]]+)\]", substituir_referencia, texto)

    def numerar_tabela(match: re.Match[str]) -> str:
        prefixo, legenda, identificador = match.groups()
        return f"{prefixo} Tabela {mapa[identificador]} - {legenda.strip()}"

    return re.sub(
        r"(?m)^(:|Table:)[ \t]*(.*?)[ \t]*\{#tbl:([^#]+)#\}[ \t]*$",
        numerar_tabela,
        texto,
    )


def processar_quebras_pagina(texto: str) -> str:
    openxml = (
        "\n```{=openxml}\n"
        '<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n'
        "```\n"
    )
    return re.sub(r"\\newpage(?:\{\})?", openxml, texto)


def substituir_underline_pandoc(texto: str) -> str:
    return re.sub(r"__(.+?)__", r"[\1]{.underline}", texto)


def aplicar_estilo_tabelas(
    docx_path: Path,
    font_name: str = "Calibri",
    header_size: int = 10,
    body_size: int = 9,
) -> None:
    documento = Document(docx_path)
    for tabela in documento.tables:
        tabela.autofit = True
        for indice, linha in enumerate(tabela.rows):
            tamanho = Pt(header_size if indice == 0 else body_size)
            for celula in linha.cells:
                for paragrafo in celula.paragraphs:
                    paragrafo.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
                    for trecho in paragrafo.runs:
                        trecho.font.name = font_name
                        trecho.font.size = tamanho
    documento.save(docx_path)


def carregar_contexto(path: Path) -> dict[str, Any]:
    dados = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(dados, dict):
        raise ValueError("O arquivo de contexto deve conter um objeto JSON na raiz.")
    return dados


def construir_contexto(dados: dict[str, Any]) -> dict[str, Any]:
    dados = dict(dados)
    auditado_dados = dados.pop("auditado")
    achados = [Achado(**item) for item in auditado_dados.get("achados", [])]
    auditado = Auditado(
        nome=auditado_dados["nome"],
        sigla=auditado_dados["sigla"],
        achados=achados,
    )
    return {
        **dados,
        "auditado": auditado,
        "data_hoje": data_hoje(),
        "data_hoje_abnt": data_hoje_abnt(),
    }


def pior_caso() -> dict[str, Any]:
    definicoes = [
        (
            1,
            "Estrutura de TIC insuficiente para coordenar, gerir e sustentar a tecnologia da informação.",
            [
                "Ausência de formalização da área, unidade, setor ou função de TIC da organização.",
                "Área de TIC sem atribuições formais suficientes para planejamento, coordenação, gestão, execução, monitoramento e controle da TIC.",
                "Posicionamento organizacional inadequado da área de TIC.",
            ],
            "formalize a estrutura de TIC, suas atribuições e seu posicionamento organizacional.",
        ),
        (
            2,
            "Governança de TIC insuficiente para avaliar, dirigir e monitorar a tecnologia da informação.",
            [
                "Ausência ou insuficiência de modelo básico de governança e gestão de TIC com papéis, responsabilidades, objetivos, indicadores, metas ou acompanhamento.",
                "Comitê de TIC ou instância equivalente não instituído formalmente.",
                "Comitê de TIC ou instância equivalente sem evidências suficientes de atuação efetiva.",
            ],
            "institua e mantenha modelo de governança e comitê de TIC com atuação efetiva.",
        ),
        (
            3,
            "Planejamento de TIC inexistente, insuficiente, desatualizado ou desconectado da gestão, do orçamento e das contratações",
            [
                "Inexistência ou fragilidade do processo formal de planejamento de TIC.",
                "Ausência de aprovação formal do plano de TIC.",
                "Plano de TIC sem alinhamento adequado ao planejamento institucional.",
                "Plano de TIC sem integração adequada com orçamento, plano de contratações, projetos ou contratações de TIC.",
                "Ausência de acompanhamento, revisão ou atualização periódica do plano de TIC.",
            ],
            "formalize, aprove, execute e acompanhe plano de TIC alinhado à estratégia, ao orçamento e às contratações.",
        ),
        (
            4,
            "Capacidade institucional insuficiente para sustentar a gestão de TIC e segurança da informação",
            [
                "Ausência de força de trabalho mínima dedicada à TIC ou à segurança da informação.",
                "A organização não definiu o quantitativo necessário de pessoal de TIC e segurança da informação.",
                "Ausência de cargos, funções, perfis ou ocupações específicas de TIC e segurança da informação.",
                "Perfis profissionais de TIC e segurança da informação inexistentes, insuficientes ou não utilizados na escolha de gestores.",
                "Lacunas de competências dos colaboradores e gestores de TIC e segurança da informação não são identificadas ou tratadas.",
                "Modelo de operação de TIC predominantemente terceirizado ou externo, sem capacidade interna mínima declarada para coordenação, planejamento, aprovação técnica ou fiscalização das atividades críticas de TIC.",
            ],
            "dimensione e desenvolva capacidade institucional mínima de TIC e segurança da informação.",
        ),
        (
            5,
            "Gestão de serviços de TIC incipiente, sem controle mínimo sobre serviços, ativos e incidentes",
            [
                "Inexistência, desatualização, indisponibilidade ou insuficiência do catálogo de serviços de TIC.",
                "Inexistência de ANS, metas mínimas ou monitoramento de níveis de serviço para os principais serviços de TIC.",
                "Inexistência ou fragilidade do inventário de ativos de TIC.",
                "Ausência ou fragilidade do processo de gestão de configuração.",
                "Inexistência ou fragilidade do processo de gestão de incidentes de TIC.",
            ],
            "implemente processos mínimos de gestão de serviços, ativos, configuração e incidentes de TIC.",
        ),
        (
            6,
            "Contratações de TIC sem governança técnica e controle de resultados",
            [
                "Inexistência ou fragilidade de processo formal e padronizado para contratações de TIC.",
                "Contratações de TIC sem análise prévia e aprovação técnica obrigatória da área de TIC.",
                "Contratações de TIC sem aderência ao plano de TIC, ao plano de contratações ou à proposta orçamentária.",
                "Contratações de TIC sem equipe de planejamento formalmente designada e com participação técnica de TIC.",
                "Artefatos de planejamento das contratações de TIC sem requisitos técnicos, análise de riscos, segurança da informação, proteção de dados ou critérios objetivos de aceite.",
                "Contratações de TIC sem níveis mínimos de serviço, métricas de desempenho ou critérios objetivos de fiscalização e recebimento.",
            ],
            "formalize a governança técnica das contratações de TIC e vincule pagamentos a resultados mensuráveis.",
        ),
    ]

    achados = []
    for numero, nome, situacoes, encaminhamento in definicoes:
        achados.append(
            {
                "numero": numero,
                "nome": nome,
                "situacoes_encontradas": situacoes,
                "evidencias": [
                    "Resposta declarada no questionário iGovTI 2026 indicando não adoção das práticas avaliadas",
                    "Ausência de documentação comprobatória suficiente nos arquivos encaminhados pelo auditado",
                    "Resultado dos procedimentos de auditoria aplicáveis à questão",
                ],
                "encaminhamentos": [
                    {
                        "tipo": "Recomendação",
                        "encaminhamento": encaminhamento,
                    }
                ],
            }
        )

    return {
        "auditado": {
            "nome": "Organização Pública de Teste - Pior Caso",
            "sigla": "PIORCASO",
            "achados": achados,
        },
        "iGovTI": 0.0,
        "iGovTI_maturidade": "Inexpressivo",
        "GovernancaTI": 0.0,
        "GovernancaTI_maturidade": "Inexpressivo",
        "iGestTI": 0.0,
        "iGestTI_maturidade": "Inexpressivo",
        "teve_ajuste": True,
        "ajustes_respostas": [
            {
                "codigo_questao": "q1001",
                "de": "Adota em maior parte ou totalmente",
                "para": "Não adota",
                "justificativa": "Não foram apresentadas evidências suficientes da prática declarada.",
            },
            {
                "codigo_questao": "q2102",
                "de": "Adota parcialmente",
                "para": "Não adota",
                "justificativa": "Não foi apresentado plano de TIC vigente e formalmente aprovado.",
            },
        ],
    }


def criar_imagens_teste(template_dir: Path) -> None:
    fontes = {
        "PIORCASO_comparativo_distribuicao_iGovTI.png": "igovti_2026_distribuicao_maturidade.png",
        "PIORCASO_componentes_iGovTI.png": "igovti_2026_distribuicao_componentes.png",
        "PIORCASO_comparativo_distribuicao_GovernancaTI.png": "igovti_2026_distribuicao_componentes.png",
        "PIORCASO_comparativo_distribuicao_iGestTI.png": "igovti_2026_distribuicao_componentes.png",
        "PIORCASO_perfil_dimensoes_iGestTI.png": "igovti_2026_distribuicao_dimensoes_gestao.png",
    }
    for destino, fonte in fontes.items():
        origem = template_dir / fonte
        caminho_destino = template_dir / destino
        if origem.exists() and not caminho_destino.exists():
            shutil.copyfile(origem, caminho_destino)


def renderizar(template_path: Path, contexto: dict[str, Any]) -> str:
    ambiente = Environment(
        loader=FileSystemLoader(str(template_path.parent)),
        undefined=StrictUndefined,
        autoescape=False,
        keep_trailing_newline=True,
        comment_start_string="{##",
        comment_end_string="##}",
    )
    template = ambiente.get_template(template_path.name)
    texto = template.render(contexto)
    texto = cross_ref_figuras(texto)
    texto = cross_ref_tabelas(texto)
    texto = processar_quebras_pagina(texto)
    return substituir_underline_pandoc(texto).rstrip() + "\n"


def converter_docx(
    markdown_path: Path,
    output_path: Path,
    reference_doc: Path | None,
    resource_paths: list[Path],
) -> None:
    argumentos = ["--figure-caption-position=above"]
    if reference_doc:
        argumentos.append(f"--reference-doc={reference_doc}")
    argumentos.append(f"--resource-path={os.pathsep.join(str(path) for path in resource_paths)}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    resultado = pypandoc.convert_file(
        str(markdown_path),
        to="docx",
        format="markdown",
        outputfile=str(output_path),
        extra_args=argumentos,
    )
    if resultado:
        raise RuntimeError(f"Pandoc retornou saída inesperada: {resultado}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE, help="Template Markdown/Jinja principal.")
    grupo = parser.add_mutually_exclusive_group(required=True)
    grupo.add_argument("--context", type=Path, help="Arquivo JSON com o contexto do relatório.")
    grupo.add_argument("--worst-case", action="store_true", help="Usa um auditado com resultado zero e todos os achados/situações.")
    parser.add_argument("--output", type=Path, required=True, help="Arquivo DOCX de saída.")
    parser.add_argument("--rendered-md", type=Path, help="Caminho opcional para preservar o Markdown renderizado.")
    parser.add_argument("--export-context", type=Path, help="Exporta o contexto efetivamente utilizado em JSON.")
    parser.add_argument("--reference-doc", type=Path, help="Documento DOCX de referência para estilos Pandoc.")
    parser.add_argument("--resource-path", type=Path, action="append", default=[], help="Diretório adicional de imagens; pode ser repetido.")
    parser.add_argument("--font-name", default="Calibri", help="Fonte das tabelas após a conversão.")
    parser.add_argument("--header-size", type=int, default=10, help="Tamanho da fonte do cabeçalho das tabelas.")
    parser.add_argument("--body-size", type=int, default=9, help="Tamanho da fonte do corpo das tabelas.")
    parser.add_argument("--skip-table-style", action="store_true", help="Não aplica o pós-processamento das tabelas.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    template_path = args.template.resolve()
    if not template_path.exists():
        raise FileNotFoundError(f"Template não encontrado: {template_path}")

    dados = pior_caso() if args.worst_case else carregar_contexto(args.context.resolve())
    if args.export_context:
        args.export_context.parent.mkdir(parents=True, exist_ok=True)
        args.export_context.write_text(json.dumps(dados, ensure_ascii=False, indent=2), encoding="utf-8")

    contexto = construir_contexto(dados)
    if args.worst_case:
        criar_imagens_teste(template_path.parent)

    reference_doc = args.reference_doc
    if reference_doc is None and ARGOS_REFERENCE_DOC.exists():
        reference_doc = ARGOS_REFERENCE_DOC
    if reference_doc is not None:
        reference_doc = reference_doc.resolve()
        if not reference_doc.exists():
            raise FileNotFoundError(f"Documento de referência não encontrado: {reference_doc}")

    markdown = renderizar(template_path, contexto)
    output_path = args.output.resolve()
    rendered_path = args.rendered_md.resolve() if args.rendered_md else output_path.with_suffix(".rendered.md")
    rendered_path.parent.mkdir(parents=True, exist_ok=True)
    rendered_path.write_text(markdown, encoding="utf-8")

    resource_paths = [template_path.parent, rendered_path.parent, Path.cwd(), *args.resource_path]
    resource_paths = [path.resolve() for path in resource_paths]
    converter_docx(rendered_path, output_path, reference_doc, resource_paths)

    if not args.skip_table_style:
        aplicar_estilo_tabelas(
            output_path,
            font_name=args.font_name,
            header_size=args.header_size,
            body_size=args.body_size,
        )

    print(f"Relatório gerado: {output_path}")
    print(f"Markdown renderizado: {rendered_path}")
    if reference_doc:
        print(f"Documento de referência: {reference_doc}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"Erro: {exc}", file=sys.stderr)
        raise SystemExit(1) from exc
