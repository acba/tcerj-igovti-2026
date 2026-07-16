"""Utilitários para gravar XLSX sem churn binário desnecessário.

Arquivos XLSX podem mudar em bytes por metadados internos mesmo quando as
células permanecem iguais. Estes helpers gravam em arquivo temporário e só
substituem o destino quando há diferença lógica no conteúdo das células.
"""

from __future__ import annotations

import math
import os
import tempfile
from pathlib import Path
from typing import Callable

import pandas as pd
from openpyxl import load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE


def _normalizar_valor_celula(valor):
    if valor is None:
        return None
    if isinstance(valor, float) and math.isnan(valor):
        return None
    if pd.isna(valor) and not isinstance(valor, (str, bytes)):
        return None
    return valor


def sanitizar_workbook_para_excel(workbook) -> None:
    """Remove controles inválidos para XML somente do artefato XLSX."""
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    cell.value = ILLEGAL_CHARACTERS_RE.sub("", cell.value)


def xlsx_logicamente_iguais(a: Path | str, b: Path | str) -> bool:
    """Compara duas planilhas XLSX por abas, dimensões e valores/fórmulas."""
    caminho_a = Path(a)
    caminho_b = Path(b)
    if not caminho_a.exists() or not caminho_b.exists():
        return False

    wb_a = load_workbook(caminho_a, read_only=True, data_only=False)
    wb_b = load_workbook(caminho_b, read_only=True, data_only=False)
    try:
        if wb_a.sheetnames != wb_b.sheetnames:
            return False

        for nome_aba in wb_a.sheetnames:
            ws_a = wb_a[nome_aba]
            ws_b = wb_b[nome_aba]
            if ws_a.max_row != ws_b.max_row or ws_a.max_column != ws_b.max_column:
                return False

            rows_a = ws_a.iter_rows(values_only=True)
            rows_b = ws_b.iter_rows(values_only=True)
            for row_a, row_b in zip(rows_a, rows_b):
                if len(row_a) != len(row_b):
                    return False
                for valor_a, valor_b in zip(row_a, row_b):
                    if _normalizar_valor_celula(valor_a) != _normalizar_valor_celula(valor_b):
                        return False

        return True
    finally:
        wb_a.close()
        wb_b.close()


def substituir_xlsx_se_diferente(temporario: Path | str, destino: Path | str) -> bool:
    """Substitui o destino apenas quando o XLSX temporário for diferente.

    Retorna ``True`` quando o destino foi criado/substituído e ``False`` quando
    o conteúdo lógico era igual e o arquivo existente foi preservado.
    """
    temp_path = Path(temporario)
    destino_path = Path(destino)
    destino_path.parent.mkdir(parents=True, exist_ok=True)

    if destino_path.exists() and xlsx_logicamente_iguais(temp_path, destino_path):
        temp_path.unlink()
        return False

    os.replace(temp_path, destino_path)
    return True


def escrever_xlsx_se_diferente(destino: Path | str, writer: Callable[[Path], None]) -> bool:
    """Executa ``writer(temp_path)`` e substitui ``destino`` se houver diferença."""
    destino_path = Path(destino)
    destino_path.parent.mkdir(parents=True, exist_ok=True)

    with tempfile.NamedTemporaryFile(
        suffix=".xlsx",
        prefix=f".{destino_path.stem}-",
        dir=destino_path.parent,
        delete=False,
    ) as handle:
        temp_path = Path(handle.name)

    try:
        writer(temp_path)
        return substituir_xlsx_se_diferente(temp_path, destino_path)
    except Exception:
        if temp_path.exists():
            temp_path.unlink()
        raise


def dataframe_to_xlsx_se_diferente(
    df: pd.DataFrame,
    destino: Path | str,
    *,
    sheet_name: str = "Sheet1",
    index: bool = False,
    engine: str = "openpyxl",
) -> bool:
    """Grava um DataFrame em XLSX apenas se o conteúdo lógico mudar."""

    def _writer(temp_path: Path) -> None:
        df.to_excel(temp_path, sheet_name=sheet_name, index=index, engine=engine)

    return escrever_xlsx_se_diferente(destino, _writer)
