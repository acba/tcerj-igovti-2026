#!/usr/bin/env python3
"""Sincroniza a camada jurídica da matriz com o mapa de verificação."""

from __future__ import annotations

import argparse
from copy import copy
from pathlib import Path
import sys

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent.parent
RESOURCES = ROOT / "scripts/resources"
if str(RESOURCES) not in sys.path:
    sys.path.insert(0, str(RESOURCES))
if str(ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(ROOT / "scripts"))

from matriz_aplicabilidade import carregar_catalogo_matriz  # noqa: E402
from gerar_matriz_planejamento import parse_matrix  # noqa: E402


DEFAULT_MATRIZ = ROOT / "01-Planejamento/03-Estrategia_e_Plano/04-Matriz_Planejamento/matriz_planejamento-pos-comentarios-gestor.md"
DEFAULT_MAPA = ROOT / "02-Execucao/03-Execucao_Procedimentos/01-Insumos/mapa-verificacao-achados-pos-comentarios-gestor.xlsx"


def _lista_texto(valores) -> str:
    return ";".join(valores)


def _preparar_aba(workbook, nome: str, titulo: str, cabecalhos: list[str]):
    if nome in workbook.sheetnames:
        del workbook[nome]
    ws = workbook.create_sheet(nome)
    ws.cell(1, 1, titulo)
    ws.cell(1, 1).font = Font(bold=True, size=14, color="FFFFFF")
    ws.cell(1, 1).fill = PatternFill("solid", fgColor="1F4E78")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cabecalhos))
    for coluna, cabecalho in enumerate(cabecalhos, 1):
        cell = ws.cell(3, coluna, cabecalho)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{ws.cell(3, len(cabecalhos)).coordinate}"
    return ws


def _situacoes_por_descricao(matriz_path: Path) -> dict[str, str]:
    matriz = parse_matrix(matriz_path.read_text(encoding="utf-8"))
    resultado = {}
    for questao in matriz.questions:
        for achado in questao.achados:
            for situacao in achado.situacoes:
                chave = situacao.descricao.strip().rstrip(".").casefold()
                if chave in resultado and resultado[chave] != situacao.id:
                    raise ValueError(f"Descrição de situação duplicada na matriz: {situacao.descricao}")
                resultado[chave] = situacao.id
    return resultado


def sincronizar(matriz_path: Path, mapa_path: Path, output_path: Path) -> None:
    catalogo = carregar_catalogo_matriz(matriz_path)
    descricoes = _situacoes_por_descricao(matriz_path)
    workbook = load_workbook(mapa_path)
    acoes = workbook["Ações de Verificação"]
    cabecalhos = {str(cell.value).strip(): cell.column for cell in acoes[3] if cell.value}
    if "id_situacao" not in cabecalhos:
        coluna_referencia = cabecalhos["descricao_situacao_inconforme"]
        acoes.insert_cols(coluna_referencia + 1)
        acoes.cell(3, coluna_referencia + 1, "id_situacao")
        origem = acoes.cell(3, coluna_referencia)
        destino = acoes.cell(3, coluna_referencia + 1)
        for atributo in ("font", "fill", "border", "alignment", "protection"):
            setattr(destino, atributo, copy(getattr(origem, atributo)))
        destino.number_format = origem.number_format
        cabecalhos = {str(cell.value).strip(): cell.column for cell in acoes[3] if cell.value}

    col_descricao = cabecalhos["descricao_situacao_inconforme"]
    col_id = cabecalhos["id_situacao"]
    ids_catalogo = {variante.id_situacao for variante in catalogo.variantes}
    for linha in range(4, acoes.max_row + 1):
        descricao = str(acoes.cell(linha, col_descricao).value or "").strip()
        if not descricao:
            continue
        chave = descricao.rstrip(".").casefold()
        id_situacao = descricoes.get(chave)
        if not id_situacao:
            raise ValueError(f"Ação {acoes.cell(linha, 1).value}: situação não localizada na matriz: {descricao}")
        if id_situacao not in ids_catalogo:
            raise ValueError(f"{id_situacao}: situação sem variante jurídica.")
        acoes.cell(linha, col_id, id_situacao)

    headers_criterios = [
        "id_criterio", "id_exibicao", "questao", "publico", "descricao", "natureza_fundamento",
        "apto_a_fundamentar_determinacao", "segmentos", "naturezas", "tags_todas",
        "tags_alguma", "tags_excluidas",
    ]
    ws_criterios = _preparar_aba(
        workbook, "Critérios de Auditoria", "Critérios de auditoria sincronizados da matriz", headers_criterios
    )
    for linha, criterio in enumerate(catalogo.criterios, 4):
        questao, id_exibicao = criterio.id.split(".", 1)
        valores = [
            criterio.id, id_exibicao, questao, criterio.rotulo_publico, criterio.descricao,
            criterio.natureza_fundamento,
            criterio.apto_a_fundamentar_determinacao,
            _lista_texto(criterio.seletor.segmentos), _lista_texto(criterio.seletor.naturezas),
            _lista_texto(criterio.seletor.tags_todas), _lista_texto(criterio.seletor.tags_alguma),
            _lista_texto(criterio.seletor.tags_excluidas),
        ]
        for coluna, valor in enumerate(valores, 1):
            ws_criterios.cell(linha, coluna, valor)

    headers_variantes = [
        "id_variante", "id_situacao", "geral", "publico", "segmentos", "naturezas",
        "tags_todas", "tags_alguma", "tags_excluidas", "criterios", "tipo_encaminhamento",
        "encaminhamento",
    ]
    ws_variantes = _preparar_aba(
        workbook, "Variantes de Encaminhamento", "Variantes jurídicas sincronizadas da matriz", headers_variantes
    )
    for linha, variante in enumerate(catalogo.variantes, 4):
        valores = [
            variante.id, variante.id_situacao, variante.geral, variante.rotulo_publico,
            _lista_texto(variante.seletor.segmentos), _lista_texto(variante.seletor.naturezas),
            _lista_texto(variante.seletor.tags_todas), _lista_texto(variante.seletor.tags_alguma),
            _lista_texto(variante.seletor.tags_excluidas), _lista_texto(variante.criterios),
            variante.tipo_encaminhamento, variante.encaminhamento,
        ]
        for coluna, valor in enumerate(valores, 1):
            ws_variantes.cell(linha, coluna, valor)

    for ws in (ws_criterios, ws_variantes):
        for indice, column in enumerate(ws.columns, 1):
            letra = get_column_letter(indice)
            maior = max(len(str(cell.value or "")) for cell in column)
            ws.column_dimensions[letra].width = min(max(maior + 2, 12), 60)
        for row in ws.iter_rows(min_row=4):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--matriz", type=Path, default=DEFAULT_MATRIZ)
    parser.add_argument("--mapa", type=Path, default=DEFAULT_MAPA)
    parser.add_argument("--output", type=Path, default=DEFAULT_MAPA)
    args = parser.parse_args()
    sincronizar(args.matriz, args.mapa, args.output)
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
