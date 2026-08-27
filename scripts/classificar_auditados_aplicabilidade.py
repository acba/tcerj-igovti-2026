#!/usr/bin/env python3
"""Materializa classificações explícitas usadas pela aplicabilidade jurídica."""

from __future__ import annotations

import argparse
from copy import copy
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent.parent
DEFAULT_BD = ROOT / "02-Execucao/03-Execucao_Procedimentos/01-Insumos/bd_auditados.xlsx"

SEGMENTOS_ESPECIAIS = {
    "ALERJ": "LEGISLATIVO_ESTADUAL",
    "DPGE": "DEFENSORIA_PUBLICA_ESTADUAL",
    "MPERJ": "MINISTERIO_PUBLICO_ESTADUAL",
    "TCE-RJ": "CONTROLE_EXTERNO_ESTADUAL",
    "TJRJ": "JUDICIARIO_ESTADUAL",
}

ADMINISTRACAO_DIRETA_ESTADUAL = {
    "CGE", "DEGASE", "GSI", "PGE", "SEAP", "SEAPPA", "SECC", "SECEC",
    "SECID", "SECTI", "SEDCON", "SEDEC", "SEDEICS", "SEDSDH", "SEEDUC",
    "SEEL", "SEENEMAR", "SEFAZ", "SEGOV", "SEHAB", "SEIJES", "SEINFRA",
    "SEPLAG", "SEPM", "SEPOL", "SERGB", "SES", "SESP", "SETD", "SETRAB",
    "SETRANS", "SETUR",
}

AUTARQUIAS = {
    "AGENERSA", "AGETRANSP", "DERRJ", "DETRAN", "DETRO", "DRM", "IEEA", "INEA",
    "IPEM", "IRM", "ISP", "ITERJ", "JUCERJA", "LOTERJ", "PROCON", "PRODERJ",
    "RIOPREVIDENCIA", "SUDERJ", "UENF", "UERJ",
}
FUNDACOES = {
    "CECIERJ", "CEPERJ", "FAETEC", "FAPERJ", "FIA", "FIPERJ", "FLXIII", "FMIS",
    "FS", "FSC", "FTM", "FUNARJ", "RJPREV",
}
EMPRESAS_ESTATAIS = {
    "AGERIO", "CEASA", "CEDAE", "CEHAB", "CENTRAL", "CODERTE", "CODIN", "EMATER",
    "EMOP", "IOERJ", "IVB", "PESAGRO", "RIOTRILHOS", "TURISRIO",
}

PREFEITURAS_MUNICIPAIS = {
    "ANGRA DOS REIS", "ARARUAMA", "ARMAÇÃO DOS BÚZIOS", "ARRAIAL DO CABO",
    "BARRA DO PIRAÍ", "BELFORD ROXO", "CABO FRIO", "CAMPOS DOS GOYTACAZES",
    "CASIMIRO DE ABREU", "DUQUE DE CAXIAS", "GUAPIMIRIM", "ITAGUAÍ", "JAPERI",
    "MACAÉ", "MAGÉ", "MARICÁ", "MESQUITA", "NITERÓI", "NOVA FRIBURGO",
    "NOVA IGUAÇU", "PARATY", "PETRÓPOLIS", "PORTO REAL", "QUATIS", "QUEIMADOS",
    "QUISSAMÃ", "RIO DAS OSTRAS", "SAQUAREMA", "SEROPÉDICA", "SÃO GONÇALO",
    "SÃO JOÃO DA BARRA", "SÃO JOÃO DE MERITI", "SÃO PEDRO DA ALDEIA",
    "TERESÓPOLIS", "VOLTA REDONDA",
}


def classificar(sigla: str, esfera: str) -> tuple[str, str, str]:
    sigla = sigla.strip().upper()
    esfera = esfera.strip().upper()
    if sigla in PREFEITURAS_MUNICIPAIS and esfera == "M":
        return "EXECUTIVO_MUNICIPAL", "ADMINISTRACAO_DIRETA", ""
    if sigla in SEGMENTOS_ESPECIAIS and esfera == "E":
        segmento = SEGMENTOS_ESPECIAIS[sigla]
        tag = {
            "TJRJ": "CNJ",
            "MPERJ": "CNMP",
        }.get(sigla, "")
        return segmento, "ORGAO_AUTONOMO", tag
    if sigla in ADMINISTRACAO_DIRETA_ESTADUAL and esfera == "E":
        natureza = "ADMINISTRACAO_DIRETA"
    elif sigla in AUTARQUIAS and esfera == "E":
        natureza = "AUTARQUIA"
    elif sigla in FUNDACOES and esfera == "E":
        natureza = "FUNDACAO"
    elif sigla in EMPRESAS_ESTATAIS and esfera == "E":
        natureza = "EMPRESA_ESTATAL"
    else:
        raise ValueError(
            f"{sigla}: classificação não declarada explicitamente para a esfera {esfera!r}."
        )
    return "EXECUTIVO_ESTADUAL", natureza, "SETIC"


def materializar(entrada: Path, saida: Path) -> None:
    workbook = load_workbook(entrada)
    ws = workbook["auditados"] if "auditados" in workbook.sheetnames else workbook.active
    headers = {str(cell.value).strip(): cell.column for cell in ws[1] if cell.value}
    for obrigatoria in ("sigla", "esfera"):
        if obrigatoria not in headers:
            raise ValueError(f"Coluna obrigatória ausente: {obrigatoria}")
    novas = ["segmento_institucional", "natureza_administrativa", "tags_aplicabilidade"]
    modelo = ws.cell(1, max(headers.values()))
    for nome in novas:
        if nome not in headers:
            coluna = ws.max_column + 1
            cell = ws.cell(1, coluna, nome)
            for atributo in ("font", "fill", "border", "alignment", "protection"):
                setattr(cell, atributo, copy(getattr(modelo, atributo)))
            cell.number_format = modelo.number_format
            headers[nome] = coluna
    for linha in range(2, ws.max_row + 1):
        sigla = str(ws.cell(linha, headers["sigla"]).value or "").strip()
        esfera = str(ws.cell(linha, headers["esfera"]).value or "").strip()
        if not sigla:
            continue
        segmento, natureza, tags = classificar(sigla, esfera)
        ws.cell(linha, headers["segmento_institucional"], segmento)
        ws.cell(linha, headers["natureza_administrativa"], natureza)
        ws.cell(linha, headers["tags_aplicabilidade"], tags)
    for nome in novas:
        ws.column_dimensions[ws.cell(1, headers[nome]).column_letter].width = 30
    saida.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(saida)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--entrada", type=Path, default=DEFAULT_BD)
    parser.add_argument("--saida", type=Path, default=DEFAULT_BD)
    args = parser.parse_args()
    materializar(args.entrada, args.saida)
    print(args.saida)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
