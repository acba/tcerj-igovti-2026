"""Preflight de integridade e manifesto de reparo dos comentários do gestor."""
from __future__ import annotations

import json
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from scripts.avaliacao_evidencias.scope_validation import formatar_violacoes, validar_resultado_no_escopo
from scripts.comentarios_gestor_pipeline import (
    registro_tem_evidencia_documental,
    validar_justificativa_publicavel,
)


def _ler_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.is_file():
        return []
    return [json.loads(linha) for linha in path.read_text(encoding="utf-8").splitlines() if linha.strip()]


def _mais_recente(registros: list[dict[str, Any]], chave_modelo: bool) -> dict[tuple[str, str], dict[str, Any]]:
    atuais: dict[tuple[str, str], dict[str, Any]] = {}
    for registro in registros:
        cid = str(registro.get("case_id") or "")
        modelo = str(registro.get("model_key") or registro.get("model") or "") if chave_modelo else ""
        if not cid:
            continue
        chave = (cid, modelo)
        anterior = atuais.get(chave)
        if anterior is None or str(registro.get("finished_at") or "") >= str(anterior.get("finished_at") or ""):
            atuais[chave] = registro
    return atuais


def _agrupar_registros(
    registros: list[dict[str, Any]],
    *,
    chave_modelo: bool,
) -> dict[tuple[str, str], list[dict[str, Any]]]:
    grupos: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for registro in registros:
        cid = str(registro.get("case_id") or "")
        modelo = str(registro.get("model_key") or registro.get("model") or "") if chave_modelo else ""
        if cid:
            grupos[(cid, modelo)].append(registro)
    for candidatos in grupos.values():
        candidatos.sort(key=lambda item: str(item.get("finished_at") or ""), reverse=True)
    return grupos


def _selecionar_registro_valido(
    candidatos: list[dict[str, Any]],
    *,
    caso: dict[str, Any],
    validar_temporal: bool,
    validar_publicacao_secao: str | None = None,
    preservar_valido_anterior: bool = True,
) -> tuple[dict[str, Any] | None, str]:
    """Preserva o registro valido mais recente diante de falhas posteriores."""
    for registro in candidatos:
        if registro.get("status") != "completed":
            continue
        violacoes = formatar_violacoes(
            validar_resultado_no_escopo(
                {**caso, "result": registro.get("result")},
                registro.get("result"),
                validar_temporal=validar_temporal,
            )
        )
        publicacao = (
            "; ".join(
                validar_justificativa_publicavel(
                    validar_publicacao_secao,
                    registro.get("result") or {},
                    evidencia_documental_disponivel=registro_tem_evidencia_documental(registro),
                )
            )
            if validar_publicacao_secao
            else ""
        )
        violacao = " | ".join(parte for parte in (violacoes, publicacao) if parte)
        if not violacao:
            return registro, ""
        # Respostas externas ao caso continuam podendo ceder lugar ao último
        # checkpoint de escopo válido. Já um parecer do próprio caso com texto
        # impróprio para publicação precisa ser reparado, não mascarado por uma
        # consolidação histórica anterior.
        if not preservar_valido_anterior and not violacoes and publicacao:
            return registro, violacao

    if not candidatos:
        return None, ""
    registro = candidatos[0]
    if registro.get("status") != "completed":
        return registro, f"status {registro.get('status')}: {registro.get('error', '')}"
    violacao = formatar_violacoes(
        validar_resultado_no_escopo(
            {**caso, "result": registro.get("result")},
            registro.get("result"),
            validar_temporal=validar_temporal,
        )
    )
    return registro, violacao


def _gravar_xlsx(path: Path, linhas: list[dict[str, Any]], resumo: dict[str, Any]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Violações"
    headers = ["Seção", "Nível", "Auditado", "Código", "Case ID", "Modelo", "Status", "Violação"]
    ws.append(headers)
    for linha in linhas:
        ws.append([linha.get(campo, "") for campo in headers])
    rs = wb.create_sheet("Resumo")
    rs.append(["Indicador", "Valor"])
    for chave, valor in resumo.items():
        if not isinstance(valor, (dict, list)):
            rs.append([chave, valor])
    for aba in wb.worksheets:
        aba.freeze_panes = "A2"
        aba.auto_filter.ref = aba.dimensions
        for cell in aba[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for row in aba.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def validar_integridade_comentarios(
    *,
    out_dir: Path,
    casos_por_secao: dict[str, list[dict[str, Any]]],
    deterministicos_por_secao: dict[str, list[dict[str, Any]]],
    modelos_esperados: list[str],
    quorum: int,
) -> dict[str, Any]:
    """Valida avaliações e pareceres contra os casos preparados e grava memória auditável."""
    linhas: list[dict[str, Any]] = []
    casos_reprocessamento: dict[str, list[dict[str, Any]]] = {"1": [], "2": []}
    total_casos = total_validos = 0

    for secao in ("1", "2"):
        casos = casos_por_secao.get(secao, [])
        deterministas = {str(item.get("case_id")) for item in deterministicos_por_secao.get(secao, [])}
        por_id = {str(caso["case_id"]): caso for caso in casos}
        total_casos += len(casos) + len(deterministas)
        individual_dir = out_dir / "individuais" / ("secao-1" if secao == "1" else "secao-2")
        registros_individuais: list[dict[str, Any]] = []
        for arquivo in individual_dir.glob("analyses_model_*.jsonl"):
            registros_individuais.extend(_ler_jsonl(arquivo))
        registros_por_caso_modelo = _agrupar_registros(registros_individuais, chave_modelo=True)
        afetados: set[str] = set()
        validos_por_caso: dict[str, int] = defaultdict(int)
        for cid, caso in por_id.items():
            for modelo in modelos_esperados:
                registro, violacao = _selecionar_registro_valido(
                    registros_por_caso_modelo.get((cid, modelo), []),
                    caso=caso,
                    validar_temporal=secao == "1",
                )
                if registro is None:
                    # A configuração define o conjunto possível de avaliadores,
                    # mas o método admite consolidação com quórum. Ausência isolada
                    # não torna o caso inválido; será tratada pela regra de quórum.
                    continue
                if violacao:
                    linhas.append({
                        "Seção": secao, "Nível": "avaliação individual",
                        "Auditado": caso.get("auditado", ""), "Código": caso.get("codigo", ""),
                        "Case ID": cid, "Modelo": modelo,
                        "Status": registro.get("status", "ausente") if registro else "ausente",
                        "Violação": violacao,
                    })
                else:
                    validos_por_caso[cid] += 1
            if validos_por_caso[cid] < quorum:
                afetados.add(cid)
                linhas.append({
                    "Seção": secao, "Nível": "quórum",
                    "Auditado": caso.get("auditado", ""), "Código": caso.get("codigo", ""),
                    "Case ID": cid, "Modelo": "", "Status": "insuficiente",
                    "Violação": f"{validos_por_caso[cid]} avaliações válidas; mínimo={quorum}",
                })

        consolidado = out_dir / "consolidado" / ("secao-1" if secao == "1" else "secao-2") / "consolidated.jsonl"
        pareceres = _agrupar_registros(_ler_jsonl(consolidado), chave_modelo=False)
        for cid, caso in por_id.items():
            registro, violacao = _selecionar_registro_valido(
                pareceres.get((cid, ""), []),
                caso=caso,
                validar_temporal=secao == "1",
                validar_publicacao_secao=secao,
                preservar_valido_anterior=False,
            )
            if registro is None:
                violacao = "parecer consolidado ausente"
            if violacao:
                afetados.add(cid)
                linhas.append({
                    "Seção": secao, "Nível": "parecer consolidado",
                    "Auditado": caso.get("auditado", ""), "Código": caso.get("codigo", ""),
                    "Case ID": cid, "Modelo": registro.get("model", "") if registro else "",
                    "Status": registro.get("status", "ausente") if registro else "ausente",
                    "Violação": violacao,
                })
            elif cid not in afetados:
                total_validos += 1

        casos_reprocessamento[secao] = [
            {
                "case_id": cid,
                "auditado": por_id[cid].get("auditado", ""),
                "codigo": por_id[cid].get("codigo", ""),
            }
            for cid in sorted(afetados)
        ]

    total_reprocessamento = sum(len(itens) for itens in casos_reprocessamento.values())
    resumo = {
        "schema_version": 1,
        "status": "conforme" if total_reprocessamento == 0 else "reparo_necessario",
        "total_casos": total_casos,
        "casos_integralmente_validos": total_validos,
        "casos_reprocessamento_secao_1": len(casos_reprocessamento["1"]),
        "casos_reprocessamento_secao_2": len(casos_reprocessamento["2"]),
        "avisos_avaliacoes_individuais": sum(
            1 for linha in linhas if linha.get("Nível") == "avaliação individual"
        ),
        "violacoes": linhas,
        "casos_reprocessamento": casos_reprocessamento,
    }
    preflight = out_dir / "preflight"
    preflight.mkdir(parents=True, exist_ok=True)
    json_path = preflight / "validacao-integridade.json"
    xlsx_path = preflight / "validacao-integridade.xlsx"
    reparo_path = preflight / "casos-reprocessamento.json"
    json_path.write_text(json.dumps(resumo, ensure_ascii=False, indent=2), encoding="utf-8")
    reparo_path.write_text(
        json.dumps({"schema_version": 1, "secoes": casos_reprocessamento}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    _gravar_xlsx(xlsx_path, linhas, resumo)
    return {**resumo, "json": str(json_path), "xlsx": str(xlsx_path), "manifesto_reparo": str(reparo_path)}
