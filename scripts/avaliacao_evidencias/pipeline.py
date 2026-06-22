from __future__ import annotations

import argparse
import difflib
import hashlib
import io
import json
import datetime as dt
import mimetypes
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import unicodedata
import zipfile
from urllib.parse import unquote
from dataclasses import dataclass
from contextlib import contextmanager
from pathlib import Path
from typing import Any, Callable, Iterable

from scripts import md2lss
from openpyxl import Workbook, load_workbook

from .providers_ai_service import _conteudo_provider_textual, estimar_tokens_payload, executar_provider


@dataclass(frozen=True)
class AnaliseCandidata:
    auditado: str
    coluna_evidencia: str
    nome_original_evidencia: str
    upload: dict[str, Any]
    resposta_id: Any = None
    evidence_index: int | None = None
    erro: str = ""
    evidencia_ausente: bool = False


@dataclass(frozen=True)
class ResolucaoEvidencia:
    caminho: Path | None
    nome_decodificado: str
    erro: str = ""


@dataclass(frozen=True)
class ItemAfirmado:
    codigo: str
    texto: str
    afirmacao: str


@dataclass(frozen=True)
class QuestaoContexto:
    codigo: str
    tipo: str
    texto: str
    itens: dict[str, str]


@dataclass(frozen=True)
class ContextoQuestionario:
    questoes: dict[str, QuestaoContexto]


@dataclass(frozen=True)
class ChecklistResolvido:
    nome: str
    caminho: Path | None
    conteudo: str
    hash_conteudo: str
    erro: str = ""


@dataclass(frozen=True)
class PromptResolvido:
    nome: str
    caminho: Path | None
    conteudo: str
    hash_conteudo: str
    erro: str = ""


@dataclass(frozen=True)
class PacoteEvidencia:
    caminho: Path
    tipo: str
    documentos: list[dict[str, Any]]
    inventario: list[str]
    erro: str = ""
    duplicados_ignorados: tuple[str, ...] = ()


REMOTE_PROVIDERS = {"gemini", "openrouter", "opencodego", "openai"}


def log_event(event: str, message: str, *, quiet: bool = False, level: str = "info", **fields: Any) -> None:
    if quiet:
        return
    payload = {
        "ts": dt.datetime.now(dt.timezone.utc).isoformat(),
        "level": level,
        "event": event,
        "message": message,
        **fields,
    }
    print(json.dumps(payload, ensure_ascii=False, default=str), file=sys.stdout, flush=True)


def validar_rpm(valor: str) -> int:
    try:
        rpm = int(valor)
    except ValueError as exc:
        raise argparse.ArgumentTypeError("--rpm deve ser um numero inteiro maior ou igual a zero") from exc
    if rpm < 0:
        raise argparse.ArgumentTypeError("--rpm deve ser maior ou igual a zero")
    return rpm


@dataclass
class RequestsPerMinuteLimiter:
    rpm: int
    clock: Callable[[], float] = time.monotonic
    sleeper: Callable[[float], None] = time.sleep
    last_started_at: float | None = None

    def wait_seconds(self) -> float:
        if self.rpm <= 0 or self.last_started_at is None:
            return 0.0
        elapsed = self.clock() - self.last_started_at
        return max(0.0, (60.0 / self.rpm) - elapsed)

    def wait_and_mark(self, wait_seconds: float | None = None) -> float:
        if wait_seconds is None:
            wait_seconds = self.wait_seconds()
        if wait_seconds > 0:
            self.sleeper(wait_seconds)
        self.last_started_at = self.clock()
        return wait_seconds


def coluna_evidencia(nome: str) -> bool:
    return "evi" in nome and not nome.endswith("[filecount]")


def _parse_upload(valor: Any) -> tuple[dict[str, Any] | None, str]:
    if valor in (None, ""):
        return None, ""
    if isinstance(valor, str):
        try:
            parsed = json.loads(valor)
        except json.JSONDecodeError as exc:
            return None, f"metadados de upload invalidos: {exc.msg}"
    else:
        parsed = valor
    if not isinstance(parsed, list):
        return None, "metadados de upload devem ser uma lista"
    if len(parsed) != 1:
        return None, "coluna de evidencia deve conter exatamente um arquivo"
    upload = parsed[0]
    if not isinstance(upload, dict):
        return None, "metadado de upload deve ser um objeto"
    name = upload.get("name")
    if not isinstance(name, str) or not name:
        return None, "metadado de upload sem atributo name"
    return upload, ""


def _rows_from_xlsx(caminho_xlsx: Path) -> Iterable[dict[str, Any]]:
    workbook = load_workbook(caminho_xlsx, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value) if value is not None else "" for value in next(rows)]
        for values in rows:
            yield dict(zip(headers, values))
    finally:
        workbook.close()


def inventariar_analises(
    caminho_xlsx: str | Path,
    raiz_evidencias: str | Path,
    caminho_questionario: str | Path,
    *,
    include_unsubmitted: bool = False,
    colunas_evidencia_permitidas: set[str] | None = None,
    include_missing_evidence: bool = False,
) -> list[AnaliseCandidata]:
    del raiz_evidencias, caminho_questionario
    linhas = list(_rows_from_xlsx(Path(caminho_xlsx)))
    if not linhas:
        return []
    colunas = list(linhas[0].keys())
    colunas_evidencia = [
        (evidence_index, col)
        for evidence_index, col in enumerate((col for col in colunas if coluna_evidencia(col)), start=1)
        if colunas_evidencia_permitidas is None or col in colunas_evidencia_permitidas
    ]
    analises: list[AnaliseCandidata] = []
    for linha in linhas:
        if not include_unsubmitted and not linha.get("submitdate"):
            continue
        auditado = str(linha.get("firstname") or "").strip()
        for evidence_index, coluna in colunas_evidencia:
            upload, erro = _parse_upload(linha.get(coluna))
            if upload is None and not erro:
                if include_missing_evidence:
                    analises.append(
                        AnaliseCandidata(
                            auditado=auditado,
                            coluna_evidencia=coluna,
                            nome_original_evidencia="",
                            upload={},
                            resposta_id=linha.get("id"),
                            evidence_index=evidence_index,
                            evidencia_ausente=True,
                        )
                    )
                continue
            analises.append(
                AnaliseCandidata(
                    auditado=auditado,
                    coluna_evidencia=coluna,
                    nome_original_evidencia=upload.get("name", "") if upload else "",
                    upload=upload or {},
                    resposta_id=linha.get("id"),
                    evidence_index=evidence_index,
                    erro=erro,
                )
            )
    return analises


def resolver_evidencia(
    auditado: str,
    raiz_evidencias: str | Path,
    upload: dict[str, Any],
    *,
    resposta_id: Any = None,
    evidence_index: int | None = None,
) -> ResolucaoEvidencia:
    nome_original = upload.get("name")
    if not isinstance(nome_original, str) or not nome_original:
        return ResolucaoEvidencia(caminho=None, nome_decodificado="", erro="metadado de upload sem atributo name")
    nome_decodificado = unquote(nome_original)
    auditado_dir = Path(raiz_evidencias) / auditado
    caminho = auditado_dir / nome_decodificado
    if not caminho.is_file():
        exportado = _resolver_evidencia_exportada_limesurvey(
            auditado_dir,
            nome_decodificado,
            resposta_id=resposta_id,
            evidence_index=evidence_index,
        )
        if exportado:
            return ResolucaoEvidencia(caminho=exportado, nome_decodificado=nome_decodificado)
        return ResolucaoEvidencia(
            caminho=None,
            nome_decodificado=nome_decodificado,
            erro=f"arquivo de evidencia nao encontrado: {caminho}",
        )
    return ResolucaoEvidencia(caminho=caminho, nome_decodificado=nome_decodificado)


def _resolver_evidencia_exportada_limesurvey(
    auditado_dir: Path,
    nome_decodificado: str,
    *,
    resposta_id: Any = None,
    evidence_index: int | None = None,
) -> Path | None:
    if not auditado_dir.is_dir():
        return None
    candidatos = [path for path in auditado_dir.iterdir() if path.is_file()]
    if not candidatos:
        return None

    if resposta_id is not None and evidence_index is not None:
        prefixo = _prefixo_exportacao_limesurvey(resposta_id, evidence_index)
        if prefixo:
            candidatos_prefixo = [path for path in candidatos if path.name.startswith(prefixo)]
            match = _melhor_match_nome_exportado(nome_decodificado, candidatos_prefixo)
            if match:
                return match

    return _melhor_match_nome_exportado(nome_decodificado, candidatos)


def _prefixo_exportacao_limesurvey(resposta_id: Any, evidence_index: int) -> str:
    try:
        return f"{int(resposta_id):05d}_{int(evidence_index):02d}_"
    except (TypeError, ValueError):
        return ""


def _melhor_match_nome_exportado(nome_decodificado: str, candidatos: list[Path]) -> Path | None:
    nome_key, suffix = _nome_exportado_key(nome_decodificado)
    matches: list[tuple[float, str, Path]] = []
    for candidato in candidatos:
        candidato_sem_prefixo = _remover_prefixo_exportacao_limesurvey(candidato.name)
        candidato_key, candidato_suffix = _nome_exportado_key(candidato_sem_prefixo)
        if candidato_suffix != suffix:
            continue
        if candidato_key == nome_key:
            score = 1.0
        else:
            score = difflib.SequenceMatcher(None, nome_key, candidato_key).ratio()
        if score >= 0.92:
            matches.append((score, candidato.name, candidato))
    if not matches:
        return None
    matches.sort(key=lambda item: (-item[0], item[1]))
    return matches[0][2]


def _remover_prefixo_exportacao_limesurvey(nome: str) -> str:
    return re.sub(r"^\d+_\d+_", "", nome)


def _nome_exportado_key(nome: str) -> tuple[str, str]:
    path = Path(nome)
    stem = unicodedata.normalize("NFKD", path.stem.casefold())
    stem = "".join(char for char in stem if not unicodedata.combining(char))
    stem = re.sub(r"[^a-z0-9]+", "-", stem)
    stem = re.sub(r"-+", "-", stem).strip("-")
    return stem, path.suffix.casefold()


def carregar_contexto_questionario(caminho_questionario: str | Path) -> ContextoQuestionario:
    survey = md2lss.parse_markdown(Path(caminho_questionario))
    questoes: dict[str, QuestaoContexto] = {}
    for group in survey.groups:
        for question in group.questions:
            if question.type == "upload":
                continue
            itens = {option.code: option.text for option in (question.subquestions or question.alternatives)}
            questoes[question.code] = QuestaoContexto(
                codigo=question.code,
                tipo=question.type,
                texto=question.text(),
                itens=itens,
            )
    return ContextoQuestionario(questoes=questoes)


def _base_coluna_evidencia(coluna_evidencia: str) -> tuple[str, str | None]:
    marker = coluna_evidencia.find("evi")
    base = coluna_evidencia[:marker]
    suffix = coluna_evidencia[marker + 3 :]
    return base, suffix or None


def _valor_afirmativo(valor: Any) -> bool:
    if isinstance(valor, str):
        return valor.strip().casefold() in {"sim", "y", "yes", "true", "1"}
    return valor is True or valor == 1


ADOPTION_VALUES = {"naoad", "adfor", "admen", "adpar", "admai", "naoap"}
ADOPTION_VALUE_ALIASES = {
    "nao adota.": "naoad",
    "adota em menor parte.": "admen",
    "adota parcialmente.": "adpar",
    "adota em maior parte ou totalmente.": "admai",
    "nao se aplica.": "naoap",
}


def _normalizar_texto_resposta(valor: Any) -> str:
    if valor is None:
        return ""
    texto = unicodedata.normalize("NFKD", str(valor).strip().casefold())
    texto = "".join(char for char in texto if not unicodedata.combining(char))
    texto = re.sub(r"\s+", " ", texto)
    return texto


def _codigo_resposta_unica(questao: QuestaoContexto, valor: Any) -> str:
    raw = str(valor or "").strip()
    if not raw:
        return ""
    if raw in questao.itens:
        return raw
    normalizado = _normalizar_texto_resposta(raw)
    for codigo, texto in questao.itens.items():
        if normalizado == _normalizar_texto_resposta(texto):
            return codigo
    match = re.match(r"^([A-Za-z])\)", raw.strip())
    if match and match.group(1).upper() in questao.itens:
        return match.group(1).upper()
    return raw


def _valor_adocao(valor: Any) -> str:
    raw = str(valor or "").strip()
    if raw in ADOPTION_VALUES:
        return raw
    return ADOPTION_VALUE_ALIASES.get(_normalizar_texto_resposta(raw), raw)


def selecionar_itens_afirmados(
    contexto: ContextoQuestionario,
    coluna_evidencia: str,
    resposta: dict[str, Any],
) -> list[ItemAfirmado]:
    base, item_especifico = _base_coluna_evidencia(coluna_evidencia)
    questao = contexto.questoes.get(base)
    if not questao:
        return []

    if item_especifico:
        valor = resposta.get(f"{base}[{item_especifico}]")
        if _valor_afirmativo(valor):
            return [
                ItemAfirmado(
                    codigo=f"{base}[{item_especifico}]",
                    texto=questao.itens.get(item_especifico, item_especifico),
                    afirmacao=str(valor),
                )
            ]
        return []

    valor_base = resposta.get(base)
    valor_adocao = _valor_adocao(valor_base)
    if questao.tipo in {"single", "adoption"} and valor_adocao in {"adpar", "admai"}:
        itens = [ItemAfirmado(codigo=base, texto=questao.texto, afirmacao=str(valor_base))]
        prefixo_ext = f"{base}ext["
        detalhe = contexto.questoes.get(f"{base}ext")
        for chave, valor in resposta.items():
            if chave.startswith(prefixo_ext) and chave.endswith("]") and _valor_afirmativo(valor):
                codigo_item = chave[len(prefixo_ext) : -1]
                itens.append(
                    ItemAfirmado(
                        codigo=chave,
                        texto=(detalhe.itens if detalhe else {}).get(codigo_item, codigo_item),
                        afirmacao="Y",
                    )
                )
        return itens

    if questao.tipo == "single" and valor_adocao in ADOPTION_VALUES:
        return []

    if questao.tipo == "single" and valor_base not in (None, ""):
        valor_codigo = _codigo_resposta_unica(questao, valor_base)
        return [
            ItemAfirmado(
                codigo=f"{base}[{valor_codigo}]",
                texto=questao.itens.get(valor_codigo, questao.texto),
                afirmacao=valor_codigo,
            )
        ]

    itens = []
    for codigo_item, texto in questao.itens.items():
        chave = f"{base}[{codigo_item}]"
        valor = resposta.get(chave)
        if _valor_afirmativo(valor):
            itens.append(ItemAfirmado(codigo=chave, texto=texto, afirmacao=str(valor)))
    return itens


def resolver_checklist(checklists_dir: str | Path, coluna_evidencia: str) -> ChecklistResolvido:
    base, item_especifico = _base_coluna_evidencia(coluna_evidencia)
    raiz = Path(checklists_dir)
    candidatos = []
    if item_especifico:
        candidatos.append(raiz / f"{base}_{item_especifico}.md")
    candidatos.append(raiz / f"{base}.md")
    for caminho in candidatos:
        if caminho.is_file():
            conteudo = caminho.read_text(encoding="utf-8")
            digest = hashlib.sha256(conteudo.encode("utf-8")).hexdigest()
            return ChecklistResolvido(
                nome=caminho.name,
                caminho=caminho,
                conteudo=conteudo,
                hash_conteudo=digest,
            )
    return ChecklistResolvido(
        nome="",
        caminho=None,
        conteudo="",
        hash_conteudo="",
        erro=f"checklist de analise nao encontrado para {coluna_evidencia}",
    )


def resolver_prompt(prompts_dir: str | Path, coluna_evidencia: str) -> PromptResolvido:
    base, item_especifico = _base_coluna_evidencia(coluna_evidencia)
    raiz = Path(prompts_dir)
    candidatos = []
    if item_especifico:
        candidatos.append(raiz / f"{base}_{item_especifico}.md")
    candidatos.append(raiz / f"{base}.md")
    for caminho in candidatos:
        if caminho.is_file():
            conteudo = caminho.read_text(encoding="utf-8")
            digest = hashlib.sha256(conteudo.encode("utf-8")).hexdigest()
            return PromptResolvido(
                nome=caminho.name,
                caminho=caminho,
                conteudo=conteudo,
                hash_conteudo=digest,
            )
    return PromptResolvido(
        nome="",
        caminho=None,
        conteudo="",
        hash_conteudo="",
        erro=f"prompt de analise nao encontrado para {coluna_evidencia}",
    )


def colunas_com_prompt(
    prompts_dir: str | Path,
    caminho_questionario: str | Path,
    *,
    only_achados: bool = False,
    catalog: str | Path | None = None,
) -> set[str]:
    achados_set = carregar_achados_set(prompts_dir, catalog) if only_achados else set()
    survey = md2lss.parse_markdown(Path(caminho_questionario))
    colunas: set[str] = set()
    for group in survey.groups:
        for question in group.questions:
            if question.type == "upload" and coluna_evidencia(question.code):
                prompt = resolver_prompt(prompts_dir, question.code)
                if prompt.caminho is None:
                    continue
                if only_achados:
                    base, _ = _base_coluna_evidencia(question.code)
                    if base not in achados_set:
                        continue
                colunas.add(question.code)
    return colunas


def itens_avaliaveis_prompt(conteudo: str) -> set[str]:
    match = re.search(r"<!--\s*itens_avaliaveis:\s*(.*?)\s*-->", conteudo)
    if not match:
        return set()
    return {item.strip() for item in match.group(1).split(",") if item.strip()}


def filtrar_itens_por_prompt(itens: list[ItemAfirmado], prompt: PromptResolvido) -> list[ItemAfirmado]:
    permitidos = itens_avaliaveis_prompt(prompt.conteudo)
    if not permitidos:
        return itens
    return [item for item in itens if item.codigo in permitidos]


def prompt_exibe_texto_itens(conteudo: str) -> bool:
    match = re.search(r"<!--\s*exibir_texto_itens:\s*(.*?)\s*-->", conteudo)
    if not match:
        return True
    valor = match.group(1).strip().lower()
    return valor not in {"nao", "não", "false", "0", "no"}


def prompt_gera_achado(conteudo: str) -> bool:
    match = re.search(r"<!--\s*gera_achado:\s*(.*?)\s*-->", conteudo)
    if not match:
        return False
    valor = match.group(1).strip().lower()
    return valor in {"sim", "true", "1", "yes"}


_NOME_PROMPT_ACHADO_RE = re.compile(r"^(q\d{4})(?:_[A-Z])?\.md$", re.IGNORECASE)


def carregar_achados_set(
    prompts_dir: str | Path | None = None,
    catalog: str | Path | None = None,
) -> set[str]:
    """Conjunto de questoes-raiz marcadas como gera_achado.

    Fonte primaria: catlogo YAML informado em ``catalog`` (le o atributo
    ``gera_achado`` de cada entrada de ``prompts``). Fallback: marcadores
    ``<!-- gera_achado: sim -->`` nos prompts markdown de ``prompts_dir``,
    para diretorios deployados antes do atributo no YAML.
    """
    if catalog:
        return _achados_de_catalogo(catalog)
    if prompts_dir:
        return _achados_de_prompts_dir(prompts_dir)
    return set()


def _achados_de_catalogo(caminho: str | Path) -> set[str]:
    import yaml

    with Path(caminho).open(encoding="utf-8") as file:
        data = yaml.safe_load(file)
    if not isinstance(data, dict):
        return set()
    questoes: set[str] = set()
    for entrada in data.get("prompts", []) or []:
        if not isinstance(entrada, dict) or not entrada.get("gera_achado"):
            continue
        arquivo = str(entrada.get("arquivo") or "")
        match = _NOME_PROMPT_ACHADO_RE.match(arquivo)
        if match:
            questoes.add(match.group(1))
    return questoes


def _achados_de_prompts_dir(prompts_dir: str | Path) -> set[str]:
    raiz = Path(prompts_dir)
    questoes: set[str] = set()
    for caminho in raiz.glob("*.md"):
        match = _NOME_PROMPT_ACHADO_RE.match(caminho.name)
        if not match:
            continue
        try:
            conteudo = caminho.read_text(encoding="utf-8")
        except OSError:
            continue
        if prompt_gera_achado(conteudo):
            questoes.add(match.group(1))
    return questoes


def preparar_itens_para_prompt(itens: list[ItemAfirmado], prompt: PromptResolvido) -> list[ItemAfirmado]:
    if prompt_exibe_texto_itens(prompt.conteudo):
        return itens
    return [
        ItemAfirmado(
            codigo=item.codigo,
            texto=item.codigo,
            afirmacao=item.afirmacao,
        )
        for item in itens
    ]


def resultado_evidencia_ausente(itens: list[ItemAfirmado]) -> dict[str, Any]:
    return {
        "status": "completed",
        "conclusoes": [
            {
                "item_codigo": item.codigo,
                "item_texto": item.texto,
                "afirmacao_auditado": item.afirmacao,
                "estado": "nao_conforme",
                "justificativa": "O auditado afirmou o item, mas nao enviou evidencia para sustentar a afirmacao.",
                "lacunas": ["Evidencia nao enviada."],
                "arquivos_referenciados": [],
                "trechos_ou_elementos": [],
                "paginas_ou_localizacao": [],
            }
            for item in itens
        ],
    }


def hash_arquivo(caminho: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(caminho).open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _zip_member_safe(name: str) -> bool:
    path = Path(name)
    return not path.is_absolute() and ".." not in path.parts


def normalizar_evidencia(caminho: str | Path) -> PacoteEvidencia:
    path = Path(caminho)
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".csv"}:
        try:
            texto = path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            texto = path.read_text(encoding="latin-1")
        return PacoteEvidencia(
            caminho=path,
            tipo=suffix.lstrip("."),
            documentos=[{"nome": path.name, "texto": texto}],
            inventario=[path.name],
        )
    if suffix == ".zip":
        documentos: list[dict[str, Any]] = []
        try:
            with zipfile.ZipFile(path) as archive:
                names = archive.namelist()
                for name in names:
                    if not _zip_member_safe(name):
                        return PacoteEvidencia(
                            caminho=path,
                            tipo="zip",
                            documentos=[],
                            inventario=names,
                            erro=f"zip contem path traversal: {name}",
                        )
                hashes_vistos: set[str] = set()
                duplicados_ignorados: list[str] = []
                for name in names:
                    if name.endswith("/"):
                        continue
                    suffix_interno = Path(name).suffix.lower()
                    if suffix_interno in {".txt", ".md", ".csv"}:
                        data = archive.read(name)
                        try:
                            texto = data.decode("utf-8")
                        except UnicodeDecodeError:
                            texto = data.decode("latin-1")
                        documentos.append({"nome": name, "texto": texto})
                    elif suffix_interno == ".pdf":
                        data = archive.read(name)
                        hash_pdf = hashlib.sha256(data).hexdigest()
                        if hash_pdf in hashes_vistos:
                            duplicados_ignorados.append(name)
                            continue
                        hashes_vistos.add(hash_pdf)
                        pdf_docs, erro_pdf = _extrair_texto_pdf_bytes(name, data)
                        if pdf_docs:
                            documentos.extend(pdf_docs)
                        else:
                            documentos.append({"nome": name, "erro": erro_pdf})
                    elif suffix_interno == ".xlsx":
                        data = archive.read(name)
                        hash_xlsx = hashlib.sha256(data).hexdigest()
                        if hash_xlsx in hashes_vistos:
                            duplicados_ignorados.append(name)
                            continue
                        hashes_vistos.add(hash_xlsx)
                        texto, erro_md = _extrair_texto_bytes_com_markitdown(name, data, suffix_interno)
                        if erro_md:
                            documentos.append({"nome": name, "erro": erro_md})
                        else:
                            documentos.append({"nome": name, "texto": texto})
                    elif suffix_interno in EXTENSOES_WORD_PARA_PDF:
                        documentos.append(
                            {
                                "nome": name,
                                "tipo_convertido": "pdf",
                                "observacao": "Documento Word preservado para avaliacao visual por conversao para PDF no upload.",
                            }
                        )
                    else:
                        documentos.append({"nome": name, "nao_suportado": True})
                return PacoteEvidencia(
                    caminho=path,
                    tipo="zip",
                    documentos=documentos,
                    inventario=names,
                    duplicados_ignorados=tuple(duplicados_ignorados),
                )
        except zipfile.BadZipFile:
            return PacoteEvidencia(caminho=path, tipo="zip", documentos=[], inventario=[], erro="zip invalido")
    if suffix == ".pdf":
        documentos, erro = _extrair_texto_pdf(path)
        return PacoteEvidencia(
            caminho=path,
            tipo="pdf",
            documentos=documentos,
            inventario=[path.name],
            erro=erro,
        )
    if suffix == ".xlsx":
        texto, erro = _extrair_texto_com_markitdown(path)
        if erro:
            return PacoteEvidencia(
                caminho=path,
                tipo="xlsx",
                documentos=[],
                inventario=[path.name],
                erro=erro,
            )
        return PacoteEvidencia(
            caminho=path,
            tipo="xlsx",
            documentos=[{"nome": path.name, "texto": texto}],
            inventario=[path.name],
        )
    if suffix in EXTENSOES_WORD_PARA_PDF:
        return PacoteEvidencia(
            caminho=path,
            tipo=suffix.lstrip("."),
            documentos=[
                {
                    "nome": path.name,
                    "tipo_convertido": "pdf",
                    "observacao": "Documento Word preservado para avaliacao visual por conversao para PDF no upload.",
                }
            ],
            inventario=[path.name],
        )
    return PacoteEvidencia(
        caminho=path,
        tipo=suffix.lstrip(".") or "desconhecido",
        documentos=[],
        inventario=[path.name],
        erro=f"tipo de evidencia nao suportado: {suffix or path.name}",
    )


def _erros_documentos(pacote: PacoteEvidencia) -> list[str]:
    erros = []
    for documento in pacote.documentos:
        erro = documento.get("erro") if isinstance(documento, dict) else None
        nome = documento.get("nome") if isinstance(documento, dict) else ""
        if erro:
            erros.append(f"{nome}: {erro}" if nome else str(erro))
    return erros


def _erro_pdf_mitigado_por_upload(erro: str, arquivos_upload: list[str]) -> bool:
    texto = erro.casefold()
    return bool(arquivos_upload) and "pdf" in texto and (
        "sem texto extraivel" in texto
        or "normalizar pdf" in texto
        or "extrair texto" in texto
    )


def erro_tecnico_bloqueante_pacote(pacote: PacoteEvidencia, arquivos_upload: list[str]) -> str:
    if pacote.erro and not _erro_pdf_mitigado_por_upload(pacote.erro, arquivos_upload):
        return pacote.erro

    erros_bloqueantes = [
        erro for erro in _erros_documentos(pacote) if not _erro_pdf_mitigado_por_upload(erro, arquivos_upload)
    ]
    if erros_bloqueantes:
        return "; ".join(erros_bloqueantes)

    if not pacote.documentos and not arquivos_upload:
        return "evidencia sem conteudo processavel para avaliacao"

    if all(documento.get("nao_suportado") for documento in pacote.documentos) and not arquivos_upload:
        return "evidencia contem apenas arquivos de tipo nao suportado"

    return ""


def resultado_indica_erro_tecnico(resultado: dict[str, Any]) -> str:
    if resultado.get("status") != "completed":
        return ""
    termos = [
        "markitdown",
        "erro tecnico",
        "erro técnico",
        "nao possui acesso ao conteudo",
        "não possui acesso ao conteúdo",
        "conteudo do arquivo inacessivel",
        "conteúdo do arquivo inacessível",
    ]
    conclusoes = resultado.get("conclusoes")
    if not isinstance(conclusoes, list):
        return ""
    for conclusao in conclusoes:
        if not isinstance(conclusao, dict):
            continue
        campos = [
            conclusao.get("justificativa", ""),
            " ".join(str(valor) for valor in conclusao.get("lacunas", []) if valor is not None)
            if isinstance(conclusao.get("lacunas"), list)
            else str(conclusao.get("lacunas", "")),
        ]
        texto = " ".join(campos).casefold()
        if any(termo in texto for termo in termos):
            return "modelo indicou erro tecnico de acesso/processamento da evidencia"
    return ""


EXTENSOES_UPLOAD_DIRETO = {".pdf", ".txt", ".md", ".csv", ".png", ".jpg", ".jpeg"}
EXTENSOES_WORD_PARA_PDF = {".doc", ".docx"}
EXTENSOES_UPLOAD_PREPARAVEIS = EXTENSOES_UPLOAD_DIRETO | EXTENSOES_WORD_PARA_PDF | {".xlsx"}


def _extrair_texto_com_markitdown(caminho: Path) -> tuple[str, str]:
    if caminho.suffix.lower() == ".xlsx":
        return _extrair_xlsx_abas_visiveis(caminho)
    try:
        from markitdown import MarkItDown
    except ModuleNotFoundError:
        return "", "markitdown nao instalado no ambiente virtual"
    try:
        md = MarkItDown()
        result = md.convert(str(caminho))
        return result.text_content, ""
    except Exception as exc:
        return "", f"erro ao extrair texto com markitdown: {exc}"


def _extrair_xlsx_abas_visiveis(caminho: Path) -> tuple[str, str]:
    """Extrai markdown de um XLSX considerando apenas abas visiveis.

    Reproduce o formato do MarkItDown (## nome_aba + tabela em markdown), mas
    filtra abas ocultas (sheet_state != 'visible') usando openpyxl, pois o
    pandas/markitdown le todas as abas sem distincao.
    """
    try:
        import pandas as pd
        from openpyxl import load_workbook
    except ModuleNotFoundError as exc:
        return "", f"dependencia para xlsx nao disponivel: {exc}"
    try:
        wb = load_workbook(caminho, read_only=True, data_only=True)
        abas_visiveis = [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"]
        wb.close()
        if not abas_visiveis:
            return "", "xlsx sem abas visiveis"
        partes: list[str] = []
        for nome_aba in abas_visiveis:
            df = pd.read_excel(caminho, sheet_name=nome_aba, engine="openpyxl")
            partes.append(f"## {nome_aba}\n")
            partes.append(df.to_markdown(index=False))
            partes.append("\n\n")
        return "".join(partes).strip(), ""
    except Exception as exc:
        return "", f"erro ao extrair xlsx com abas visiveis: {exc}"


def _extrair_texto_bytes_com_markitdown(nome: str, data: bytes, sufixo: str) -> tuple[str, str]:
    with tempfile.TemporaryDirectory() as tmp_dir:
        temp_file = Path(tmp_dir) / f"temp_{_nome_upload_seguro(nome, sufixo)}"
        temp_file.write_bytes(data)
        return _extrair_texto_com_markitdown(temp_file)


def _converter_word_para_pdf(origem: Path, destino: Path, nome_original: str | None = None) -> tuple[Path | None, str]:
    conversor = shutil.which("soffice") or shutil.which("libreoffice")
    if not conversor:
        return _converter_word_para_pdf_com_microsoft_word(origem, destino, nome_original)
    destino.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp_dir:
        tmp_path = Path(tmp_dir)
        try:
            result = subprocess.run(
                [
                    conversor,
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(tmp_path),
                    str(origem),
                ],
                check=False,
                capture_output=True,
                text=True,
                timeout=120,
            )
        except subprocess.TimeoutExpired:
            convertido, erro_word = _converter_word_para_pdf_com_microsoft_word(origem, destino, nome_original)
            if convertido:
                return convertido, ""
            return None, f"timeout ao converter documento Word em PDF; fallback Microsoft Word: {erro_word}"
        except OSError as exc:
            convertido, erro_word = _converter_word_para_pdf_com_microsoft_word(origem, destino, nome_original)
            if convertido:
                return convertido, ""
            return None, f"erro ao executar conversor de documento Word para PDF: {exc}; fallback Microsoft Word: {erro_word}"
        if result.returncode != 0:
            detalhe = (result.stderr or result.stdout or "").strip()
            convertido, erro_word = _converter_word_para_pdf_com_microsoft_word(origem, destino, nome_original)
            if convertido:
                return convertido, ""
            return None, f"erro ao converter documento Word em PDF: {detalhe or result.returncode}; fallback Microsoft Word: {erro_word}"
        pdf_convertido = tmp_path / f"{origem.stem}.pdf"
        if not pdf_convertido.is_file():
            candidatos = sorted(tmp_path.glob("*.pdf"))
            if not candidatos:
                convertido, erro_word = _converter_word_para_pdf_com_microsoft_word(origem, destino, nome_original)
                if convertido:
                    return convertido, ""
                return None, f"conversor de documento Word para PDF nao gerou arquivo PDF; fallback Microsoft Word: {erro_word}"
            pdf_convertido = candidatos[0]
        nome_pdf = _nome_upload_seguro(nome_original or origem.name, ".pdf")
        target = _caminho_unico(destino, nome_pdf)
        shutil.copyfile(pdf_convertido, target)
        return target, ""


def _converter_word_para_pdf_com_microsoft_word(
    origem: Path,
    destino: Path,
    nome_original: str | None = None,
) -> tuple[Path | None, str]:
    if os.name != "nt":
        return None, "Microsoft Word COM disponivel apenas no Windows"
    if not shutil.which("powershell.exe"):
        return None, "powershell.exe nao encontrado para acionar Microsoft Word"
    destino.mkdir(parents=True, exist_ok=True)
    nome_pdf = _nome_upload_seguro(nome_original or origem.name, ".pdf")
    target = _caminho_unico(destino, nome_pdf)
    script = r"""
param([string]$Source, [string]$Target)
$ErrorActionPreference = 'Stop'
$word = $null
$doc = $null
try {
    $word = New-Object -ComObject Word.Application
    $word.Visible = $false
    $word.DisplayAlerts = 0
    $sourcePath = (Resolve-Path -LiteralPath $Source).Path
    $doc = $word.Documents.Open([ref]$sourcePath, [ref]$false, [ref]$true, [ref]$false)
    if ($null -eq $doc) {
        $doc = $word.ActiveDocument
    }
    if ($null -eq $doc) {
        throw 'Microsoft Word nao abriu o documento'
    }
    $doc.ExportAsFixedFormat($Target, 17)
} finally {
    if ($null -ne $doc) {
        $doc.Close($false) | Out-Null
        [System.Runtime.InteropServices.Marshal]::FinalReleaseComObject($doc) | Out-Null
    }
    if ($null -ne $word) {
        $word.Quit() | Out-Null
        [System.Runtime.InteropServices.Marshal]::FinalReleaseComObject($word) | Out-Null
    }
    [GC]::Collect()
    [GC]::WaitForPendingFinalizers()
}
"""
    try:
        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_path = Path(tmp_dir)
            local_source = tmp_path / f"source{origem.suffix.lower()}"
            local_target = tmp_path / "converted.pdf"
            script_path = tmp_path / "convert-word-to-pdf.ps1"
            shutil.copyfile(origem, local_source)
            script_path.write_text(script, encoding="utf-8")
            result = subprocess.run(
                [
                    "powershell.exe",
                    "-NoProfile",
                    "-ExecutionPolicy",
                    "Bypass",
                    "-File",
                    str(script_path),
                    "-Source",
                    str(local_source),
                    "-Target",
                    str(local_target),
                ],
                check=False,
                capture_output=True,
                text=True,
                timeout=180,
            )
            if result.returncode != 0:
                detalhe = (result.stderr or result.stdout or "").strip()
                return None, f"erro ao converter documento Word em PDF com Microsoft Word: {detalhe or result.returncode}"
            if not local_target.is_file():
                return None, "Microsoft Word nao gerou arquivo PDF"
            shutil.copyfile(local_target, target)
    except subprocess.TimeoutExpired:
        return None, "timeout ao converter documento Word em PDF com Microsoft Word"
    except OSError as exc:
        return None, f"erro ao acionar Microsoft Word para converter documento Word em PDF: {exc}"
    if not target.is_file():
        return None, "Microsoft Word nao gerou arquivo PDF"
    return target, ""


@contextmanager
def _suprimir_stderr():
    """Silencia warnings do pypdf (ex.: 'Multiple definitions in dictionary')."""
    saved_fd = os.dup(2)
    devnull = os.open(os.devnull, os.O_WRONLY)
    try:
        os.dup2(devnull, 2)
        yield
    finally:
        os.dup2(saved_fd, 2)
        os.close(devnull)
        os.close(saved_fd)


MAX_PAGINAS_PDF = 50


def _extrair_texto_pdf_reader(nome: str, reader: Any) -> tuple[list[dict[str, Any]], str]:
    documentos: list[dict[str, Any]] = []
    total_paginas = len(reader.pages)
    limite = min(total_paginas, MAX_PAGINAS_PDF)
    for index, page in enumerate(reader.pages[:limite], start=1):
        texto = (page.extract_text() or "").strip()
        if texto:
            documentos.append({"nome": nome, "pagina": index, "texto": texto})
    if not documentos:
        return [], "pdf sem texto extraivel"
    return documentos, ""


def _extrair_texto_pdf(caminho: str | Path) -> tuple[list[dict[str, Any]], str]:
    try:
        from pypdf import PdfReader
    except ModuleNotFoundError:
        return [], "pypdf nao instalado para normalizar PDF"
    try:
        path = Path(caminho)
        with _suprimir_stderr():
            return _extrair_texto_pdf_reader(path.name, PdfReader(path))
    except Exception as exc:
        return [], f"erro ao normalizar PDF: {exc}"


def _extrair_texto_pdf_bytes(nome: str, data: bytes) -> tuple[list[dict[str, Any]], str]:
    try:
        from pypdf import PdfReader
    except ModuleNotFoundError:
        return [], "pypdf nao instalado para normalizar PDF"
    try:
        with _suprimir_stderr():
            return _extrair_texto_pdf_reader(nome, PdfReader(io.BytesIO(data)))
    except Exception as exc:
        return [], f"erro ao normalizar PDF: {exc}"


def _tentar_extrair_markdown_pdf(
    caminho: Path,
    raiz: Path,
    imagens_dir: Path,
    slug: str,
    dpi: int,
) -> tuple[str, str]:
    """Tenta extrair markdown de um PDF com pymupdf4llm. Retorna (markdown, erro)."""
    try:
        import pymupdf4llm
    except ModuleNotFoundError:
        return "", "pymupdf4llm nao instalado no ambiente virtual"
    # Limitar a MAX_PAGINAS_PDF primeiras paginas
    try:
        import fitz

        doc = fitz.open(str(caminho))
        total_paginas = doc.page_count
        doc.close()
    except Exception:
        total_paginas = 0
    pages = list(range(1, min(total_paginas, MAX_PAGINAS_PDF) + 1)) if total_paginas > 0 else None
    caminho_absoluto = caminho.resolve()
    cwd = Path.cwd()
    try:
        os.chdir(raiz)
        kwargs: dict[str, Any] = {
            "write_images": True,
            "image_path": "images",
            "image_format": "png",
            "dpi": dpi,
        }
        if pages is not None:
            kwargs["pages"] = pages
        markdown = pymupdf4llm.to_markdown(str(caminho_absoluto), **kwargs)
        return markdown, ""
    except Exception as exc:
        return "", f"erro ao extrair PDF com pymupdf4llm: {exc}"
    finally:
        os.chdir(cwd)


def _recuperar_pdf(caminho: Path, destino_dir: Path) -> tuple[Path | None, str, str]:
    """Aplica protocolo de recuperacao de PDF corrompido em camadas.

    Retorna (caminho_recuperado, camada_aplicada, erro).
    Se todas as camadas falharem, retorna (None, ultima_camada_tentada, erro).
    """
    data = caminho.read_bytes()
    ultima_camada = ""
    ultimo_erro = ""

    # Fase 0 — Diagnostico rapido
    tem_eof = b"%%EOF" in data[-1024:]
    tem_startxref = b"startxref" in data[-1024:]
    header_ok = data[:5] == b"%PDF-"
    if header_ok and tem_eof and tem_startxref:
        # Provavelmente nao esta truncado; ainda assim tenta camada 1.2
        pass

    # Camada 1.1 — Adicionar %%EOF no final
    ultima_camada = "1.1_eof"
    if not tem_eof:
        recuperado = destino_dir / f"{caminho.stem}_rec_eof.pdf"
        recuperado.write_bytes(data + b"\n%%EOF\n")
        if _pdf_tem_paginas(recuperado):
            return recuperado, ultima_camada, ""

    # Camada 1.2 — PyMuPDF: abrir e re-salvar
    ultima_camada = "1.2_fitz_resave"
    recuperado = _tentar_recuperar_fitz(caminho, destino_dir)
    if recuperado is not None:
        return recuperado, ultima_camada, ""

    # Camada 1.3 — pikepdf com flags permissivas
    ultima_camada = "1.3_pikepdf"
    recuperado = _tentar_recuperar_pikepdf(caminho, destino_dir)
    if recuperado is not None:
        return recuperado, ultima_camada, ""

    # Camada 4 — Reconstrucao manual (Catalog/Pages/xref/trailer)
    ultima_camada = "4_reconstrucao_manual"
    recuperado = _tentar_reconstrucao_manual(caminho, destino_dir, data)
    if recuperado is not None:
        return recuperado, ultima_camada, ""

    # Camada 5 — Reconstrucao a partir de imagens extraidas
    ultima_camada = "5_imagens_extraidas"
    recuperado = _tentar_reconstrucao_imagens(caminho, destino_dir, data)
    if recuperado is not None:
        return recuperado, ultima_camada, ""

    ultimo_erro = "todas as camadas de recuperacao falharam"
    return None, ultima_camada, ultimo_erro


def _pdf_tem_paginas(caminho: Path) -> bool:
    """Verifica se um PDF tem pelo menos 1 pagina legivel."""
    try:
        import fitz

        doc = fitz.open(str(caminho))
        n = doc.page_count
        doc.close()
        return n > 0
    except Exception:
        return False


def _tentar_recuperar_fitz(caminho: Path, destino_dir: Path) -> Path | None:
    """Camada 1.2: Abre com PyMuPDF e re-salva com garbage collection."""
    try:
        import fitz

        doc = fitz.open(str(caminho))
        if doc.page_count == 0:
            doc.close()
            return None
        recuperado = destino_dir / f"{caminho.stem}_rec_fitz.pdf"
        doc.save(str(recuperado), garbage=4, deflate=True, clean=True)
        doc.close()
        if _pdf_tem_paginas(recuperado):
            return recuperado
    except Exception:
        pass
    return None


def _tentar_recuperar_pikepdf(caminho: Path, destino_dir: Path) -> Path | None:
    """Camada 1.3/2.1: Abre com pikepdf com flags permissivas e re-salva."""
    try:
        import pikepdf

        recuperado = destino_dir / f"{caminho.stem}_rec_pikepdf.pdf"
        with pikepdf.open(str(caminho), ignore_xref_streams=True) as pdf:
            pdf.save(str(recuperado))
        if _pdf_tem_paginas(recuperado):
            return recuperado
    except Exception:
        pass
    return None


def _tentar_reconstrucao_manual(
    caminho: Path,
    destino_dir: Path,
    data: bytes,
) -> Path | None:
    """Camada 4: Reconstruir manualmente Catalog, Pages, xref e trailer.

    Localiza objetos /Page no binario, recria Catalog e Pages, e escreve um
    novo PDF com offsets e xref corretos.
    """
    try:
        # Localizar todos os objetos PDF: "N M obj"
        objs = list(re.finditer(rb"(\d+)\s+(\d+)\s+obj\b", data))
        if not objs:
            return None

        # Identificar objetos que sao /Page (mas nao /Pages)
        page_objs: list[int] = []
        for m in objs:
            obj_num = int(m.group(1))
            # Procurar /Type /Page nos bytes seguintes (mas nao /Pages)
            trecho = data[m.end():m.end() + 200]
            if rb"/Type" in trecho and rb"/Page" in trecho and rb"/Pages" not in trecho:
                page_objs.append(obj_num)

        if not page_objs:
            return None

        # Determinar numeros de objetos para Catalog e Pages (nao usados)
        nums_usados = {int(m.group(1)) for m in objs}
        catalog_num = 1 if 1 not in nums_usados else max(nums_usados) + 1
        pages_num = catalog_num + 1 if catalog_num + 1 not in nums_usados else max(nums_usados) + 2

        # Construir Catalog e Pages
        catalog_obj = (
            f"{catalog_num} 0 obj\n<< /Type /Catalog /Pages {pages_num} 0 R >>\nendobj\n"
        ).encode()
        kids = " ".join(f"{n} 0 R" for n in page_objs)
        pages_obj = (
            f"{pages_num} 0 obj\n<< /Type /Pages /Kids [ {kids} ] /Count {len(page_objs)} >>\nendobj\n"
        ).encode()

        # Reescrever PDF: header + catalog + pages + objetos originais + xref + trailer
        output = bytearray()
        output.extend(b"%PDF-1.7\n%\xe2\xe3\xcf\xd3\n")

        offsets: dict[int, int] = {}
        # Escrever Catalog
        offsets[catalog_num] = len(output)
        output.extend(catalog_obj)
        # Escrever Pages
        offsets[pages_num] = len(output)
        output.extend(pages_obj)
        # Escrever objetos originais
        for m in objs:
            obj_num = int(m.group(1))
            if obj_num in (catalog_num, pages_num):
                continue
            # Encontrar endobj apos o objeto
            end = data.find(b"endobj", m.end())
            if end == -1:
                continue
            conteudo = data[m.start():end + 6] + b"\n"
            offsets[obj_num] = len(output)
            output.extend(conteudo)

        # Construir xref
        xref_offset = len(output)
        output.extend(b"xref\n")
        output.extend(f"0 {max(offsets.keys()) + 1}\n".encode())
        output.extend(b"0000000000 65535 f \n")
        for i in range(1, max(offsets.keys()) + 1):
            if i in offsets:
                output.extend(f"{offsets[i]:010d} 00000 n \n".encode())
            else:
                output.extend(b"0000000000 00000 f \n")

        # Trailer
        output.extend(b"trailer\n")
        output.extend(f"<< /Size {max(offsets.keys()) + 1} /Root {catalog_num} 0 R >>\n".encode())
        output.extend(b"startxref\n")
        output.extend(f"{xref_offset}\n".encode())
        output.extend(b"%%EOF\n")

        recuperado = destino_dir / f"{caminho.stem}_rec_manual.pdf"
        recuperado.write_bytes(bytes(output))

        # Otimizar com fitz
        try:
            import fitz

            doc = fitz.open(str(recuperado))
            if doc.page_count > 0:
                otimizado = destino_dir / f"{caminho.stem}_rec_manual_opt.pdf"
                doc.save(str(otimizado), garbage=4, deflate=True, clean=True)
                doc.close()
                return otimizado
            doc.close()
        except Exception:
            pass

        if _pdf_tem_paginas(recuperado):
            return recuperado
    except Exception:
        pass
    return None


def _tentar_reconstrucao_imagens(
    caminho: Path,
    destino_dir: Path,
    data: bytes,
) -> Path | None:
    """Camada 5: Extrai imagens JPEG embarcadas e cria novo PDF."""
    try:
        import fitz
        from PIL import Image
    except ModuleNotFoundError:
        return None

    # Extrair JPEGs do binario
    imagens: list[bytes] = []
    start = 0
    while True:
        pos = data.find(b"\xff\xd8\xff", start)
        if pos == -1:
            break
        end = data.find(b"\xff\xd9", pos)
        if end == -1:
            break
        imagens.append(data[pos:end + 2])
        start = end + 2

    if not imagens:
        return None

    recuperado = destino_dir / f"{caminho.stem}_rec_imagens.pdf"
    imagens_dir = destino_dir / f"{caminho.stem}_imgs"
    imagens_dir.mkdir(parents=True, exist_ok=True)

    doc = fitz.open()
    for i, img_data in enumerate(imagens):
        img_path = imagens_dir / f"img_{i:03d}.jpg"
        img_path.write_bytes(img_data)
        try:
            img = Image.open(img_path)
            img.verify()
            img = Image.open(img_path)
            page = doc.new_page(width=img.width, height=img.height)
            page.insert_image(fitz.Rect(0, 0, img.width, img.height), filename=str(img_path))
        except Exception:
            continue

    if doc.page_count == 0:
        doc.close()
        return None

    doc.save(str(recuperado))
    doc.close()
    return recuperado


def _extrair_pdf_markdown_imagens(
    caminho: Path,
    destino: Path,
    nome_original: str | None = None,
    *,
    dpi: int = 150,
) -> tuple[PacoteEvidencia | None, list[str], str]:
    try:
        import pymupdf4llm
    except ModuleNotFoundError:
        return None, [], "pymupdf4llm nao instalado no ambiente virtual"

    destino.mkdir(parents=True, exist_ok=True)
    nome_base = nome_original or caminho.name
    slug = _slug_ascii(Path(nome_base).stem, fallback="pdf")
    raiz = _caminho_unico(destino, f"{slug}-pymupdf4llm")
    raiz.mkdir(parents=True, exist_ok=True)
    imagens_dir = raiz / "images"
    imagens_dir.mkdir(parents=True, exist_ok=True)
    markdown_nome = f"{slug}.md"
    markdown_path = raiz / markdown_nome

    # Tentativa inicial; em caso de erro, tentar recuperar o PDF e repetir
    markdown, erro = _tentar_extrair_markdown_pdf(caminho, raiz, imagens_dir, slug, dpi)
    if erro:
        pdf_recuperado, camada, erro_rec = _recuperar_pdf(caminho, destino)
        if pdf_recuperado is not None:
            markdown, erro = _tentar_extrair_markdown_pdf(
                pdf_recuperado, raiz, imagens_dir, slug, dpi
            )
            if erro:
                return None, [], f"erro ao extrair PDF com pymupdf4llm apos recuperacao (camada {camada}): {erro}"
        else:
            return None, [], f"erro ao extrair PDF com pymupdf4llm: {erro}. recuperacao tentada ate camada {camada} sem sucesso: {erro_rec}"

    markdown_path.write_text(markdown, encoding="utf-8")
    imagens = sorted(path for path in imagens_dir.glob("*") if path.is_file())

    # Dedup de imagens por hash de conteudo: se duas imagens sao identicas,
    # manter apenas uma e ajustar as referencias no markdown.
    hashes_imagens: dict[str, Path] = {}
    dedup_map: dict[str, Path] = {}
    duplicadas_removidas: list[str] = []
    for img in imagens:
        h = hash_arquivo(img)
        if h in hashes_imagens:
            dedup_map[str(img)] = hashes_imagens[h]
            duplicadas_removidas.append(img.name)
        else:
            hashes_imagens[h] = img

    if dedup_map:
        texto_md = markdown_path.read_text(encoding="utf-8")
        for caminho_dup, caminho_orig in dedup_map.items():
            nome_dup = Path(caminho_dup).name
            nome_orig = caminho_orig.name
            texto_md = texto_md.replace(nome_dup, nome_orig)
        markdown_path.write_text(texto_md, encoding="utf-8")
        markdown = texto_md
        for img_dup in dedup_map:
            Path(img_dup).unlink()

    imagens_unicas = sorted(path for path in imagens_dir.glob("*") if path.is_file())
    imagens_relativas = [str(path.relative_to(raiz).as_posix()) for path in imagens_unicas]
    documentos = [
        {
            "nome": markdown_nome,
            "origem_pdf": nome_base,
            "tipo_extraido": "pdf_markdown_imagens",
            "texto": markdown,
            "imagens_extraidas": imagens_relativas,
            "duplicadas_removidas": duplicadas_removidas,
        }
    ]
    inventario = [nome_base, markdown_nome, *imagens_relativas]
    arquivos_upload = [str(markdown_path), *(str(path) for path in imagens_unicas)]
    pacote = PacoteEvidencia(
        caminho=caminho,
        tipo="pdf_markdown_imagens",
        documentos=documentos,
        inventario=inventario,
    )
    return pacote, arquivos_upload, ""


def _extensao_imagem_mammoth(content_type: str) -> str:
    extensao = mimetypes.guess_extension(content_type.split(";")[0].strip())
    if extensao == ".jpe":
        return ".jpg"
    return extensao or ".bin"


def _extrair_docx_html_imagens(
    caminho: Path,
    destino: Path,
    nome_original: str | None = None,
) -> tuple[PacoteEvidencia | None, list[str], str]:
    try:
        import mammoth
    except ModuleNotFoundError:
        return None, [], "mammoth nao instalado no ambiente virtual"

    destino.mkdir(parents=True, exist_ok=True)
    nome_base = nome_original or caminho.name
    slug = _slug_ascii(Path(nome_base).stem, fallback="docx")
    raiz = _caminho_unico(destino, f"{slug}-mammoth")
    raiz.mkdir(parents=True, exist_ok=True)
    imagens_dir = raiz / "images"
    imagens_dir.mkdir(parents=True, exist_ok=True)
    html_nome = f"{slug}.html"
    html_path = raiz / html_nome
    imagens: list[Path] = []
    contador = 0

    def salvar_imagem(image: Any) -> dict[str, str]:
        nonlocal contador
        contador += 1
        suffix = _extensao_imagem_mammoth(str(getattr(image, "content_type", "")))
        image_path = imagens_dir / f"image_{contador:03d}{suffix}"
        with image.open() as image_bytes:
            image_path.write_bytes(image_bytes.read())
        imagens.append(image_path)
        return {"src": f"images/{image_path.name}"}

    try:
        with caminho.open("rb") as handle:
            resultado = mammoth.convert_to_html(
                handle,
                convert_image=mammoth.images.img_element(salvar_imagem),
            )
    except Exception as exc:
        return None, [], f"erro ao extrair DOCX com mammoth: {exc}"

    html_path.write_text(resultado.value, encoding="utf-8")
    imagens = sorted(path for path in imagens if path.is_file())
    imagens_relativas = [str(path.relative_to(raiz).as_posix()) for path in imagens]
    avisos = [str(message) for message in getattr(resultado, "messages", [])]
    documentos = [
        {
            "nome": html_nome,
            "origem_docx": nome_base,
            "tipo_extraido": "docx_html_imagens",
            "formato": "html",
            "texto": resultado.value,
            "imagens_extraidas": imagens_relativas,
            "avisos_conversao": avisos,
        }
    ]
    inventario = [nome_base, html_nome, *imagens_relativas]
    imagens_upload = [path for path in imagens if path.suffix.casefold() in {".png", ".jpg", ".jpeg"}]
    arquivos_upload = [str(path) for path in imagens_upload]
    pacote = PacoteEvidencia(
        caminho=caminho,
        tipo="docx_html_imagens",
        documentos=documentos,
        inventario=inventario,
    )
    return pacote, arquivos_upload, ""


def _slug_ascii(valor: str, *, fallback: str = "evidencia", max_len: int = 80) -> str:
    normalizado = unicodedata.normalize("NFKD", valor)
    ascii_text = normalizado.encode("ascii", "ignore").decode("ascii")
    slug = re.sub(r"[^A-Za-z0-9._-]+", "-", ascii_text).strip("._-").lower()
    slug = re.sub(r"-{2,}", "-", slug)
    if not slug:
        slug = fallback
    return slug[:max_len].strip("._-") or fallback


def _nome_upload_seguro(nome_original: str, sufixo: str) -> str:
    suffix = sufixo.lower()
    stem = Path(nome_original).stem or "evidencia"
    slug = _slug_ascii(stem)
    digest = hashlib.sha256(nome_original.encode("utf-8", errors="ignore")).hexdigest()[:10]
    return f"{slug}-{digest}{suffix}"


def _caminho_unico(destino: Path, nome: str) -> Path:
    candidato = destino / nome
    if not candidato.exists():
        return candidato
    stem = candidato.stem
    suffix = candidato.suffix
    contador = 2
    while True:
        proximo = destino / f"{stem}-{contador}{suffix}"
        if not proximo.exists():
            return proximo
        contador += 1


def _copiar_upload_seguro(origem: Path, destino: Path, nome_original: str | None = None) -> Path:
    destino.mkdir(parents=True, exist_ok=True)
    nome = _nome_upload_seguro(nome_original or origem.name, origem.suffix)
    target = _caminho_unico(destino, nome)
    shutil.copyfile(origem, target)
    return target


def _gravar_upload_texto(destino: Path, nome_original: str, texto: str) -> Path:
    destino.mkdir(parents=True, exist_ok=True)
    nome = _nome_upload_seguro(nome_original, ".txt")
    target = _caminho_unico(destino, nome)
    target.write_text(texto, encoding="utf-8")
    return target


def _preparar_documento_para_upload(caminho: Path, destino: Path, nome_original: str | None = None) -> Path | None:
    texto, erro = _extrair_texto_com_markitdown(caminho)
    if erro:
        return None
    return _gravar_upload_texto(destino, nome_original or caminho.name, texto)


def arquivos_compativeis_upload(caminho: str | Path, destino_zip: str | Path | None = None) -> list[str]:
    path = Path(caminho)
    if path.is_file() and path.suffix.lower() in EXTENSOES_UPLOAD_PREPARAVEIS and path.suffix.lower() != ".zip":
        if destino_zip is None:
            return [str(path)] if path.suffix.lower() in EXTENSOES_UPLOAD_DIRETO else []
        destino = Path(destino_zip)
        if path.suffix.lower() in EXTENSOES_WORD_PARA_PDF:
            preparado, erro = _converter_word_para_pdf(path, destino)
            if erro:
                raise RuntimeError(erro)
            return [str(preparado)] if preparado else []
        if path.suffix.lower() == ".xlsx":
            preparado = _preparar_documento_para_upload(path, destino)
            return [str(preparado)] if preparado else []
        return [str(_copiar_upload_seguro(path, destino))]
    if path.suffix.lower() == ".zip" and destino_zip is not None:
        destino = Path(destino_zip)
        destino.mkdir(parents=True, exist_ok=True)
        arquivos = []
        hashes_vistos: set[str] = set()
        with zipfile.ZipFile(path) as archive:
            for name in archive.namelist():
                if name.endswith("/") or not _zip_member_safe(name):
                    continue
                member = Path(name)
                suffix = member.suffix.lower()
                if suffix not in EXTENSOES_UPLOAD_PREPARAVEIS:
                    continue
                # Dedup por conteudo: pular arquivos duplicados (mesmo hash)
                data = archive.read(name)
                hash_conteudo = hashlib.sha256(data).hexdigest()
                if hash_conteudo in hashes_vistos:
                    continue
                hashes_vistos.add(hash_conteudo)
                if suffix in EXTENSOES_WORD_PARA_PDF:
                    staging_dir = destino / "_word_sources"
                    staging_dir.mkdir(parents=True, exist_ok=True)
                    staging = _caminho_unico(staging_dir, _nome_upload_seguro(name, suffix))
                    staging.write_bytes(data)
                    preparado, erro = _converter_word_para_pdf(staging, destino, name)
                    if erro:
                        raise RuntimeError(f"{name}: {erro}")
                    if preparado:
                        arquivos.append(str(preparado))
                    continue
                if suffix == ".xlsx":
                    staging_dir = destino / f"_{suffix[1:]}_sources"
                    staging_dir.mkdir(parents=True, exist_ok=True)
                    staging = _caminho_unico(staging_dir, _nome_upload_seguro(name, suffix))
                    staging.write_bytes(data)
                    preparado = _preparar_documento_para_upload(staging, destino, name)
                    if preparado:
                        arquivos.append(str(preparado))
                    continue
                target = _caminho_unico(destino, _nome_upload_seguro(name, suffix))
                target.write_bytes(data)
                arquivos.append(str(target))
        return arquivos
    return []


def calcular_identidade_analise(
    *,
    auditado: str,
    coluna_evidencia: str,
    nome_original_evidencia: str,
    hash_conteudo: str,
    provider: str,
    model: str,
    prompt_hash: str = "",
    checklist_hash: str = "",
    prompt_version: str,
    reasoning_effort: str = "",
    evidence_processing_mode: str = "",
) -> str:
    artifact_hash = prompt_hash or checklist_hash
    payload = {
        "auditado": auditado,
        "coluna_evidencia": coluna_evidencia,
        "nome_original_evidencia": nome_original_evidencia,
        "hash_conteudo": hash_conteudo,
        "provider": provider,
        "model": model,
        "prompt_hash": artifact_hash,
        "prompt_version": prompt_version,
        "reasoning_effort": reasoning_effort,
        "evidence_processing_mode": evidence_processing_mode,
    }
    encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def gravar_registro_analise(checkpoint: str | Path, registro: dict[str, Any]) -> None:
    path = Path(checkpoint)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as file:
        file.write(json.dumps(registro, ensure_ascii=False, sort_keys=True, default=str))
        file.write("\n")


def _componente_nome_arquivo(valor: str) -> str:
    normalizado = unicodedata.normalize("NFKD", str(valor or "")).encode("ascii", "ignore").decode("ascii")
    seguro = re.sub(r"[^A-Za-z0-9._-]+", "_", normalizado).strip("._-")
    return seguro or "nao_informado"


def nome_checkpoint_padrao(provider: str, model: str) -> str:
    return f"analyses_{_componente_nome_arquivo(provider)}_{_componente_nome_arquivo(model)}.jsonl"


def resolver_checkpoint_saida(out_dir: str | Path, out_file: str | Path | None, provider: str, model: str) -> Path:
    caminho = Path(out_file or nome_checkpoint_padrao(provider, model))
    if caminho.is_absolute():
        return caminho
    return Path(out_dir) / caminho


def montar_registro_analise(
    *,
    identity: str,
    status: str,
    analise: AnaliseCandidata,
    questao_base: str,
    provider: str,
    model: str,
    started_at: dt.datetime,
    reasoning_effort: str = "",
    evidence_processing_mode: str = "",
    result: dict[str, Any] | None = None,
    error: str = "",
    prompt_payload: str | None = None,
    gera_achado: bool = False,
    payload_tokens: dict[str, int] | None = None,
) -> dict[str, Any]:
    finished_at = dt.datetime.now(dt.timezone.utc)
    registro: dict[str, Any] = {
        "identity": identity,
        "status": status,
        "auditado": analise.auditado,
        "questao": questao_base,
        "coluna_evidencia": analise.coluna_evidencia,
        "evidencia": analise.nome_original_evidencia,
        "provider": provider,
        "model": model,
        "started_at": started_at.isoformat(),
        "finished_at": finished_at.isoformat(),
        "duration_seconds": round((finished_at - started_at).total_seconds(), 3),
        "reasoning_effort": reasoning_effort,
        "evidence_processing_mode": evidence_processing_mode,
        "gera_achado": gera_achado,
        "error": error,
    }
    if result is not None:
        registro["result"] = result
    if prompt_payload is not None:
        registro["prompt_payload"] = prompt_payload
    if payload_tokens is not None:
        registro["payload_tokens"] = payload_tokens
    return registro


def carregar_registros_analise(checkpoint: str | Path) -> dict[str, dict[str, Any]]:
    path = Path(checkpoint)
    if not path.is_file():
        return {}
    registros: dict[str, dict[str, Any]] = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        registro = json.loads(line)
        identity = registro.get("identity")
        if isinstance(identity, str):
            registros[identity] = registro
    return registros


def deve_processar_identidade(
    registros: dict[str, dict[str, Any]],
    identity: str,
    *,
    skip_errors: bool = False,
) -> bool:
    registro = registros.get(identity)
    if not registro:
        return True
    status = registro.get("status")
    if status == "completed":
        return False
    if status == "error" and skip_errors:
        return False
    return True


def _join_value(value: Any) -> str:
    if isinstance(value, list):
        return "; ".join(str(item) for item in value)
    if value is None:
        return ""
    return str(value)


def gerar_checkpoint_limpo(checkpoint: str | Path, destino: str | Path) -> int:
    """Copia para ``destino`` apenas os registros ``completed`` do checkpoint.

    O checkpoint original (``analyses_*.jsonl``) mantem todos os registros,
    inclusive ``error``, para permitir reprocessamento via dedup. O arquivo
    limpo contem apenas os ``completed``, pronto para consolidacao e relatorios.
    """
    path = Path(checkpoint)
    if not path.is_file():
        return 0
    destino_path = Path(destino)
    destino_path.parent.mkdir(parents=True, exist_ok=True)
    total = 0
    with path.open(encoding="utf-8") as f_in, destino_path.open("w", encoding="utf-8") as f_out:
        for linha in f_in:
            if not linha.strip():
                continue
            registro = json.loads(linha)
            if registro.get("status") == "completed":
                f_out.write(json.dumps(registro, ensure_ascii=False) + "\n")
                total += 1
    return total


def gerar_relatorio_conformidade(checkpoint: str | Path, destino: str | Path) -> int:
    registros = carregar_registros_analise(checkpoint)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Conformidade"
    headers = [
        "auditado",
        "questao",
        "item",
        "afirmacao_auditado",
        "estado",
        "justificativa",
        "lacunas",
        "referencias",
        "evidencia",
        "provider",
        "model",
        "data_analise",
        "status_revisao_humana",
    ]
    sheet.append(headers)
    total_linhas = 0
    for registro in registros.values():
        result = registro.get("result") if isinstance(registro.get("result"), dict) else {}
        conclusoes = result.get("conclusoes") if isinstance(result, dict) else None
        if not conclusoes and registro.get("status") == "error":
            conclusoes = [
                {
                    "item_codigo": "",
                    "afirmacao_auditado": "",
                    "estado": "erro",
                    "justificativa": registro.get("error", "Erro de analise"),
                    "lacunas": [],
                    "arquivos_referenciados": [],
                }
            ]
        for conclusao in conclusoes or []:
            total_linhas += 1
            sheet.append(
                [
                    registro.get("auditado", ""),
                    registro.get("questao", ""),
                    conclusao.get("item_codigo", ""),
                    conclusao.get("afirmacao_auditado", ""),
                    conclusao.get("estado", ""),
                    conclusao.get("justificativa", ""),
                    _join_value(conclusao.get("lacunas")),
                    _join_value(conclusao.get("arquivos_referenciados")),
                    registro.get("evidencia", ""),
                    registro.get("provider", ""),
                    registro.get("model", ""),
                    registro.get("finished_at", ""),
                    None,
                ]
            )
    destino_path = Path(destino)
    destino_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destino_path)
    return total_linhas


def testar_conexao_provider(provider: str, model: str, api_key: str, reasoning_effort: str = "") -> bool:
    print(f"=== Testando conexao com provider: {provider} | modelo: {model} ===")
    env_key = {"gemini": "GEMINI_API_KEY", "openrouter": "OPENROUTER_API_KEY", "opencodego": "OPENCODEGO_API_KEY", "openai": "OPENAI_API_KEY"}.get(provider, "")
    if provider in REMOTE_PROVIDERS and provider != "openai" and not api_key:
        print(f"[ERRO] A chave de API ({env_key}) nao esta configurada no ambiente!")
        return False

    prompt = "Teste de conexao de IA. Valide que voce recebeu esta mensagem."
    itens = [
        {
            "codigo": "test_01",
            "texto": "Teste de conexao",
            "afirmacao": "Mensagem de teste de conexao",
        }
    ]
    pacote = {
        "documentos": [],
        "inventario": [],
        "erro": "",
        "arquivos_upload": [],
    }
    
    try:
        resultado = executar_provider(
            provider=provider,
            model=model,
            api_key=api_key,
            prompt=prompt,
            auditado="TESTE",
            questao_base="qtest",
            coluna_evidencia="qtestevi",
            itens_afirmados=itens,
            pacote=pacote,
            reasoning_effort=reasoning_effort,
        )
        print("Resultado da chamada:")
        print(json.dumps(resultado, ensure_ascii=False, indent=2))
        if isinstance(resultado, dict) and resultado.get("status") == "completed":
            print("\n[OK] Conexao e comunicacao com o modelo funcionando com sucesso!")
            return True
        else:
            print("\n[ERRO] O provider retornou um status inesperado ou erro.")
            return False
    except Exception as exc:
        print(f"\n[FALHA] Erro ao tentar se comunicar com o provider: {exc}")
        return False


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Pre-analisa evidencias enviadas por auditados.")
    parser.add_argument("respostas", nargs="?", default=None)
    parser.add_argument("evidencias", nargs="?", default=None)
    parser.add_argument("--questionario", default=None)
    parser.add_argument("--provider", default="fake")
    parser.add_argument("--model", default="fake")
    parser.add_argument(
        "--reasoning",
        "--reasoning-effort",
        dest="reasoning_effort",
        choices=["low", "medium", "high"],
        default="",
        help="Nivel de reasoning a solicitar ao provider quando suportado: low, medium ou high.",
    )
    parser.add_argument("--test-connection", action="store_true", help="Testa a conexao com o provider e modelo configurados e encerra.")
    parser.add_argument("--prompts-dir", default=None, help="Diretorio com Prompts de analise por questao.")
    parser.add_argument("--checklists-dir", default=None, help="Alias legado para --prompts-dir.")
    parser.add_argument("--out-dir", default=".saida_analise")
    parser.add_argument(
        "--out-file",
        default=None,
        help="Nome ou caminho do JSONL incremental de saida. Padrao: analyses_<provider>_<model>.jsonl dentro de --out-dir.",
    )
    parser.add_argument("--prompt-version", default="v1")
    parser.add_argument(
        "--only-prompts-present",
        action="store_true",
        help="Processa somente colunas de evidencia com prompt existente e registra item afirmado sem anexo como nao_conforme.",
    )
    parser.add_argument(
        "--only-achados",
        action="store_true",
        help="Processa somente colunas de evidencia cujo prompt esteja marcado como gera_achado: sim "
        "(questoes que ensejam achado no mapa de verificacao). Implicita o comportamento de "
        "--only-prompts-present e registra item afirmado sem anexo como nao_conforme.",
    )
    parser.add_argument(
        "--catalog",
        default=None,
        help="Caminho do catalogo YAML de prompts. Usado para derivar o conjunto de questoes-achado "
        "quando --only-achados esta ativo (le o atributo gera_achado do YAML).",
    )
    parser.add_argument(
        "--rpm",
        type=validar_rpm,
        default=12,
        help="Limita requests por minuto para providers remotos; padrao 12, 0 desativa.",
    )
    parser.add_argument("--skip-errors", action="store_true")
    parser.add_argument(
        "--include-unsubmitted",
        action="store_true",
        help="Inclui respostas sem submitdate. Por padrao, somente respostas submetidas sao analisadas.",
    )
    parser.add_argument("--list-only", action="store_true")
    parser.add_argument("--quiet", action="store_true", help="Nao emite logs estruturados durante a execucao.")
    parser.add_argument("--store-prompts", action="store_true", help="Armazena o payload textual completo enviado ao provider em cada registro JSONL.")
    parser.add_argument(
        "--pdf2md",
        "--pdf-markdown-images",
        dest="pdf2md",
        action="store_true",
        help="Para evidencias PDF, extrai Markdown e imagens com pymupdf4llm e envia esses artefatos ao modelo em vez do PDF original.",
    )
    parser.add_argument(
        "--pdf2md-dpi",
        "--pdf-markdown-images-dpi",
        dest="pdf2md_dpi",
        type=int,
        default=150,
        help="DPI usado nas imagens extraidas de PDF quando --pdf2md estiver ativo. Padrao: 150.",
    )
    parser.add_argument(
        "--docx2html",
        action="store_true",
        help="Para evidencias DOCX, extrai HTML e imagens com Mammoth e envia esses artefatos ao modelo em vez de converter o DOCX para PDF.",
    )
    parser.add_argument(
        "--auditados",
        nargs="+",
        default=None,
        help="Nomes de auditados (firstname) específicos para avaliar.",
    )
    args = parser.parse_args(argv)
    if args.auditados:
        auditados_flat = []
        for item in args.auditados:
            auditados_flat.extend([a.strip() for a in item.split(",") if a.strip()])
        args.auditados = auditados_flat
    if args.test_connection:
        env_key = {"gemini": "GEMINI_API_KEY", "openrouter": "OPENROUTER_API_KEY", "opencodego": "OPENCODEGO_API_KEY", "openai": "OPENAI_API_KEY"}.get(args.provider, "")
        api_key = os.environ.get(env_key, "") if env_key else ""
        sucesso = testar_conexao_provider(args.provider, args.model, api_key, args.reasoning_effort)
        return 0 if sucesso else 1

    if not args.respostas or not args.evidencias or not args.questionario:
        parser.error("Os argumentos respostas, evidencias e --questionario sao obrigatorios quando nao for --test-connection")
    if args.pdf2md_dpi <= 0:
        parser.error("--pdf2md-dpi deve ser maior que zero")

    prompts_dir = args.prompts_dir or args.checklists_dir or "checklists"
    modos_processamento = []
    if args.pdf2md:
        modos_processamento.append("pdf2md")
    if args.docx2html:
        modos_processamento.append("docx2html")
    evidence_processing_mode = "+".join(modos_processamento)
    filtrar_por_prompt = args.only_prompts_present or args.only_achados
    colunas_permitidas = (
        colunas_com_prompt(
            prompts_dir,
            args.questionario,
            only_achados=args.only_achados,
            catalog=args.catalog,
        )
        if filtrar_por_prompt
        else None
    )
    achados_set = carregar_achados_set(prompts_dir, args.catalog)
    rate_limiter = RequestsPerMinuteLimiter(args.rpm)
    log_event(
        "pipeline_started",
        "Inicio da avaliacao de evidencias.",
        quiet=args.quiet,
        respostas=args.respostas,
        evidencias=args.evidencias,
        questionario=args.questionario,
        prompts_dir=prompts_dir,
        provider=args.provider,
        model=args.model,
        reasoning_effort=args.reasoning_effort,
        out_dir=args.out_dir,
        out_file=args.out_file or nome_checkpoint_padrao(args.provider, args.model),
        prompt_version=args.prompt_version,
        rpm=args.rpm,
        include_unsubmitted=args.include_unsubmitted,
        skip_errors=args.skip_errors,
        list_only=args.list_only,
        store_prompts=args.store_prompts,
        pdf2md=args.pdf2md,
        pdf2md_dpi=args.pdf2md_dpi,
        docx2html=args.docx2html,
        auditados=args.auditados,
        only_prompts_present=args.only_prompts_present,
        only_achados=args.only_achados,
        colunas_com_prompt=sorted(colunas_permitidas or []),
    )
    analises = inventariar_analises(
        args.respostas,
        args.evidencias,
        args.questionario,
        include_unsubmitted=args.include_unsubmitted,
        colunas_evidencia_permitidas=colunas_permitidas,
        include_missing_evidence=filtrar_por_prompt,
    )
    if args.auditados:
        auditados_selecionados = {str(a).strip().upper() for a in args.auditados}
        analises = [a for a in analises if str(a.auditado).strip().upper() in auditados_selecionados]
    log_event(
        "inventory_completed",
        "Inventario de evidencias concluido.",
        quiet=args.quiet,
        total_analises=len(analises),
        analises_com_erro=sum(1 for analise in analises if analise.erro),
        auditados=sorted({analise.auditado for analise in analises if analise.auditado}),
    )
    if args.list_only:
        log_event(
            "list_only_started",
            "Listando analises candidatas sem processar evidencias.",
            quiet=args.quiet,
            total_analises=len(analises),
        )
        for analise in analises:
            print(json.dumps(analise.__dict__, ensure_ascii=False, default=str))
        log_event(
            "pipeline_finished",
            "Execucao finalizada em modo list-only.",
            quiet=args.quiet,
            total_analises=len(analises),
        )
        return 0
    contexto = carregar_contexto_questionario(args.questionario)
    checkpoint = resolver_checkpoint_saida(args.out_dir, args.out_file, args.provider, args.model)
    out_dir = checkpoint.parent
    registros = carregar_registros_analise(checkpoint)
    log_event(
        "checkpoint_loaded",
        "Checkpoint carregado para deduplicacao.",
        quiet=args.quiet,
        checkpoint=str(checkpoint),
        registros=len(registros),
    )
    total_processadas = 0
    total_puladas = 0
    total_erros = 0
    total_concluidas = 0
    for index, analise in enumerate(analises, start=1):
        analysis_started_at = dt.datetime.now(dt.timezone.utc)
        questao_base, _ = _base_coluna_evidencia(analise.coluna_evidencia)
        base_log = {
            "index": index,
            "total": len(analises),
            "auditado": analise.auditado,
            "questao": questao_base,
            "coluna_evidencia": analise.coluna_evidencia,
            "evidencia": analise.nome_original_evidencia,
        }
        log_event(
            "analysis_started",
            "Analise candidata iniciada.",
            quiet=args.quiet,
            **base_log,
        )
        prompt = resolver_prompt(prompts_dir, analise.coluna_evidencia)
        gera_achado = (
            questao_base in achados_set
            if achados_set
            else (prompt_gera_achado(prompt.conteudo) if prompt.caminho else False)
        )
        if analise.erro:
            identity = hashlib.sha256(
                json.dumps(analise.__dict__, ensure_ascii=False, sort_keys=True, default=str).encode("utf-8")
            ).hexdigest()
            if not deve_processar_identidade(registros, identity, skip_errors=args.skip_errors):
                total_puladas += 1
                log_event(
                    "analysis_skipped",
                    "Analise com erro de inventario ignorada por checkpoint.",
                    quiet=args.quiet,
                    identity=identity,
                    reason="checkpoint",
                    **base_log,
                )
            else:
                gravar_registro_analise(
                    checkpoint,
                    montar_registro_analise(
                        identity=identity,
                        status="error",
                        analise=analise,
                        questao_base=questao_base,
                        provider=args.provider,
                        model=args.model,
                        started_at=analysis_started_at,
                        reasoning_effort=args.reasoning_effort,
                        evidence_processing_mode=evidence_processing_mode,
                        gera_achado=gera_achado,
                        error=analise.erro,
                    ),
                )
                registros[identity] = {"identity": identity, "status": "error"}
                total_processadas += 1
                total_erros += 1
                log_event(
                    "analysis_recorded_error",
                    "Erro de inventario registrado no checkpoint.",
                    quiet=args.quiet,
                    level="error",
                    identity=identity,
                    error=analise.erro,
                    **base_log,
                )
            continue
        log_event(
            "prompt_resolved",
            "Resolucao do prompt de analise concluida.",
            quiet=args.quiet,
            level="error" if prompt.erro else "info",
            prompt=prompt.nome,
            prompt_hash=prompt.hash_conteudo,
            error=prompt.erro,
            **base_log,
        )
        hash_conteudo = "" if analise.evidencia_ausente else None
        resolucao: ResolucaoEvidencia | None = None
        if hash_conteudo is None and not prompt.erro:
            resolucao = resolver_evidencia(
                analise.auditado,
                args.evidencias,
                analise.upload,
                resposta_id=analise.resposta_id,
                evidence_index=analise.evidence_index,
            )
            log_event(
                "evidence_resolved",
                "Resolucao do arquivo de evidencia concluida.",
                quiet=args.quiet,
                level="error" if resolucao.erro else "info",
                caminho=str(resolucao.caminho) if resolucao.caminho else "",
                nome_decodificado=resolucao.nome_decodificado,
                error=resolucao.erro,
                **base_log,
            )
            hash_conteudo = "" if resolucao.erro else hash_arquivo(resolucao.caminho)
        identity = calcular_identidade_analise(
            auditado=analise.auditado,
            coluna_evidencia=analise.coluna_evidencia,
            nome_original_evidencia=analise.nome_original_evidencia,
            hash_conteudo=hash_conteudo or "",
            provider=args.provider,
            model=args.model,
            prompt_hash=prompt.hash_conteudo,
            prompt_version=args.prompt_version,
            reasoning_effort=args.reasoning_effort,
            evidence_processing_mode=evidence_processing_mode,
        )
        if not deve_processar_identidade(registros, identity, skip_errors=args.skip_errors):
            total_puladas += 1
            log_event(
                "analysis_skipped",
                "Analise ignorada por ja existir no checkpoint.",
                quiet=args.quiet,
                identity=identity,
                reason="checkpoint",
                **base_log,
            )
            continue
        erro = prompt.erro or (resolucao.erro if resolucao else "")
        if erro:
            gravar_registro_analise(
                checkpoint,
                montar_registro_analise(
                    identity=identity,
                    status="error",
                    analise=analise,
                    questao_base=questao_base,
                    provider=args.provider,
                    model=args.model,
                    started_at=analysis_started_at,
                    reasoning_effort=args.reasoning_effort,
                    evidence_processing_mode=evidence_processing_mode,
                    gera_achado=gera_achado,
                    error=erro,
                ),
            )
            registros[identity] = {"identity": identity, "status": "error"}
            total_processadas += 1
            total_erros += 1
            log_event(
                "analysis_recorded_error",
                "Erro antes da chamada ao provider registrado no checkpoint.",
                quiet=args.quiet,
                level="error",
                identity=identity,
                error=erro,
                **base_log,
            )
            continue
        itens = selecionar_itens_afirmados(contexto, analise.coluna_evidencia, _linha_por_id(args.respostas, analise.resposta_id))
        itens = filtrar_itens_por_prompt(itens, prompt)
        itens_provider = preparar_itens_para_prompt(itens, prompt)
        log_event(
            "items_selected",
            "Itens afirmados pelo auditado selecionados para avaliacao.",
            quiet=args.quiet,
            identity=identity,
            total_itens=len(itens),
            itens=[item.codigo for item in itens],
            textos_itens_visiveis=prompt_exibe_texto_itens(prompt.conteudo),
            evidencia_ausente=analise.evidencia_ausente,
            **base_log,
        )
        if analise.evidencia_ausente and not itens:
            total_puladas += 1
            log_event(
                "analysis_skipped_missing_evidence_not_applicable",
                "Analise sem anexo ignorada porque nenhum item avaliavel foi afirmado.",
                quiet=args.quiet,
                identity=identity,
                **base_log,
            )
            continue
        if not itens:
            result = {
                "status": "completed",
                "conclusoes": [],
                "skip_reason": "nenhum_item_afirmado",
            }
            prompt_payload = _conteudo_provider_textual(
                prompt=prompt.conteudo,
                auditado=analise.auditado,
                questao_base=questao_base,
                coluna_evidencia=analise.coluna_evidencia,
                itens_afirmados=itens_provider,
                pacote={"documentos": [], "inventario": [], "erro": "", "arquivos_upload": []},
            )
            gravar_registro_analise(
                checkpoint,
                montar_registro_analise(
                    identity=identity,
                    status=result["status"],
                    analise=analise,
                    questao_base=questao_base,
                    provider=args.provider,
                    model=args.model,
                    started_at=analysis_started_at,
                    reasoning_effort=args.reasoning_effort,
                    evidence_processing_mode=evidence_processing_mode,
                    gera_achado=gera_achado,
                    result=result,
                    error="",
                    prompt_payload=prompt_payload if args.store_prompts else None,
                ),
            )
            registros[identity] = {"identity": identity, "status": result["status"]}
            total_processadas += 1
            total_concluidas += 1
            log_event(
                "analysis_skipped_no_items",
                "Analise concluida sem chamada ao provider porque nenhum item foi afirmado.",
                quiet=args.quiet,
                identity=identity,
                checkpoint=str(checkpoint),
                **base_log,
            )
            continue
        if analise.evidencia_ausente:
            result = resultado_evidencia_ausente(itens_provider)
            prompt_payload = _conteudo_provider_textual(
                prompt=prompt.conteudo,
                auditado=analise.auditado,
                questao_base=questao_base,
                coluna_evidencia=analise.coluna_evidencia,
                itens_afirmados=itens_provider,
                pacote={"documentos": [], "inventario": [], "erro": "evidencia ausente", "arquivos_upload": []},
            )
            gravar_registro_analise(
                checkpoint,
                montar_registro_analise(
                    identity=identity,
                    status=result["status"],
                    analise=analise,
                    questao_base=questao_base,
                    provider=args.provider,
                    model=args.model,
                    started_at=analysis_started_at,
                    reasoning_effort=args.reasoning_effort,
                    evidence_processing_mode=evidence_processing_mode,
                    gera_achado=gera_achado,
                    result=result,
                    error="",
                    prompt_payload=prompt_payload if args.store_prompts else None,
                ),
            )
            registros[identity] = {"identity": identity, "status": result["status"]}
            total_processadas += 1
            total_concluidas += 1
            log_event(
                "analysis_recorded_missing_evidence",
                "Item afirmado sem evidencia registrado como nao_conforme.",
                quiet=args.quiet,
                identity=identity,
                checkpoint=str(checkpoint),
                conclusoes=len(result["conclusoes"]),
                **base_log,
            )
            continue
        if resolucao is None or resolucao.caminho is None:
            raise RuntimeError("resolucao de evidencia ausente em analise com anexo")
        pacote = normalizar_evidencia(resolucao.caminho)
        log_event(
            "evidence_normalized",
            "Evidencia normalizada para envio ao provider.",
            quiet=args.quiet,
            level="warning" if pacote.erro else "info",
            identity=identity,
            tipo=pacote.tipo,
            documentos=len(pacote.documentos),
            inventario=len(pacote.inventario),
            duplicados_ignorados=list(pacote.duplicados_ignorados),
            error=pacote.erro,
            **base_log,
        )
        env_key = {"gemini": "GEMINI_API_KEY", "openrouter": "OPENROUTER_API_KEY", "opencodego": "OPENCODEGO_API_KEY", "openai": "OPENAI_API_KEY"}.get(args.provider, "")
        api_key = os.environ.get(env_key, "") if env_key else ""
        with tempfile.TemporaryDirectory() as upload_tmp:
            try:
                suffix_evidencia = resolucao.caminho.suffix.lower()
                if args.pdf2md and suffix_evidencia == ".pdf":
                    pacote_extraido, arquivos_upload, erro_extracao = _extrair_pdf_markdown_imagens(
                        resolucao.caminho,
                        Path(upload_tmp),
                        resolucao.nome_decodificado or resolucao.caminho.name,
                        dpi=args.pdf2md_dpi,
                    )
                    if erro_extracao:
                        raise RuntimeError(erro_extracao)
                    if pacote_extraido is not None:
                        pacote = pacote_extraido
                elif args.docx2html and suffix_evidencia == ".docx":
                    pacote_extraido, arquivos_upload, erro_extracao = _extrair_docx_html_imagens(
                        resolucao.caminho,
                        Path(upload_tmp),
                        resolucao.nome_decodificado or resolucao.caminho.name,
                    )
                    if erro_extracao:
                        raise RuntimeError(erro_extracao)
                    if pacote_extraido is not None:
                        pacote = pacote_extraido
                else:
                    arquivos_upload = arquivos_compativeis_upload(resolucao.caminho, upload_tmp)
            except RuntimeError as exc:
                result = {"status": "error", "error": f"erro ao preparar evidencia para upload: {exc}"}
                gravar_registro_analise(
                    checkpoint,
                    montar_registro_analise(
                        identity=identity,
                        status=result["status"],
                        analise=analise,
                        questao_base=questao_base,
                        provider=args.provider,
                        model=args.model,
                        started_at=analysis_started_at,
                        reasoning_effort=args.reasoning_effort,
                        evidence_processing_mode=evidence_processing_mode,
                        gera_achado=gera_achado,
                        result=result,
                        error=result["error"],
                    ),
                )
                registros[identity] = {"identity": identity, "status": result["status"]}
                total_processadas += 1
                total_erros += 1
                log_event(
                    "upload_prepare_error",
                    "Erro ao preparar evidencia para upload ao provider.",
                    quiet=args.quiet,
                    level="error",
                    identity=identity,
                    error=result["error"],
                    **base_log,
                )
                continue
            log_event(
                "upload_prepared",
                "Arquivos preparados para upload ao provider.",
                quiet=args.quiet,
                identity=identity,
                total_arquivos=len(arquivos_upload),
                arquivos=[Path(arquivo).name for arquivo in arquivos_upload],
                modo_processamento=evidence_processing_mode,
                **base_log,
            )
            erro_tecnico = erro_tecnico_bloqueante_pacote(pacote, arquivos_upload)
            if erro_tecnico:
                result = {"status": "error", "error": f"erro tecnico ao processar evidencia: {erro_tecnico}"}
                gravar_registro_analise(
                    checkpoint,
                    montar_registro_analise(
                        identity=identity,
                        status=result["status"],
                        analise=analise,
                        questao_base=questao_base,
                        provider=args.provider,
                        model=args.model,
                        started_at=analysis_started_at,
                        reasoning_effort=args.reasoning_effort,
                        evidence_processing_mode=evidence_processing_mode,
                        gera_achado=gera_achado,
                        result=result,
                        error=result["error"],
                    ),
                )
                registros[identity] = {"identity": identity, "status": result["status"]}
                total_processadas += 1
                total_erros += 1
                log_event(
                    "evidence_processing_error",
                    "Erro tecnico de processamento da evidencia registrado antes da chamada ao provider.",
                    quiet=args.quiet,
                    level="error",
                    identity=identity,
                    error=result["error"],
                    **base_log,
                )
                continue
            tokens_info: dict[str, int] | None = None
            if args.provider in REMOTE_PROVIDERS and (api_key or args.provider == "openai"):
                wait_seconds = rate_limiter.wait_seconds()
                if wait_seconds > 0:
                    log_event(
                        "rate_limit_wait",
                        "Aguardando limite de requests por minuto antes da chamada ao provider.",
                        quiet=args.quiet,
                        identity=identity,
                        provider=args.provider,
                        model=args.model,
                        rpm=args.rpm,
                        wait_seconds=round(wait_seconds, 3),
                        resposta_id=analise.resposta_id,
                        itens=[item.codigo for item in itens],
                        **base_log,
                    )
                rate_limiter.wait_and_mark(wait_seconds)
            log_event(
                "provider_started",
                "Chamada ao provider iniciada.",
                quiet=args.quiet,
                identity=identity,
                provider=args.provider,
                model=args.model,
                **base_log,
            )
            pacote_provider = {
                "documentos": pacote.documentos,
                "inventario": pacote.inventario,
                "erro": pacote.erro,
                "arquivos_upload": arquivos_upload,
            }
            prompt_payload = _conteudo_provider_textual(
                prompt=prompt.conteudo,
                auditado=analise.auditado,
                questao_base=questao_base,
                coluna_evidencia=analise.coluna_evidencia,
                itens_afirmados=itens_provider,
                pacote={k: v for k, v in pacote_provider.items() if k != "arquivos_upload"},
            )

            def on_provider_event(event: str, fields: dict[str, Any]) -> None:
                log_event(
                    event,
                    "Evento do provider durante a chamada.",
                    quiet=args.quiet,
                    level="warning",
                    provider=args.provider,
                    model=args.model,
                    identity=identity,
                    **base_log,
                    **fields,
                )

            tokens_info = estimar_tokens_payload(
                prompt=prompt.conteudo,
                auditado=analise.auditado,
                questao_base=questao_base,
                coluna_evidencia=analise.coluna_evidencia,
                itens_afirmados=itens_provider,
                pacote=pacote_provider,
                provider=args.provider,
            )
            log_event(
                "payload_tokens_estimated",
                "Tokens do payload estimados antes da chamada ao provider.",
                quiet=args.quiet,
                identity=identity,
                **tokens_info,
                **base_log,
            )
            # Bloquear payload que excede limite de tokens do provider.
            # Gemini: pular bloqueio da estimativa — a contagem exata e feita
            # via client.models.count_tokens dentro do executar_julgamento_gemini_genai
            limite_tokens = _limite_tokens_provider(args.provider, args.model)
            if limite_tokens and args.provider != "gemini" and tokens_info["tokens_total"] > limite_tokens:
                result = {
                    "status": "error",
                    "error": (
                        f"payload excede limite de tokens do provider: "
                        f"{tokens_info['tokens_total']:,} > {limite_tokens:,} "
                        f"(texto={tokens_info['tokens_texto']:,}, "
                        f"pdfs={tokens_info['n_pdfs']}({tokens_info['tokens_pdfs']:,}), "
                        f"imagens={tokens_info['n_imagens']}({tokens_info['tokens_imagens']:,}))"
                    ),
                }
                log_event(
                    "payload_tokens_exceeded",
                    "Payload bloqueado por exceder limite de tokens do provider.",
                    quiet=args.quiet,
                    level="warning",
                    identity=identity,
                    tokens_total=tokens_info["tokens_total"],
                    limite=limite_tokens,
                    provider=args.provider,
                    model=args.model,
                    **base_log,
                )
            else:
                result = executar_provider(
                    provider=args.provider,
                    model=args.model,
                    api_key=api_key,
                    prompt=prompt.conteudo,
                    auditado=analise.auditado,
                    questao_base=questao_base,
                    coluna_evidencia=analise.coluna_evidencia,
                    itens_afirmados=itens_provider,
                    pacote=pacote_provider,
                    reasoning_effort=args.reasoning_effort,
                    on_event=on_provider_event,
                )
            # Rota 2: todas as chaves exauridas - registrar erro e passar para o proximo item
            if isinstance(result, dict) and result.get("all_keys_exhausted"):
                wait_seconds = float(result.get("retry_after_seconds", 60))
                log_event(
                    "all_keys_exhausted",
                    f"Todas as chaves {args.provider} exauridas. Registrando erro e passando para o proximo item.",
                    quiet=args.quiet,
                    level="warning",
                    provider=args.provider,
                    model=args.model,
                    retry_after_seconds=round(wait_seconds, 1),
                    available_keys=0,
                    **base_log,
                )
                result = {
                    "status": "error",
                    "error": f"todas as chaves {args.provider} exauridas (retry_after {wait_seconds:.0f}s)",
                }
            erro_tecnico_resultado = resultado_indica_erro_tecnico(result)
            if erro_tecnico_resultado:
                result = {
                    "status": "error",
                    "error": erro_tecnico_resultado,
                    "raw_completed_result": result,
                }
        conclusoes = result.get("conclusoes") if isinstance(result, dict) else None
        log_event(
            "provider_finished",
            "Chamada ao provider finalizada.",
            quiet=args.quiet,
            level="error" if result.get("status") == "error" else "info",
            identity=identity,
            provider=args.provider,
            model=args.model,
            status=result.get("status"),
            conclusoes=len(conclusoes or []),
            error=result.get("error", ""),
            **base_log,
        )
        gravar_registro_analise(
            checkpoint,
            montar_registro_analise(
                identity=identity,
                status=result["status"],
                analise=analise,
                questao_base=questao_base,
                provider=args.provider,
                model=args.model,
                started_at=analysis_started_at,
                reasoning_effort=args.reasoning_effort,
                evidence_processing_mode=evidence_processing_mode,
                gera_achado=gera_achado,
                result=result,
                error=result.get("error", ""),
                prompt_payload=prompt_payload if args.store_prompts else None,
                payload_tokens=tokens_info,
            ),
        )
        registros[identity] = {"identity": identity, "status": result["status"]}
        total_processadas += 1
        if result["status"] == "error":
            total_erros += 1
        else:
            total_concluidas += 1
        log_event(
            "analysis_recorded",
            "Resultado da analise gravado no checkpoint.",
            quiet=args.quiet,
            identity=identity,
            status=result["status"],
            checkpoint=str(checkpoint),
            **base_log,
        )
    relatorio = out_dir / "relatorio_conformidade.xlsx"
    linhas_relatorio = gerar_relatorio_conformidade(checkpoint, relatorio)
    log_event(
        "report_generated",
        "Relatorio de conformidade gerado.",
        quiet=args.quiet,
        relatorio=str(relatorio),
        linhas=linhas_relatorio,
    )
    checkpoint_limpo = out_dir / f"analyses_clean_{_componente_nome_arquivo(args.provider)}_{_componente_nome_arquivo(args.model)}.jsonl"
    registros_limpos = gerar_checkpoint_limpo(checkpoint, checkpoint_limpo)
    log_event(
        "clean_checkpoint_generated",
        "Checkpoint limpo (somente completed) gerado.",
        quiet=args.quiet,
        checkpoint_limpo=str(checkpoint_limpo),
        registros=registros_limpos,
    )
    log_event(
        "pipeline_finished",
        "Execucao do pipeline finalizada.",
        quiet=args.quiet,
        total_analises=len(analises),
        processadas=total_processadas,
        puladas=total_puladas,
        concluidas=total_concluidas,
        erros=total_erros,
        checkpoint=str(checkpoint),
        checkpoint_limpo=str(checkpoint_limpo),
        relatorio=str(relatorio),
    )
    return 0


def _limite_tokens_provider(provider: str, model: str) -> int:
    """Retorna o limite aproximado de tokens de entrada por provider/model.

    Fontes: documentacao oficial de cada provider (valores conservadores).
    Retorna 0 se desconhecido (nao bloqueia).
    """
    provider = provider.lower()
    model = model.lower()
    if provider == "gemini":
        return 1_048_576
    if provider == "openrouter":
        # OpenRouter roteia para varios modelos; usar limite conservador
        if "gpt-5.4" in model or "gpt-4" in model:
            return 400_000
        if "minimax" in model:
            return 1_000_000
        return 400_000
    if provider == "openai":
        return 400_000
    if provider == "opencodego":
        return 1_000_000
    return 0


def _linha_por_id(caminho_xlsx: str | Path, resposta_id: Any) -> dict[str, Any]:
    for linha in _rows_from_xlsx(Path(caminho_xlsx)):
        if linha.get("id") == resposta_id:
            return linha
    return {}
