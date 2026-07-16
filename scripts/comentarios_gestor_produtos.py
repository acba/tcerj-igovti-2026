#!/usr/bin/env python3
"""Materializa pareceres e contexto de relatório dos comentários do gestor."""
from __future__ import annotations

import argparse
import json
import re
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts.comentarios_gestor_pipeline import consolidar_submissoes, ler_xlsx, texto


ESTADOS_SANEADOS = {"afastada_na_data_base", "corrigida_posteriormente"}


def ler_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.is_file():
        return []
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def normalizar_data_referencia(value: str) -> str:
    try:
        return datetime.strptime(value, "%d/%m/%Y").strftime("%d/%m/%Y")
    except ValueError as exc:
        raise ValueError("--data-referencia deve estar no formato DD/MM/AAAA") from exc


def _limpar_fundamentacao(value: Any) -> str:
    value = re.sub(r"\s+", " ", texto(value)).strip()
    frases = re.split(r"(?<=[.!?])\s+", value)
    marcas_historicas = re.compile(
        r"data[- ]base|período da auditoria|época da auditoria|correç(?:ão|ao) posterior|regulariza(?:ção|cao) após",
        flags=re.I,
    )
    value = " ".join(frase for frase in frases if not marcas_historicas.search(frase)).strip()
    value = re.sub(r"\b(?:modelo|provider|provedor|juiz de ia)\b", "análise técnica", value, flags=re.I)
    return value or "Não foi registrada fundamentação conclusiva."


def _manifestacoes(respostas: Path) -> tuple[dict[str, dict[str, Any]], set[str]]:
    _, rows = ler_xlsx(respostas)
    validas, _ = consolidar_submissoes(rows)
    por_orgao = {texto(row.get("firstname")).upper(): row for row in validas if texto(row.get("firstname"))}
    return por_orgao, set(por_orgao)


def _decisao_secao1(conclusao: dict[str, Any]) -> tuple[str, str]:
    temporal = texto(conclusao.get("estado_temporal"))
    if temporal in ESTADOS_SANEADOS:
        return "acolhida", "Situação não caracterizada na data de referência"
    if temporal == "mantida":
        motivos = conclusao.get("conclusoes_motivos") or []
        if motivos and any(texto(m.get("estado_motivo")) == "afastado" for m in motivos if isinstance(m, dict)):
            return "parcialmente acolhida", "Situação mantida com alteração parcial de fundamentos"
        return "não acolhida", "Situação mantida"
    return "inconclusiva", "Elementos insuficientes para conclusão"


def _decisao_secao2(conclusoes: list[dict[str, Any]]) -> str:
    estados = [texto(item.get("estado")) for item in conclusoes]
    if estados and all(estado == "conforme" for estado in estados):
        return "acolhida"
    if "conforme" in estados:
        return "parcialmente acolhida"
    if estados and all(estado == "nao_conforme" for estado in estados):
        return "não acolhida"
    return "inconclusiva"


def _texto_equipe_secao1(decisao: str, situacao_atual: str, fundamento: str, data: str) -> str:
    inicio = {
        "acolhida": "A manifestação foi acolhida.",
        "parcialmente acolhida": "A manifestação foi parcialmente acolhida.",
        "não acolhida": "A manifestação não foi acolhida.",
        "inconclusiva": "A manifestação não permitiu conclusão definitiva.",
    }[decisao]
    return f"{inicio} Em {data}, a avaliação concluiu: {situacao_atual.lower()}. {fundamento}"


def _texto_equipe_secao2(decisao: str, itens: list[dict[str, Any]], fundamento: str, data: str) -> str:
    conformes = [texto(item.get("item_codigo")) for item in itens if texto(item.get("estado")) == "conforme"]
    nao_conformes = [texto(item.get("item_codigo")) for item in itens if texto(item.get("estado")) == "nao_conforme"]
    partes = [{
        "acolhida": "A reavaliação foi acolhida.",
        "parcialmente acolhida": "A reavaliação foi parcialmente acolhida.",
        "não acolhida": "A reavaliação não foi acolhida.",
        "inconclusiva": "A reavaliação não permitiu conclusão definitiva.",
    }[decisao], f"A posição em {data} foi aferida sem elevar respostas além do valor originalmente declarado."]
    if conformes:
        partes.append("Itens considerados conformes: " + ", ".join(conformes) + ".")
    if nao_conformes:
        partes.append("Itens que permaneceram não conformes: " + ", ".join(nao_conformes) + ".")
    partes.append(fundamento)
    return " ".join(partes)


def _carregar_revisoes(path: Path | None) -> dict[tuple[str, str, str], dict[str, str]]:
    if not path or not path.is_file():
        return {}
    wb = load_workbook(path, read_only=True, data_only=True)
    revisoes: dict[tuple[str, str, str], dict[str, str]] = {}
    try:
        for ws in wb.worksheets[:2]:
            rows = ws.iter_rows(values_only=True)
            headers = [texto(v) for v in next(rows)]
            for values in rows:
                row = dict(zip(headers, values))
                decisao = texto(row.get("Decisão revisada"))
                fundamento = texto(row.get("Manifestação revisada da equipe"))
                if decisao or fundamento:
                    key = (texto(row.get("Auditado")).upper(), texto(row.get("Seção")), texto(row.get("Código")))
                    revisoes[key] = {"decisao": decisao, "manifestacao_equipe": fundamento}
    finally:
        wb.close()
    return revisoes


def _situacoes_auditoria(path: Path | None) -> dict[str, dict[str, set[str]]]:
    if not path or not path.is_file():
        return {}
    dados = json.loads(path.read_text(encoding="utf-8"))
    result: dict[str, dict[str, set[str]]] = {}
    for sigla, orgao in dados.items():
        achados: set[str] = set()
        situacoes: set[str] = set()
        for proc in orgao.get("procedimentos_executados") or []:
            achado = proc.get("achado") or {}
            if not proc.get("achado_ocorreu") and not achado:
                continue
            numero = texto(achado.get("numero") or proc.get("numero_achado"))
            if numero:
                achados.add(numero)
            situacoes.update(texto(item) for item in achado.get("situacoes_encontradas") or [] if texto(item))
        result[str(sigla).upper()] = {"achados": achados, "situacoes": situacoes}
    return result


def _indices(path: Path | None) -> dict[str, float]:
    if not path or not path.is_file():
        return {}
    df = pd.read_excel(path)
    sigla = next((c for c in df.columns if str(c).casefold() == "sigla"), None)
    indice = next((c for c in df.columns if str(c).casefold() == "igovti"), None)
    if sigla is None or indice is None:
        return {}
    return {texto(row[sigla]).upper(): float(row[indice]) for _, row in df.iterrows() if texto(row[sigla]) and pd.notna(row[indice])}


def materializar(
    *,
    respostas: Path,
    consolidado_secao1: Path,
    consolidado_secao2: Path,
    data_referencia: str,
    output_xlsx: Path,
    output_json: Path,
    revisoes_xlsx: Path | None = None,
    auditoria_anterior: Path | None = None,
    auditoria_atual: Path | None = None,
    contexto_igovti_anterior: Path | None = None,
    contexto_igovti_atual: Path | None = None,
) -> dict[str, Any]:
    data_referencia = normalizar_data_referencia(data_referencia)
    respostas_orgao, respondentes = _manifestacoes(respostas)
    revisoes = _carregar_revisoes(revisoes_xlsx)
    por_orgao: dict[str, dict[str, Any]] = defaultdict(lambda: {
        "data_referencia": data_referencia,
        "status_produto": "minuta para revisão final da equipe de auditoria",
        "respondeu": False,
        "situacoes": [],
        "itens_questionario": [],
        "resumo_impacto": {},
    })
    linhas_s1: list[dict[str, Any]] = []
    linhas_s2: list[dict[str, Any]] = []

    for registro in ler_jsonl(consolidado_secao1):
        if registro.get("status") != "completed":
            continue
        sigla = texto(registro.get("auditado")).upper()
        codigo = texto(registro.get("codigo"))
        conclusoes = (registro.get("result") or {}).get("conclusoes") or []
        if not conclusoes:
            continue
        conclusao = conclusoes[0]
        decisao, situacao_atual = _decisao_secao1(conclusao)
        fundamento = _limpar_fundamentacao(conclusao.get("justificativa"))
        manifestacao_equipe = _texto_equipe_secao1(decisao, situacao_atual, fundamento, data_referencia)
        revisao = revisoes.get((sigla, "1", codigo), {})
        decisao = revisao.get("decisao") or decisao
        manifestacao_equipe = revisao.get("manifestacao_equipe") or manifestacao_equipe
        row_raw = respostas_orgao.get(sigla, {})
        manifestacao_gestor = " ".join(filter(None, [
            texto(row_raw.get(f"{codigo}Conc")), texto(row_raw.get(f"{codigo}Com")),
            texto(row_raw.get(f"{codigo}Jus")),
        ]))
        contexto = registro.get("contexto") or {}
        item = {
            "codigo": codigo,
            "achado": texto(contexto.get("achado")),
            "situacao": texto(contexto.get("situacao") or conclusao.get("item_texto")),
            "manifestacao_gestor": manifestacao_gestor or "Sem texto adicional apresentado.",
            "decisao": decisao,
            "situacao_atual": situacao_atual,
            "manifestacao_equipe": manifestacao_equipe,
            "case_id": texto(registro.get("case_id")),
            "identidade_parecer": texto(registro.get("identity")),
        }
        por_orgao[sigla]["situacoes"].append(item)
        linhas_s1.append({"Auditado": sigla, "Seção": "1", "Código": codigo, **item,
                          "Decisão revisada": "", "Manifestação revisada da equipe": ""})

    for registro in ler_jsonl(consolidado_secao2):
        if registro.get("status") != "completed":
            continue
        sigla = texto(registro.get("auditado")).upper()
        codigo = texto(registro.get("codigo"))
        conclusoes = (registro.get("result") or {}).get("conclusoes") or []
        if not conclusoes:
            continue
        decisao = _decisao_secao2(conclusoes)
        fundamentos = " ".join(dict.fromkeys(_limpar_fundamentacao(c.get("justificativa")) for c in conclusoes))
        manifestacao_equipe = _texto_equipe_secao2(decisao, conclusoes, fundamentos, data_referencia)
        revisao = revisoes.get((sigla, "2", codigo), {})
        decisao = revisao.get("decisao") or decisao
        manifestacao_equipe = revisao.get("manifestacao_equipe") or manifestacao_equipe
        item = {
            "codigo": codigo,
            "itens": ", ".join(texto(c.get("item_codigo")) for c in conclusoes),
            "manifestacao_gestor": texto((registro.get("contexto") or {}).get("comentario")) or "Manifestação baseada exclusivamente nos anexos apresentados.",
            "decisao": decisao,
            "manifestacao_equipe": manifestacao_equipe,
            "case_id": texto(registro.get("case_id")),
            "identidade_parecer": texto(registro.get("identity")),
        }
        por_orgao[sigla]["itens_questionario"].append(item)
        linhas_s2.append({"Auditado": sigla, "Seção": "2", "Código": codigo, **item,
                          "Decisão revisada": "", "Manifestação revisada da equipe": ""})

    anterior, atual = _situacoes_auditoria(auditoria_anterior), _situacoes_auditoria(auditoria_atual)
    idx_anterior, idx_atual = _indices(contexto_igovti_anterior), _indices(contexto_igovti_atual)
    universo = set(por_orgao) | respondentes | set(anterior) | set(atual) | set(idx_anterior) | set(idx_atual)
    for sigla in universo:
        produto = por_orgao[sigla]
        produto["respondeu"] = sigla in respondentes
        antes = anterior.get(sigla, {"achados": set(), "situacoes": set()})
        depois = atual.get(sigla, antes)
        ia, idp = idx_anterior.get(sigla), idx_atual.get(sigla, idx_anterior.get(sigla))
        produto["resumo_impacto"] = {
            "situacoes_antes": len(antes["situacoes"]), "situacoes_atuais": len(depois["situacoes"]),
            "situacoes_removidas": len(antes["situacoes"] - depois["situacoes"]),
            "achados_antes": len(antes["achados"]), "achados_atuais": len(depois["achados"]),
            "achados_removidos": len(antes["achados"] - depois["achados"]),
            "igovti_anterior": ia, "igovti_atual": idp,
            "variacao_igovti": (idp - ia) if ia is not None and idp is not None else None,
        }
        produto["situacoes"].sort(key=lambda x: (x["achado"], x["codigo"]))
        produto["itens_questionario"].sort(key=lambda x: x["codigo"])

    output_json.parent.mkdir(parents=True, exist_ok=True)
    output_json.write_text(json.dumps(dict(sorted(por_orgao.items())), ensure_ascii=False, indent=2), encoding="utf-8")
    _gravar_xlsx(output_xlsx, linhas_s1, linhas_s2, por_orgao)
    return {"auditados": len(por_orgao), "situacoes": len(linhas_s1), "reavaliacoes": len(linhas_s2),
            "xlsx": str(output_xlsx), "json": str(output_json)}


def _gravar_xlsx(path: Path, s1: list[dict[str, Any]], s2: list[dict[str, Any]], produtos: dict[str, Any]) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for nome, rows in (("Seção 1 - situações", s1), ("Seção 2 - itens", s2)):
        ws = wb.create_sheet(nome)
        headers = list(rows[0]) if rows else ["Auditado", "Seção", "Código", "Decisão revisada", "Manifestação revisada da equipe"]
        ws.append(headers)
        for row in rows:
            ws.append([row.get(h, "") for h in headers])
        _formatar(ws)
    ws = wb.create_sheet("Resumo de impactos")
    headers = ["Auditado", "Situações antes", "Situações atuais", "Situações removidas", "Achados antes", "Achados atuais", "Achados removidos", "iGovTI anterior", "iGovTI atual", "Variação iGovTI"]
    ws.append(headers)
    for sigla, produto in sorted(produtos.items()):
        r = produto["resumo_impacto"]
        ws.append([sigla, r["situacoes_antes"], r["situacoes_atuais"], r["situacoes_removidas"], r["achados_antes"], r["achados_atuais"], r["achados_removidos"], r["igovti_anterior"], r["igovti_atual"], r["variacao_igovti"]])
    _formatar(ws)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _formatar(ws: Any) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col in ws.columns:
        letra = col[0].column_letter
        ws.column_dimensions[letra].width = min(70, max(12, max(len(texto(c.value)) for c in col) + 2))


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--respostas-comentarios", type=Path, required=True)
    parser.add_argument("--consolidado-secao1", type=Path, required=True)
    parser.add_argument("--consolidado-secao2", type=Path, required=True)
    parser.add_argument("--data-referencia", required=True)
    parser.add_argument("--output-xlsx", type=Path, required=True)
    parser.add_argument("--output-json", type=Path, required=True)
    parser.add_argument("--revisoes-xlsx", type=Path)
    parser.add_argument("--auditoria-anterior", type=Path)
    parser.add_argument("--auditoria-atual", type=Path)
    parser.add_argument("--contexto-igovti-anterior", type=Path)
    parser.add_argument("--contexto-igovti-atual", type=Path)
    return parser


if __name__ == "__main__":
    args = build_parser().parse_args()
    print(json.dumps(materializar(
        respostas=args.respostas_comentarios, consolidado_secao1=args.consolidado_secao1,
        consolidado_secao2=args.consolidado_secao2, data_referencia=args.data_referencia,
        output_xlsx=args.output_xlsx, output_json=args.output_json, revisoes_xlsx=args.revisoes_xlsx,
        auditoria_anterior=args.auditoria_anterior, auditoria_atual=args.auditoria_atual,
        contexto_igovti_anterior=args.contexto_igovti_anterior, contexto_igovti_atual=args.contexto_igovti_atual,
    ), ensure_ascii=False))
