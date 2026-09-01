import unittest
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_ORIENT

from scripts.gerar_matriz_achados import (
    atualizar_cabecalho,
    atualizar_documento,
    construir_painel_ocorrencias,
    montar_dados,
)
from scripts.gerar_matriz_planejamento import (
    AUDITED_ENTITIES_HEADER,
    AUDIT_OBJECTIVE_HEADER,
    W,
)


MATRIX_PATH = Path("02-Execucao/04-Matriz_Achados/AN06 – Matriz de achados.docx")


class MatrizAchadosHeaderTest(unittest.TestCase):
    def test_cabecalho_usa_escopo_e_objetivo_padronizados(self):
        with ZipFile(MATRIX_PATH) as archive:
            updated = atualizar_cabecalho(archive.read("word/header1.xml"))
        root = ET.fromstring(updated)
        text = "".join(element.text or "" for element in root.iter(f"{W}t"))

        self.assertIn(f"JURISDICIONADOS: {AUDITED_ENTITIES_HEADER}", text)
        self.assertIn(f"OBJETIVO DA AUDITORIA: {AUDIT_OBJECTIVE_HEADER}", text)


class MatrizAchadosConteudoTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dados = montar_dados(Path(".").resolve())

    def test_criterios_usam_ids_sequenciais_da_visao_da_matriz(self):
        achado = self.dados[0]

        self.assertEqual(
            [item["id_exibicao"] for item in achado["criterios"]],
            ["C1", "C2", "C3", "C4", "C5"],
        )
        self.assertEqual(
            {
                publico: [item["id_exibicao"] for item in criterios]
                for publico, criterios in achado["criterios_especificos"].items()
            },
            {
                "Poder Executivo Estadual": ["C6"],
            },
        )

    def test_evidencias_sao_agrupadas_por_situacao(self):
        evidencias = self.dados[0]["evidencias"]

        self.assertEqual(
            evidencias[0]["situacao"],
            "S1.1 - Ausência de área, unidade, setor ou função de TIC formalmente instituída",
        )
        self.assertEqual(evidencias[0]["quantidade"], 5)
        self.assertEqual(evidencias[0]["itens"], ["alternativa f) em q0101"])

    def test_variantes_comuns_sao_agrupadas_com_seus_criterios(self):
        grupos = self.dados[0]["encaminhamentos"][0]["grupos"]

        por_publico = {grupo["publico"]: grupo for grupo in grupos}
        self.assertEqual(por_publico["Demais jurisdicionados"]["criterios"], ["C1", "C5"])
        self.assertEqual(por_publico["Demais jurisdicionados"]["quantidade"], 2)
        self.assertEqual(por_publico["Poder Executivo Estadual"]["criterios"], ["C1", "C6"])
        self.assertEqual(por_publico["Poder Executivo Estadual"]["quantidade"], 3)

    def test_todos_os_encaminhamentos_referenciam_criterios_exibidos(self):
        for achado in self.dados:
            exibidos = {item["id_exibicao"] for item in achado["criterios"]}
            exibidos.update(
                item["id_exibicao"]
                for criterios in achado["criterios_especificos"].values()
                for item in criterios
            )
            for situacao in achado["encaminhamentos"]:
                for grupo in situacao["grupos"]:
                    self.assertTrue(grupo["criterios"])
                    self.assertTrue(set(grupo["criterios"]).issubset(exibidos))
                    self.assertGreater(grupo["quantidade"], 0)

    def test_situacoes_e_variantes_sem_ocorrencia_sao_omitidas(self):
        achado_4 = next(item for item in self.dados if item["achado"].codigo == "A4")

        self.assertNotIn(
            "S4.1",
            {situacao.codigo for situacao in achado_4["achado"].situacoes},
        )
        self.assertEqual(achado_4["criterios_especificos"], {})
        self.assertTrue(
            all(
                grupo["quantidade"] > 0
                for situacao in achado_4["encaminhamentos"]
                for grupo in situacao["grupos"]
            )
        )

    def test_painel_consolida_tipo_resolvido_por_auditado(self):
        situacoes, auditados, valores = construir_painel_ocorrencias(self.dados)

        self.assertEqual(len(situacoes), 23)
        self.assertEqual(len(auditados), 113)
        self.assertEqual(valores[("ARARUAMA", "S1.1")], "R")
        self.assertEqual(valores[("FIA", "S1.1")], "D")
        self.assertEqual(valores[("TJRJ", "S5.3")], "D")
        self.assertNotIn(("TJRJ", "S1.1"), valores)

    def test_documento_final_contem_painel_em_paisagem(self):
        document = Document(str(MATRIX_PATH))
        painel = document.tables[-1]
        headers = [cell.text for cell in painel.rows[0].cells]
        linhas = {
            row.cells[0].text: dict(zip(headers[1:], (cell.text for cell in row.cells[1:])))
            for row in painel.rows[1:]
        }

        self.assertEqual(document.sections[-1].orientation, WD_ORIENT.LANDSCAPE)
        self.assertEqual(len(document.tables), 7)
        self.assertEqual(len(painel.rows), 114)
        self.assertEqual(len(painel.columns), 24)
        self.assertEqual(linhas["ARARUAMA"]["S1.1"], "R")
        self.assertEqual(linhas["FIA"]["S1.1"], "D")
        self.assertEqual(linhas["TJRJ"]["S5.3"], "D")

    def test_documento_exibe_dispositivo_em_negrito_e_vinculos(self):
        with ZipFile(MATRIX_PATH) as archive:
            updated = atualizar_documento(archive.read("word/document.xml"), self.dados)
        root = ET.fromstring(updated)
        namespace = {"w": W.removesuffix("}").removeprefix("{")}
        first_row = root.findall(".//w:tbl", namespace)[0].findall("w:tr", namespace)[1]
        cells = first_row.findall("w:tc", namespace)

        criterion_runs = cells[1].findall(".//w:r", namespace)
        bold_text = [
            "".join(node.text or "" for node in run.findall("w:t", namespace))
            for run in criterion_runs
            if run.find("w:rPr/w:b", namespace) is not None
        ]
        evidence_text = "".join(node.text or "" for node in cells[2].iter(f"{W}t"))
        referral_text = "".join(node.text or "" for node in cells[5].iter(f"{W}t"))

        self.assertIn("C1: COBIT 2019, APO01.04", bold_text)
        self.assertIn("C6: Decreto Estadual nº 48.997/2024, art. 4º", bold_text)
        self.assertIn("S1.1 - Ausência de área", evidence_text)
        self.assertIn("• alternativa f) em q0101.", evidence_text)
        self.assertIn(
            "Poder Executivo Estadual — 3 organizações com ocorrência",
            referral_text,
        )
        self.assertIn("[S1.1, C1, C6]", referral_text)

    def test_colunas_achado_e_encaminhamento_seguem_composicao_visual(self):
        with ZipFile(MATRIX_PATH) as archive:
            updated = atualizar_documento(archive.read("word/document.xml"), self.dados)
        root = ET.fromstring(updated)
        namespace = {"w": W.removesuffix("}").removeprefix("{")}
        first_row = root.findall(".//w:tbl", namespace)[0].findall("w:tr", namespace)[1]
        cells = first_row.findall("w:tc", namespace)
        finding_paragraphs = [
            "".join(node.text or "" for node in paragraph.iter(f"{W}t"))
            for paragraph in cells[0].findall("w:p", namespace)
        ]
        referral_paragraphs = [
            "".join(node.text or "" for node in paragraph.iter(f"{W}t"))
            for paragraph in cells[5].findall("w:p", namespace)
        ]

        self.assertIn(
            "Achado composto pelas seguintes situações efetivamente identificadas:",
            finding_paragraphs,
        )
        self.assertIn(
            "• S1.1 - Ausência de área, unidade, setor ou função de TIC formalmente instituída (5 ocorrências)",
            finding_paragraphs,
        )
        self.assertIn(
            "Ausência de área, unidade, setor ou função de TIC formalmente instituída",
            referral_paragraphs,
        )
        self.assertTrue(any(paragraph.startswith("• Comunicação com Recomendação")
                            for paragraph in referral_paragraphs))
        self.assertTrue(any("[S1.1, C1, C5]" in paragraph for paragraph in referral_paragraphs))
        self.assertIn("Demais jurisdicionados — 2 organizações com ocorrência", referral_paragraphs)
        self.assertNotIn("Todos os grupos alcançados", referral_paragraphs)
        self.assertNotIn(
            "S1.1 - Ausência de área, unidade, setor ou função de TIC formalmente instituída",
            referral_paragraphs,
        )

    def test_achados_segundo_ao_sexto_comecam_em_nova_pagina(self):
        with ZipFile(MATRIX_PATH) as archive:
            updated = atualizar_documento(archive.read("word/document.xml"), self.dados)
        root = ET.fromstring(updated)
        namespace = {"w": W.removesuffix("}").removeprefix("{")}
        body = root.find("w:body", namespace)
        children = list(body)
        tables = body.findall("w:tbl", namespace)

        self.assertEqual(len(tables), 6)
        for table in tables[1:]:
            position = children.index(table)
            self.assertGreater(position, 0)
            previous = children[position - 1]
            self.assertEqual(previous.tag, f"{W}p")
            self.assertIsNotNone(
                previous.find(".//w:br[@w:type='page']", namespace),
                "Cada achado posterior ao primeiro deve ter uma quebra de página estrutural.",
            )


if __name__ == "__main__":
    unittest.main()
