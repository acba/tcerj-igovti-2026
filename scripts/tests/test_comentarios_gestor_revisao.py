from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from scripts.comentarios_gestor_revisao import (
    detectar_prioridades,
    sincronizar_planilha_revisao,
    validar_aprovacoes,
)


def registro(case_id: str, model: str, estado: str, *, judge: bool = False) -> dict:
    return {
        "identity": f"id-{case_id}-{model}",
        "case_id": case_id,
        "secao": "reavaliacao",
        "auditado": "ORG",
        "codigo": "q1001",
        "model": model,
        "model_key": model,
        "provider": "fake",
        "status": "completed",
        "finished_at": "2026-07-23T10:00:00+00:00",
        "result": {
            "status": "completed",
            "conclusoes": [{"item_codigo": "q1001[A]", "estado": estado, "justificativa": "Texto."}],
        },
        "opinioes_validas": 4 if judge else 0,
    }


class ComentariosGestorRevisaoTest(unittest.TestCase):
    def _detectar(self, estados: list[str], estado_juiz: str):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            analyses = root / "analyses.jsonl"
            analyses.write_text(
                "\n".join(json.dumps(registro("case-1", f"m{i}", estado)) for i, estado in enumerate(estados)) + "\n",
                encoding="utf-8",
            )
            consolidado = root / "consolidated_clean.jsonl"
            consolidado.write_text(json.dumps(registro("case-1", "juiz", estado_juiz, judge=True)) + "\n", encoding="utf-8")
            return detectar_prioridades(secoes={"2": ([analyses], consolidado)})

    def test_empate_dois_a_dois_e_prioritario(self) -> None:
        prioridades, divergencias, _ = self._detectar(
            ["conforme", "conforme", "nao_conforme", "nao_conforme"],
            "conforme",
        )
        self.assertEqual(len(prioridades), 1)
        self.assertEqual(divergencias[0]["Tipo da prioridade"], "empate")

    def test_juiz_contra_maioria_e_prioritario(self) -> None:
        prioridades, divergencias, _ = self._detectar(
            ["conforme", "conforme", "conforme", "nao_conforme"],
            "nao_conforme",
        )
        self.assertEqual(len(prioridades), 1)
        self.assertEqual(divergencias[0]["Tipo da prioridade"], "juiz_contra_maioria")

    def test_juiz_alinhado_a_maioria_nao_e_prioritario(self) -> None:
        prioridades, divergencias, universo = self._detectar(
            ["conforme", "conforme", "conforme", "nao_conforme"],
            "conforme",
        )
        self.assertEqual(prioridades, [])
        self.assertEqual(divergencias, [])
        self.assertEqual(universo[0]["Prioritário"], "Não")

    def test_aprovacao_e_preservada_apenas_para_mesmo_hash(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "revisao.xlsx"
            prioridade = {
                "Seção": "2", "Auditado": "ORG", "Código": "q1001", "Motivos da prioridade": "empate",
                "Case ID": "case-1", "Identidade do parecer": "identidade", "Hash do parecer": "hash-1",
                "Origem da decisão": "juiz_ia", "Status da revisão": "Pendente", "Revisor": "",
                "Data da revisão": "", "Fundamento da revisão": "",
            }
            sincronizar_planilha_revisao(path=path, prioridades=[prioridade], divergencias=[], universo=[])
            wb = load_workbook(path)
            ws = wb["Casos prioritários"]
            headers = [cell.value for cell in ws[1]]
            ws.cell(2, headers.index("Status da revisão") + 1, "Aprovado")
            ws.cell(2, headers.index("Revisor") + 1, "Auditor")
            ws.cell(2, headers.index("Data da revisão") + 1, "23/07/2026")
            wb.save(path)
            sincronizar_planilha_revisao(path=path, prioridades=[dict(prioridade)], divergencias=[], universo=[])
            self.assertEqual(validar_aprovacoes(path, [prioridade])["status"], "approved")
            alterada = dict(prioridade, **{"Hash do parecer": "hash-2"})
            sincronizar_planilha_revisao(path=path, prioridades=[alterada], divergencias=[], universo=[])
            self.assertEqual(validar_aprovacoes(path, [alterada])["status"], "awaiting_review")


if __name__ == "__main__":
    unittest.main()

