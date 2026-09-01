import unittest
from pathlib import Path
from types import SimpleNamespace
from tempfile import TemporaryDirectory

from docx import Document
from jinja2 import Environment, FileSystemLoader

from scripts.resources.argos_utils import aplicar_fonte_justificativas_avaliacao

from scripts.gerar_relatorio_consolidado import (
    SITUACOES_POR_ACHADO,
    _nota_criterios_especificos_achado,
)


ROOT = Path(".").resolve()
MAPA = ROOT / "02-Execucao/03-Execucao_Procedimentos/01-Insumos/mapa-verificacao-achados-pos-comentarios-gestor.xlsx"
RESULTADO = ROOT / "02-Execucao/03-Execucao_Procedimentos/02-Resultados_Auditoria/03-pos-comentarios-gestor/resultado_auditoria.json"
AUDITADOS = ROOT / "02-Execucao/03-Execucao_Procedimentos/01-Insumos/bd_auditados.xlsx"
TEMPLATES = ROOT / "03-Relatorios/03-Relatorios_Individuais_Finais"


class CriteriosEspecificosConsolidadoTest(unittest.TestCase):
    def nota(self, numero_achado: int) -> str:
        return _nota_criterios_especificos_achado(
            str(MAPA),
            str(RESULTADO),
            str(AUDITADOS),
            SITUACOES_POR_ACHADO[numero_achado],
        )

    def test_omite_grupo_sem_ocorrencia_de_situacao_especifica(self):
        self.assertEqual(self.nota(4), "")

    def test_menciona_apenas_grupos_e_destinatarios_com_ocorrencia(self):
        nota_1 = self.nota(1)
        self.assertIn("Poder Executivo Estadual", nota_1)
        self.assertIn("S1.1 — determinação para FIA, FTM, TURISRIO", nota_1)
        self.assertNotIn("Poder Judiciário Estadual", nota_1)
        self.assertNotIn("Ministério Público Estadual", nota_1)

        nota_5 = self.nota(5)
        self.assertIn("Poder Executivo Estadual", nota_5)
        self.assertIn("Poder Judiciário Estadual", nota_5)
        self.assertIn("determinação para TJRJ", nota_5)
        self.assertNotIn("Ministério Público Estadual", nota_5)


class BlocosRelatorioIndividualTest(unittest.TestCase):
    def setUp(self):
        self.env = Environment(loader=FileSystemLoader(TEMPLATES))

    def test_criterios_sao_exibidos_sem_ids_e_sem_nota(self):
        criterios = [
            {"id_exibicao": "C2", "descricao": "COBIT 2019, APO01.05.", "especifico": False},
            {
                "id_exibicao": "C8",
                "descricao": "Decreto Estadual nº 48.997/2024, art. 4º.",
                "especifico": True,
                "publico": "Poder Executivo Estadual",
            },
        ]
        enquadramentos = [{
            "publico": "Poder Executivo Estadual",
            "determinacao_por_criterio_especifico": True,
        }]
        auditado = SimpleNamespace(
            sigla="SEEL",
            get_criterios_achado=lambda _: criterios,
            get_enquadramentos_especificos_achado=lambda _: enquadramentos,
        )
        texto = self.env.get_template("bloco_criterios_achado.md").render(
            auditado=auditado,
            nome_achado="Achado",
            achado=SimpleNamespace(numero=1),
        )
        self.assertIn("Decreto Estadual nº 48.997/2024, art. 4º", texto)
        self.assertNotIn("C2", texto)
        self.assertNotIn("C8", texto)
        self.assertNotIn("[^", texto)

    def test_encaminhamento_incorpora_fundamentacao_sem_sufixo_generico(self):
        achado = SimpleNamespace(encaminhamentos=[{
            "tipo": "Determinação",
            "fundamentacao_encaminhamento": (
                "alinhando-se à prática APO01.05 do COBIT 2019 e conforme o art. 4º "
                "do Decreto Estadual nº 48.997/2024"
            ),
            "encaminhamento": "defina formalmente as atribuições da área de TIC",
        }])
        texto = self.env.get_template("bloco_encaminhamentos_achado.md").render(
            auditado=SimpleNamespace(sigla="SEEL"), achado=achado
        )
        self.assertIn("para que, alinhando-se à prática APO01.05", texto)
        self.assertIn("defina formalmente as atribuições", texto)
        self.assertNotIn("A natureza determinativa decorre", texto)
        self.assertNotIn("Foram considerados os critérios específicos", texto)

    def test_evidencias_sao_exibidas_sem_ids(self):
        auditado = SimpleNamespace(
            get_evidencias_numeradas=lambda _: [{
                "ref": "E1",
                "descricao": "Resposta negativa sobre a formalização.",
                "complemento": "",
            }]
        )
        texto = self.env.get_template("bloco_evidencias_achado.md").render(
            auditado=auditado, nome_achado="Achado"
        )
        self.assertIn("Resposta negativa sobre a formalização.", texto)
        self.assertNotIn("E1", texto)

    def test_justificativa_da_avaliacao_recebe_fonte_oito_no_docx(self):
        with TemporaryDirectory() as diretorio:
            caminho = Path(diretorio) / "relatorio.docx"
            documento = Document()
            normal = documento.add_paragraph("Evidência principal.")
            justificativa = documento.add_paragraph(
                "Justificativa da avaliação: documento insuficiente."
            )
            documento.save(caminho)

            aplicar_fonte_justificativas_avaliacao(caminho)

            resultado = Document(caminho)
            self.assertIsNone(resultado.paragraphs[0].runs[0].font.size)
            self.assertEqual(resultado.paragraphs[1].runs[0].font.size.pt, 8)
            self.assertEqual(normal.text, "Evidência principal.")
            self.assertEqual(justificativa.text, "Justificativa da avaliação: documento insuficiente.")


if __name__ == "__main__":
    unittest.main()
