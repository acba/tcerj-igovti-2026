from __future__ import annotations

import hashlib
import json
import tempfile
import unittest
from argparse import Namespace
from pathlib import Path
from unittest.mock import patch

import yaml

from scripts.avaliacao_evidencias.providers.response import (
    json_schema_response_format,
    validar_resultado_ia,
)
from scripts.avaliacao_evidencias.evidence_processing import (
    MAX_LINHAS_XLSX_POR_ABA,
    PacoteEvidencia,
    _dataframe_para_markdown_compacto,
    _extrair_xlsx_abas_visiveis,
    _motivo_filtro_imagem_pdf,
    _remover_referencia_imagem_markdown,
    _tentar_extrair_markdown_pdf,
)
from scripts.avaliacao_evidencias.providers.base import (
    GenericProvider,
    conteudo_provider_textual,
    pacote_textual_sem_documentos_de_arquivos_nativos,
)
from scripts.comentarios_gestor_pipeline import (
    DEFAULT_AJUSTES,
    DEFAULT_CATALOG,
    DEFAULT_LSS,
    DEFAULT_MAPA,
    DEFAULT_RESPOSTAS_BASE,
    DEFAULT_RESULTADO,
    ItemAnterior,
    _agregar_propostas_ajuste,
    _identidade_logica_analise,
    _latest_by_case_model,
    _materializar_consolidado,
    _resolver_valor_positivo,
    _selecionar_situacao,
    avaliar_casos,
    carregar_rotas_catalogo,
    capturar_links_manifestacao,
    consolidar_casos,
    executar_caso,
    extrair_urls_manifestacao,
    filtrar_itens_saneados_secao1,
    gerar_painel_evidencias_pos_comentarios,
    gravar_avaliacoes_modelos_xlsx,
    mapear_ajustes_secao1,
    materializar_checkpoint_limpo_modelo,
    parse_uploads,
    validar_justificativa_publicavel,
)
from scripts.run_comentarios_gestor import DEFAULT_MODELS_CONFIG, carregar_configuracao_modelos, models, run


def opinion(case_id: str, provider: str, model: str) -> dict:
    return {
        "identity": f"{provider}-{model}",
        "case_id": case_id,
        "secao": "reavaliacao",
        "auditado": "ORG",
        "codigo": "q1001",
        "questao": "q1001",
        "coluna_evidencia": "q1001evi",
        "evidencia": "",
        "provider": provider,
        "model": model,
        "model_key": model,
        "status": "completed",
        "finished_at": "2026-07-13T10:00:00+00:00",
        "evidence_paths": [],
        "pacote_contexto_consolidacao": {"documentos": [], "inventario": [], "erro": ""},
        "result": {
            "status": "completed",
            "conclusoes": [
                {
                    "item_codigo": "q1001[A]",
                    "item_texto": "Prática",
                    "afirmacao_auditado": "Sim",
                    "estado": "nao_conforme",
                    "justificativa": "Teste",
                    "lacunas": [],
                    "arquivos_referenciados": [],
                    "trechos_ou_elementos": [],
                    "paginas_ou_localizacao": [],
                }
            ],
        },
    }


class ComentariosGestorPipelineTest(unittest.TestCase):
    def test_complete_pipeline_succeeds_when_only_redundant_evaluators_failed(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            out_dir = Path(tmp)
            args = Namespace(
                action="completo",
                fake=False,
                out_dir=out_dir,
                revisao_humana=None,
                models_config=Path("models.json"),
                preflight_only=False,
            )
            avaliacao_com_aviso = {"modelos": [{"erros": 1}], "secao": "1"}
            consolidacao_valida = {"pendentes_quorum": 0, "secao": "1"}
            integridade = {
                "status": "conforme",
                "casos_integralmente_validos": 598,
                "avisos_avaliacoes_individuais": 146,
            }
            revisao = {"status": "approved", "pendentes": []}
            ajustes = {"ajustes": 160, "pendencias": 0}

            with (
                patch("scripts.run_comentarios_gestor.carregar_configuracao_modelos", return_value={}),
                patch("scripts.run_comentarios_gestor.evaluate_section", side_effect=[
                    avaliacao_com_aviso,
                    {**avaliacao_com_aviso, "secao": "2"},
                ]),
                patch("scripts.run_comentarios_gestor.consolidate_section", side_effect=[
                    consolidacao_valida,
                    {**consolidacao_valida, "secao": "2"},
                ]),
                patch("scripts.run_comentarios_gestor._carregar_saneados_secao1", return_value=({}, [], [])),
                patch("scripts.run_comentarios_gestor.validate_integrity", return_value=integridade),
                patch("scripts.run_comentarios_gestor.review_gate", return_value=revisao),
                patch("scripts.run_comentarios_gestor.generate_adjustments", return_value=ajustes),
            ):
                returncode = run(args)

            resumo = json.loads((out_dir / "resumo-execucao.json").read_text(encoding="utf-8"))

        self.assertEqual(returncode, 0)
        self.assertEqual(resumo["status"], "success")
        self.assertEqual(resumo["integridade"]["avisos_avaliacoes_individuais"], 146)

    def test_captured_manager_link_is_reused_with_its_original_hash(self) -> None:
        documentos = [{"tipo": "comentario_gestor", "texto": "https://www.exemplo.gov.br/pedtic"}]
        with tempfile.TemporaryDirectory() as tmp:
            destino = Path(tmp) / "links"

            def baixar(url: str, pasta: Path):
                caminho = pasta / "pedtic.html"
                caminho.parent.mkdir(parents=True, exist_ok=True)
                caminho.write_text('<script data-rpid="primeiro"></script>', encoding="utf-8")
                return caminho, "", {
                    "url_original": url,
                    "url_final": url,
                    "content_type": "text/html",
                }

            with patch("scripts.comentarios_gestor_pipeline.baixar_recurso_url", side_effect=baixar) as download:
                primeiros, registros_primeiros, erros_primeiros = capturar_links_manifestacao(documentos, destino)
                segundos, registros_segundos, erros_segundos = capturar_links_manifestacao(documentos, destino)

        self.assertEqual(download.call_count, 1)
        self.assertEqual(erros_primeiros, [])
        self.assertEqual(erros_segundos, [])
        self.assertEqual(primeiros, segundos)
        self.assertEqual(registros_primeiros, registros_segundos)

    def test_captured_manager_link_is_refreshed_only_when_requested(self) -> None:
        documentos = [{"tipo": "comentario_gestor", "texto": "https://www.exemplo.gov.br/pedtic"}]
        with tempfile.TemporaryDirectory() as tmp:
            destino = Path(tmp) / "links"
            chamadas = 0

            def baixar(url: str, pasta: Path):
                nonlocal chamadas
                chamadas += 1
                caminho = pasta / f"pedtic-{chamadas}.html"
                caminho.parent.mkdir(parents=True, exist_ok=True)
                caminho.write_text(f'<script data-rpid="{chamadas}"></script>', encoding="utf-8")
                return caminho, "", {
                    "url_original": url,
                    "url_final": url,
                    "content_type": "text/html",
                }

            with patch("scripts.comentarios_gestor_pipeline.baixar_recurso_url", side_effect=baixar):
                _, primeiros, _ = capturar_links_manifestacao(documentos, destino)
                _, segundos, _ = capturar_links_manifestacao(documentos, destino, refresh=True)

        self.assertEqual(chamadas, 2)
        self.assertNotEqual(primeiros[0]["sha256"], segundos[0]["sha256"])

    def test_invalid_cached_manager_link_is_captured_again(self) -> None:
        documentos = [{"tipo": "comentario_gestor", "texto": "https://www.exemplo.gov.br/pedtic"}]
        with tempfile.TemporaryDirectory() as tmp:
            destino = Path(tmp) / "links"
            destino.mkdir(parents=True)
            captura = destino / "pedtic.html"
            captura.write_text("conteudo adulterado", encoding="utf-8")
            (destino / "manifesto-links.json").write_text(
                json.dumps([
                    {
                        "url_original": "https://www.exemplo.gov.br/pedtic",
                        "url_final": "https://www.exemplo.gov.br/pedtic",
                        "content_type": "text/html",
                        "status": "capturado",
                        "caminho": str(captura),
                        "sha256": "hash-incorreto",
                        "capturado_em": "2026-07-22T00:00:00+00:00",
                    }
                ]),
                encoding="utf-8",
            )

            def baixar(url: str, pasta: Path):
                caminho = pasta / "pedtic-novo.html"
                caminho.write_text("nova captura", encoding="utf-8")
                return caminho, "", {
                    "url_original": url,
                    "url_final": url,
                    "content_type": "text/html",
                }

            with patch("scripts.comentarios_gestor_pipeline.baixar_recurso_url", side_effect=baixar) as download:
                caminhos, registros, erros = capturar_links_manifestacao(documentos, destino)

        self.assertEqual(download.call_count, 1)
        self.assertEqual(erros, [])
        self.assertEqual(caminhos[0].name, "pedtic-novo.html")
        self.assertNotEqual(registros[0]["sha256"], "hash-incorreto")

    def test_cached_manager_link_outside_case_directory_is_not_reused(self) -> None:
        documentos = [{"tipo": "comentario_gestor", "texto": "https://www.exemplo.gov.br/pedtic"}]
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            destino = root / "links"
            destino.mkdir()
            externo = root / "fora-do-caso.html"
            externo.write_text("conteudo externo", encoding="utf-8")
            hash_externo = hashlib.sha256(externo.read_bytes()).hexdigest()
            (destino / "manifesto-links.json").write_text(
                json.dumps([
                    {
                        "url_original": "https://www.exemplo.gov.br/pedtic",
                        "url_final": "https://www.exemplo.gov.br/pedtic",
                        "content_type": "text/html",
                        "status": "capturado",
                        "caminho": str(externo),
                        "sha256": hash_externo,
                    }
                ]),
                encoding="utf-8",
            )

            def baixar(url: str, pasta: Path):
                caminho = pasta / "captura-valida.html"
                caminho.write_text("captura valida", encoding="utf-8")
                return caminho, "", {
                    "url_original": url,
                    "url_final": url,
                    "content_type": "text/html",
                }

            with patch("scripts.comentarios_gestor_pipeline.baixar_recurso_url", side_effect=baixar) as download:
                caminhos, _, erros = capturar_links_manifestacao(documentos, destino)

        self.assertEqual(download.call_count, 1)
        self.assertEqual(erros, [])
        self.assertEqual(caminhos[0].name, "captura-valida.html")

    def test_unchanged_captured_link_skips_second_judge_consolidation(self) -> None:
        url = "https://www.exemplo.gov.br/pedtic"
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            paths = []
            for provider_name, model in (("p1", "m1"), ("p2", "m2")):
                record = opinion("case-link-cache", provider_name, model)
                record["pacote_contexto_consolidacao"] = {
                    "documentos": [{"nome": "comentario_gestor.txt", "tipo": "comentario_gestor", "texto": url}],
                    "inventario": ["comentario_gestor.txt"],
                    "erro": "",
                }
                path = root / f"{provider_name}.jsonl"
                path.write_text(json.dumps(record) + "\n", encoding="utf-8")
                paths.append(path)

            def baixar(link: str, pasta: Path):
                caminho = pasta / "pedtic.html"
                caminho.parent.mkdir(parents=True, exist_ok=True)
                caminho.write_text('<script data-rpid="dinamico"></script>', encoding="utf-8")
                return caminho, "", {
                    "url_original": link,
                    "url_final": link,
                    "content_type": "text/html",
                }

            with patch("scripts.comentarios_gestor_pipeline.baixar_recurso_url", side_effect=baixar) as download:
                first = consolidar_casos(
                    analyses_files=paths, out_dir=root / "out", secao="2",
                    expected_models=["m1", "m2"], judge_provider="fake", judge_model="judge",
                    min_opinions=2, quiet=True,
                )
                second = consolidar_casos(
                    analyses_files=paths, out_dir=root / "out", secao="2",
                    expected_models=["m1", "m2"], judge_provider="fake", judge_model="judge",
                    min_opinions=2, quiet=True,
                )

        self.assertEqual(download.call_count, 1)
        self.assertEqual(first["concluidos"], 1)
        self.assertEqual(second["pulados"], 1)
        self.assertEqual(second["concluidos"], 0)

    def test_catalog_v3_preserves_evaluator_prompts_and_changes_only_judge(self) -> None:
        catalogs = Path("scripts/avaliacao_evidencias/prompt_catalogs")
        v2 = yaml.safe_load((catalogs / "igovti_2026_comentarios_gestor_atual_v2.yml").read_text(encoding="utf-8"))
        v3 = yaml.safe_load((catalogs / "igovti_2026_comentarios_gestor_atual_v3.yml").read_text(encoding="utf-8"))
        self.assertEqual(v3["prompt_situacoes"], v2["prompt_situacoes"])
        self.assertEqual(v3["prompt_reavaliacao_prefixo"], v2["prompt_reavaliacao_prefixo"])
        self.assertNotEqual(v3["prompt_juiz_situacoes"], v2["prompt_juiz_situacoes"])
        self.assertNotEqual(v3["prompt_juiz_reavaliacao"], v2["prompt_juiz_reavaliacao"])

    def test_xlsx_markdown_is_compact_and_escapes_cell_content(self) -> None:
        import pandas as pd

        dataframe = pd.DataFrame(
            [
                ["curto", "a|b", "linha 1\nlinha 2"],
                ["x", "y", "z"],
            ],
            columns=["codigo", "com|pipe", "multilinha"],
        )

        markdown = _dataframe_para_markdown_compacto(dataframe, pd)

        self.assertEqual(
            markdown.splitlines(),
            [
                r"|codigo|com\|pipe|multilinha|",
                "|---|---|---|",
                r"|curto|a\|b|linha 1<br>linha 2|",
                "|x|y|z|",
            ],
        )
        self.assertNotIn("| codigo ", markdown)

    def test_xlsx_markdown_limits_visible_sheet_to_first_thousand_rows(self) -> None:
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "planilha-grande.xlsx"
            wb = Workbook()
            visible = wb.active
            visible.title = "Visivel"
            visible.append(["codigo", "valor"])
            for index in range(MAX_LINHAS_XLSX_POR_ABA + 5):
                visible.append([index, f"linha-{index}"])
            hidden = wb.create_sheet("Oculta")
            hidden.sheet_state = "hidden"
            hidden.append(["segredo"])
            hidden.append(["nao-deve-aparecer"])
            wb.save(path)
            wb.close()

            markdown, error = _extrair_xlsx_abas_visiveis(path)

        self.assertEqual(error, "")
        self.assertIn("## Visivel", markdown)
        self.assertIn("Conteúdo truncado", markdown)
        self.assertIn("linha-998", markdown)
        self.assertNotIn("linha-999", markdown)
        self.assertNotIn("Oculta", markdown)
        self.assertNotIn("nao-deve-aparecer", markdown)

    def test_individual_evaluation_uses_context_when_attachment_is_unprocessable(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            attachment = Path(tmp) / "anexo-vazio.pdf"
            attachment.write_bytes(b"")
            case = {
                "case_id": "case-warning",
                "secao": "situacoes",
                "auditado": "ORG",
                "codigo": "A1G1",
                "coluna_evidencia": "A1G1Evi",
                "resposta_id": 1,
                "prompt": "Avalie o comentário e não presuma o conteúdo de anexos ilegíveis.",
                "prompt_hash": "prompt-hash",
                "response_profile": "manager_comments_temporal",
                "itens": [],
                "contexto": {"motivos": []},
                "documentos_contexto": [
                    {"nome": "manifestacao_gestor.json", "tipo": "manifestacao_gestor", "texto": "Compromisso futuro."},
                    {"nome": "contexto_situacao.json", "tipo": "contexto_auditoria", "texto": "Situação encontrada."},
                ],
                "evidence_paths": [attachment],
                "evidencias": [{"name": attachment.name}],
            }
            pacote_invalido = PacoteEvidencia(
                caminho=attachment,
                tipo="pdf",
                documentos=[],
                inventario=[],
                erro="evidencia sem conteudo processavel para avaliacao",
            )
            resultado_provider = {
                "status": "completed",
                "conclusoes": [{
                    "item_codigo": "A1G1",
                    "estado": "nao_conforme",
                    "justificativa": (
                        "O conteúdo do arquivo está inacessível, mas o comentário descreve "
                        "somente compromisso futuro e não afasta a situação."
                    ),
                }],
            }
            with patch(
                "scripts.comentarios_gestor_pipeline.preparar_evidencia_para_provider",
                return_value=(pacote_invalido, [], ""),
            ), patch(
                "scripts.comentarios_gestor_pipeline.executar_provider",
                return_value=resultado_provider,
            ) as provider:
                registro = executar_caso(
                    case,
                    provider="fake",
                    model="fake",
                    reasoning="high",
                    pdf2md=True,
                    docx2html=True,
                )

        self.assertEqual(registro["status"], "completed")
        self.assertEqual(len(registro["avisos_processamento_evidencias"]), 1)
        pacote_enviado = provider.call_args.kwargs["pacote"]
        nomes = [documento["nome"] for documento in pacote_enviado["documentos"]]
        self.assertIn("manifestacao_gestor.json", nomes)
        self.assertIn("avisos_processamento_anexos.json", nomes)
        self.assertEqual(pacote_enviado["arquivos_upload"], [])

    def test_judge_uses_valid_opinions_when_attachment_is_unprocessable(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            attachment = root / "anexo-vazio.pdf"
            attachment.write_bytes(b"")
            paths = []
            for provider_name, model in (("p1", "m1"), ("p2", "m2")):
                record = opinion("case-warning", provider_name, model)
                record["evidence_paths"] = [str(attachment)]
                path = root / f"{provider_name}.jsonl"
                path.write_text(json.dumps(record) + "\n", encoding="utf-8")
                paths.append(path)
            pacote_invalido = PacoteEvidencia(
                caminho=attachment,
                tipo="pdf",
                documentos=[],
                inventario=[],
                erro="evidencia sem conteudo processavel para avaliacao",
            )
            resultado_provider = {
                "status": "completed",
                "conclusoes": [{"item_codigo": "q1001[A]", "estado": "nao_conforme"}],
            }
            with patch(
                "scripts.comentarios_gestor_pipeline.preparar_evidencia_para_provider",
                return_value=(pacote_invalido, [], ""),
            ), patch(
                "scripts.comentarios_gestor_pipeline.executar_provider",
                return_value=resultado_provider,
            ) as provider:
                result = consolidar_casos(
                    analyses_files=paths,
                    out_dir=root / "out",
                    secao="2",
                    expected_models=["m1", "m2"],
                    judge_provider="fake",
                    judge_model="judge",
                    min_opinions=2,
                    quiet=True,
                )
            clean = [
                json.loads(line)
                for line in (root / "out/consolidated_clean.jsonl").read_text(encoding="utf-8").splitlines()
            ]

        self.assertEqual(result["concluidos"], 1)
        self.assertEqual(clean[0]["status"], "completed")
        self.assertEqual(len(clean[0]["avisos_processamento_evidencias"]), 1)
        pacote_enviado = provider.call_args.kwargs["pacote"]
        nomes = [documento["nome"] for documento in pacote_enviado["documentos"]]
        self.assertIn("avaliacoes_individuais.json", nomes)
        self.assertIn("avisos_processamento_anexos.json", nomes)

    def test_judge_prefers_existing_attachment_path_from_any_opinion(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            missing = root / "extracao-antiga.pdf"
            current = root / "extracao-atual.pdf"
            current.write_bytes(b"pdf")
            first = opinion("case-path", "p1", "m1")
            second = opinion("case-path", "p2", "m2")
            first["evidence_paths"] = [str(missing)]
            second["evidence_paths"] = [str(current)]
            paths = []
            for index, record in enumerate([first, second], start=1):
                path = root / f"p{index}.jsonl"
                path.write_text(json.dumps(record) + "\n", encoding="utf-8")
                paths.append(path)
            pacote = PacoteEvidencia(
                caminho=current,
                tipo="pdf",
                documentos=[{"nome": "atual", "tipo": "texto", "conteudo": "ok"}],
                inventario=["atual"],
                erro="",
            )
            with patch(
                "scripts.comentarios_gestor_pipeline.preparar_evidencia_para_provider",
                return_value=(pacote, [], ""),
            ) as preparar:
                result = consolidar_casos(
                    analyses_files=paths,
                    out_dir=root / "out",
                    secao="2",
                    expected_models=["m1", "m2"],
                    judge_provider="fake",
                    judge_model="judge",
                    min_opinions=2,
                    quiet=True,
                )
        self.assertEqual(result["concluidos"], 1)
        self.assertEqual(preparar.call_args.args[0], current)

    def test_saneamento_ajusta_detalhamento_sem_resposta_original_para_sim(self) -> None:
        valor, motivo = _resolver_valor_positivo(
            auditado="ORG", codigo="q1001ext[A]", fonte="questionario", ajustes_evidencias={},
        )
        self.assertEqual(valor, "Sim")
        self.assertIn("saneada", motivo)

    def test_saneamento_converte_detalhamento_e_subitem_binario(self) -> None:
        detalhe, _ = _resolver_valor_positivo(
            auditado="ORG", codigo="q2101ext[C]", fonte="questionario",
            ajustes_evidencias={}, resposta_original="Não",
        )
        subitem, _ = _resolver_valor_positivo(
            auditado="ORG", codigo="q2708[D]", fonte="questionario",
            ajustes_evidencias={}, resposta_original="Não",
        )
        self.assertEqual(detalhe, "Sim")
        self.assertEqual(subitem, "Sim")

    def test_q0101_preserva_original_e_q0103_aplica_regras_especificas(self) -> None:
        q0101, _ = _resolver_valor_positivo(
            auditado="ORG", codigo="q0101", fonte="questionario",
            ajustes_evidencias={}, resposta_original="e) Híbrida",
        )
        q0103d, _ = _resolver_valor_positivo(
            auditado="ORG", codigo="q0103[D]", fonte="questionario",
            ajustes_evidencias={}, resposta_original="Não",
        )
        q0103g, _ = _resolver_valor_positivo(
            auditado="ORG", codigo="q0103[G]", fonte="questionario",
            ajustes_evidencias={}, resposta_original="Sim",
        )
        self.assertEqual(q0101, "e) Híbrida")
        self.assertEqual(q0103d, "Sim")
        self.assertEqual(q0103g, "Não")

    def test_checkpoint_limpo_exclui_identidade_de_prompt_anterior(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            checkpoint = root / "analyses_model_m.jsonl"
            registros = [
                {"identity": "antiga", "logical_identity": "antiga", "case_id": "c", "model_key": "m", "status": "completed"},
                {"identity": "nova", "logical_identity": "nova", "case_id": "c", "model_key": "m", "status": "completed"},
            ]
            checkpoint.write_text("\n".join(json.dumps(r) for r in registros) + "\n", encoding="utf-8")
            clean = materializar_checkpoint_limpo_modelo(
                root, model_key="m", routes=[], expected_identities={"nova"},
            )
            vigentes = [json.loads(line) for line in clean.read_text(encoding="utf-8").splitlines()]
        self.assertEqual([r["logical_identity"] for r in vigentes], ["nova"])

    def test_opiniao_sem_conclusao_nao_conta_para_quorum(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "analyses.jsonl"
            valida = opinion("caso", "gemini", "modelo-valido")
            vazia = opinion("caso", "opencodego", "modelo-vazio")
            vazia["result"]["conclusoes"] = []
            path.write_text(
                "\n".join(json.dumps(r) for r in [valida, vazia]) + "\n",
                encoding="utf-8",
            )
            grupos = _latest_by_case_model([path])
        self.assertEqual(set(grupos["caso"]), {"modelo-valido"})

    def test_checkpoint_preserva_parecer_valido_anterior_a_resposta_vazia(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            valida = opinion("caso", "opencodego", "qwen")
            valida.update({"identity": "identidade", "logical_identity": "identidade", "model_key": "qwen"})
            vazia = json.loads(json.dumps(valida))
            vazia["finished_at"] = "2026-07-14T10:00:00+00:00"
            vazia["result"]["conclusoes"] = []
            (root / "analyses_model_qwen.jsonl").write_text(
                "\n".join(json.dumps(r) for r in [valida, vazia]) + "\n",
                encoding="utf-8",
            )
            clean = materializar_checkpoint_limpo_modelo(
                root,
                model_key="qwen",
                routes=[],
                expected_identities={"identidade"},
            )
            vigente = json.loads(clean.read_text(encoding="utf-8").strip())
        self.assertTrue(vigente["result"]["conclusoes"])
        self.assertEqual(vigente["finished_at"], valida["finished_at"])

    def test_consolidado_limpo_exclui_caso_historico_fora_do_manifesto(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            checkpoint = root / "consolidated.jsonl"
            clean = root / "consolidated_clean.jsonl"
            registros = [
                {"case_id": "vigente", "status": "completed", "finished_at": "2026-07-01"},
                {"case_id": "obsoleto", "status": "completed", "finished_at": "2026-07-02"},
            ]
            checkpoint.write_text(
                "\n".join(json.dumps(r) for r in registros) + "\n",
                encoding="utf-8",
            )
            vigentes = _materializar_consolidado(
                checkpoint,
                clean,
                expected_case_ids={"vigente"},
            )
        self.assertEqual([r["case_id"] for r in vigentes], ["vigente"])

    def test_default_models_configuration_is_external_and_includes_glm(self) -> None:
        config = carregar_configuracao_modelos(DEFAULT_MODELS_CONFIG)
        pairs = {(item["provider"], item["model"]) for item in config["evaluators"]}
        self.assertIn(("openrouter", "z-ai/glm-5.2"), pairs)
        self.assertEqual(config["judge"]["min_valid_opinions"], 3)

    def test_models_configuration_rejects_duplicate_pairs(self) -> None:
        config = json.loads(DEFAULT_MODELS_CONFIG.read_text(encoding="utf-8"))
        config["evaluators"].append(dict(config["evaluators"][0], name="duplicado"))
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "models.json"
            path.write_text(json.dumps(config), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "provider/model duplicado"):
                carregar_configuracao_modelos(path)

    def test_models_configuration_filters_disabled_routes(self) -> None:
        config = json.loads(DEFAULT_MODELS_CONFIG.read_text(encoding="utf-8"))
        config["evaluators"][0]["enabled"] = False
        config["judge"]["min_valid_opinions"] = 1
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "models.json"
            path.write_text(json.dumps(config), encoding="utf-8")
            loaded = carregar_configuracao_modelos(path)
        self.assertEqual(
            len(models(loaded, False)),
            sum(item.get("enabled", True) for item in config["evaluators"]),
        )

    def test_models_configuration_defaults_enabled_and_model_key(self) -> None:
        config = json.loads(DEFAULT_MODELS_CONFIG.read_text(encoding="utf-8"))
        config["evaluators"][0].pop("enabled")
        config["evaluators"][0].pop("model_key")
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "models.json"
            path.write_text(json.dumps(config), encoding="utf-8")
            loaded = carregar_configuracao_modelos(path)
        self.assertTrue(loaded["evaluators"][0]["enabled"])
        self.assertEqual(loaded["evaluators"][0]["model_key"], loaded["evaluators"][0]["model"])
        self.assertEqual(loaded["evaluators"][0]["pdf_detail"], "auto")
        self.assertEqual(loaded["judge"]["pdf_detail"], "auto")

    def test_models_configuration_keeps_historical_model_without_executing_it(self) -> None:
        config = json.loads(DEFAULT_MODELS_CONFIG.read_text(encoding="utf-8"))
        target = next(item for item in config["evaluators"] if item.get("enabled", True))
        target["execute"] = False
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "models.json"
            path.write_text(json.dumps(config), encoding="utf-8")
            loaded = carregar_configuracao_modelos(path)
        self.assertIn(target["model_key"], {item["model_key"] for item in models(loaded, False)})
        self.assertNotIn(
            target["model_key"],
            {item["model_key"] for item in models(loaded, False, execute_only=True)},
        )

    def test_models_configuration_rejects_execution_of_disabled_model(self) -> None:
        config = json.loads(DEFAULT_MODELS_CONFIG.read_text(encoding="utf-8"))
        config["evaluators"][0]["enabled"] = False
        config["evaluators"][0]["execute"] = True
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "models.json"
            path.write_text(json.dumps(config), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "execute exige enabled=true"):
                carregar_configuracao_modelos(path)

    def test_models_configuration_rejects_invalid_pdf_detail(self) -> None:
        config = json.loads(DEFAULT_MODELS_CONFIG.read_text(encoding="utf-8"))
        config["evaluators"][0]["pdf_detail"] = "medium"
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "models.json"
            path.write_text(json.dumps(config), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "pdf_detail"):
                carregar_configuracao_modelos(path)

    def test_models_configuration_allows_alternative_disabled_route(self) -> None:
        config = json.loads(DEFAULT_MODELS_CONFIG.read_text(encoding="utf-8"))
        first = config["evaluators"][0]
        config["evaluators"].append(
            {
                **first,
                "name": "gemini-alternativo",
                "provider": "openrouter",
                "model": "google/gemini-3.1-flash-lite",
                "enabled": False,
            }
        )
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "models.json"
            path.write_text(json.dumps(config), encoding="utf-8")
            loaded = carregar_configuracao_modelos(path)
        self.assertEqual(
            len(models(loaded, False)),
            sum(item.get("enabled", True) for item in config["evaluators"]),
        )

    def test_models_configuration_rejects_two_active_routes_for_model_key(self) -> None:
        config = json.loads(DEFAULT_MODELS_CONFIG.read_text(encoding="utf-8"))
        first = config["evaluators"][0]
        config["evaluators"].append(
            {
                **first,
                "name": "gemini-alternativo",
                "provider": "openrouter",
                "model": "google/gemini-3.1-flash-lite",
                "enabled": True,
            }
        )
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "models.json"
            path.write_text(json.dumps(config), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "uma rota.*model_key"):
                carregar_configuracao_modelos(path)

    def test_logical_identity_ignores_provider_but_tracks_payload(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            evidence = Path(tmp) / "A.pdf"
            evidence.write_bytes(b"conteudo")
            kwargs = {
                "case_id_value": "case",
                "model_key": "gpt-5.4-mini",
                "prompt_hash": "prompt",
                "reasoning": "high",
                "pdf2md": False,
                "docx2html": False,
                "contexto": {"situacao": "A"},
                "evidence_paths": [evidence],
            }
            first = _identidade_logica_analise(**kwargs)
            second = _identidade_logica_analise(**kwargs)
            changed = _identidade_logica_analise(**{**kwargs, "pdf2md": True})
            changed_detail = _identidade_logica_analise(**{**kwargs, "pdf_detail": "low"})
        self.assertEqual(first, second)
        self.assertNotEqual(first, changed)
        self.assertNotEqual(first, changed_detail)

    def test_provider_switch_reuses_canonical_model_checkpoint(self) -> None:
        case = {
            "case_id": "case", "secao": "reavaliacao", "auditado": "ORG", "codigo": "q1",
            "coluna_evidencia": "q1evi", "prompt_hash": "prompt", "prompt_version": "v1",
            "contexto": {}, "evidence_paths": [], "evidencias": [], "documentos_contexto": [],
        }
        routes = [
            {"provider": "openai", "model": "gpt-5.4-mini", "model_key": "gpt-5.4-mini"},
            {"provider": "openrouter", "model": "openai/gpt-5.4-mini", "model_key": "gpt-5.4-mini"},
        ]

        def completed(caso, *, provider, model, **_):
            record = opinion(caso["case_id"], provider, model)
            record.update({"contexto": {}, "evidence_paths": [], "prompt_hash": "prompt"})
            return record

        with tempfile.TemporaryDirectory() as tmp, patch(
            "scripts.comentarios_gestor_pipeline.executar_caso", side_effect=completed
        ):
            out = Path(tmp)
            first = avaliar_casos(
                [case], provider="openai", model="gpt-5.4-mini", model_key="gpt-5.4-mini",
                out_dir=out, routes=routes, quiet=True,
            )
            second = avaliar_casos(
                [case], provider="openrouter", model="openai/gpt-5.4-mini", model_key="gpt-5.4-mini",
                out_dir=out, routes=routes, quiet=True,
            )
        self.assertEqual(first["processados"], 1)
        self.assertEqual(second["processados"], 0)
        self.assertEqual(second["pulados"], 1)

    def test_missing_evidence_hash_fails_only_its_case(self) -> None:
        def case(case_id: str, evidence_paths: list[Path]) -> dict:
            return {
                "case_id": case_id,
                "secao": "reavaliacao",
                "auditado": "ORG",
                "codigo": case_id,
                "coluna_evidencia": "q1evi",
                "prompt_hash": "prompt",
                "prompt_version": "v1",
                "contexto": {},
                "evidence_paths": evidence_paths,
                "evidencias": [],
                "documentos_contexto": [],
            }

        def completed(caso, *, provider, model, **_):
            record = opinion(caso["case_id"], provider, model)
            record.update({"contexto": {}, "evidence_paths": [], "prompt_hash": "prompt"})
            return record

        with tempfile.TemporaryDirectory() as tmp, patch(
            "scripts.comentarios_gestor_pipeline.executar_caso", side_effect=completed
        ) as execute:
            root = Path(tmp)
            result = avaliar_casos(
                [case("ausente", [root / "nao-existe.pdf"]), case("valido", [])],
                provider="fake",
                model="fake",
                model_key="fake",
                out_dir=root / "out",
                quiet=True,
            )
        self.assertEqual(result["processados"], 2)
        self.assertEqual(result["erros"], 1)
        self.assertEqual(result["concluidos"], 1)
        self.assertEqual(execute.call_count, 1)

    def test_pdf2md_does_not_change_process_working_directory(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "origem.pdf"
            source.write_bytes(b"pdf sintetico")
            images = root / "saida" / "img"
            images.mkdir(parents=True)
            cwd_before = Path.cwd()
            absolute_image = (images / "pagina.png").resolve().as_posix()
            with patch(
                "pymupdf4llm.to_markdown",
                return_value=f"![imagem]({absolute_image})",
            ) as convert:
                markdown, error = _tentar_extrair_markdown_pdf(
                    source,
                    images.parent,
                    images,
                    "origem",
                    150,
                )
            self.assertEqual(Path.cwd(), cwd_before)
            self.assertEqual(error, "")
            self.assertEqual(markdown, "![imagem](img/pagina.png)")
            self.assertTrue(Path(convert.call_args.kwargs["image_path"]).is_absolute())

    def test_pdf2md_filters_small_files_and_tiny_fragments(self) -> None:
        from PIL import Image

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            tiny = root / "clausula.png"
            small_file = root / "icone.png"
            narrow_but_useful = root / "cabecalho.png"
            Image.new("RGB", (78, 22), "white").save(tiny)
            Image.new("RGB", (200, 200), "white").save(small_file)
            Image.effect_noise((2000, 100), 100).save(narrow_but_useful)

            self.assertIn("arquivo_derivado_ate_25kb", _motivo_filtro_imagem_pdf(tiny))
            self.assertIn("arquivo_derivado_ate_25kb", _motivo_filtro_imagem_pdf(small_file))
            self.assertGreater(narrow_but_useful.stat().st_size, 25_000)
            self.assertEqual(_motivo_filtro_imagem_pdf(narrow_but_useful), "")

            markdown = "antes ![cláusula](img/clausula.png) depois"
            self.assertEqual(
                _remover_referencia_imagem_markdown(markdown, tiny.name),
                "antes  depois",
            )

    def test_provider_switch_reuses_legacy_provider_checkpoint(self) -> None:
        case = {
            "case_id": "case", "secao": "reavaliacao", "auditado": "ORG", "codigo": "q1",
            "coluna_evidencia": "q1evi", "prompt_hash": "prompt", "prompt_version": "v1",
            "contexto": {}, "evidence_paths": [], "evidencias": [], "documentos_contexto": [],
        }
        routes = [
            {"provider": "openai", "model": "gpt-5.4-mini", "model_key": "gpt-5.4-mini"},
            {"provider": "openrouter", "model": "openai/gpt-5.4-mini", "model_key": "gpt-5.4-mini"},
        ]
        legacy = opinion("case", "openai", "gpt-5.4-mini")
        legacy.pop("model_key")
        legacy.update(
            {
                "identity": "identidade-antiga-com-provider",
                "contexto": {}, "evidence_paths": [], "prompt_hash": "prompt",
                "reasoning_effort": "high", "evidence_processing_mode": "",
            }
        )
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp)
            (out / "analyses_openai_gpt-5.4-mini.jsonl").write_text(
                json.dumps(legacy) + "\n", encoding="utf-8"
            )
            with patch("scripts.comentarios_gestor_pipeline.executar_caso") as execute:
                result = avaliar_casos(
                    [case], provider="openrouter", model="openai/gpt-5.4-mini",
                    model_key="gpt-5.4-mini", out_dir=out, routes=routes, quiet=True,
                )
            clean = [
                json.loads(line)
                for line in (out / "analyses_clean_model_gpt-5.4-mini.jsonl").read_text().splitlines()
            ]
        execute.assert_not_called()
        self.assertEqual(result["pulados"], 1)
        self.assertEqual(clean[0]["source_identity"], "identidade-antiga-com-provider")
        self.assertEqual(clean[0]["model_key"], "gpt-5.4-mini")

    def test_empty_limesurvey_upload_is_absence(self) -> None:
        self.assertEqual(parse_uploads("[]"), [])

    def test_section_one_selection_rule(self) -> None:
        self.assertFalse(_selecionar_situacao("Concorda, mas ainda não adotou nenhuma medida", "", "", []))
        self.assertTrue(_selecionar_situacao("Discorda da sinalização de inadequação", "Justificativa", "", []))
        self.assertFalse(_selecionar_situacao("Concorda e já está atendendo", "", "", []))
        self.assertTrue(_selecionar_situacao("Concorda e já atendeu às propostas de encaminhamento", "", "", []))
        self.assertFalse(
            _selecionar_situacao(
                "Concorda, mas ainda não adotou nenhuma medida",
                "Comentário adicional",
                "Justificativa adicional",
                [{"name": "evidencia.pdf"}],
            )
        )

    def test_justificativa_publicavel_secao1_rejeita_temporalidade(self) -> None:
        result = {
            "conclusoes": [{
                "item_codigo": "A2G4",
                "justificativa": (
                    "A manifestação foi acolhida. A documentação é suficiente. "
                    "A correção ocorreu posteriormente à data-base. Assim, a situação é considerada sanada."
                ),
                "conclusoes_motivos": [{"id_motivo": "MR016", "estado_motivo": "afastado"}],
            }]
        }
        self.assertTrue(validar_justificativa_publicavel("1", result))

    def test_justificativa_publicavel_secao1_aceita_padrao_coerente(self) -> None:
        result = {
            "conclusoes": [{
                "item_codigo": "A2G4",
                "justificativa": (
                    "A manifestação foi acolhida. A organização apresentou como evidência a Portaria 1. "
                    "A documentação é suficiente para demonstrar a formalização do comitê, pois comprova "
                    "sua instituição. Assim, a situação é considerada sanada."
                ),
                "conclusoes_motivos": [{"id_motivo": "MR016", "estado_motivo": "afastado"}],
            }]
        }
        self.assertEqual(validar_justificativa_publicavel("1", result), [])

    def test_justificativa_publicavel_rejeita_artefato_interno_como_evidencia(self) -> None:
        result = {
            "conclusoes": [{
                "item_codigo": "A5G20",
                "justificativa": (
                    "A manifestação não foi acolhida. A organização apresentou como evidência o(a) "
                    "manifestacao_gestor.json. A documentação é insuficiente para demonstrar a "
                    "existência do controle. Assim, a inconformidade permanece."
                ),
                "conclusoes_motivos": [{"id_motivo": "MR075", "estado_motivo": "mantido"}],
            }]
        }
        violacoes = validar_justificativa_publicavel("1", result)
        self.assertIn(
            "conclusão destinada ao auditado cita artefato interno",
            violacoes,
        )

    def test_justificativa_publicavel_rejeita_artefato_interno_em_metadados(self) -> None:
        result = {
            "conclusoes": [{
                "item_codigo": "A5G20",
                "justificativa": (
                    "A manifestação não foi acolhida. A organização não apresentou documentação "
                    "comprobatória. Assim, a inconformidade permanece."
                ),
                "arquivos_referenciados": ["contexto_situacao.json"],
                "conclusoes_motivos": [{"id_motivo": "MR075", "estado_motivo": "mantido"}],
            }]
        }
        self.assertIn(
            "conclusão destinada ao auditado cita artefato interno",
            validar_justificativa_publicavel("1", result),
        )

    def test_justificativa_publicavel_sem_anexo_exige_declaracao_expressa(self) -> None:
        result = {
            "conclusoes": [{
                "item_codigo": "A5G22",
                "justificativa": (
                    "A manifestação não foi acolhida. A organização apresentou como evidência a "
                    "documentação comprobatória. A documentação é insuficiente. "
                    "Assim, a inconformidade permanece."
                ),
                "arquivos_referenciados": [],
                "conclusoes_motivos": [{"id_motivo": "MR066", "estado_motivo": "mantido"}],
            }]
        }
        violacoes = validar_justificativa_publicavel(
            "1", result, evidencia_documental_disponivel=False
        )
        self.assertIn(
            "caso sem anexo ou link deve informar que a organização não apresentou documentação comprobatória",
            violacoes,
        )
        self.assertIn("caso sem anexo ou link não pode declarar evidência apresentada", violacoes)

    def test_justificativa_publicavel_sem_anexo_rejeita_arquivo_referenciado(self) -> None:
        result = {
            "conclusoes": [{
                "item_codigo": "A3G11",
                "justificativa": (
                    "A manifestação não foi acolhida. A organização não apresentou documentação "
                    "comprobatória. Assim, a inconformidade permanece."
                ),
                "arquivos_referenciados": ["Portaria 1/2026"],
                "conclusoes_motivos": [{"id_motivo": "MR034", "estado_motivo": "mantido"}],
            }]
        }
        self.assertIn(
            "conclusão referencia documentos sem anexo ou link comprobatório",
            validar_justificativa_publicavel(
                "1", result, evidencia_documental_disponivel=False
            ),
        )

    def test_justificativa_publicavel_rejeita_marcador_e_nome_tecnico(self) -> None:
        result = {
            "conclusoes": [{
                "item_codigo": "A5G20",
                "justificativa": (
                    "A manifestação não foi acolhida. A organização apresentou como evidência o(a) "
                    "documento 00036_02_imagens-gestao-de-ativos.pdf. A documentação é insuficiente. "
                    "Assim, a inconformidade permanece."
                ),
                "conclusoes_motivos": [{"id_motivo": "MR075", "estado_motivo": "mantido"}],
            }]
        }
        violacoes = validar_justificativa_publicavel(
            "1", result, evidencia_documental_disponivel=True
        )
        self.assertIn("justificativa contém marcador de gênero não adaptado", violacoes)
        self.assertIn("justificativa expõe nome técnico de armazenamento do anexo", violacoes)

    def test_justificativa_publicavel_parcial_exige_insuficiencia_integral(self) -> None:
        result = {
            "conclusoes": [{
                "item_codigo": "A2G6",
                "justificativa": (
                    "A manifestação foi parcialmente acolhida. A organização apresentou como evidência "
                    "o Decreto 1/2026. A documentação é suficiente para demonstrar a formalização. "
                    "Assim, a inconformidade permanece."
                ),
                "conclusoes_motivos": [
                    {"id_motivo": "MR008", "estado_motivo": "afastado"},
                    {"id_motivo": "MR010", "estado_motivo": "mantido"},
                ],
            }]
        }
        self.assertIn(
            "acolhimento parcial deve declarar que a documentação é insuficiente para demonstrar integralmente a prática",
            validar_justificativa_publicavel(
                "1", result, evidencia_documental_disponivel=True
            ),
        )

    def test_links_sao_extraidos_somente_da_manifestacao_do_gestor(self) -> None:
        documentos = [
            {"tipo": "manifestacao_gestor", "texto": "Consulte https://www.rj.gov.br/ato.pdf."},
            {"tipo": "comentario_gestor", "texto": "Também https://dados.gov.br/pagina?q=1"},
            {"tipo": "contexto_auditoria", "texto": "Ignore https://criterio.gov.br/referencia"},
        ]
        self.assertEqual(
            extrair_urls_manifestacao(documentos),
            ["https://www.rj.gov.br/ato.pdf", "https://dados.gov.br/pagina?q=1"],
        )

    def test_q2804_uses_specific_evidence_route(self) -> None:
        routes = carregar_rotas_catalogo(DEFAULT_CATALOG)
        self.assertEqual(routes["q2804"][0]["coluna_evidencia"], "q2804eviA")

    def test_temporal_schema_requires_temporal_fields(self) -> None:
        schema = json_schema_response_format(response_profile="manager_comments_temporal")
        properties = schema["json_schema"]["schema"]["properties"]["conclusoes"]["items"]["properties"]
        self.assertIn("estado_temporal", properties)
        result = opinion("case", "fake", "m")["result"]
        with self.assertRaises(ValueError):
            validar_resultado_ia(result, response_profile="manager_comments_temporal")

    def test_temporal_reason_accepts_extras_and_normalizes_safe_aliases(self) -> None:
        result = opinion("case", "fake", "m")["result"]
        conclusion = result["conclusoes"][0]
        conclusion.update({
            "estado_temporal": "mantida",
            "conclusoes_motivos": [{
                "motivo_id": "MR001",
                "estado": "mantido",
                "justificativa_motivo": "A manifestação não afasta o motivo.",
                "evidencias_consideradas": ["anexo.pdf"],
            }],
            "providencias_informadas": [],
            "comentarios_encaminhamento": "",
            "consequencias_praticas": [],
            "alternativas_propostas": [],
        })

        validated = validar_resultado_ia(result, response_profile="manager_comments_temporal")
        reason = validated["conclusoes"][0]["conclusoes_motivos"][0]
        self.assertEqual(reason["id_motivo"], "MR001")
        self.assertEqual(reason["estado_motivo"], "mantido")
        self.assertEqual(reason["justificativa"], "A manifestação não afasta o motivo.")
        self.assertEqual(reason["evidencias_consideradas"], ["anexo.pdf"])
        self.assertIn("avisos_estrutura_motivos", validated["conclusoes"][0])

    def test_temporal_reason_reports_the_specific_missing_fields(self) -> None:
        result = opinion("case", "fake", "m")["result"]
        conclusion = result["conclusoes"][0]
        conclusion.update({
            "estado_temporal": "mantida",
            "conclusoes_motivos": [{"id_motivo": "MR001"}],
            "providencias_informadas": [],
            "comentarios_encaminhamento": "",
            "consequencias_praticas": [],
            "alternativas_propostas": [],
        })

        with self.assertRaisesRegex(
            ValueError,
            r"motivo 0 sem campos: estado_motivo, justificativa",
        ):
            validar_resultado_ia(result, response_profile="manager_comments_temporal")

    def test_invalid_provider_response_preserves_raw_content(self) -> None:
        raw = '{"status":"completed","conclusoes":[{}]}'
        parsed = GenericProvider("modelo")._interpretar_resposta(
            raw,
            response_profile="manager_comments_temporal",
        )
        self.assertEqual(parsed["status"], "error")
        self.assertEqual(parsed["raw_response"], raw)
        self.assertIn("conclusao 0 sem campos", parsed["error"])

    def test_temporal_prompt_contains_reason_object_example(self) -> None:
        payload = json.loads(conteudo_provider_textual(
            prompt="Avalie.",
            auditado="ORG",
            questao_base="A1G1",
            coluna_evidencia="comentario",
            itens_afirmados=[],
            pacote={},
            response_profile="manager_comments_temporal",
        ))
        reason = payload["saida_obrigatoria"]["conclusoes"][0]["conclusoes_motivos"][0]
        self.assertEqual(set(reason), {"id_motivo", "estado_motivo", "justificativa"})

    def test_native_upload_filter_preserves_manager_context(self) -> None:
        package = {
            "arquivos_upload": ["/tmp/nova-evidencia.pdf"],
            "documentos": [
                {"nome": "nova-evidencia.pdf", "texto": "conteúdo extraído"},
                {"nome": "comentario-gestor", "texto": "esclarecimento do auditado"},
            ],
        }
        filtered = pacote_textual_sem_documentos_de_arquivos_nativos(package, {".pdf"})
        self.assertEqual([doc["nome"] for doc in filtered["documentos"]], ["comentario-gestor"])

    def test_section_one_maps_binary_adjustment_and_flags_derived_field(self) -> None:
        def record(case: str, code: str, item: str, source: str) -> dict:
            return {
                "identity": f"id-{case}", "case_id": case, "secao": "situacoes",
                "auditado": "AGERIO", "codigo": code, "status": "completed",
                "finished_at": "2026-07-14T10:00:00+00:00",
                "contexto": {
                    "achado": 2, "situacao": f"Situação {code}",
                    "motivos": [{
                        "id": f"MR-{code}",
                        "acoes": [{
                            "id": f"AV-{code}", "informacao_requerida": item,
                            "id_fonte_informacao": source,
                        }],
                    }],
                },
                "result": {"status": "completed", "error": "", "conclusoes": [{
                    "item_codigo": code, "item_texto": f"Situação {code}",
                    "afirmacao_auditado": "Discorda", "estado": "nao_conforme",
                    "estado_temporal": "corrigida_posteriormente", "justificativa": "Sanada.",
                    "conclusoes_motivos": [{
                        "id_motivo": f"MR-{code}", "estado_motivo": "afastado",
                        "justificativa": "Fundamento afastado.",
                    }],
                }]},
            }

        with tempfile.TemporaryDirectory() as tmp:
            consolidated = Path(tmp) / "consolidated_clean.jsonl"
            consolidated.write_text(
                "\n".join([
                    json.dumps(record("c1", "A2G4", "q1001ext[E]", "avaliacao_evidencias_ajustes")),
                    json.dumps(record("c2", "A4G12", "total_SI", "questionario")),
                ]) + "\n",
                encoding="utf-8",
            )
            ajustes, pendencias = mapear_ajustes_secao1(
                consolidado=consolidated,
                resultado_auditoria=DEFAULT_RESULTADO,
                mapa=DEFAULT_MAPA,
                lss=DEFAULT_LSS,
                ajustes_pos_avaliacao_evidencias=DEFAULT_AJUSTES,
                respostas_base=DEFAULT_RESPOSTAS_BASE,
            )
        self.assertEqual([(a["Código do item avaliado"], a["Resposta ajustada"]) for a in ajustes], [("q1001ext[E]", "Sim")])
        self.assertEqual(pendencias[0]["Código do item avaliado"], "total_SI")

    def test_section_two_filter_removes_only_sane_item_from_base_group(self) -> None:
        first = ItemAnterior("ORG", "q1001ext[E]", "q1001", "Sim", "Não conforme", "", "E")
        second = ItemAnterior("ORG", "q1001ext[F]", "q1001", "Sim", "Não conforme", "", "F")
        saneamento = {("ORG", "q1001ext[E]"): {"Auditado": "ORG", "Código do item avaliado": "q1001ext[E]", "Resposta ajustada": "Sim"}}
        restantes, excluidos = filtrar_itens_saneados_secao1({("ORG", "q1001"): [first, second]}, saneamento)
        self.assertEqual([item.codigo for item in restantes[("ORG", "q1001")]], ["q1001ext[F]"])
        self.assertEqual(excluidos[0]["Código do item avaliado"], "q1001ext[E]")

    def test_conflicting_adjustments_become_pending(self) -> None:
        base = {"Auditado": "ORG", "Código do item avaliado": "q1001ext[A]"}
        finais, pendencias = _agregar_propostas_ajuste(
            [{**base, "Resposta ajustada": "Sim"}, {**base, "Resposta ajustada": "Não"}], []
        )
        self.assertEqual(finais, [])
        self.assertIn("conflitantes", pendencias[0]["Motivo da pendência"])

    def test_individual_evaluations_workbook_has_human_readable_sheets(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            checkpoint = root / "analyses.jsonl"
            checkpoint.write_text(json.dumps(opinion("case", "provider", "model")) + "\n", encoding="utf-8")
            output = root / "avaliacoes_modelos.xlsx"
            gravar_avaliacoes_modelos_xlsx(
                output, [checkpoint], expected_models=["model"], expected_case_ids={"case"},
            )
            from openpyxl import load_workbook
            wb = load_workbook(output, read_only=True)
            sheets = wb.sheetnames
            wb.close()
        self.assertEqual(sheets, ["Avaliações", "Motivos", "Erros e ausências", "Escopo seção 2"])

    def test_revised_evidence_panel_blanks_sane_item_and_keeps_schema(self) -> None:
        from openpyxl import Workbook, load_workbook
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "painel.xlsx"
            output = root / "painel-final.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.append(["Auditado", "q1001ext[E]", "q1001ext[E]__justificativa"])
            ws.append(["ORG", "Não conforme", "Lacuna"])
            wb.save(source)
            count = gerar_painel_evidencias_pos_comentarios(
                painel_origem=source,
                output=output,
                ajustes=[{"Auditado": "ORG", "Código do item avaliado": "q1001ext[E]"}],
            )
            result = load_workbook(output, data_only=True)
            values = list(result.active.iter_rows(values_only=True))
            result.close()
        self.assertEqual(count, 1)
        self.assertEqual(values[0], ("Auditado", "q1001ext[E]", "q1001ext[E]__justificativa"))
        self.assertEqual(values[1], ("ORG", None, None))

    def test_reconsolidates_when_third_and_fourth_opinions_arrive(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            paths = []
            pairs = [("p1", "m1"), ("p2", "m2"), ("p3", "m3"), ("p4", "m4")]
            for provider, model in pairs:
                path = root / f"{provider}.jsonl"
                path.write_text(json.dumps(opinion("case", provider, model)) + "\n", encoding="utf-8")
                paths.append(path)
            first = consolidar_casos(
                analyses_files=paths[:2], out_dir=root / "out", secao="2",
                expected_models=[model for _, model in pairs],
                judge_provider="fake", judge_model="judge", min_opinions=2, quiet=True,
            )
            second = consolidar_casos(
                analyses_files=paths[:3], out_dir=root / "out", secao="2",
                expected_models=[model for _, model in pairs],
                judge_provider="fake", judge_model="judge", min_opinions=2, quiet=True,
            )
            third = consolidar_casos(
                analyses_files=paths, out_dir=root / "out", secao="2",
                expected_models=[model for _, model in pairs],
                judge_provider="fake", judge_model="judge", min_opinions=2, quiet=True,
            )
            raw = [json.loads(line) for line in (root / "out/consolidated.jsonl").read_text().splitlines()]
            clean = [json.loads(line) for line in (root / "out/consolidated_clean.jsonl").read_text().splitlines()]
            self.assertEqual(first["concluidos"], 1)
            self.assertEqual(second["concluidos"], 1)
            self.assertEqual(third["concluidos"], 1)
            self.assertEqual([r["opinioes_validas"] for r in raw], [2, 3, 4])
            self.assertEqual(len(clean), 1)
            self.assertEqual(clean[0]["opinioes_validas"], 4)
            self.assertTrue(clean[0]["supersedes_identity"])

    def test_consolidation_counts_one_opinion_per_model_key(self) -> None:
        first = opinion("case", "openai", "gpt-5.4-mini")
        first["model_key"] = "gpt-5.4-mini"
        first["identity"] = "logical-identity"
        second = opinion("case", "openrouter", "openai/gpt-5.4-mini")
        second["model_key"] = "gpt-5.4-mini"
        second["identity"] = "logical-identity"
        second["finished_at"] = "2026-07-13T11:00:00+00:00"
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            paths = [root / "openai.jsonl", root / "openrouter.jsonl"]
            paths[0].write_text(json.dumps(first) + "\n", encoding="utf-8")
            paths[1].write_text(json.dumps(second) + "\n", encoding="utf-8")
            result = consolidar_casos(
                analyses_files=paths,
                out_dir=root / "out",
                secao="2",
                expected_models=["gpt-5.4-mini"],
                judge_provider="fake",
                judge_model="judge",
                min_opinions=1,
                quiet=True,
            )
            clean = [json.loads(line) for line in (root / "out/consolidated_clean.jsonl").read_text().splitlines()]
        self.assertEqual(result["concluidos"], 1)
        self.assertEqual(clean[0]["opinioes_validas"], 1)
        self.assertEqual(clean[0]["opinioes"][0]["provider"], "openrouter")


if __name__ == "__main__":
    unittest.main()
