import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from scripts.gerar_relatorio_consolidado import (
    LARGURA_CAIXA_CABECALHO_EMU,
    LARGURA_CAIXA_CABECALHO_PT,
    MARCADOR_CORPO_INSTITUCIONAL,
    NUMERO_PROCESSO,
    _reassociar_notas_tabela_institucional,
    ajustar_tabela_lista_anexos,
    ajustar_tabela_modelo_plano_acao,
    aplicar_paginas_institucionais,
)
from scripts.resources.argos_utils import cross_ref_tabelas, inserir_campo_sumario_docx


ROOT = Path(__file__).resolve().parents[2]
MODELO = ROOT / "scripts/resources/template-relatorio-consolidado-institucional.docx"


def _texto_xml(elemento) -> str:
    return "".join(
        no.text or ""
        for no in elemento.iter()
        if no.tag == qn("w:t")
    )


class PaginasInstitucionaisConsolidadoTest(unittest.TestCase):
    def test_titulo_sumario_repete_visual_do_titulo_1_sem_ser_titulo_1(self):
        conteudo = inserir_campo_sumario_docx("::: {.toc}\n:::")

        self.assertIn('<w:pStyle w:val="TOCHeading"/>', conteudo)
        self.assertNotIn('<w:pStyle w:val="Heading1"/>', conteudo)
        self.assertIn(
            '<w:rFonts w:ascii="Cambria" w:hAnsi="Cambria" '
            'w:eastAsia="Cambria" w:cs="Cambria"/>',
            conteudo,
        )
        self.assertIn('<w:sz w:val="32"/>', conteudo)
        self.assertIn('<w:szCs w:val="32"/>', conteudo)
        self.assertIn('<w:b/>', conteudo)
        self.assertIn('<w:color w:val="2F5496"/>', conteudo)

    def _criar_conteudo(self, caminho: Path) -> None:
        documento = Document()
        documento.add_heading("RELATÓRIO DE AUDITORIA GOVERNAMENTAL", 0)
        documento.add_heading("DADOS DA FISCALIZAÇÃO", 1)
        tabela = documento.add_table(rows=11, cols=2)
        for indice, linha in enumerate(tabela.rows):
            linha.cells[0].text = f"Campo {indice + 1}:"
            linha.cells[1].text = f"Valor atualizado {indice + 1}"
        documento.add_heading("SUMÁRIO", 1)
        campo_sumario = documento.add_paragraph()
        for tipo, texto in (
            ("begin", None),
            (None, 'TOC \\o "1-3" \\h \\z \\u'),
            ("separate", None),
            (None, "O sumário será atualizado ao abrir o documento no Word."),
            ("end", None),
        ):
            run = OxmlElement("w:r")
            if tipo:
                campo = OxmlElement("w:fldChar")
                campo.set(qn("w:fldCharType"), tipo)
                run.append(campo)
            else:
                no_texto = OxmlElement(
                    "w:instrText" if texto.startswith("TOC ") else "w:t"
                )
                no_texto.text = texto
                run.append(no_texto)
            campo_sumario._element.append(run)
        documento.add_heading("LISTA DE ANEXOS", 1)
        documento.add_paragraph("CONTEÚDO ATUAL DA LISTA")
        documento.add_heading("1. RESUMO", 1)
        documento.add_paragraph("CONTEÚDO ATUAL DO RELATÓRIO")
        documento.add_heading("7. PROPOSTA DE ENCAMINHAMENTO", 1)
        documento.add_paragraph("ARQUIVAMENTO do presente processo.")
        documento.save(caminho)

    def test_compoe_frontispicio_corpo_e_encerramento_sem_duplicacoes(self):
        with tempfile.TemporaryDirectory() as diretorio:
            caminho = Path(diretorio) / "relatorio.docx"
            self._criar_conteudo(caminho)

            aplicar_paginas_institucionais(
                caminho,
                MODELO,
                data_relatorio="11/09/2026",
            )

            documento = Document(caminho)
            textos = [paragrafo.text.strip() for paragrafo in documento.paragraphs]
            texto_total = "\n".join(textos)

            self.assertEqual(
                documento.tables[0].cell(0, 0).paragraphs[0].text,
                f"Processo: {NUMERO_PROCESSO}",
            )
            self.assertNotIn("Origem: SIGILOSO", documento.tables[0].cell(0, 0).text)
            self.assertEqual(documento.tables[1].cell(0, 1).text, "Valor atualizado 1")
            for paragrafo in documento.tables[0].cell(0, 0).paragraphs[:3]:
                for run in paragrafo.runs:
                    self.assertEqual(run.font.name, "Arial")
                    self.assertEqual(run.font.size.pt, 12)
            for linha in documento.tables[1].rows:
                for celula in linha.cells:
                    for paragrafo in celula.paragraphs:
                        self.assertEqual(paragrafo.paragraph_format.line_spacing, 1.0)
                        for run in paragrafo.runs:
                            self.assertEqual(run.font.name, "Arial")
                            self.assertEqual(run.font.size.pt, 10)
            self.assertEqual(textos.count("LISTA DE ANEXOS"), 1)
            self.assertEqual(textos.count("7. PROPOSTA DE ENCAMINHAMENTO"), 1)
            self.assertEqual(
                sum(texto.startswith("O presente relatório foi objeto de supervisão") for texto in textos),
                1,
            )
            self.assertEqual(textos.count("CAD-TI, 11/09/2026"), 2)
            self.assertNotIn(MARCADOR_CORPO_INSTITUCIONAL, texto_total)
            self.assertIn("CONTEÚDO ATUAL DO RELATÓRIO", texto_total)

            tabela_equipe = documento.tables[-1]
            self.assertEqual(len(tabela_equipe.rows), 3)
            self.assertIn("AUGUSTO CÉSAR BENVENUTO DE ALMEIDA", tabela_equipe.cell(0, 0).text)
            self.assertIn("JOÃO PAULO DE FREITAS RAMIREZ", tabela_equipe.cell(1, 0).text)
            self.assertIn("BRUNO MATTOS SOUZA DE SOUZA MELO", tabela_equipe.cell(2, 0).text)
            self.assertIn("ALBERTO DE FONTES TAVARES NETO", texto_total)

    def test_preserva_cabecalho_sumario_e_campos_sem_revisoes(self):
        with tempfile.TemporaryDirectory() as diretorio:
            caminho = Path(diretorio) / "relatorio.docx"
            self._criar_conteudo(caminho)

            aplicar_paginas_institucionais(
                caminho,
                MODELO,
                data_relatorio="11/09/2026",
            )

            documento = Document(caminho)
            cabecalho = _texto_xml(documento.sections[0].header._element)
            self.assertIn("TCE-RJ", cabecalho)
            self.assertIn(f"Processo nº {NUMERO_PROCESSO}", cabecalho)
            self.assertIn("Rubrica", cabecalho)

            with zipfile.ZipFile(caminho) as arquivo:
                documento_xml = arquivo.read("word/document.xml")
                cabecalhos_xml = b"".join(
                    arquivo.read(nome)
                    for nome in arquivo.namelist()
                    if nome.startswith("word/header") and nome.endswith(".xml")
                )
                cabecalhos_brutos = [
                    arquivo.read(nome)
                    for nome in arquivo.namelist()
                    if nome.startswith("word/header") and nome.endswith(".xml")
                ]
                settings_xml = arquivo.read("word/settings.xml")
                xml_completo = b"".join(
                    arquivo.read(nome)
                    for nome in arquivo.namelist()
                    if nome.startswith("word/") and nome.endswith(".xml")
                )

            self.assertIn(b'TOC \\o "1-3"', documento_xml)
            for opcao in (b"\\h", b"\\z", b"\\u"):
                self.assertIn(opcao, documento_xml)
            self.assertIn(b" PAGE ", cabecalhos_xml)
            raiz_cabecalhos = [etree.fromstring(xml) for xml in cabecalhos_brutos]
            caixas_processo = []
            for raiz in raiz_cabecalhos:
                for caixa in raiz.xpath("//*[local-name()='txbxContent']"):
                    if f"Processo nº {NUMERO_PROCESSO}" in _texto_xml(caixa):
                        caixas_processo.append(caixa)
            self.assertTrue(caixas_processo)
            for caixa in caixas_processo:
                tamanhos = caixa.xpath(".//*[local-name()='sz' or local-name()='szCs']/@w:val", namespaces={"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"})
                self.assertTrue(tamanhos)
                self.assertEqual(set(tamanhos), {"16"})
            larguras_drawing = {
                ancora.xpath("./*[local-name()='extent']/@cx")[0]
                for raiz in raiz_cabecalhos
                for ancora in raiz.xpath("//*[local-name()='anchor']")
                if f"Processo nº {NUMERO_PROCESSO}" in _texto_xml(ancora)
            }
            self.assertEqual(larguras_drawing, {str(LARGURA_CAIXA_CABECALHO_EMU)})
            estilos_vml = [
                estilo
                for raiz in raiz_cabecalhos
                for forma in raiz.xpath("//*[local-name()='shape']")
                if f"Processo nº {NUMERO_PROCESSO}" in _texto_xml(forma)
                for estilo in forma.xpath("./@style")
            ]
            self.assertTrue(any(f"width:{LARGURA_CAIXA_CABECALHO_PT}" in estilo for estilo in estilos_vml))
            self.assertIn(b'<w:updateFields w:val="true"', settings_xml)
            self.assertNotIn(b"<w:trackRevisions", xml_completo)
            self.assertNotIn(b"<w:ins ", xml_completo)
            self.assertNotIn(b"<w:del ", xml_completo)

    def test_reassocia_notas_da_tabela_sem_colidir_com_notas_legadas(self):
        with tempfile.TemporaryDirectory() as diretorio:
            caminho = Path(diretorio) / "relatorio.docx"
            documento = Document(MODELO)
            paragrafo = documento.tables[1].cell(4, 1).paragraphs[0]
            notas_origem = {}
            for id_nota, texto in (("2", "RELAÇÃO ESTADUAL"), ("3", "RELAÇÃO MUNICIPAL")):
                run = OxmlElement("w:r")
                referencia = OxmlElement("w:footnoteReference")
                referencia.set(qn("w:id"), id_nota)
                run.append(referencia)
                paragrafo._element.append(run)

                nota = OxmlElement("w:footnote")
                nota.set(qn("w:id"), id_nota)
                paragrafo_nota = OxmlElement("w:p")
                run_nota = OxmlElement("w:r")
                texto_nota = OxmlElement("w:t")
                texto_nota.text = texto
                run_nota.append(texto_nota)
                paragrafo_nota.append(run_nota)
                nota.append(paragrafo_nota)
                notas_origem[id_nota] = nota
            documento.save(caminho)

            _reassociar_notas_tabela_institucional(caminho, notas_origem)

            with zipfile.ZipFile(caminho) as arquivo:
                documento_xml = etree.fromstring(arquivo.read("word/document.xml"))
                notas_xml = etree.fromstring(arquivo.read("word/footnotes.xml"))
            tabelas = documento_xml.findall(f".//{qn('w:tbl')}")
            ids_tabela = [
                no.get(qn("w:id"))
                for no in tabelas[1].findall(f".//{qn('w:footnoteReference')}")
            ]
            self.assertEqual(ids_tabela, ["2", "3"])
            notas_usadas = [
                _texto_xml(nota)
                for nota in notas_xml.findall(qn("w:footnote"))
                if nota.get(qn("w:type")) is None
            ]
            self.assertEqual(notas_usadas, ["RELAÇÃO ESTADUAL", "RELAÇÃO MUNICIPAL"])

    def test_preserva_recuo_ao_numerar_legenda_de_tabela(self):
        markdown = (
            "    : Modelo referencial {#tbl:modelo#}\n\n"
            "    | Coluna |\n    |---|\n    | Valor |\n"
        )
        processado = cross_ref_tabelas(markdown)
        self.assertIn("    : Tabela 1 - Modelo referencial", processado)
        self.assertNotIn("{#tbl:modelo#}", processado)

    def test_alinha_tabela_referencial_ao_texto_da_lista(self):
        with tempfile.TemporaryDirectory() as diretorio:
            caminho = Path(diretorio) / "relatorio.docx"
            documento = Document()
            documento.add_paragraph("Tabela 16 - Modelo referencial de plano de ação")
            tabela = documento.add_table(rows=2, cols=2)
            tabela.cell(0, 0).text = "Determinação ou recomendação"
            tabela.cell(0, 1).text = "Medida a adotar"
            tabela.cell(1, 0).text = "Item"
            tabela.cell(1, 1).text = "Medida"
            documento.add_paragraph("(Fonte: elaboração própria)")
            documento.save(caminho)

            ajustar_tabela_modelo_plano_acao(caminho)

            ajustado = Document(caminho)
            tabela = ajustado.tables[0]
            tbl_pr = tabela._tbl.tblPr
            self.assertEqual(tbl_pr.find(qn("w:tblInd")).get(qn("w:w")), "720")
            self.assertEqual(tbl_pr.find(qn("w:tblW")).get(qn("w:w")), "8352")
            self.assertEqual(
                sum(
                    int(coluna.get(qn("w:w")))
                    for coluna in tabela._tbl.tblGrid.findall(qn("w:gridCol"))
                ),
                8352,
            )
            self.assertEqual(ajustado.paragraphs[0].paragraph_format.left_indent.twips, 720)
            self.assertEqual(ajustado.paragraphs[1].paragraph_format.left_indent.twips, 720)

    def test_distribui_colunas_da_lista_de_anexos_pelo_conteudo(self):
        with tempfile.TemporaryDirectory() as diretorio:
            caminho = Path(diretorio) / "relatorio.docx"
            documento = Document()
            tabela = documento.add_table(rows=2, cols=2)
            tabela.cell(0, 0).text = "Documento nº"
            tabela.cell(0, 1).text = "Descrição"
            tabela.cell(1, 0).text = "AN01"
            tabela.cell(1, 1).text = "Descrição extensa do documento anexado"
            documento.save(caminho)

            ajustar_tabela_lista_anexos(caminho)

            ajustado = Document(caminho)
            tabela = ajustado.tables[0]
            larguras = [
                int(coluna.get(qn("w:w")))
                for coluna in tabela._tbl.tblGrid.findall(qn("w:gridCol"))
            ]
            self.assertEqual(larguras, [1600, 7472])
            self.assertGreater(larguras[1], larguras[0] * 4)
            self.assertEqual(tabela._tbl.tblPr.find(qn("w:tblW")).get(qn("w:w")), "9072")


if __name__ == "__main__":
    unittest.main()
