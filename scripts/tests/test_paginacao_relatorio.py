"""Testes automatizados de otimização editorial e de paginação dos relatórios individuais."""

import os
import re
import sys
import unittest
from pathlib import Path
import docx

REPO_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO_ROOT / "scripts"))
sys.path.insert(0, str(REPO_ROOT / "scripts" / "resources"))

from argos_utils import (
    otimizar_paginacao_relatorio,
    is_tabela_plano_acao,
    configurar_tabela_plano_acao,
)


class TestPaginacaoRelatorio(unittest.TestCase):
    """Bateria de testes cobrindo os 10 requisitos da Seção 35."""

    def test_01_nenhuma_quebra_manual_antes_achados(self):
        """Item 1: Nenhuma quebra manual \newpage imediatamente antes dos seis achados."""
        achado_files = [
            "achado_questao_1_estrutura_tic.md",
            "achado_questao_2_governanca_comite_tic.md",
            "achado_questao_3_planejamento_tic.md",
            "achado_questao_4_capacidade_institucional_tic_si.md",
            "achado_questao_5_gestao_servicos_tic.md",
            "achado_questao_6_contratacoes_tic.md",
        ]
        templates_dir = REPO_ROOT / "03-Relatorios" / "03-Relatorios_Individuais_Finais"
        for filename in achado_files:
            path = templates_dir / filename
            self.assertTrue(path.exists(), f"Arquivo {filename} não encontrado.")
            content = path.read_text(encoding="utf-8")
            match = re.search(r"\\newpage\s*\n+\s*##\s+Achado\s+\{\{\s*achado\.numero\s*\}\}", content)
            self.assertIsNone(
                match,
                f"Quebra \\newpage indevida encontrada antes do título em {filename}",
            )

    def test_02_headings_com_keep_with_next(self):
        """Item 2: Headings relevantes com keep_with_next ativado."""
        sample_docx = "/tmp/igovti-pipeline-test/Relatório Individual - SEEL.docx"
        if not os.path.exists(sample_docx):
            self.skipTest("DOCX de teste não disponível para validação de headings")

        doc = docx.Document(sample_docx)
        for p in doc.paragraphs:
            style = p.style.name or ""
            is_heading = style.startswith(("Heading", "heading", "Título", "Title", "Subtitle")) or bool(
                p._element.xpath("./w:pPr/w:outlineLvl")
            )
            if is_heading and p.text.strip():
                self.assertTrue(
                    p.paragraph_format.keep_with_next,
                    f"Heading '{p.text[:40]}' sem keep_with_next ativado",
                )

    def test_03_captions_figura_vinculadas(self):
        """Item 3: Captions de figura vinculadas às imagens (keep_with_next)."""
        sample_docx = "/tmp/igovti-pipeline-test/Relatório Individual - SEEL.docx"
        if not os.path.exists(sample_docx):
            self.skipTest("DOCX de teste não disponível")

        doc = docx.Document(sample_docx)
        for i, p in enumerate(doc.paragraphs):
            has_drawing = bool(p._element.xpath(".//w:drawing") or p._element.xpath(".//w:graphic"))
            if has_drawing and i > 0:
                prev_p = doc.paragraphs[i - 1]
                if prev_p.text.strip().startswith("Figura ") or "Caption" in (prev_p.style.name or ""):
                    self.assertTrue(
                        prev_p.paragraph_format.keep_with_next,
                        f"Legenda '{prev_p.text[:40]}' sem keep_with_next antes da imagem",
                    )

    def test_04_imagem_vinculada_a_fonte(self):
        """Item 4: Imagem vinculada à fonte quando aplicável e fonte com keep_with_next=False."""
        sample_docx = "/tmp/igovti-pipeline-test/Relatório Individual - SEEL.docx"
        if not os.path.exists(sample_docx):
            self.skipTest("DOCX de teste não disponível")

        doc = docx.Document(sample_docx)
        for i, p in enumerate(doc.paragraphs):
            has_drawing = bool(p._element.xpath(".//w:drawing") or p._element.xpath(".//w:graphic"))
            if has_drawing and i < len(doc.paragraphs) - 1:
                next_p = doc.paragraphs[i + 1]
                if "FonteImagem" in (next_p.style.name or "") or next_p.text.strip().startswith("(Fonte:"):
                    self.assertTrue(p.paragraph_format.keep_with_next, "Imagem não vinculada à fonte posterior")
                    self.assertFalse(next_p.paragraph_format.keep_with_next, "Fonte não deve ter keep_with_next=True")

    def test_05_captions_tabela_vinculadas(self):
        """Item 5: Captions de tabela vinculadas à tabela (keepNext no parágrafo anterior)."""
        sample_docx = "/tmp/igovti-pipeline-test/Relatório Individual - SEEL.docx"
        if not os.path.exists(sample_docx):
            self.skipTest("DOCX de teste não disponível")

        doc = docx.Document(sample_docx)
        for tbl in doc.tables:
            prev_el = tbl._element.getprevious()
            if prev_el is not None and prev_el.tag.endswith("p"):
                has_keep = bool(prev_el.xpath(".//w:keepNext"))
                self.assertTrue(has_keep, "Parágrafo imediatamente anterior à tabela sem keepNext")

    def test_06_repeticao_cabecalho_tabelas(self):
        """Item 6: Primeira linha de tabelas com repetição de cabeçalho (<w:tblHeader/>)."""
        sample_docx = "/tmp/igovti-pipeline-test/Relatório Individual - SEEL.docx"
        if not os.path.exists(sample_docx):
            self.skipTest("DOCX de teste não disponível")

        doc = docx.Document(sample_docx)
        for tbl in doc.tables:
            if len(tbl.rows) > 1:
                has_header = bool(tbl.rows[0]._element.xpath(".//w:tblHeader"))
                self.assertTrue(has_header, "Primeira linha da tabela sem tblHeader")

    def test_07_plano_acao_layout_fixo_e_larguras(self):
        """Item 7: Plano de Ação com layout fixo e larguras esperadas."""
        sample_docx = "/tmp/igovti-pipeline-test/Relatório Individual - SEEL.docx"
        if not os.path.exists(sample_docx):
            self.skipTest("DOCX de teste não disponível")

        doc = docx.Document(sample_docx)
        plano_tbl = None
        for tbl in doc.tables:
            if is_tabela_plano_acao(tbl):
                plano_tbl = tbl
                break

        self.assertIsNotNone(plano_tbl, "Tabela de Plano de Ação não encontrada")
        self.assertFalse(plano_tbl.autofit, "Tabela de Plano de Ação não deve ter autofit ativado")
        self.assertTrue(bool(plano_tbl._element.xpath('w:tblPr/w:tblLayout[@w:type="fixed"]')))

        cell_medida = plano_tbl.rows[0].cells[2]
        tcW = cell_medida._element.xpath(".//w:tcW")
        self.assertTrue(bool(tcW))
        w_val = int(tcW[0].get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w", 0))
        self.assertGreaterEqual(w_val, 3000, f"Largura da coluna Medida proposta ({w_val} dxa) insuficiente")

    def test_08_ausencia_cantsplit_indiscriminado(self):
        """Item 8: Ausência de cantSplit indiscriminado em linhas longas."""
        sample_docx = "/tmp/igovti-pipeline-test/Relatório Individual - SEEL.docx"
        if not os.path.exists(sample_docx):
            self.skipTest("DOCX de teste não disponível")

        doc = docx.Document(sample_docx)
        for tbl in doc.tables:
            for r_idx, row in enumerate(tbl.rows):
                if r_idx == 0:
                    continue
                total_chars = sum(len(c.text.strip()) for c in row.cells)
                total_paras = sum(len(c.paragraphs) for c in row.cells)
                is_long = total_chars > 250 or total_paras > 3
                if is_long:
                    has_cant_split = bool(row._element.xpath("w:trPr/w:cantSplit"))
                    self.assertFalse(
                        has_cant_split,
                        f"Linha longa ({total_chars} caracteres) com cantSplit indevido",
                    )

    def test_09_idempotencia_pos_processamento(self):
        """Item 9: Idempotência do pós-processamento."""
        sample_docx = "/tmp/igovti-pipeline-test/Relatório Individual - TJRJ.docx"
        if not os.path.exists(sample_docx):
            self.skipTest("DOCX de teste não disponível")

        import shutil
        import tempfile

        with tempfile.TemporaryDirectory() as tmpdir:
            test_file = os.path.join(tmpdir, "test_idempotencia.docx")
            shutil.copyfile(sample_docx, test_file)

            otimizar_paginacao_relatorio(test_file)
            doc1 = docx.Document(test_file)
            n_para1 = len(doc1.paragraphs)
            n_tbl1 = len(doc1.tables)

            otimizar_paginacao_relatorio(test_file)
            doc2 = docx.Document(test_file)
            n_para2 = len(doc2.paragraphs)
            n_tbl2 = len(doc2.tables)

            self.assertEqual(n_para1, n_para2, "Segunda execução alterou quantidade de parágrafos")
            self.assertEqual(n_tbl1, n_tbl2, "Segunda execução alterou quantidade de tabelas")

    def test_10_preservacao_quantidade_elementos(self):
        """Item 10: Preservação de quantidade de tabelas e imagens antes e depois."""
        for name in ["SEEL", "TJRJ", "MPERJ", "DUQUE_DE_CAXIAS"]:
            base = f"/tmp/igovti-pagination-baseline/{name}.docx"
            opt_name = name.replace("_", " ")
            opt = f"/tmp/igovti-pipeline-test/Relatório Individual - {opt_name}.docx"
            if not os.path.exists(base) or not os.path.exists(opt):
                continue

            doc_b = docx.Document(base)
            doc_o = docx.Document(opt)

            n_tbl_b = len(doc_b.tables)
            n_tbl_o = len(doc_o.tables)
            self.assertEqual(n_tbl_b, n_tbl_o, f"Diferença no total de tabelas em {name}: {n_tbl_b} vs {n_tbl_o}")

            n_img_b = len(doc_b._element.xpath(".//w:drawing") + doc_b._element.xpath(".//w:graphic"))
            n_img_o = len(doc_o._element.xpath(".//w:drawing") + doc_o._element.xpath(".//w:graphic"))
            self.assertEqual(n_img_b, n_img_o, f"Diferença no total de imagens em {name}: {n_img_b} vs {n_img_o}")


if __name__ == "__main__":
    unittest.main()
