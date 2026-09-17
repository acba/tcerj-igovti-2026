#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Gera o arquivo DOCX do Relatório Consolidado a partir do Markdown.

Este script processa o Markdown aplicando renderizações de template Jinja2,
referências cruzadas, quebras de página, substituição de underlines do Pandoc
e formatação de tabelas, utilizando o template de estilos do Argos.
"""

from __future__ import annotations

import argparse
from copy import deepcopy
import json
import logging
import os
import re
import sys
import tempfile
import zipfile
from pathlib import Path

from lxml import etree

# Configura path para importar utilitários da pasta resources
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "resources"))
from argos_utils import (
    cross_ref_figuras,
    processar_quebras_pagina,
    cross_ref_tabelas,
    substituir_underline_pandoc,
    aplicar_estilo_tabelas,
    evitar_quebra_elementos,
    inserir_campo_sumario_docx,
    marcar_atualizacao_campos_docx,
    data_hoje_abnt,
    data_hoje,
)

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")
logger = logging.getLogger(__name__)
TOC_MARKER_FILTER = Path(__file__).resolve().parent / "resources" / "toc-marker.lua"
DEFAULT_CONTEXT_JSON = (
    Path(__file__).resolve().parent.parent
    / "03-Relatorios/01-Relatorio_Consolidado/dados/diagnostico-transversal-igovti-2026.json"
)
DEFAULT_MODELO_INSTITUCIONAL = (
    Path(__file__).resolve().parent
    / "resources/template-relatorio-consolidado-institucional.docx"
)
MARCADOR_CORPO_INSTITUCIONAL = "[[CORPO_RELATORIO_CONSOLIDADO]]"
NUMERO_PROCESSO = "101.088-0/2026"
LARGURA_CAIXA_CABECALHO_EMU = 1_408_176  # 1,54 polegada
LARGURA_CAIXA_CABECALHO_PT = "110.88pt"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = f"{{{W_NS}}}"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
V_NS = "urn:schemas-microsoft-com:vml"

SITUACOES_POR_ACHADO = {
    1: ("S1.1", "S1.2", "S1.3"),
    2: ("S2.1", "S2.2", "S2.3"),
    3: ("S3.1", "S3.2", "S3.4", "S3.5", "S3.6"),
    4: ("S4.1", "S4.2", "S4.3", "S4.6"),
    5: ("S5.1", "S5.2", "S5.3", "S5.4", "S5.5"),
    6: ("S6.1", "S6.2", "S6.3", "S6.4"),
}


def _resumir_dispositivo(descricao: str) -> str:
    """Mantém a identificação do ato e do dispositivo, sem reproduzir sua explicação."""
    return re.split(r"\s+[—–]\s+", str(descricao or "").strip(), maxsplit=1)[0].rstrip(".;")


def _destinatarios_variantes(
    resultado_path: str | None,
    auditados_path: str | None,
    resolvedor,
) -> dict[str, set[str]]:
    if (
        not resultado_path
        or not Path(resultado_path).is_file()
        or not auditados_path
        or not Path(auditados_path).is_file()
    ):
        return {}
    import pandas as pd
    from aplicabilidade_juridica import PerfilAuditado

    dados = json.loads(Path(resultado_path).read_text(encoding="utf-8"))
    cadastro = pd.read_excel(auditados_path, sheet_name="auditados", keep_default_na=False)
    perfis = {
        perfil.sigla: perfil
        for perfil in (
            PerfilAuditado.from_mapping(row.to_dict())
            for _, row in cadastro.iterrows()
        )
        if perfil.sigla
    }
    destinatarios: dict[str, set[str]] = {}
    for auditado in dados.values():
        if not isinstance(auditado, dict):
            continue
        sigla = str(auditado.get("sigla") or "").strip().upper()
        perfil = perfis.get(sigla)
        if perfil is None:
            raise ValueError(f"{sigla}: auditado do resultado final ausente do cadastro institucional.")
        for procedimento in auditado.get("procedimentos_executados", []):
            for situacao in (procedimento.get("achado") or {}).get("situacoes_detalhadas", []):
                id_situacao = str(situacao.get("id_situacao") or "").strip()
                if not id_situacao:
                    raise ValueError(f"{sigla}: situação final sem identificador estável.")
                variante = resolvedor.resolver(perfil, id_situacao).variante
                if not variante.geral:
                    destinatarios.setdefault(variante.id, set()).add(sigla)
    return destinatarios


def _resumir_destinatarios(auditados: set[str]) -> str:
    ordenados = sorted(auditados)
    if len(ordenados) <= 5:
        return ", ".join(ordenados)
    return f"{len(ordenados)} organizações desse grupo institucional"


def _nota_criterios_especificos_achado(
    mapa_path: str | None,
    resultado_path: str | None,
    auditados_path: str | None,
    ids_situacoes: tuple[str, ...],
) -> str:
    if not mapa_path or not Path(mapa_path).is_file():
        return ""

    from aplicabilidade_juridica import carregar_catalogo_planilha

    resolvedor = carregar_catalogo_planilha(mapa_path)
    destinatarios = _destinatarios_variantes(
        resultado_path,
        auditados_path,
        resolvedor,
    )
    grupos: dict[str, dict] = {}

    for id_situacao in ids_situacoes:
        for variante in resolvedor.variantes_por_situacao.get(id_situacao, []):
            if variante.geral:
                continue
            auditados_variante = destinatarios.get(variante.id, set())
            if not auditados_variante:
                continue
            grupo = grupos.setdefault(
                variante.rotulo_publico,
                {"criterios": {}, "aplicacoes": []},
            )
            for id_criterio in variante.criterios:
                criterio = resolvedor.criterios[id_criterio]
                if criterio.seletor.vazio:
                    continue
                registro = grupo["criterios"].setdefault(
                    id_criterio,
                    {"criterio": criterio, "situacoes": []},
                )
                if id_situacao not in registro["situacoes"]:
                    registro["situacoes"].append(id_situacao)
            grupo["aplicacoes"].append(
                (id_situacao, variante.tipo_encaminhamento, auditados_variante)
            )

    partes = []
    for publico, grupo in grupos.items():
        criterios = []
        for id_criterio, registro in grupo["criterios"].items():
            criterio = registro["criterio"]
            id_exibicao = id_criterio.split(".", 1)[-1]
            situacoes = ", ".join(registro["situacoes"])
            criterios.append(
                f"{id_exibicao} — {_resumir_dispositivo(criterio.descricao)} [{situacoes}]"
            )
        texto = f"**{publico}:** " + "; ".join(criterios) + "."
        if grupo["aplicacoes"]:
            aplicacoes = "; ".join(
                f"{id_situacao} — {tipo.lower()} para {_resumir_destinatarios(auditados)}"
                for id_situacao, tipo, auditados in grupo["aplicacoes"]
            )
            texto += f" Aplicação no resultado final: {aplicacoes}."
        partes.append(texto)
    return " ".join(partes)


def _texto_elemento_docx(elemento) -> str:
    return "".join(
        no.text or ""
        for no in elemento.iter()
        if no.tag.endswith("}t")
    )


def _localizar_ultimo_elemento_docx(documento, texto: str) -> int:
    correspondencias = [
        indice
        for indice, elemento in enumerate(documento.element.body)
        if _texto_elemento_docx(elemento).strip() == texto
    ]
    if not correspondencias:
        raise ValueError(f"Elemento {texto!r} não encontrado no DOCX.")
    return correspondencias[-1]


def _remover_corpo_antes_de(documento, texto: str) -> None:
    indice = _localizar_ultimo_elemento_docx(documento, texto)
    body = documento.element.body
    for elemento in list(body)[:indice]:
        body.remove(elemento)


def _remover_corpo_a_partir_de(documento, texto: str) -> None:
    from docx.oxml.ns import qn

    indice = _localizar_ultimo_elemento_docx(documento, texto)
    body = documento.element.body
    for elemento in list(body)[indice:]:
        if elemento.tag != qn("w:sectPr"):
            body.remove(elemento)


def _localizar_tabela_apos_texto(documento, texto: str):
    """Localiza a primeira tabela posterior a um parágrafo-âncora."""
    from docx.table import Table

    indice = _localizar_ultimo_elemento_docx(documento, texto)
    body = documento.element.body
    for elemento in list(body)[indice + 1:]:
        if elemento.tag == f"{W}tbl":
            return Table(elemento, documento)
        if _texto_elemento_docx(elemento).strip():
            break
    raise ValueError(f"Tabela posterior a {texto!r} não encontrada no DOCX.")


def _substituir_texto_paragrafo(paragrafo, texto: str, *, negrito_rotulo: str | None = None) -> None:
    for run in paragrafo.runs:
        run._element.getparent().remove(run._element)
    if negrito_rotulo and texto.startswith(negrito_rotulo):
        run_rotulo = paragrafo.add_run(negrito_rotulo)
        run_rotulo.bold = True
        paragrafo.add_run(texto[len(negrito_rotulo):])
    else:
        paragrafo.add_run(texto)


def _preencher_celula_com_paragrafos(celula, paragrafos_origem) -> None:
    for paragrafo in list(celula.paragraphs):
        paragrafo._element.getparent().remove(paragrafo._element)
    for paragrafo in paragrafos_origem:
        celula._tc.append(deepcopy(paragrafo._element))


def _formatar_paragrafo_docx(paragrafo, *, fonte: str, tamanho_pt: float) -> None:
    """Aplica fonte diretamente aos runs, inclusive para scripts alternativos."""
    from docx.oxml.ns import qn
    from docx.shared import Pt

    for run in paragrafo.runs:
        run.font.name = fonte
        run.font.size = Pt(tamanho_pt)
        fontes = run._element.get_or_add_rPr().get_or_add_rFonts()
        for atributo in ("ascii", "hAnsi", "eastAsia", "cs"):
            fontes.set(qn(f"w:{atributo}"), fonte)


def _ajustar_caixa_cabecalho(no_texto) -> None:
    """Define a caixa TCE-RJ em 1,54 polegada e todo o seu texto em 8 pt."""
    conteudo_caixa = no_texto.getparent()
    while conteudo_caixa is not None and conteudo_caixa.tag != f"{W}txbxContent":
        conteudo_caixa = conteudo_caixa.getparent()
    if conteudo_caixa is not None:
        for tamanho in conteudo_caixa.findall(f".//{W}sz") + conteudo_caixa.findall(
            f".//{W}szCs"
        ):
            tamanho.set(f"{W}val", "16")

    ancestral = no_texto.getparent()
    while ancestral is not None:
        if ancestral.tag == f"{{{WP_NS}}}anchor":
            extensao = ancestral.find(f"{{{WP_NS}}}extent")
            if extensao is not None:
                extensao.set("cx", str(LARGURA_CAIXA_CABECALHO_EMU))
            for extensao_forma in ancestral.findall(f".//{{{A_NS}}}ext"):
                extensao_forma.set("cx", str(LARGURA_CAIXA_CABECALHO_EMU))
            return
        if ancestral.tag == f"{{{V_NS}}}shape":
            estilo = ancestral.get("style", "")
            estilo, quantidade = re.subn(
                r"(?<=;)width:[^;]+",
                f"width:{LARGURA_CAIXA_CABECALHO_PT}",
                estilo,
                count=1,
            )
            if not quantidade and estilo.startswith("width:"):
                estilo = re.sub(
                    r"^width:[^;]+",
                    f"width:{LARGURA_CAIXA_CABECALHO_PT}",
                    estilo,
                    count=1,
                )
            ancestral.set("style", estilo)
            return
        ancestral = ancestral.getparent()


def _atualizar_paginas_modelo(modelo, conteudo, data_relatorio: str) -> None:
    """Atualiza dados variáveis preservando o desenho institucional do modelo."""
    tabela_processo = modelo.tables[0]
    dados_processo = [
        ("Processo:", NUMERO_PROCESSO),
        ("Natureza:", "RELATÓRIO DE AUDITORIA GOVERNAMENTAL - AUDITORIA DE CONFORMIDADE"),
        (
            "Observação:",
            "Avaliar o grau de maturidade e a conformidade das práticas de governança e gestão de "
            "Tecnologia da Informação e Comunicação (TIC) adotadas pelos jurisdicionados, identificando "
            "evoluções em relação ao ciclo de 2023 e induzindo a adoção de boas práticas internacionais.",
        ),
    ]
    celula_processo = tabela_processo.cell(0, 0)
    while len(celula_processo.paragraphs) < len(dados_processo):
        celula_processo.add_paragraph()
    for paragrafo, (rotulo, valor) in zip(celula_processo.paragraphs, dados_processo):
        _substituir_texto_paragrafo(
            paragrafo,
            f"{rotulo} {valor}",
            negrito_rotulo=rotulo,
        )
        _formatar_paragrafo_docx(paragrafo, fonte="Arial", tamanho_pt=12)
    for paragrafo in list(celula_processo.paragraphs[len(dados_processo):]):
        paragrafo._element.getparent().remove(paragrafo._element)

    tabela_modelo = modelo.tables[1]
    tabela_conteudo = _localizar_tabela_apos_texto(conteudo, "DADOS DA FISCALIZAÇÃO")
    if len(tabela_modelo.rows) != len(tabela_conteudo.rows):
        raise ValueError("A tabela de dados da fiscalização diverge da estrutura do modelo.")
    for linha_modelo, linha_conteudo in zip(tabela_modelo.rows, tabela_conteudo.rows):
        for celula_modelo, celula_conteudo in zip(linha_modelo.cells, linha_conteudo.cells):
            _preencher_celula_com_paragrafos(celula_modelo, celula_conteudo.paragraphs)
            for paragrafo in celula_modelo.paragraphs:
                _formatar_paragrafo_docx(paragrafo, fonte="Arial", tamanho_pt=10)
                paragrafo.paragraph_format.line_spacing = 1.0

    # O cabeçalho do modelo inclui o número do processo em caixas de texto VML,
    # que não são expostas como runs pelo python-docx.
    partes_cabecalho = {}
    for secao in modelo.sections:
        for cabecalho in (
            secao.header,
            secao.even_page_header,
            secao.first_page_header,
        ):
            partes_cabecalho[cabecalho.part.partname] = cabecalho
    substituicoes_processo = 0
    for cabecalho in partes_cabecalho.values():
        for no in cabecalho._element.iter():
            if no.tag.endswith("}t") and no.text and "Processo nº" in no.text:
                novo_texto, quantidade = re.subn(
                    r"(?<=Processo nº )\d{3}\.\d{3}-\d/\d{2,4}",
                    NUMERO_PROCESSO,
                    no.text,
                )
                no.text = novo_texto
                substituicoes_processo += quantidade
                _ajustar_caixa_cabecalho(no)
    if not substituicoes_processo:
        raise ValueError("Número do processo não encontrado no cabeçalho institucional.")

    datas_atualizadas = 0
    for paragrafo in modelo.paragraphs:
        if not paragrafo.text.strip().startswith("CAD-TI,"):
            continue
        for filho in list(paragrafo._element):
            if filho.tag != f"{W}pPr":
                paragrafo._element.remove(filho)
        run = paragrafo.add_run(f"CAD-TI, {data_relatorio}")
        run.bold = True
        datas_atualizadas += 1
    if datas_atualizadas != 2:
        raise ValueError(
            "Quantidade inesperada de datas CAD-TI no encerramento institucional: "
            f"esperadas 2, encontradas {datas_atualizadas}."
        )


def _mapear_estilos_equivalentes(documento_origem, documento_destino) -> dict[str, str]:
    """Mapeia IDs de estilos pelo nome, pois cada DOCX pode usar IDs distintos."""
    destino_por_nome = {
        estilo.name.casefold(): estilo.style_id
        for estilo in documento_destino.styles
        if estilo.name
    }
    return {
        estilo.style_id: destino_por_nome[estilo.name.casefold()]
        for estilo in documento_origem.styles
        if estilo.name and estilo.name.casefold() in destino_por_nome
    }


def _normalizar_estilos_notas_rodape(
    docx_path: Path,
    estilos_equivalentes: dict[str, str],
) -> None:
    """Reassocia estilos das notas após a composição com o modelo institucional.

    O docxcompose converte as referências no corpo, mas pode manter em
    ``footnotes.xml`` os IDs do documento anexado. Se esses IDs não existem no
    documento principal, o Word exibe a nota com a formatação padrão.
    """
    with zipfile.ZipFile(docx_path, "r") as entrada:
        footnotes = etree.fromstring(entrada.read("word/footnotes.xml"))
        alteracoes = 0
        for atributo_estilo in footnotes.findall(f".//{W}pStyle") + footnotes.findall(
            f".//{W}rStyle"
        ):
            estilo_origem = atributo_estilo.get(f"{W}val")
            estilo_destino = estilos_equivalentes.get(estilo_origem)
            if estilo_destino and estilo_destino != estilo_origem:
                atributo_estilo.set(f"{W}val", estilo_destino)
                alteracoes += 1

        if not alteracoes:
            return

        itens_docx = [
            (item, entrada.read(item.filename))
            for item in entrada.infolist()
        ]

    with tempfile.NamedTemporaryFile(
        suffix=".docx",
        dir=docx_path.parent,
        delete=False,
    ) as arquivo_temporario:
        caminho_temporario = Path(arquivo_temporario.name)

    try:
        with zipfile.ZipFile(caminho_temporario, "w") as saida:
            for item, dados in itens_docx:
                if item.filename == "word/footnotes.xml":
                    dados = etree.tostring(
                        footnotes,
                        xml_declaration=True,
                        encoding="UTF-8",
                        standalone=True,
                    )
                saida.writestr(item, dados)
        os.replace(caminho_temporario, docx_path)
    finally:
        if caminho_temporario.exists():
            caminho_temporario.unlink()


def _capturar_notas_primeira_tabela(docx_path: Path) -> dict[str, etree._Element]:
    """Captura as notas referenciadas pela tabela DADOS DA FISCALIZAÇÃO."""
    with zipfile.ZipFile(docx_path, "r") as arquivo:
        documento_xml = etree.fromstring(arquivo.read("word/document.xml"))
        if "word/footnotes.xml" not in arquivo.namelist():
            return {}
        notas_xml = etree.fromstring(arquivo.read("word/footnotes.xml"))

    tabelas = documento_xml.findall(f".//{W}tbl")
    if not tabelas:
        raise ValueError("Tabela DADOS DA FISCALIZAÇÃO ausente no DOCX gerado.")
    ids = {
        referencia.get(f"{W}id")
        for referencia in tabelas[0].findall(f".//{W}footnoteReference")
    }
    notas_por_id = {
        nota.get(f"{W}id"): deepcopy(nota)
        for nota in notas_xml.findall(f"{W}footnote")
        if nota.get(f"{W}id") in ids
    }
    if ids != set(notas_por_id):
        raise ValueError("Definição de nota da tabela de fiscalização não encontrada.")
    return notas_por_id


def _reassociar_notas_tabela_institucional(
    docx_path: Path,
    notas_origem: dict[str, etree._Element],
) -> None:
    """Copia notas da tabela institucional com IDs livres no DOCX composto."""
    if not notas_origem:
        return

    with zipfile.ZipFile(docx_path, "r") as entrada:
        itens = [(item, entrada.read(item.filename)) for item in entrada.infolist()]
    dados_por_nome = {item.filename: dados for item, dados in itens}
    documento_xml = etree.fromstring(dados_por_nome["word/document.xml"])
    notas_xml = etree.fromstring(dados_por_nome["word/footnotes.xml"])

    tabelas = documento_xml.findall(f".//{W}tbl")
    if len(tabelas) < 2:
        raise ValueError("Tabela institucional DADOS DA FISCALIZAÇÃO não encontrada.")
    referencias = tabelas[1].findall(f".//{W}footnoteReference")
    ids_referenciados = {referencia.get(f"{W}id") for referencia in referencias}
    if not set(notas_origem).issubset(ids_referenciados):
        raise ValueError("Referências de notas da tabela institucional foram perdidas na composição.")

    ids_existentes = [
        int(nota.get(f"{W}id"))
        for nota in notas_xml.findall(f"{W}footnote")
        if nota.get(f"{W}id", "").lstrip("-").isdigit()
    ]
    proximo_id = max(ids_existentes, default=0) + 1
    novos_ids: dict[str, str] = {}
    for id_origem, nota_origem in notas_origem.items():
        novo_id = str(proximo_id)
        proximo_id += 1
        novos_ids[id_origem] = novo_id
        nova_nota = deepcopy(nota_origem)
        nova_nota.set(f"{W}id", novo_id)
        notas_xml.append(nova_nota)

    for referencia in referencias:
        id_origem = referencia.get(f"{W}id")
        if id_origem in novos_ids:
            referencia.set(f"{W}id", novos_ids[id_origem])

    # O modelo institucional conserva definições de notas de seu documento
    # original, embora elas não sejam mais referenciadas. O LibreOffice associa
    # notas também pela ordem física; por isso, eliminamos as definições órfãs
    # e renumeramos as notas usadas na mesma ordem das referências no corpo.
    referencias_documento = documento_xml.findall(f".//{W}footnoteReference")
    ids_em_ordem = list(
        dict.fromkeys(referencia.get(f"{W}id") for referencia in referencias_documento)
    )
    definicoes = {
        nota.get(f"{W}id"): nota
        for nota in notas_xml.findall(f"{W}footnote")
        if nota.get(f"{W}type") is None
    }
    ids_sem_definicao = [id_nota for id_nota in ids_em_ordem if id_nota not in definicoes]
    if ids_sem_definicao:
        raise ValueError(
            "Notas referenciadas sem definição após a composição: "
            + ", ".join(ids_sem_definicao)
        )
    mapa_compacto = {
        id_antigo: str(indice)
        for indice, id_antigo in enumerate(ids_em_ordem, start=2)
    }
    for referencia in referencias_documento:
        referencia.set(f"{W}id", mapa_compacto[referencia.get(f"{W}id")])
    for nota in list(notas_xml.findall(f"{W}footnote")):
        if nota.get(f"{W}type") is None:
            notas_xml.remove(nota)
    for id_antigo in ids_em_ordem:
        nota = definicoes[id_antigo]
        nota.set(f"{W}id", mapa_compacto[id_antigo])
        notas_xml.append(nota)

    dados_por_nome["word/document.xml"] = etree.tostring(
        documento_xml,
        xml_declaration=True,
        encoding="UTF-8",
        standalone=True,
    )
    dados_por_nome["word/footnotes.xml"] = etree.tostring(
        notas_xml,
        xml_declaration=True,
        encoding="UTF-8",
        standalone=True,
    )

    with tempfile.NamedTemporaryFile(
        suffix=".docx",
        dir=docx_path.parent,
        delete=False,
    ) as arquivo_temporario:
        caminho_temporario = Path(arquivo_temporario.name)
    try:
        with zipfile.ZipFile(
            caminho_temporario,
            "w",
            compression=zipfile.ZIP_DEFLATED,
        ) as saida:
            for item, _ in itens:
                saida.writestr(item, dados_por_nome[item.filename])
        os.replace(caminho_temporario, docx_path)
    finally:
        if caminho_temporario.exists():
            caminho_temporario.unlink()


def _validar_estilos_notas_rodape(docx_path: Path) -> int:
    """Valida os estilos das notas efetivamente referenciadas no relatório."""
    with zipfile.ZipFile(docx_path, "r") as arquivo:
        estilos = etree.fromstring(arquivo.read("word/styles.xml"))
        documento = etree.fromstring(arquivo.read("word/document.xml"))
        footnotes = etree.fromstring(arquivo.read("word/footnotes.xml"))

    estilos_definidos = {
        estilo.get(f"{W}styleId") for estilo in estilos.findall(f"{W}style")
    }
    notas_referenciadas = {
        referencia.get(f"{W}id")
        for referencia in documento.findall(f".//{W}footnoteReference")
    }
    estilos_ausentes: set[tuple[str, str]] = set()
    notas_encontradas = 0
    for nota in footnotes.findall(f"{W}footnote"):
        nota_id = nota.get(f"{W}id")
        if nota_id not in notas_referenciadas:
            continue
        notas_encontradas += 1
        for no_estilo in nota.findall(f".//{W}pStyle") + nota.findall(f".//{W}rStyle"):
            estilo_id = no_estilo.get(f"{W}val")
            if estilo_id and estilo_id not in estilos_definidos:
                estilos_ausentes.add((nota_id, estilo_id))

    if notas_encontradas != len(notas_referenciadas):
        raise ValueError(
            "Há referências de nota de rodapé sem definição correspondente: "
            f"{len(notas_referenciadas) - notas_encontradas}."
        )
    if estilos_ausentes:
        detalhes = ", ".join(
            f"nota {nota_id}: {estilo_id}"
            for nota_id, estilo_id in sorted(estilos_ausentes)
        )
        raise ValueError(f"Notas de rodapé referenciam estilos inexistentes: {detalhes}")
    return notas_encontradas


def _mover_notas_isoladas_para_legendas_figuras(docx_path: Path) -> int:
    """Move notas declaradas após figuras para suas legendas no DOCX.

    O escritor DOCX do Pandoc omite a legenda quando uma nota de rodapé é
    inserida diretamente no texto alternativo da imagem. No Markdown, a nota é
    portanto declarada em um parágrafo isolado logo após a figura. Esta etapa
    incorpora a referência na legenda, depois de "situações encontradas", e
    remove o parágrafo intermediário vazio.
    """
    with zipfile.ZipFile(docx_path, "r") as entrada:
        documento = etree.fromstring(entrada.read("word/document.xml"))
        itens_docx = [(item, entrada.read(item.filename)) for item in entrada.infolist()]

    marcador_legenda = "Manifestações sobre as situações encontradas"
    ponto_insercao = "situações encontradas"
    notas_movidas = 0

    for paragrafo_nota in list(documento.findall(f".//{W}p")):
        referencias = paragrafo_nota.findall(f".//{W}footnoteReference")
        texto_nota = "".join(
            no.text or "" for no in paragrafo_nota.findall(f".//{W}t")
        ).strip()
        if texto_nota or len(referencias) != 1:
            continue

        paragrafo_imagem = paragrafo_nota.getprevious()
        paragrafo_legenda = (
            paragrafo_imagem.getprevious() if paragrafo_imagem is not None else None
        )
        if (
            paragrafo_imagem is None
            or paragrafo_legenda is None
            or paragrafo_imagem.find(f".//{W}drawing") is None
        ):
            continue

        texto_legenda = "".join(
            no.text or "" for no in paragrafo_legenda.findall(f".//{W}t")
        )
        if marcador_legenda not in texto_legenda:
            continue

        no_texto = next(
            (
                no
                for no in paragrafo_legenda.findall(f".//{W}t")
                if ponto_insercao in (no.text or "")
            ),
            None,
        )
        if no_texto is None:
            continue

        texto_original = no_texto.text or ""
        posicao = texto_original.index(ponto_insercao) + len(ponto_insercao)
        prefixo, sufixo = texto_original[:posicao], texto_original[posicao:]
        no_texto.text = prefixo

        run_legenda = no_texto.getparent()
        run_referencia = referencias[0].getparent()
        paragrafo_nota.remove(run_referencia)

        indice_run = paragrafo_legenda.index(run_legenda)
        paragrafo_legenda.insert(indice_run + 1, run_referencia)
        if sufixo:
            run_sufixo = etree.Element(f"{W}r")
            propriedades = run_legenda.find(f"{W}rPr")
            if propriedades is not None:
                run_sufixo.append(deepcopy(propriedades))
            texto_sufixo = etree.SubElement(run_sufixo, f"{W}t")
            texto_sufixo.set(
                "{http://www.w3.org/XML/1998/namespace}space",
                "preserve",
            )
            texto_sufixo.text = sufixo
            paragrafo_legenda.insert(indice_run + 2, run_sufixo)

        paragrafo_nota.getparent().remove(paragrafo_nota)
        notas_movidas += 1

    if not notas_movidas:
        return 0

    with tempfile.NamedTemporaryFile(
        suffix=".docx",
        dir=docx_path.parent,
        delete=False,
    ) as arquivo_temporario:
        caminho_temporario = Path(arquivo_temporario.name)

    try:
        with zipfile.ZipFile(
            caminho_temporario,
            "w",
            compression=zipfile.ZIP_DEFLATED,
        ) as saida:
            for item, dados in itens_docx:
                if item.filename == "word/document.xml":
                    dados = etree.tostring(
                        documento,
                        xml_declaration=True,
                        encoding="UTF-8",
                        standalone=True,
                    )
                saida.writestr(item, dados)
        os.replace(caminho_temporario, docx_path)
    finally:
        if caminho_temporario.exists():
            caminho_temporario.unlink()

    return notas_movidas


def _configurar_faixas_horizontais_tabelas(documento) -> int:
    """Mantém faixas horizontais e desativa faixas verticais nas tabelas."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    atributos_bits = {
        "firstRow": 0x0020,
        "lastRow": 0x0040,
        "firstColumn": 0x0080,
        "lastColumn": 0x0100,
        "noHBand": 0x0200,
        "noVBand": 0x0400,
    }
    for tabela in documento.tables:
        propriedades = tabela._tbl.tblPr
        tbl_look = propriedades.find(qn("w:tblLook"))
        if tbl_look is None:
            tbl_look = OxmlElement("w:tblLook")
            propriedades.append(tbl_look)

        tbl_look.set(qn("w:noHBand"), "0")
        tbl_look.set(qn("w:noVBand"), "1")

        # w:val duplica os atributos individuais em uma máscara hexadecimal.
        # Mantê-los coerentes evita interpretações diferentes entre editores.
        mascara = 0
        for atributo, bit in atributos_bits.items():
            if tbl_look.get(qn(f"w:{atributo}"), "0") == "1":
                mascara |= bit
        tbl_look.set(qn("w:val"), f"{mascara:04X}")

    return len(documento.tables)


def _validar_faixas_horizontais_tabelas(
    docx_path: Path,
    quantidade_esperada: int,
) -> int:
    """Confirma que tabelas de relatório usam linhas, não colunas, alternadas."""
    with zipfile.ZipFile(docx_path, "r") as arquivo:
        estilos = etree.fromstring(arquivo.read("word/styles.xml"))
        documento = etree.fromstring(arquivo.read("word/document.xml"))

    estilos_tabela = set()
    for estilo in estilos.findall(f"{W}style"):
        nome = estilo.find(f"{W}name")
        if nome is not None and nome.get(f"{W}val", "").casefold() == "table":
            estilos_tabela.add(estilo.get(f"{W}styleId"))

    tabelas_relatorio = []
    configuracoes_incorretas = []
    for indice, tabela in enumerate(documento.findall(f".//{W}tbl")):
        estilo = tabela.find(f"./{W}tblPr/{W}tblStyle")
        if estilo is None or estilo.get(f"{W}val") not in estilos_tabela:
            continue
        tabelas_relatorio.append(indice)
        tbl_look = tabela.find(f"./{W}tblPr/{W}tblLook")
        if (
            tbl_look is None
            or tbl_look.get(f"{W}noHBand") != "0"
            or tbl_look.get(f"{W}noVBand") != "1"
        ):
            configuracoes_incorretas.append(indice)

    if len(tabelas_relatorio) != quantidade_esperada:
        raise ValueError(
            "Quantidade inesperada de tabelas do relatório após a composição: "
            f"esperadas {quantidade_esperada}, encontradas {len(tabelas_relatorio)}."
        )
    if configuracoes_incorretas:
        raise ValueError(
            "Há tabelas com faixas verticais habilitadas ou faixas horizontais "
            f"desabilitadas: {configuracoes_incorretas}."
        )
    return len(tabelas_relatorio)


def _remover_revisoes_controladas(docx_path: Path) -> None:
    """Aceita inserções, descarta exclusões e desativa o controle de alterações."""
    with zipfile.ZipFile(docx_path, "r") as entrada:
        entradas = {item.filename: entrada.read(item.filename) for item in entrada.infolist()}

    alterado = False
    for nome, dados in list(entradas.items()):
        if not nome.startswith("word/") or not nome.endswith(".xml"):
            continue
        try:
            raiz = etree.fromstring(dados)
        except etree.XMLSyntaxError:
            continue

        parte_alterada = False
        for tag in (f"{W}ins", f"{W}moveTo"):
            for no in list(raiz.findall(f".//{tag}")):
                pai = no.getparent()
                indice = pai.index(no)
                for filho in list(no):
                    pai.insert(indice, filho)
                    indice += 1
                pai.remove(no)
                parte_alterada = True
        for tag in (f"{W}del", f"{W}moveFrom"):
            for no in list(raiz.findall(f".//{tag}")):
                no.getparent().remove(no)
                parte_alterada = True
        for no in list(raiz.findall(f".//{W}trackRevisions")):
            no.getparent().remove(no)
            parte_alterada = True

        if parte_alterada:
            entradas[nome] = etree.tostring(
                raiz,
                xml_declaration=True,
                encoding="UTF-8",
                standalone=True,
            )
            alterado = True

    if not alterado:
        return

    with tempfile.NamedTemporaryFile(
        suffix=".docx",
        dir=docx_path.parent,
        delete=False,
    ) as arquivo_temporario:
        caminho_temporario = Path(arquivo_temporario.name)
    try:
        with zipfile.ZipFile(
            caminho_temporario,
            "w",
            compression=zipfile.ZIP_DEFLATED,
        ) as saida:
            for nome, dados in entradas.items():
                saida.writestr(nome, dados)
        os.replace(caminho_temporario, docx_path)
    finally:
        if caminho_temporario.exists():
            caminho_temporario.unlink()


def aplicar_paginas_institucionais(
    docx_gerado: Path,
    modelo_path: Path,
    *,
    data_relatorio: str,
) -> None:
    """Compõe frontispício, sumário, corpo atual e encerramento institucional.

    O reference-docx do Pandoc transfere estilos, mas não o conteúdo pré-textual.
    Esta etapa preserva as páginas institucionais do modelo e incorpora, do
    documento recém-gerado, a lista de anexos e o corpo até a seção 7.
    """
    try:
        from docxcompose.composer import Composer
    except ImportError as exc:
        raise RuntimeError(
            "docxcompose é necessário para aplicar as páginas institucionais do modelo."
        ) from exc

    if not modelo_path.exists():
        raise FileNotFoundError(f"Modelo institucional não encontrado: {modelo_path}")

    documento_modelo = docx.Document(str(modelo_path))
    documento_conteudo = docx.Document(str(docx_gerado))
    notas_tabela_institucional = _capturar_notas_primeira_tabela(docx_gerado)
    estilos_equivalentes = _mapear_estilos_equivalentes(
        documento_conteudo,
        documento_modelo,
    )
    _atualizar_paginas_modelo(
        documento_modelo,
        documento_conteudo,
        data_relatorio,
    )

    indice_marcador = _localizar_ultimo_elemento_docx(
        documento_modelo,
        MARCADOR_CORPO_INSTITUCIONAL,
    )
    elementos_modelo = list(documento_modelo.element.body)
    elementos_encerramento = [
        deepcopy(elemento)
        for elemento in elementos_modelo[indice_marcador + 1:]
        if elemento.tag != f"{W}sectPr"
    ]
    if not any(
        _texto_elemento_docx(elemento).strip().startswith(
            "O presente relatório foi objeto de supervisão"
        )
        for elemento in elementos_encerramento
    ):
        raise ValueError("Página de encerramento ausente no modelo institucional.")

    # O frontispício vem do modelo; o sumário, a lista de anexos e o conteúdo
    # técnico vêm do Markdown atual para que o campo TOC não retenha paginação
    # em cache do documento usado como referência institucional.
    _remover_corpo_a_partir_de(documento_modelo, MARCADOR_CORPO_INSTITUCIONAL)
    _remover_corpo_antes_de(documento_conteudo, "SUMÁRIO")
    total_tabelas = _configurar_faixas_horizontais_tabelas(documento_conteudo)

    with tempfile.TemporaryDirectory() as temp_dir:
        temp_dir_path = Path(temp_dir)
        modelo_recortado = temp_dir_path / "modelo-frontispicio.docx"
        conteudo_recortado = temp_dir_path / "conteudo-relatorio.docx"
        documento_modelo.save(modelo_recortado)
        documento_conteudo.save(conteudo_recortado)

        principal = docx.Document(str(modelo_recortado))
        compositor = Composer(principal)
        compositor.append(docx.Document(str(conteudo_recortado)))
        compositor.save(str(docx_gerado))

    documento_final = docx.Document(str(docx_gerado))
    body_final = documento_final.element.body
    sect_pr = body_final.sectPr
    for elemento in elementos_encerramento:
        body_final.insert(body_final.index(sect_pr), elemento)
    documento_final.save(str(docx_gerado))

    _reassociar_notas_tabela_institucional(docx_gerado, notas_tabela_institucional)
    _remover_revisoes_controladas(docx_gerado)
    _normalizar_estilos_notas_rodape(docx_gerado, estilos_equivalentes)
    total_notas = _validar_estilos_notas_rodape(docx_gerado)
    logger.info("Notas de rodapé validadas após a composição: %d", total_notas)
    total_tabelas = _validar_faixas_horizontais_tabelas(
        docx_gerado,
        total_tabelas,
    )
    logger.info(
        "Tabelas validadas com faixas somente horizontais: %d",
        total_tabelas,
    )
    marcar_atualizacao_campos_docx(str(docx_gerado))


def materializar_sumario(docx_path: Path) -> None:
    """Calcula as entradas e páginas do sumário no editor disponível."""
    if os.name == "nt":
        from atualizar_sumario_word import atualizar_sumario_word

        sucesso, _ = atualizar_sumario_word(docx_path, exportar_pdf=False)
        if not sucesso:
            raise RuntimeError("O Word não conseguiu materializar o sumário do relatório.")
        logger.info("Sumário materializado pelo Microsoft Word.")
        return

    from atualizar_sumario_libreoffice import atualizar_sumario_libreoffice

    quantidade = atualizar_sumario_libreoffice(docx_path)
    logger.info(
        "Sumário materializado pelo LibreOffice (%d índice).",
        quantidade,
    )


def ajustar_tabela_modelo_plano_acao(docx_path: Path) -> None:
    """Alinha o quadro referencial ao início do texto do encaminhamento 2."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Twips
    from docx.text.paragraph import Paragraph

    recuo_dxa = 720  # 1,27 cm, mesmo recuo do texto da lista numerada.
    largura_dxa = 8352  # 16 cm de área útil menos o recuo da lista.
    documento = docx.Document(str(docx_path))
    tabelas_encontradas = 0
    for tabela in documento.tables:
        cabecalho = [celula.text.strip().casefold() for celula in tabela.rows[0].cells]
        if not cabecalho or "determinação ou recomendação" not in cabecalho[0]:
            continue
        tabelas_encontradas += 1
        tabela.alignment = WD_TABLE_ALIGNMENT.LEFT
        tabela.autofit = False
        tbl_pr = tabela._tbl.tblPr
        tbl_ind = tbl_pr.find(qn("w:tblInd"))
        if tbl_ind is None:
            tbl_ind = OxmlElement("w:tblInd")
            tbl_pr.append(tbl_ind)
        tbl_ind.set(qn("w:w"), str(recuo_dxa))
        tbl_ind.set(qn("w:type"), "dxa")
        tbl_w = tbl_pr.find(qn("w:tblW"))
        if tbl_w is None:
            tbl_w = OxmlElement("w:tblW")
            tbl_pr.append(tbl_w)
        tbl_w.set(qn("w:w"), str(largura_dxa))
        tbl_w.set(qn("w:type"), "dxa")

        colunas_grid = tabela._tbl.tblGrid.findall(qn("w:gridCol"))
        larguras_atuais = [int(coluna.get(qn("w:w"), "0")) for coluna in colunas_grid]
        total_atual = sum(larguras_atuais)
        if total_atual <= 0:
            raise ValueError("Larguras das colunas do modelo de plano de ação inválidas.")
        larguras_novas = [
            round(largura * largura_dxa / total_atual)
            for largura in larguras_atuais
        ]
        larguras_novas[-1] += largura_dxa - sum(larguras_novas)
        for coluna, largura in zip(colunas_grid, larguras_novas):
            coluna.set(qn("w:w"), str(largura))
        for linha in tabela.rows:
            for celula, largura in zip(linha.cells, larguras_novas):
                celula.width = Twips(largura)

        anterior = tabela._element.getprevious()
        posterior = tabela._element.getnext()
        for elemento, nome_estilo in (
            (anterior, "Table Caption"),
            (posterior, "FonteImagem"),
        ):
            if elemento is None or elemento.tag != qn("w:p"):
                continue
            paragrafo = Paragraph(elemento, tabela._parent)
            if nome_estilo in documento.styles:
                paragrafo.style = documento.styles[nome_estilo]
            paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragrafo.paragraph_format.left_indent = Twips(recuo_dxa)
            paragrafo.paragraph_format.right_indent = Twips(0)

    if tabelas_encontradas != 1:
        raise ValueError(
            "Quantidade inesperada de tabelas do modelo referencial de plano de ação: "
            f"esperada 1, encontrada {tabelas_encontradas}."
        )
    documento.save(str(docx_path))


def ajustar_tabela_lista_anexos(docx_path: Path) -> None:
    """Distribui a tabela de anexos em coluna curta de código e descrição ampla."""
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Twips

    larguras_dxa = (1600, 7472)  # 17,6% para o documento e 82,4% para a descrição.
    largura_total = sum(larguras_dxa)
    documento = docx.Document(str(docx_path))
    tabelas_encontradas = 0
    for tabela in documento.tables:
        if len(tabela.columns) != 2 or not tabela.rows:
            continue
        cabecalho = [celula.text.strip().casefold() for celula in tabela.rows[0].cells]
        if cabecalho != ["documento nº", "descrição"]:
            continue
        tabelas_encontradas += 1
        tabela.alignment = WD_TABLE_ALIGNMENT.LEFT
        tabela.autofit = False
        tbl_pr = tabela._tbl.tblPr
        tbl_ind = tbl_pr.find(qn("w:tblInd"))
        if tbl_ind is None:
            tbl_ind = OxmlElement("w:tblInd")
            tbl_pr.append(tbl_ind)
        tbl_ind.set(qn("w:w"), "0")
        tbl_ind.set(qn("w:type"), "dxa")
        tbl_w = tbl_pr.find(qn("w:tblW"))
        if tbl_w is None:
            tbl_w = OxmlElement("w:tblW")
            tbl_pr.append(tbl_w)
        tbl_w.set(qn("w:w"), str(largura_total))
        tbl_w.set(qn("w:type"), "dxa")

        colunas_grid = tabela._tbl.tblGrid.findall(qn("w:gridCol"))
        if len(colunas_grid) != 2:
            raise ValueError("Grade inesperada na tabela da Lista de Anexos.")
        for coluna, largura in zip(colunas_grid, larguras_dxa):
            coluna.set(qn("w:w"), str(largura))
        for linha in tabela.rows:
            for celula, largura in zip(linha.cells, larguras_dxa):
                celula.width = Twips(largura)

    if tabelas_encontradas != 1:
        raise ValueError(
            "Quantidade inesperada de tabelas da Lista de Anexos: "
            f"esperada 1, encontrada {tabelas_encontradas}."
        )
    documento.save(str(docx_path))


REFERENCE_LABEL_RE = re.compile(r"\{#(?:fig|tbl):[^#\r\n]+#\}")


def proteger_rotulos_referencias(conteudo: str) -> tuple[str, dict[str, str]]:
    """Protege labels de figuras/tabelas para que o Jinja não os trate como comentários."""
    rotulos: dict[str, str] = {}

    def substituir(match: re.Match[str]) -> str:
        token = f"@@ARGOS_REFERENCE_LABEL_{len(rotulos)}@@"
        if token in conteudo:
            raise ValueError(f"Token interno de proteção já existe no Markdown: {token}")
        rotulos[token] = match.group(0)
        return token

    return REFERENCE_LABEL_RE.sub(substituir, conteudo), rotulos


def restaurar_rotulos_referencias(conteudo: str, rotulos: dict[str, str]) -> str:
    """Restaura os labels que permaneceram após a renderização Jinja."""
    for token, rotulo in rotulos.items():
        conteudo = conteudo.replace(token, rotulo)
    return conteudo

# Tenta importar as bibliotecas necessárias
try:
    import pypandoc
except ImportError:
    logger.error(
        "pypandoc não está instalado neste ambiente Python. Execute o script utilizando o ambiente virtual do projeto: "
        "  scripts/.venv/bin/python scripts/gerar_relatorio_consolidado.py"
    )
    sys.exit(1)

try:
    import docx
except ImportError:
    logger.error(
        "python-docx não está instalado neste ambiente Python. Execute o script utilizando o ambiente virtual do projeto: "
        "  scripts/.venv/bin/python scripts/gerar_relatorio_consolidado.py"
    )
    sys.exit(1)



def main() -> int:
    parser = argparse.ArgumentParser(
        description="Gera o arquivo DOCX do Relatório Consolidado a partir do Markdown."
    )
    parser.add_argument(
        "input_positional",
        nargs="?",
        default=None,
        help="Caminho do arquivo Markdown de entrada (posicional)."
    )
    parser.add_argument(
        "--input",
        default=None,
        help="Caminho do arquivo Markdown de entrada."
    )
    parser.add_argument(
        "--output",
        default=None,
        help="Caminho do arquivo DOCX de saída gerado. Se omitido, usa o mesmo nome do arquivo de entrada com extensão .docx."
    )
    parser.add_argument(
        "--reference-docx",
        default="scripts/resources/template-base-estilos.docx",
        help="Caminho para o documento de estilos do Word usado como referência."
    )
    parser.add_argument(
        "--modelo-institucional",
        default=str(DEFAULT_MODELO_INSTITUCIONAL),
        help=(
            "DOCX com frontispício, cabeçalho, sumário e encerramento institucionais. "
            "Por padrão, utiliza o modelo consolidado versionado no repositório."
        ),
    )
    parser.add_argument(
        "--resource-files", "--resource-dirs",
        nargs="*",
        default=[],
        dest="resource_files",
        help="Diretórios adicionais, arquivos ou padrões glob (wildcards) contendo recursos de contexto."
    )
    parser.add_argument(
        "--context-json",
        default=str(DEFAULT_CONTEXT_JSON),
        help=(
            "Caminho para um arquivo JSON contendo variáveis de contexto para renderização Jinja2. "
            "Por padrão, utiliza o diagnóstico transversal oficial do iGovTI 2026."
        ),
    )
    parser.add_argument(
        "--context-vars",
        nargs="*",
        default=[],
        help="Parâmetros de contexto adicionais chave=valor (ex: sigla=FTM ano=2026)."
    )
    parser.add_argument("--resultados-2026", default=None, help="XLSX de resultados iGovTI 2026 usado para gerar gráficos do relatório consolidado.")
    parser.add_argument("--respostas-2026", default=None, help="XLSX de respostas do questionário usado para gerar gráficos do relatório consolidado.")
    parser.add_argument("--comparavel-2026", default=None, help="XLSX iGovTI 2026 ajustado comparável usado para gerar gráficos do relatório consolidado.")
    parser.add_argument("--setic-2023", default=None, help="XLSX comparável do SETIC 2023 usado para gerar gráficos do relatório consolidado.")
    parser.add_argument("--municipios-2023", default=None, help="XLSX comparável dos municípios 2023 usado para gerar gráficos do relatório consolidado.")
    parser.add_argument("--auditados-xlsx", default=None, help="Base de auditados XLSX usada para gerar gráficos consolidados de achados.")
    parser.add_argument("--resultado-auditoria-json", default=None, help="Resultado estruturado da auditoria em JSON usado para gerar gráficos consolidados de achados.")
    parser.add_argument("--mapa", default=None, help="Mapa de verificação revisado usado nos perfis das questões Q1 a Q6.")
    args = parser.parse_args()

    input_str = args.input_positional or args.input or "03-Relatorios/01-Relatorio_Consolidado/Relatório_altaresolucao_novo.md"
    input_path = Path(input_str)
    
    if args.output:
        output_path = Path(args.output)
    else:
        output_path = input_path.with_suffix(".docx")

    ref_path = Path(args.reference_docx)

    if not input_path.exists():
        logger.error("Arquivo de entrada não encontrado: %s", input_path)
        return 1

    if not ref_path.exists():
        logger.error("Documento de referência de estilos não encontrado: %s", ref_path)
        return 1

    logger.info("Lendo markdown: %s", input_path)
    conteudo = input_path.read_text(encoding="utf-8")

    # 1. Resolve variáveis de contexto para renderização Jinja2
    contexto = {
        "data_hoje_abnt": data_hoje_abnt(),
        "data_hoje": data_hoje(),
    }
    for numero_achado, ids_situacoes in SITUACOES_POR_ACHADO.items():
        contexto[f"nota_criterios_especificos_a{numero_achado}"] = (
            _nota_criterios_especificos_achado(
                args.mapa,
                args.resultado_auditoria_json,
                args.auditados_xlsx,
                ids_situacoes,
            )
        )
    if args.context_json:
        context_json_path = Path(args.context_json)
        if context_json_path.exists():
            logger.info("Carregando variáveis do arquivo JSON: %s", context_json_path)
            contexto.update(json.loads(context_json_path.read_text(encoding="utf-8")))
        else:
            logger.warning("Arquivo JSON de contexto não encontrado: %s", context_json_path)

    if args.context_vars:
        for item in args.context_vars:
            if "=" in item:
                k, v = item.split("=", 1)
                contexto[k.strip()] = v.strip()

    # 2. Renderiza como Jinja2 se houver tags de template. Os labels
    # {#fig:...#}/{#tbl:...#} precisam ser protegidos porque essa sintaxe
    # também representa comentários para o Jinja.
    if "{{" in conteudo or "{%" in conteudo:
        logger.info("Renderizando variáveis do template com Jinja2...")
        try:
            from jinja2 import Template
            conteudo_protegido, rotulos_referencias = proteger_rotulos_referencias(conteudo)
            template = Template(conteudo_protegido)
            conteudo = restaurar_rotulos_referencias(
                template.render(contexto),
                rotulos_referencias,
            )
        except Exception as e:
            logger.error("Erro ao renderizar template Jinja2: %s", e)
            return 1

    # 3. Aplica os processamentos de marcação herdados do Argos
    conteudo_processado = cross_ref_figuras(conteudo)
    conteudo_processado = cross_ref_tabelas(conteudo_processado)
    conteudo_processado = inserir_campo_sumario_docx(conteudo_processado)
    conteudo_processado = processar_quebras_pagina(conteudo_processado)
    conteudo_processado = substituir_underline_pandoc(conteudo_processado)

    # Cria arquivo temporário para o markdown intermediário
    with tempfile.NamedTemporaryFile(suffix=".md", mode="w", encoding="utf-8", delete=False) as temp_md:
        temp_md.write(conteudo_processado)
        temp_md_name = temp_md.name

    with tempfile.TemporaryDirectory() as temp_resources_dir:
        temp_img_dir = os.path.join(temp_resources_dir, "img")
        os.makedirs(temp_img_dir, exist_ok=True)
        
        # 4. Executa a geração de gráficos gerais consolidados na pasta temporária
        import subprocess
        import shutil
        
        script_gerais = os.path.join(os.path.dirname(__file__), "gerar_graficos_relatorios_consolidado_individuais_igovti.py")
        logger.info("Executando a geração de gráficos gerais consolidados no diretório temporário...")
        try:
            cmd_graficos_gerais = [sys.executable, script_gerais, "--somente-consolidados", "--output-root", temp_resources_dir]
            for option, value in [
                ("--resultados-2026", args.resultados_2026),
                ("--respostas-2026", args.respostas_2026),
                ("--comparavel-2026", args.comparavel_2026),
                ("--setic-2023", args.setic_2023),
                ("--municipios-2023", args.municipios_2023),
                ("--mapa", args.mapa),
            ]:
                if value:
                    cmd_graficos_gerais.extend([option, value])
            # Executa com o mesmo interpretador python
            subprocess.run(
                cmd_graficos_gerais,
                check=True,
                stdout=subprocess.DEVNULL
            )
        except Exception as e:
            logger.error("Erro ao gerar gráficos gerais: %s", e)
            return 1

        # Executa a geração de gráficos de achados diretamente para a pasta temporária
        script_achados = os.path.join(os.path.dirname(__file__), "gerar_graficos_achados_consolidado.py")
        logger.info("Executando a geração de gráficos de achados no diretório temporário...")
        try:
            cmd_graficos_achados = [sys.executable, script_achados, "--output-dir", temp_resources_dir]
            if args.auditados_xlsx:
                cmd_graficos_achados.extend(["--auditados", args.auditados_xlsx])
            if args.resultado_auditoria_json:
                cmd_graficos_achados.extend(["--resultado-auditoria-json", args.resultado_auditoria_json])
            subprocess.run(
                cmd_graficos_achados,
                check=True,
                stdout=subprocess.DEVNULL
            )
        except Exception as e:
            logger.error("Erro ao gerar gráficos de achados: %s", e)
            return 1
            
        # 5. Copia imagens de contexto de outros locais (se existirem / especificados) para consolidar na pasta temporária
        tmp_root = Path("C:/tmp") if sys.platform.startswith("win") else Path(tempfile.gettempdir())
        tmp_pkg = tmp_root / "tcerj-igovti-2026"
        tmp_consolidado_img = str(tmp_pkg / "relatorio-consolidado" / "img")
        tmp_individuais_img = str(tmp_pkg / "relatorios-individuais" / "img")
        
        resource_inputs = [
            tmp_consolidado_img,
            tmp_individuais_img,
            "03-Relatorios/01-Relatorio_Consolidado/img",
            "03-Relatorios/99-Analise_Longitudinal/img",
            "03-Relatorios/99-Avaliacao_IA/img",
        ]
        if args.resource_files:
            resource_inputs.extend(args.resource_files)
            
        # Resolve wildcards/globs primeiro
        import glob
        import zipfile
        
        resolved_paths = []
        for r_path in resource_inputs:
            if not r_path:
                continue
            # Verifica se contém caracteres wildcard
            if any(char in r_path for char in ['*', '?', '[']):
                expanded = glob.glob(r_path, recursive=True)
                if expanded:
                    resolved_paths.extend([p for p in expanded if os.path.exists(p)])
                else:
                    logger.warning("Nenhum recurso encontrado para o padrão: %r", r_path)
            else:
                resolved_paths.append(r_path)

        for r_path in resolved_paths:
            if not os.path.exists(r_path):
                continue

            if os.path.isdir(r_path):
                for root, _, files in os.walk(r_path):
                    for filename in files:
                        if filename.lower().endswith((".png", ".jpg", ".jpeg")):
                            src_file = os.path.join(root, filename)
                            dest_file = os.path.join(temp_img_dir, filename)
                            if not os.path.exists(dest_file):
                                shutil.copy2(src_file, dest_file)
            elif r_path.lower().endswith('.zip'):
                try:
                    with tempfile.TemporaryDirectory() as unzip_tmp:
                        with zipfile.ZipFile(r_path, 'r') as zip_ref:
                            zip_ref.extractall(unzip_tmp)
                        for root, _, files in os.walk(unzip_tmp):
                            for filename in files:
                                if filename.lower().endswith((".png", ".jpg", ".jpeg")):
                                    src_file = os.path.join(root, filename)
                                    dest_file = os.path.join(temp_img_dir, filename)
                                    if not os.path.exists(dest_file):
                                        shutil.copy2(src_file, dest_file)
                except Exception as e:
                    logger.error("Erro ao extrair zip de recursos %r: %s", r_path, e)
            else:
                # Arquivo normal
                filename = os.path.basename(r_path)
                if filename.lower().endswith((".png", ".jpg", ".jpeg")):
                    dest_file = os.path.join(temp_img_dir, filename)
                    if not os.path.exists(dest_file):
                        shutil.copy2(r_path, dest_file)

        # Planifica todas as imagens encontradas nas subpastas do diretório temporário
        # colocando-as também na raiz de temp_resources_dir para que o Pandoc as encontre
        # diretamente quando referenciadas por nome simples (sem prefixo de pasta)
        for root, _, files in os.walk(temp_resources_dir):
            if root == temp_resources_dir:
                continue
            for filename in files:
                if filename.lower().endswith((".png", ".jpg", ".jpeg")):
                    src_file = os.path.join(root, filename)
                    dest_file_root = os.path.join(temp_resources_dir, filename)
                    if not os.path.exists(dest_file_root):
                        shutil.copy2(src_file, dest_file_root)


        try:
            logger.info("Compilando com Pandoc para DOCX usando template de referência: %s", ref_path)
            
            # Garante diretórios de recursos para o Pandoc
            resource_paths = [
                ".",
                temp_resources_dir,
                str(input_path.parent),
                os.path.dirname(args.reference_docx)
            ]
            resource_path_arg = '--resource-path=' + os.pathsep.join(resource_paths)
            
            extra_args = [
                '--figure-caption-position=above',
                '--reference-doc=' + str(ref_path),
                resource_path_arg
            ]
            if TOC_MARKER_FILTER.exists():
                extra_args.append('--lua-filter=' + str(TOC_MARKER_FILTER))

            # Executa conversão
            output_path.parent.mkdir(parents=True, exist_ok=True)
            pypandoc.convert_file(
                temp_md_name,
                to="docx",
                outputfile=str(output_path),
                extra_args=extra_args
            )

            notas_em_legendas = _mover_notas_isoladas_para_legendas_figuras(output_path)
            if notas_em_legendas:
                logger.info(
                    "%d notas de rodapé vinculadas às legendas das figuras.",
                    notas_em_legendas,
                )

            logger.info("Aplicando estilos de tabela pós-conversão no DOCX...")
            aplicar_estilo_tabelas(str(output_path))
            ajustar_tabela_modelo_plano_acao(output_path)
            ajustar_tabela_lista_anexos(output_path)

            logger.info("Ajustando layout para evitar quebras órfãs de figuras, tabelas e fontes...")
            evitar_quebra_elementos(str(output_path))

            logger.info("Marcando campos do DOCX para atualização ao abrir no Word...")
            marcar_atualizacao_campos_docx(str(output_path))

            logger.info(
                "Aplicando páginas do modelo institucional: %s",
                args.modelo_institucional,
            )
            aplicar_paginas_institucionais(
                output_path,
                Path(args.modelo_institucional),
                data_relatorio=str(contexto["data_hoje"]),
            )
            materializar_sumario(output_path)

            logger.info("Relatório DOCX gerado com sucesso em: %s", output_path)
            return 0

        except Exception as e:
            logger.error("Erro durante a geração do relatório: %s", e)
            return 1
        finally:
            if os.path.exists(temp_md_name):
                os.remove(temp_md_name)


if __name__ == "__main__":
    sys.exit(main())
