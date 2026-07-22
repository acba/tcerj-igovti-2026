#!/usr/bin/env python
"""Pipeline auditável das seções 1 e 2 dos comentários do gestor."""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import html
import json
import math
import os
import re
import sys
import tempfile
import threading
import unicodedata
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from contextlib import nullcontext
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Iterable
from xml.etree import ElementTree

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts.avaliacao_evidencias.evidence_processing import (
    _url_download_permitida,
    baixar_recurso_url,
    erro_tecnico_bloqueante_pacote,
    preparar_evidencia_para_provider,
    resultado_indica_erro_tecnico,
)
from scripts.avaliacao_evidencias.inventory import resolver_evidencia
from scripts.avaliacao_evidencias.key_pool import (
    ExclusiveApiKeyPool,
    cooldown_429,
    resultado_429,
)
from scripts.avaliacao_evidencias.prompts import (
    filtrar_itens_por_prompt,
    preparar_itens_para_prompt,
    resolver_prompt,
)
from scripts.avaliacao_evidencias.providers import estimar_tokens_payload, executar_provider, limite_tokens_provider
from scripts.avaliacao_evidencias.questionnaire import (
    ItemAfirmado,
    carregar_contexto_questionario,
)
from scripts.avaliacao_evidencias.scope_validation import (
    formatar_violacoes,
    motivos_ativos_registro,
    validar_resultado_no_escopo,
)
from scripts.avaliacao_evidencias.utils import RequestsPerMinuteLimiter, hash_arquivo, log_event
from scripts.resources.xlsx_utils import sanitizar_workbook_para_excel


DEFAULT_LSS = ROOT / "02-Execucao/05-Comentarios_Gestor/questionario_comentarios_gestor.lss"
DEFAULT_RESULTADO = ROOT / "02-Execucao/03-Execucao_Procedimentos/02-Resultados_Auditoria/resultado_auditoria.json"
DEFAULT_MAPA = ROOT / "02-Execucao/03-Execucao_Procedimentos/01-Insumos/mapa-verificacao-achados.xlsx"
DEFAULT_AJUSTES = ROOT / "02-Execucao/01-Questionario/02-Ajustes_Respostas/ajustes_respostas_questionario_pos_avaliacao_evidencias.xlsx"
DEFAULT_RESPOSTAS_BASE = ROOT / "02-Execucao/01-Questionario/03-Respostas_Processadas/20260621-respostas-questionario-02-pos-avaliacao-evidencias.xlsx"
DEFAULT_RESPOSTAS_ORIGINAIS = ROOT / "02-Execucao/01-Questionario/01-Coleta_LimeSurvey/20260621-respostas-questionario-bruto.xlsx"
DEFAULT_PAINEL_EVIDENCIAS = ROOT / "02-Execucao/03-Execucao_Procedimentos/99-Avaliacao_Evidencias/painel-avaliacao-evidencias.xlsx"
DEFAULT_REVISOES_RESPOSTAS = ROOT / "02-Execucao/05-Comentarios_Gestor/02-Avaliacao_Comentarios_Gestor/revisoes_respostas.yml"
DEFAULT_QUESTIONARIO = ROOT / "01-Planejamento/02-Metodologia_iGovTI/igovti_2026.md"
DEFAULT_PROMPTS = ROOT / "scripts/avaliacao_evidencias/prompts/igovti_2026_achados_binario_v1"
DEFAULT_CATALOG = ROOT / "scripts/avaliacao_evidencias/prompt_catalogs/igovti_2026_achados_binario_v1.yml"
DEFAULT_CATALOG_COMENTARIOS = ROOT / "scripts/avaliacao_evidencias/prompt_catalogs/igovti_2026_comentarios_gestor_atual_v4.yml"

ENV_PROVIDER_KEYS = {
    "gemini": "GEMINI_API_KEY",
    "openrouter": "OPENROUTER_API_KEY",
    "opencodego": "OPENCODEGO_API_KEY",
    "openai": "OPENAI_API_KEY",
}

ESTADOS_TEMPORAIS = {
    "mantida",
    "afastada_na_data_base",
    "corrigida_posteriormente",
}

RESPOSTAS_SECAO1_AVALIAVEIS = {
    "discorda da sinalizacao de inadequacao",
    "concorda e ja atendeu as propostas de encaminhamento",
}

# Conversões pdf2md/docx2html são intensivas e algumas bibliotecas auxiliares
# mantêm estado global. Serializa apenas a preparação; as chamadas aos modelos
# continuam paralelas depois que os anexos de cada caso estão prontos.
_EVIDENCE_CONVERSION_LOCK = threading.Lock()


@dataclass(frozen=True)
class SituacaoDef:
    codigo: str
    achado: int
    situacao: str


@dataclass(frozen=True)
class ItemAnterior:
    auditado: str
    codigo: str
    base: str
    resposta_afirmada: str
    resultado: str
    justificativa: str
    texto: str


def texto(valor: Any) -> str:
    if valor is None or (isinstance(valor, float) and math.isnan(valor)):
        return ""
    return str(valor).strip()


def chave(valor: Any) -> str:
    bruto = unicodedata.normalize("NFKD", texto(valor).casefold())
    bruto = "".join(c for c in bruto if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", bruto).strip()


def base_item(codigo: str) -> str:
    match = re.match(r"^(q\d{4})", codigo, flags=re.I)
    return match.group(1).lower() if match else codigo.lower()


def ler_xlsx(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        # A primeira aba é a fonte publicada; a aba ativa pode ser uma aba
        # auxiliar/origin deixada por ferramentas de revisão do XLSX.
        ws = wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        headers = [texto(v) for v in next(rows)]
        dados = [dict(zip(headers, row)) for row in rows if any(v is not None for v in row)]
        return headers, dados
    finally:
        wb.close()


def consolidar_submissoes(rows: Iterable[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    validas: list[dict[str, Any]] = []
    excluidas: list[dict[str, Any]] = []
    for row in rows:
        if not row.get("submitdate"):
            excluidas.append({**row, "motivo_exclusao": "Submissão não concluída"})
        elif chave(row.get("firstname")) == "teste zip" or texto(row.get("token")) == "123":
            excluidas.append({**row, "motivo_exclusao": "Registro de teste"})
        else:
            validas.append(row)

    grupos: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in validas:
        identificador = texto(row.get("token")) or chave(row.get("firstname"))
        grupos[identificador].append(row)

    finais: list[dict[str, Any]] = []
    for grupo in grupos.values():
        ordenado = sorted(grupo, key=lambda r: texto(r.get("submitdate")))
        finais.append(ordenado[-1])
        for antiga in ordenado[:-1]:
            excluidas.append({**antiga, "motivo_exclusao": "Substituída por reenvio posterior do mesmo token"})
    finais.sort(key=lambda r: chave(r.get("firstname")))
    return finais, excluidas


def parse_uploads(valor: Any) -> list[dict[str, Any]]:
    if valor in (None, ""):
        return []
    parsed = json.loads(valor) if isinstance(valor, str) else valor
    if not isinstance(parsed, list):
        raise ValueError("metadados de upload devem ser uma lista")
    uploads: list[dict[str, Any]] = []
    for item in parsed:
        if not isinstance(item, dict) or not texto(item.get("name")):
            raise ValueError("metadado de upload inválido ou sem atributo name")
        uploads.append(item)
    return uploads


def indices_upload(headers: list[str]) -> dict[str, int]:
    colunas = [h for h in headers if "evi" in h.casefold() and not h.endswith("[filecount]")]
    return {coluna: idx for idx, coluna in enumerate(colunas, start=1)}


def limpar_html(valor: str) -> str:
    sem_tags = re.sub(r"<[^>]+>", " ", html.unescape(valor or ""))
    return re.sub(r"\s+", " ", sem_tags).strip()


def carregar_situacoes_lss(path: Path) -> dict[str, SituacaoDef]:
    root = ElementTree.parse(path).getroot()
    resultado: dict[str, SituacaoDef] = {}
    for row in root.findall("./questions/rows/row"):
        titulo = texto(row.findtext("title"))
        match = re.fullmatch(r"A(\d+)G(\d+)Conc", titulo, flags=re.I)
        if not match:
            continue
        pergunta_html = row.findtext("question") or ""
        strong = re.search(r"<strong>(.*?)</strong>", pergunta_html, flags=re.I | re.S)
        situacao = limpar_html(strong.group(1) if strong else pergunta_html)
        codigo = titulo[:-4]
        resultado[codigo] = SituacaoDef(codigo=codigo, achado=int(match.group(1)), situacao=situacao.rstrip("."))
    if not resultado:
        raise ValueError(f"nenhuma situação A#G# encontrada no LSS: {path}")
    return resultado


def carregar_acoes_mapa(path: Path) -> dict[str, dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb["Ações de Verificação"]
        headers = [texto(v) for v in next(ws.iter_rows(min_row=3, max_row=3, values_only=True))]
        return {
            texto(row[0]): dict(zip(headers, row))
            for row in ws.iter_rows(min_row=4, values_only=True)
            if row and texto(row[0])
        }
    finally:
        wb.close()


def carregar_contextos_situacoes(resultado_path: Path, mapa_path: Path) -> dict[tuple[str, int, str], dict[str, Any]]:
    dados = json.loads(resultado_path.read_text(encoding="utf-8"))
    acoes = carregar_acoes_mapa(mapa_path)
    contextos: dict[tuple[str, int, str], dict[str, Any]] = {}
    for sigla, auditado in dados.items():
        for procedimento in auditado.get("procedimentos_executados", []):
            achado = procedimento.get("achado") or {}
            numero = int(achado.get("numero") or procedimento.get("numero_achado") or 0)
            for situacao in achado.get("situacoes_encontradas", []) or []:
                motivos = (achado.get("motivos_situacoes") or {}).get(situacao, []) or []
                motivos_enriquecidos = []
                for motivo in motivos:
                    refs = motivo.get("refs_acoes") or motivo.get("acoes_referencia") or []
                    if isinstance(refs, str):
                        refs = re.findall(r"AV\d+", refs)
                    motivos_enriquecidos.append(
                        {
                            **motivo,
                            "acoes": [acoes[ref] for ref in refs if ref in acoes],
                        }
                    )
                key = (texto(sigla).upper(), numero, chave(situacao.rstrip(".")))
                if key in contextos:
                    raise ValueError(f"situação duplicada no resultado de auditoria: {key}")
                contextos[key] = {
                    "achado": numero,
                    "nome_achado": achado.get("nome", ""),
                    "situacao": situacao,
                    "motivos": motivos_enriquecidos,
                    "evidencias": achado.get("evidencias_detalhadas") or achado.get("evidencias") or [],
                    "encaminhamentos": achado.get("encaminhamentos") or [],
                }
    return contextos


def resultado_anterior(row: dict[str, Any]) -> tuple[str, str]:
    revisao = texto(row.get("Avaliação do auditor revisor"))
    if chave(revisao) not in {"", "nan", "none", "sem parecer", "sem_parecer", "nao revisado"}:
        return revisao, texto(row.get("Justificativa do auditor revisor"))
    return texto(row.get("Resultado da avaliação do juiz")), texto(row.get("Justificativa do juiz"))


def carregar_itens_nao_conformes(path: Path, questionario: Path) -> dict[tuple[str, str], list[ItemAnterior]]:
    contexto = carregar_contexto_questionario(questionario)
    _, rows = ler_xlsx(path)
    itens: dict[tuple[str, str], list[ItemAnterior]] = defaultdict(list)
    for row in rows:
        auditado = texto(row.get("Auditado")).upper()
        codigo = texto(row.get("Código do item avaliado"))
        resultado, justificativa = resultado_anterior(row)
        if not auditado or not codigo or chave(resultado) != "nao conforme":
            continue
        base = base_item(codigo)
        item_texto = codigo
        questao = contexto.questoes.get(base)
        match = re.search(r"\[([^]]+)\]$", codigo)
        if questao and match:
            item_texto = questao.itens.get(match.group(1), codigo)
        elif questao:
            item_texto = questao.texto
        itens[(auditado, base)].append(
            ItemAnterior(
                auditado=auditado,
                codigo=codigo,
                base=base,
                resposta_afirmada=texto(row.get("Resposta afirmada")),
                resultado=resultado,
                justificativa=justificativa,
                texto=item_texto,
            )
        )
    return itens


def carregar_indice_consolidado(path: Path) -> dict[tuple[str, str], list[dict[str, Any]]]:
    indice: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for linha in path.read_text(encoding="utf-8").splitlines():
        if not linha.strip():
            continue
        registro = json.loads(linha)
        if registro.get("status") != "completed":
            continue
        auditado = texto(registro.get("auditado")).upper()
        result = registro.get("result") if isinstance(registro.get("result"), dict) else {}
        for conclusao in result.get("conclusoes") or []:
            codigo = texto(conclusao.get("item_codigo"))
            if codigo:
                indice[(auditado, codigo)].append(
                    {
                        "evidencia": registro.get("evidencia", ""),
                        "estado": conclusao.get("estado", ""),
                        "justificativa": conclusao.get("justificativa", ""),
                        "lacunas": conclusao.get("lacunas") or [],
                        "arquivos_referenciados": conclusao.get("arquivos_referenciados") or [],
                        "trechos_ou_elementos": conclusao.get("trechos_ou_elementos") or [],
                        "paginas_ou_localizacao": conclusao.get("paginas_ou_localizacao") or [],
                    }
                )
    return indice


def carregar_rotas_catalogo(path: Path) -> dict[str, list[dict[str, Any]]]:
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    rotas: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for entry in data.get("prompts", []) or []:
        coluna = texto(entry.get("coluna_evidencia"))
        match = re.match(r"^(q\d{4})", coluna, flags=re.I)
        if not match:
            continue
        entry = dict(entry)
        entry["itens_avaliaveis"] = [texto(v) for v in entry.get("itens_avaliaveis", []) or []]
        rotas[match.group(1).lower()].append(entry)
    return rotas


def resolver_caminhos_upload(
    auditado: str,
    root: Path,
    uploads: list[dict[str, Any]],
    *,
    resposta_id: Any,
    evidence_index: int,
) -> list[Path]:
    caminhos: list[Path] = []
    for upload in uploads:
        resolucao = resolver_evidencia(
            auditado,
            root,
            upload,
            resposta_id=resposta_id,
            evidence_index=evidence_index,
        )
        if resolucao.erro or not resolucao.caminho:
            raise ValueError(resolucao.erro or f"arquivo não encontrado: {upload.get('name')}")
        caminhos.append(resolucao.caminho.resolve())
    return caminhos


def carregar_prompts_comentarios(path: Path) -> dict[str, str]:
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    return {
        "situacoes": texto(data.get("prompt_situacoes")),
        "reavaliacao": texto(data.get("prompt_reavaliacao_prefixo")),
        "juiz_situacoes": texto(data.get("prompt_juiz_situacoes")),
        "juiz_reavaliacao": texto(data.get("prompt_juiz_reavaliacao")),
        "version": texto(data.get("version")),
    }


def api_key(provider: str) -> str:
    return os.environ.get(ENV_PROVIDER_KEYS.get(provider, ""), "")


def case_id(secao: str, auditado: str, codigo: str, resposta_id: Any, extra: str = "") -> str:
    bruto = json.dumps([secao, auditado, codigo, resposta_id, extra], ensure_ascii=False)
    return hashlib.sha256(bruto.encode("utf-8")).hexdigest()


def _resultado_fake_temporal(result: dict[str, Any], motivos: list[dict[str, Any]]) -> dict[str, Any]:
    for conclusao in result.get("conclusoes") or []:
        conclusao["estado_temporal"] = "mantida"
        conclusao["conclusoes_motivos"] = [
            {
                "id_motivo": texto(m.get("id")),
                "estado_motivo": "mantido",
                "justificativa": "Provider fake não emite conclusão substantiva.",
            }
            for m in motivos
        ]
        conclusao["providencias_informadas"] = []
        conclusao["comentarios_encaminhamento"] = ""
        conclusao["consequencias_praticas"] = []
        conclusao["alternativas_propostas"] = []
    return result


def _registrar_anexo_nao_processado(
    avisos: list[dict[str, str]],
    caminho: Path,
    erro: str,
) -> None:
    """Registra anexo ilegível sem atribuir a ele valor probatório."""
    avisos.append(
        {
            "arquivo": caminho.name,
            "caminho": str(caminho),
            "erro": erro,
            "tratamento": (
                "O anexo não foi enviado ao modelo nem considerado como evidência. "
                "A avaliação deve usar somente os demais elementos processáveis."
            ),
        }
    )


def _adicionar_avisos_anexos_ao_pacote(
    pacote: dict[str, Any],
    avisos: list[dict[str, str]],
) -> None:
    if not avisos:
        return
    documento = _documento(
        "avisos_processamento_anexos.json",
        "aviso_tecnico_processamento",
        {
            "orientacao": (
                "Os anexos abaixo não puderam ser processados e não constituem prova. "
                "Não presuma seu conteúdo nem extraia conclusão favorável de sua mera existência."
            ),
            "anexos": avisos,
        },
    )
    pacote["documentos"].append(documento)
    pacote["inventario"].append(documento["nome"])


def executar_caso(
    caso: dict[str, Any],
    *,
    provider: str,
    model: str,
    reasoning: str,
    pdf2md: bool,
    docx2html: bool,
    pdf_detail: str = "auto",
) -> dict[str, Any]:
    started = dt.datetime.now(dt.timezone.utc)
    tokens_info: dict[str, int] = {}
    avisos_processamento: list[dict[str, str]] = []
    with tempfile.TemporaryDirectory() as tmp:
        pacote = {
            "documentos": list(caso["documentos_contexto"]),
            "inventario": [d.get("nome", "") for d in caso["documentos_contexto"]],
            "erro": "",
            "arquivos_upload": [],
        }
        conversion_context = _EVIDENCE_CONVERSION_LOCK if (pdf2md or docx2html) else nullcontext()
        with conversion_context:
            for caminho in caso["evidence_paths"]:
                preparado, arquivos, erro = preparar_evidencia_para_provider(
                    Path(caminho), tmp, pdf2md=pdf2md, docx2html=docx2html, dpi=150
                )
                if erro:
                    _registrar_anexo_nao_processado(
                        avisos_processamento, Path(caminho), f"erro ao preparar: {erro}"
                    )
                    continue
                bloqueante = erro_tecnico_bloqueante_pacote(preparado, arquivos)
                if bloqueante:
                    _registrar_anexo_nao_processado(
                        avisos_processamento, Path(caminho), bloqueante
                    )
                    continue
                pacote["documentos"].extend(preparado.documentos)
                pacote["inventario"].extend(preparado.inventario)
                pacote["arquivos_upload"].extend(arquivos)

        if not pacote["documentos"] and not pacote["arquivos_upload"]:
            raise ValueError("caso sem comentário, contexto ou evidência processável para avaliação")
        _adicionar_avisos_anexos_ao_pacote(pacote, avisos_processamento)

        tokens_info = estimar_tokens_payload(
            prompt=caso["prompt"], auditado=caso["auditado"], questao_base=caso["codigo"],
            coluna_evidencia=caso["coluna_evidencia"], itens_afirmados=caso["itens"], pacote=pacote,
            provider=provider, response_profile=caso.get("response_profile", "evidence"),
        )
        limite = limite_tokens_provider(provider, model)
        if limite and provider != "gemini" and tokens_info["tokens_total"] > limite:
            result = {"status": "error", "error": f"payload excede limite de tokens: {tokens_info['tokens_total']:,} > {limite:,}"}
        else:
            result = executar_provider(
                provider=provider,
                model=model,
                api_key=api_key(provider),
                prompt=caso["prompt"],
                auditado=caso["auditado"],
                questao_base=caso["codigo"],
                coluna_evidencia=caso["coluna_evidencia"],
                itens_afirmados=caso["itens"],
                pacote=pacote,
                reasoning_effort=reasoning,
                response_profile=caso.get("response_profile", "evidence"),
                pdf_detail=pdf_detail,
            )
        erro_tecnico = "" if avisos_processamento else resultado_indica_erro_tecnico(result)
        if erro_tecnico:
            result = {"status": "error", "error": erro_tecnico}
        if caso["secao"] == "situacoes" and provider == "fake" and result.get("status") == "completed":
            result = _resultado_fake_temporal(result, caso["contexto"].get("motivos", []))
        if result.get("status") == "completed":
            registro_escopo = {
                "secao": caso["secao"],
                "codigo": caso["codigo"],
                "itens": caso["itens"],
                "contexto": caso["contexto"],
            }
            violacoes = validar_resultado_no_escopo(
                registro_escopo,
                result,
                validar_temporal=caso["secao"] == "situacoes",
            )
            if violacoes:
                result = {
                    "status": "error",
                    "error": "resposta fora do escopo lógico do caso: " + formatar_violacoes(violacoes),
                }

    return {
        "identity": hashlib.sha256(
            json.dumps(
                [
                    caso["case_id"],
                    provider,
                    model,
                    caso["prompt_hash"],
                    [hash_arquivo(Path(p)) for p in caso["evidence_paths"]],
                    _hash_json(caso.get("contexto", {})),
                ],
                sort_keys=True,
            ).encode("utf-8")
        ).hexdigest(),
        "case_id": caso["case_id"],
        "secao": caso["secao"],
        "auditado": caso["auditado"],
        "codigo": caso["codigo"],
        "coluna_evidencia": caso["coluna_evidencia"],
        "resposta_id": caso["resposta_id"],
        "provider": provider,
        "model": model,
        "started_at": started.isoformat(),
        "finished_at": dt.datetime.now(dt.timezone.utc).isoformat(),
        "status": result.get("status", "error"),
        "error": result.get("error", ""),
        "result": result,
        "contexto": caso["contexto"],
        "documentos_contexto": caso["documentos_contexto"],
        "evidence_paths": [str(p) for p in caso["evidence_paths"]],
        "evidencias": caso["evidencias"],
        "itens": [asdict(i) if hasattr(i, "__dataclass_fields__") else dict(i) for i in caso["itens"]],
        "prompt_hash": caso["prompt_hash"],
        "payload_tokens": tokens_info,
        "avisos_processamento_evidencias": avisos_processamento,
    }


def gravar_jsonl(path: Path, registros: Iterable[dict[str, Any]]) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    registros = list(registros)
    with path.open("w", encoding="utf-8") as file:
        for registro in registros:
            file.write(json.dumps(registro, ensure_ascii=False, default=str) + "\n")
    return len(registros)

def _documento(nome: str, tipo: str, valor: Any) -> dict[str, Any]:
    conteudo = valor if isinstance(valor, str) else json.dumps(valor, ensure_ascii=False, indent=2, default=str)
    return {"nome": nome, "tipo": tipo, "texto": conteudo}


def extrair_urls_manifestacao(documentos: list[dict[str, Any]]) -> list[str]:
    urls: list[str] = []
    for documento in documentos:
        if texto(documento.get("tipo")) not in {"manifestacao_gestor", "comentario_gestor"}:
            continue
        for encontrada in re.findall(r"https?://[^\s<>\"']+", texto(documento.get("texto")), flags=re.I):
            url = encontrada.rstrip(".,;:!?)]}")
            if url and url not in urls:
                urls.append(url)
    return urls


def capturar_links_manifestacao(
    documentos: list[dict[str, Any]],
    destino: Path,
    *,
    refresh: bool = False,
) -> tuple[list[Path], list[dict[str, Any]], list[str]]:
    manifesto_path = destino / "manifesto-links.json"
    anteriores: list[dict[str, Any]] = []
    if not refresh and manifesto_path.is_file():
        try:
            carregado = json.loads(manifesto_path.read_text(encoding="utf-8"))
            if isinstance(carregado, list):
                anteriores = [item for item in carregado if isinstance(item, dict)]
        except (OSError, ValueError, TypeError):
            anteriores = []

    def captura_reutilizavel(url: str) -> tuple[Path, dict[str, Any]] | None:
        destino_resolvido = destino.resolve()
        for registro in anteriores:
            if registro.get("url_original") != url or registro.get("status") != "capturado":
                continue
            caminho_registrado = Path(texto(registro.get("caminho")))
            candidatos = [caminho_registrado]
            if not caminho_registrado.is_absolute():
                candidatos.append(destino / caminho_registrado.name)
            for candidato in candidatos:
                try:
                    resolvido = candidato.resolve()
                    resolvido.relative_to(destino_resolvido)
                except (OSError, ValueError):
                    continue
                if not resolvido.is_file():
                    continue
                hash_registrado = texto(registro.get("sha256"))
                if not hash_registrado or hash_arquivo(resolvido) != hash_registrado:
                    continue
                reutilizado = dict(registro)
                reutilizado["caminho"] = str(candidato)
                return candidato, reutilizado
        return None

    caminhos: list[Path] = []
    registros: list[dict[str, Any]] = []
    erros: list[str] = []
    for url in extrair_urls_manifestacao(documentos):
        permitido, motivo = _url_download_permitida(url)
        if not permitido:
            registros.append({"url_original": url, "status": "ignorado", "motivo": motivo})
            continue
        reutilizavel = captura_reutilizavel(url)
        if reutilizavel is not None:
            caminho, registro = reutilizavel
            registros.append(registro)
            caminhos.append(caminho)
            continue
        caminho, erro, metadados = baixar_recurso_url(url, destino)
        if erro or caminho is None:
            erros.append(f"{url}: {erro or 'download não materializado'}")
            registros.append({"url_original": url, "status": "erro", "motivo": erro})
            continue
        registro = {
            **metadados,
            "status": "capturado",
            "caminho": str(caminho),
            "sha256": hash_arquivo(caminho),
            "capturado_em": dt.datetime.now(dt.timezone.utc).isoformat(),
        }
        registros.append(registro)
        caminhos.append(caminho)
    if registros:
        destino.mkdir(parents=True, exist_ok=True)
        manifesto_path.write_text(
            json.dumps(registros, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
    return caminhos, registros, erros


def _hash_json(valor: Any) -> str:
    return hashlib.sha256(json.dumps(valor, ensure_ascii=False, sort_keys=True, default=str).encode("utf-8")).hexdigest()


def _selecionar_situacao(resposta: str, comentario: str, justificativa: str, uploads: list[dict[str, Any]]) -> bool:
    del comentario, justificativa, uploads
    return chave(resposta) in RESPOSTAS_SECAO1_AVALIAVEIS


def _itens_reavaliacao(itens: list[ItemAnterior], rota: dict[str, Any], prompt: Any) -> list[ItemAfirmado]:
    afirmados = [ItemAfirmado(codigo=i.codigo, texto=i.texto, afirmacao=i.resposta_afirmada) for i in itens]
    afirmados = filtrar_itens_por_prompt(afirmados, prompt)
    return preparar_itens_para_prompt(afirmados, prompt)


def filtrar_itens_saneados_secao1(
    itens_anteriores: dict[tuple[str, str], list[ItemAnterior]],
    saneados_secao1: dict[tuple[str, str], dict[str, Any]],
) -> tuple[dict[tuple[str, str], list[ItemAnterior]], list[dict[str, Any]]]:
    restantes_por_base: dict[tuple[str, str], list[ItemAnterior]] = {}
    excluidos: list[dict[str, Any]] = []
    for chave_itens, itens_lista in itens_anteriores.items():
        restantes: list[ItemAnterior] = []
        for item in itens_lista:
            saneamento = saneados_secao1.get((item.auditado, item.codigo))
            if saneamento:
                excluidos.append(
                    {
                        **saneamento,
                        "Questão-base": item.base,
                        "Motivo da exclusão": "Item saneado por parecer consolidado da seção 1.",
                    }
                )
            else:
                restantes.append(item)
        if restantes:
            restantes_por_base[chave_itens] = restantes
    return restantes_por_base, excluidos


def preparar_casos_comentarios(
    *,
    respostas: Path,
    evidencias_root: Path,
    secao: str,
    lss: Path = DEFAULT_LSS,
    resultado_auditoria: Path = DEFAULT_RESULTADO,
    mapa: Path = DEFAULT_MAPA,
    ajustes: Path = DEFAULT_AJUSTES,
    questionario: Path = DEFAULT_QUESTIONARIO,
    prompts_dir: Path = DEFAULT_PROMPTS,
    catalog: Path = DEFAULT_CATALOG,
    catalog_comentarios: Path = DEFAULT_CATALOG_COMENTARIOS,
    auditados: set[str] | None = None,
    saneados_secao1: dict[tuple[str, str], dict[str, Any]] | None = None,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    """Prepara casos da seção solicitada e resolve todos os anexos de forma bloqueante."""
    headers, rows = ler_xlsx(respostas)
    validas, excluidas = consolidar_submissoes(rows)
    indices = indices_upload(headers)
    prompts_comentarios = carregar_prompts_comentarios(catalog_comentarios)
    casos: list[dict[str, Any]] = []
    deterministicos: list[dict[str, Any]] = []
    fora_escopo_secao1: list[dict[str, Any]] = []
    erros: list[str] = []

    if secao == "1":
        situacoes = carregar_situacoes_lss(lss)
        contextos = carregar_contextos_situacoes(resultado_auditoria, mapa)
        for row in validas:
            auditado = texto(row.get("firstname")).upper()
            if auditados and auditado not in auditados:
                continue
            for codigo, definicao in situacoes.items():
                resposta = texto(row.get(f"{codigo}Conc"))
                if not resposta:
                    continue
                comentario = texto(row.get(f"{codigo}Com"))
                justificativa = texto(row.get(f"{codigo}Jus"))
                try:
                    uploads = parse_uploads(row.get(f"{codigo}Evi"))
                except (ValueError, json.JSONDecodeError) as exc:
                    erros.append(f"{auditado}/{codigo}: {exc}")
                    continue
                context_key = (auditado, definicao.achado, chave(definicao.situacao.rstrip(".")))
                contexto = contextos.get(context_key)
                if not contexto:
                    erros.append(f"{auditado}/{codigo}: contexto da situação não encontrado")
                    continue
                if not contexto.get("motivos"):
                    erros.append(f"{auditado}/{codigo}: situação sem motivos ativos")
                    continue
                cid = case_id("situacoes", auditado, codigo, row.get("id"))
                if not _selecionar_situacao(resposta, comentario, justificativa, uploads):
                    fora_escopo_secao1.append({
                        "case_id": cid,
                        "auditado": auditado,
                        "codigo": codigo,
                        "resposta": resposta,
                        "motivo_exclusao": (
                            "A alternativa não contesta a situação nem declara atendimento concluído."
                        ),
                    })
                    continue
                caminhos: list[Path] = []
                if uploads:
                    try:
                        caminhos = resolver_caminhos_upload(
                            auditado,
                            evidencias_root,
                            uploads,
                            resposta_id=row.get("id"),
                            evidence_index=indices[f"{codigo}Evi"],
                        )
                    except (ValueError, KeyError) as exc:
                        erros.append(f"{auditado}/{codigo}: {exc}")
                        continue
                documentos = [
                    _documento("manifestacao_gestor.json", "manifestacao_gestor", {
                        "resposta": resposta,
                        "comentario": comentario,
                        "justificativa": justificativa,
                    }),
                    _documento("contexto_situacao.json", "contexto_auditoria", contexto),
                ]
                itens = [ItemAfirmado(codigo=codigo, texto=definicao.situacao, afirmacao=resposta)]
                casos.append(
                    {
                        "case_id": cid,
                        "secao": "situacoes",
                        "auditado": auditado,
                        "codigo": codigo,
                        "coluna_evidencia": f"{codigo}Evi",
                        "resposta_id": row.get("id"),
                        "prompt": prompts_comentarios["situacoes"],
                        "prompt_hash": _hash_json(prompts_comentarios["situacoes"]),
                        "prompt_version": prompts_comentarios["version"],
                        "response_profile": "manager_comments_temporal",
                        "itens": itens,
                        "contexto": contexto,
                        "documentos_contexto": documentos,
                        "evidence_paths": caminhos,
                        "evidencias": uploads,
                    }
                )
    elif secao == "2":
        itens_anteriores = carregar_itens_nao_conformes(ajustes, questionario)
        itens_anteriores, itens_excluidos = filtrar_itens_saneados_secao1(
            itens_anteriores,
            saneados_secao1 or {},
        )
        rotas = carregar_rotas_catalogo(catalog)
        for row in validas:
            auditado = texto(row.get("firstname")).upper()
            if auditados and auditado not in auditados:
                continue
            for base in sorted({base for org, base in itens_anteriores if org == auditado}):
                comentario = texto(row.get(f"REV{base.upper()}Com"))
                try:
                    uploads = parse_uploads(row.get(f"REV{base.upper()}Evi"))
                except (ValueError, json.JSONDecodeError) as exc:
                    erros.append(f"{auditado}/{base}: {exc}")
                    continue
                if not comentario and not uploads:
                    continue
                opcoes_rota = rotas.get(base) or []
                if len(opcoes_rota) != 1:
                    erros.append(f"{auditado}/{base}: esperada uma rota de prompt; encontradas {len(opcoes_rota)}")
                    continue
                rota = opcoes_rota[0]
                prompt = resolver_prompt(prompts_dir, texto(rota.get("coluna_evidencia")))
                if prompt.erro:
                    erros.append(f"{auditado}/{base}: {prompt.erro}")
                    continue
                itens = _itens_reavaliacao(itens_anteriores[(auditado, base)], rota, prompt)
                if not itens:
                    erros.append(f"{auditado}/{base}: nenhum item não conforme avaliável pelo prompt")
                    continue
                coluna_upload = f"REV{base.upper()}Evi"
                caminhos: list[Path] = []
                if uploads:
                    try:
                        caminhos = resolver_caminhos_upload(
                            auditado,
                            evidencias_root,
                            uploads,
                            resposta_id=row.get("id"),
                            evidence_index=indices[coluna_upload],
                        )
                    except (ValueError, KeyError) as exc:
                        erros.append(f"{auditado}/{base}: {exc}")
                        continue
                anteriores = itens_anteriores[(auditado, base)]
                documentos = [
                    _documento("comentario_gestor.txt", "comentario_gestor", comentario or "(sem comentário textual)"),
                    _documento("avaliacao_consolidada_anterior.json", "avaliacao_anterior", [asdict(i) for i in anteriores]),
                ]
                prompt_completo = f"{prompts_comentarios['reavaliacao']}\n\n{prompt.conteudo}".strip()
                conjunto_itens = "|".join(sorted(i.codigo for i in itens))
                cid = case_id("reavaliacao", auditado, base, row.get("id"), conjunto_itens)
                casos.append(
                    {
                        "case_id": cid,
                        "secao": "reavaliacao",
                        "auditado": auditado,
                        "codigo": base,
                        "coluna_evidencia": texto(rota.get("coluna_evidencia")),
                        "resposta_id": row.get("id"),
                        "prompt": prompt_completo,
                        "prompt_hash": _hash_json(prompt_completo),
                        "prompt_version": prompts_comentarios["version"],
                        "response_profile": "evidence",
                        "itens": itens,
                        "contexto": {"avaliacao_anterior": [asdict(i) for i in anteriores], "comentario": comentario},
                        "documentos_contexto": documentos,
                        "evidence_paths": caminhos,
                        "evidencias": uploads,
                    }
                )
    else:
        raise ValueError("secao deve ser '1' ou '2'")

    if erros:
        detalhe = "\n".join(f"- {erro}" for erro in erros[:50])
        restante = f"\n- ... e mais {len(erros) - 50} erro(s)" if len(erros) > 50 else ""
        raise ValueError(f"preflight encontrou {len(erros)} erro(s):\n{detalhe}{restante}")
    resumo = {
        "respostas_total": len(rows),
        "respostas_validas": len(validas),
        "respostas_excluidas": len(excluidas),
        "casos_ia": len(casos),
        "casos_deterministicos": len(deterministicos),
        "anexos": sum(len(c["evidence_paths"]) for c in casos),
        "secao": secao,
    }
    if secao == "2":
        resumo["itens_excluidos_secao1"] = itens_excluidos
        resumo["total_itens_excluidos_secao1"] = len(itens_excluidos)
    else:
        resumo["casos_fora_escopo_avaliacao"] = fora_escopo_secao1
        resumo["total_casos_fora_escopo_avaliacao"] = len(fora_escopo_secao1)
    return casos, deterministicos, excluidas, resumo


def registro_deterministico_situacao(
    *, case_id_value: str, auditado: str, codigo: str, resposta_id: Any,
    situacao: str, resposta: str, contexto: dict[str, Any]
) -> dict[str, Any]:
    motivos = [
        {
            "id_motivo": texto(m.get("id")),
            "estado_motivo": "mantido",
            "justificativa": "A organização concordou com a situação e informou não ter adotado medida, sem manifestação ou evidência adicional.",
        }
        for m in contexto.get("motivos", [])
    ]
    conclusao = {
        "item_codigo": codigo,
        "item_texto": situacao,
        "afirmacao_auditado": resposta,
        "estado": "nao_conforme",
        "justificativa": "A situação permanece mantida por concordância expressa sem medida adotada e sem elementos adicionais.",
        "lacunas": [],
        "arquivos_referenciados": [],
        "trechos_ou_elementos": [],
        "paginas_ou_localizacao": [],
        "estado_temporal": "mantida",
        "conclusoes_motivos": motivos,
        "providencias_informadas": [],
        "comentarios_encaminhamento": "",
        "consequencias_praticas": [],
        "alternativas_propostas": [],
    }
    return {
        "identity": _hash_json([case_id_value, "regra_deterministica_v1"]),
        "case_id": case_id_value,
        "secao": "situacoes",
        "auditado": auditado,
        "codigo": codigo,
        "questao": codigo,
        "coluna_evidencia": f"{codigo}Evi",
        "evidencia": "",
        "resposta_id": resposta_id,
        "provider": "regra_auditoria",
        "model": "concordancia_sem_medida_v1",
        "status": "completed",
        "origem_decisao": "regra_deterministica",
        "result": {"status": "completed", "conclusoes": [conclusao], "error": ""},
        "error": "",
    }


def _ler_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.is_file():
        return []
    return [json.loads(linha) for linha in path.read_text(encoding="utf-8").splitlines() if linha.strip()]


def _append_jsonl(path: Path, registro: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as file:
        file.write(json.dumps(registro, ensure_ascii=False, default=str) + "\n")


def _nome_modelo(provider: str, model: str) -> str:
    seguro = re.sub(r"[^a-zA-Z0-9._-]+", "-", f"{provider}_{model}").strip("-")
    return seguro or "modelo"


def _nome_model_key(model_key: str) -> str:
    seguro = re.sub(r"[^a-zA-Z0-9._-]+", "-", model_key).strip("-")
    return seguro or "modelo"


def _identidade_logica_analise(
    *,
    case_id_value: str,
    model_key: str,
    prompt_hash: str,
    reasoning: str,
    pdf2md: bool,
    docx2html: bool,
    contexto: Any,
    evidence_paths: Iterable[str | Path],
    pdf_detail: str = "auto",
) -> str:
    payload = {
        "case_id": case_id_value,
        "model_key": model_key,
        "prompt_hash": prompt_hash,
        "reasoning": reasoning,
        "pdf2md": pdf2md,
        "docx2html": docx2html,
        "contexto_hash": _hash_json(contexto),
        "evidencias_hash": [hash_arquivo(Path(path)) for path in evidence_paths],
    }
    # ``auto`` ja era o comportamento implicito da API antes de o campo ser
    # exposto na configuracao. Omiti-lo preserva checkpoints historicos.
    detail_normalizado = texto(pdf_detail).lower() or "auto"
    if detail_normalizado != "auto":
        payload["pdf_detail"] = detail_normalizado
    return _hash_json(payload)


def _model_key_registro(registro: dict[str, Any], routes: Iterable[dict[str, Any]]) -> str:
    explicit = texto(registro.get("model_key"))
    if explicit:
        return explicit
    pair = (texto(registro.get("provider")), texto(registro.get("model")))
    for route in routes:
        if pair == (texto(route.get("provider")), texto(route.get("model"))):
            return texto(route.get("model_key")) or pair[1]
    return pair[1]


def _normalizar_registro_legado(
    registro: dict[str, Any],
    *,
    model_key: str,
) -> dict[str, Any] | None:
    if registro.get("status") != "completed" or not registro.get("case_id"):
        return None
    logical_identity = texto(registro.get("logical_identity"))
    # Registros gerados na curta versao que incluiu ``pdf_detail=auto`` no
    # hash precisam ser recanonizados para o formato retrocompativel.
    if not logical_identity or "pdf_detail" in registro:
        mode = {part for part in texto(registro.get("evidence_processing_mode")).split("+") if part}
        try:
            logical_identity = _identidade_logica_analise(
                case_id_value=texto(registro.get("case_id")),
                model_key=model_key,
                prompt_hash=texto(registro.get("prompt_hash")),
                reasoning=texto(registro.get("reasoning_effort")),
                pdf2md="pdf2md" in mode,
                docx2html="docx2html" in mode,
                contexto=registro.get("contexto", {}),
                evidence_paths=registro.get("evidence_paths") or [],
                pdf_detail=texto(registro.get("pdf_detail")) or "auto",
            )
        except (OSError, ValueError):
            return None
    normalized = dict(registro)
    source_identity = texto(registro.get("source_identity")) or texto(registro.get("identity"))
    normalized["identity"] = logical_identity
    normalized["logical_identity"] = logical_identity
    normalized["model_key"] = model_key
    if source_identity and source_identity != logical_identity:
        normalized["source_identity"] = source_identity
    return normalized


def _carregar_historico_modelo(
    out_dir: Path,
    *,
    model_key: str,
    routes: Iterable[dict[str, Any]],
) -> dict[str, dict[str, Any]]:
    routes = list(routes)
    canonical = out_dir / f"analyses_model_{_nome_model_key(model_key)}.jsonl"
    paths = {canonical}
    for route in routes:
        if texto(route.get("model_key")) == model_key:
            paths.add(out_dir / f"analyses_{_nome_modelo(texto(route.get('provider')), texto(route.get('model')))}.jsonl")
    records: dict[str, dict[str, Any]] = {}
    for path in sorted(paths):
        for raw in _ler_jsonl(path):
            if _model_key_registro(raw, routes) != model_key:
                continue
            normalized = _normalizar_registro_legado(raw, model_key=model_key)
            if normalized is None:
                continue
            identity = normalized["identity"]
            current = records.get(identity)
            current_valid = _registro_avaliacao_utilizavel(current) if current is not None else False
            normalized_valid = _registro_avaliacao_utilizavel(normalized)
            if (
                current is None
                or (normalized_valid and not current_valid)
                or (
                    normalized_valid == current_valid
                    and texto(normalized.get("finished_at")) >= texto(current.get("finished_at"))
                )
            ):
                records[identity] = normalized
    return records


def materializar_checkpoint_limpo_modelo(
    out_dir: Path,
    *,
    model_key: str,
    routes: Iterable[dict[str, Any]],
    expected_identities: set[str] | None = None,
) -> Path:
    records = _carregar_historico_modelo(out_dir, model_key=model_key, routes=routes)
    if expected_identities is not None:
        records = {identity: record for identity, record in records.items() if identity in expected_identities}
    current_by_case: dict[str, dict[str, Any]] = {}
    for record in records.values():
        case_id_value = texto(record.get("case_id"))
        if not case_id_value:
            continue
        current = current_by_case.get(case_id_value)
        if current is None or texto(record.get("finished_at")) >= texto(current.get("finished_at")):
            current_by_case[case_id_value] = record
    clean = out_dir / f"analyses_clean_model_{_nome_model_key(model_key)}.jsonl"
    gravar_jsonl(
        clean,
        sorted(current_by_case.values(), key=lambda record: (texto(record.get("auditado")), texto(record.get("codigo")))),
    )
    return clean


def avaliar_casos(
    casos: list[dict[str, Any]],
    *,
    provider: str,
    model: str,
    model_key: str,
    out_dir: Path,
    routes: Iterable[dict[str, Any]] = (),
    reasoning: str = "high",
    rpm: int = 0,
    pdf2md: bool = False,
    docx2html: bool = False,
    pdf_detail: str = "auto",
    skip_errors: bool = False,
    quiet: bool = False,
) -> dict[str, Any]:
    """Avalia casos com checkpoint lógico compartilhado entre providers do mesmo modelo."""
    stem = _nome_model_key(model_key)
    checkpoint = out_dir / f"analyses_model_{stem}.jsonl"
    clean = out_dir / f"analyses_clean_model_{stem}.jsonl"
    por_identidade = _carregar_historico_modelo(out_dir, model_key=model_key, routes=routes)
    limiter = RequestsPerMinuteLimiter(rpm)
    processados = concluidos = erros = pulados = 0
    identidades_esperadas: set[str] = set()
    for index, caso in enumerate(casos, start=1):
        # Identidade de contingência permite registrar uma falha de hash como
        # erro de um único caso, sem derrubar a execução inteira do modelo.
        identity = _hash_json({
            "case_id": caso.get("case_id"),
            "model_key": model_key,
            "prompt_hash": caso.get("prompt_hash"),
            "identity_fallback": True,
        })
        try:
            identity = _identidade_logica_analise(
                case_id_value=caso["case_id"],
                model_key=model_key,
                prompt_hash=caso["prompt_hash"],
                reasoning=reasoning,
                pdf2md=pdf2md,
                docx2html=docx2html,
                contexto=caso.get("contexto", {}),
                evidence_paths=caso["evidence_paths"],
                pdf_detail=pdf_detail,
            )
            identidades_esperadas.add(identity)
            anterior = por_identidade.get(identity)
            if anterior and (
                _registro_avaliacao_utilizavel(anterior)
                or (skip_errors and anterior.get("status") == "error")
            ):
                pulados += 1
                log_event(
                    "comment_analysis_skipped",
                    "Caso ignorado por checkpoint.",
                    quiet=quiet,
                    case_id=caso["case_id"],
                    index=index,
                    total=len(casos),
                    provider=provider,
                    model=model,
                )
                continue
            if provider != "fake" and rpm:
                espera = limiter.wait_seconds()
                limiter.wait_and_mark(espera)
            registro = executar_caso(
                caso,
                provider=provider,
                model=model,
                reasoning=reasoning,
                pdf2md=pdf2md,
                docx2html=docx2html,
                pdf_detail=pdf_detail,
            )
            registro["identity"] = identity
            registro["logical_identity"] = identity
            registro["model_key"] = model_key
            registro["questao"] = caso["codigo"]
            registro["evidencia"] = "; ".join(texto(u.get("name")) for u in caso.get("evidencias", []))
            registro["reasoning_effort"] = reasoning
            registro["pdf_detail"] = pdf_detail
            registro["evidence_processing_mode"] = "+".join(x for x, ativo in [("pdf2md", pdf2md), ("docx2html", docx2html)] if ativo)
            registro["prompt_version"] = caso.get("prompt_version", "")
            registro["pacote_contexto_consolidacao"] = {
                "documentos": caso["documentos_contexto"],
                "inventario": [d.get("nome", "") for d in caso["documentos_contexto"]],
                "erro": "",
            }
        except Exception as exc:
            agora = dt.datetime.now(dt.timezone.utc).isoformat()
            registro = {
                "identity": identity,
                "logical_identity": identity,
                "model_key": model_key,
                "case_id": caso["case_id"],
                "secao": caso["secao"],
                "auditado": caso["auditado"],
                "codigo": caso["codigo"],
                "questao": caso["codigo"],
                "coluna_evidencia": caso["coluna_evidencia"],
                "evidencia": "; ".join(texto(u.get("name")) for u in caso.get("evidencias", [])),
                "provider": provider,
                "model": model,
                "started_at": agora,
                "finished_at": agora,
                "status": "error",
                "error": str(exc),
                "result": {"status": "error", "error": str(exc)},
                "prompt_hash": caso.get("prompt_hash", ""),
                "reasoning_effort": reasoning,
                "pdf_detail": pdf_detail,
                "evidence_processing_mode": "+".join(
                    x for x, active in [("pdf2md", pdf2md), ("docx2html", docx2html)] if active
                ),
                "contexto": caso.get("contexto", {}),
                "evidence_paths": [str(path) for path in caso.get("evidence_paths", [])],
            }
        _append_jsonl(checkpoint, registro)
        por_identidade[identity] = registro
        processados += 1
        if registro.get("status") == "completed":
            concluidos += 1
        else:
            erros += 1
        log_event(
            "comment_analysis_recorded", "Avaliação de comentário registrada.", quiet=quiet,
            case_id=caso["case_id"], secao=caso["secao"], auditado=caso["auditado"],
            codigo=caso["codigo"], index=index, total=len(casos), provider=provider,
            model=model, status=registro.get("status"), error=registro.get("error", ""),
        )

    clean = materializar_checkpoint_limpo_modelo(
        out_dir,
        model_key=model_key,
        routes=routes,
        expected_identities=identidades_esperadas,
    )
    vigentes = _ler_jsonl(clean)
    return {
        "provider": provider, "model": model, "model_key": model_key,
        "checkpoint": str(checkpoint), "clean": str(clean),
        "casos": len(casos), "processados": processados, "concluidos": concluidos,
        "erros": erros, "pulados": pulados, "vigentes": len(vigentes),
    }


def gravar_deterministicos(registros: list[dict[str, Any]], out_dir: Path) -> Path:
    path = out_dir / "deterministicos.jsonl"
    gravar_jsonl(path, registros)
    return path


def _lookup_ajustes_evidencias(path: Path) -> dict[tuple[str, str], dict[str, Any]]:
    _, rows = ler_xlsx(path)
    lookup: dict[tuple[str, str], dict[str, Any]] = {}
    for row in rows:
        auditado = texto(row.get("Auditado")).upper()
        codigo = texto(row.get("Código do item avaliado") or row.get("Código do item"))
        if auditado and codigo:
            lookup[(auditado, codigo)] = row
    return lookup


def _lookup_respostas_base(path: Path) -> tuple[dict[str, dict[str, Any]], dict[str, str]]:
    headers, rows = ler_xlsx(path)
    colunas = {h.casefold(): h for h in headers}
    por_auditado = {
        texto(row.get("firstname")).upper(): row
        for row in rows
        if texto(row.get("firstname"))
    }
    return por_auditado, colunas


def _resolver_valor_positivo(
    *,
    auditado: str,
    codigo: str,
    fonte: str,
    ajustes_evidencias: dict[tuple[str, str], dict[str, Any]],
    resposta_original: Any = None,
) -> tuple[str, str]:
    original = texto(resposta_original)
    if codigo.casefold() == "q0101":
        if original and not original.casefold().startswith("f)"):
            return original, "Resposta categórica original preservada conforme a coleta LimeSurvey."
        return "", "A alternativa original f) exige classificação manual em uma das alternativas a) a e)."
    if codigo.casefold() == "q0103[d]":
        return "Sim", "Atribuição de governança, planejamento ou gestão de TIC reconhecida pelo saneamento da situação."
    if codigo.casefold() == "q0103[g]" and original.casefold() == "sim":
        return "Não", "Negativa de atribuições formais revertida após o saneamento da situação."
    anterior = ajustes_evidencias.get((auditado, codigo))
    resposta = texto((anterior or {}).get("Resposta afirmada"))
    if resposta:
        return resposta, "Resposta originalmente afirmada restaurada, sem majoração além da declaração do gestor."
    if re.fullmatch(r"q\d{4}ext\[[^]]+\]", codigo, flags=re.I):
        return "Sim", "Detalhamento ajustado para Sim porque a situação inconforme correspondente foi saneada."
    if re.fullmatch(r"q\d{4}\[[^]]+\]", codigo, flags=re.I) and original.casefold() in {"sim", "não", "nao"}:
        return "Sim", "Subitem binário ajustado para Sim porque a situação inconforme correspondente foi saneada."
    if fonte == "avaliacao_evidencias_ajustes":
        return "", "Resposta afirmada não localizada na planilha de avaliação de evidências."
    if re.fullmatch(r"q\d{4}\[[^]]+\]", codigo, flags=re.I):
        return "", "Subitem sem resposta originalmente afirmada que possa ser restaurada com segurança."
    if re.fullmatch(r"q\d{4}", codigo, flags=re.I):
        return "", "Questão-base categórica sem valor positivo único e seguro."
    return "", "Campo derivado, quantitativo ou não reconhecido sem valor positivo único e seguro."


def mapear_ajustes_secao1(
    *,
    consolidado: Path,
    resultado_auditoria: Path,
    mapa: Path,
    lss: Path,
    ajustes_pos_avaliacao_evidencias: Path,
    respostas_base: Path,
    respostas_originais: Path = DEFAULT_RESPOSTAS_ORIGINAIS,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Converte situações saneadas em propostas seguras de ajuste por item."""
    registros = _ler_jsonl(consolidado)
    ajustes_evidencias = _lookup_ajustes_evidencias(ajustes_pos_avaliacao_evidencias)
    respostas, colunas = _lookup_respostas_base(respostas_base)
    respostas_brutas, colunas_brutas = _lookup_respostas_base(respostas_originais)
    contextos = carregar_contextos_situacoes(resultado_auditoria, mapa)
    situacoes = carregar_situacoes_lss(lss)
    propostas: list[dict[str, Any]] = []
    pendencias: list[dict[str, Any]] = []

    for registro in registros:
        if registro.get("status") != "completed":
            continue
        violacoes = validar_resultado_no_escopo(
            registro,
            registro.get("result"),
            validar_temporal=True,
        )
        if violacoes:
            pendencias.append({
                "Auditado": texto(registro.get("auditado")).upper(),
                "Código do item avaliado": "",
                "Código da situação": texto(registro.get("codigo")),
                "Situação": "",
                "Estado temporal": "",
                "Motivo da pendência": "Parecer rejeitado por violação de escopo: " + formatar_violacoes(violacoes),
                "Case ID": registro.get("case_id", ""),
                "Identidade do parecer": registro.get("identity", ""),
            })
            continue
        conclusoes = (registro.get("result") or {}).get("conclusoes") or []
        for conclusao in conclusoes:
            estado_temporal = texto(conclusao.get("estado_temporal"))
            auditado = texto(registro.get("auditado")).upper()
            codigo_situacao = texto(registro.get("codigo"))
            if texto(conclusao.get("item_codigo")) != codigo_situacao:
                pendencias.append({
                    "Auditado": auditado,
                    "Código do item avaliado": "",
                    "Código da situação": codigo_situacao,
                    "Situação": "",
                    "Estado temporal": estado_temporal,
                    "Motivo da pendência": "Código da conclusão não pertence à situação avaliada.",
                    "Case ID": registro.get("case_id", ""),
                    "Identidade do parecer": registro.get("identity", ""),
                })
                continue
            contexto = registro.get("contexto") if isinstance(registro.get("contexto"), dict) else None
            if not contexto:
                definicao = situacoes.get(codigo_situacao)
                if definicao:
                    contexto = contextos.get((auditado, definicao.achado, chave(definicao.situacao)))
            contexto = contexto or {}
            motivos = contexto.get("motivos") or []
            if not motivos:
                pendencias.append({
                    "Auditado": auditado, "Código do item avaliado": "", "Código da situação": codigo_situacao,
                    "Situação": contexto.get("situacao") or conclusao.get("item_texto", ""),
                    "Estado temporal": estado_temporal, "Motivo da pendência": "Situação sem motivos ativos rastreáveis.",
                    "Case ID": registro.get("case_id", ""), "Identidade do parecer": registro.get("identity", ""),
                })
                continue
            conclusoes_motivos = {
                texto(motivo.get("id_motivo")): motivo
                for motivo in conclusao.get("conclusoes_motivos") or []
                if isinstance(motivo, dict) and texto(motivo.get("id_motivo"))
            }
            for motivo in motivos:
                id_motivo = texto(motivo.get("id"))
                conclusao_motivo = conclusoes_motivos.get(id_motivo)
                if not conclusao_motivo or texto(conclusao_motivo.get("estado_motivo")) != "afastado":
                    continue
                for acao in motivo.get("acoes") or []:
                    codigo_item = texto(acao.get("informacao_requerida"))
                    fonte = texto(acao.get("id_fonte_informacao"))
                    base = {
                        "Auditado": auditado,
                        "Código do item avaliado": codigo_item,
                        "Questão-base": base_item(codigo_item),
                        "Achado": contexto.get("achado", ""),
                        "Código da situação": codigo_situacao,
                        "Situação": contexto.get("situacao") or conclusao.get("item_texto", ""),
                        "Estado temporal": estado_temporal,
                        "Estado do motivo": "afastado",
                        "IDs dos motivos": texto(motivo.get("id")),
                        "IDs das ações": texto(acao.get("id")),
                        "Fonte da ação": fonte,
                        "Resultado da avaliação do juiz": conclusao.get("estado", ""),
                        "Justificativa do juiz": conclusao.get("justificativa", ""),
                        "Justificativa do motivo": conclusao_motivo.get("justificativa", ""),
                        "Case ID": registro.get("case_id", ""),
                        "Identidade do parecer": registro.get("identity", ""),
                        "Data do parecer": registro.get("finished_at", ""),
                        "Origem": "secao_1_situacao",
                    }
                    if codigo_item == "total_SI" and auditado == "SÃO JOÃO DE MERITI":
                        codigo_fonte = "q0105[SI_comissionados]"
                        coluna_fonte = colunas.get(codigo_fonte.casefold())
                        resposta_auditado = respostas.get(auditado)
                        if coluna_fonte and resposta_auditado is not None:
                            propostas.append({
                                **base,
                                "Código do item avaliado": codigo_fonte,
                                "Questão-base": base_item(codigo_fonte),
                                "Resposta anterior": texto(resposta_auditado.get(coluna_fonte)) or "0",
                                "Resposta afirmada": 5,
                                "Resposta ajustada": 5,
                                "Origem do valor positivo": (
                                    "Decomposição de total_SI: 1 Agente Estratégico e 4 Assessores Executivos "
                                    "de Segurança da Informação, cargos comissionados criados pela LC municipal 233/2026."
                                ),
                                "Justificativa": (
                                    f"Ajuste da fonte de total_SI após comentários do gestor: situação {estado_temporal}. "
                                    f"{texto(conclusao.get('justificativa'))}"
                                ),
                                "observacao": texto(conclusao.get("justificativa")),
                                "Avaliação do auditor revisor": "",
                                "Justificativa do auditor revisor": "",
                            })
                        else:
                            pendencias.append({**base, "Motivo da pendência": "Fonte q0105[SI_comissionados] não localizada."})
                        continue
                    resposta_original = None
                    coluna_original = colunas_brutas.get(codigo_item.casefold())
                    if coluna_original and respostas_brutas.get(auditado) is not None:
                        resposta_original = respostas_brutas[auditado].get(coluna_original)
                    valor, origem_valor = _resolver_valor_positivo(
                        auditado=auditado,
                        codigo=codigo_item,
                        fonte=fonte,
                        ajustes_evidencias=ajustes_evidencias,
                        resposta_original=resposta_original,
                    )
                    coluna_real = colunas.get(codigo_item.casefold())
                    resposta_auditado = respostas.get(auditado)
                    if not codigo_item or not coluna_real or resposta_auditado is None:
                        pendencias.append({
                            **base,
                            "Motivo da pendência": (
                                "Item não localizado na base de respostas."
                                if codigo_item else "Ação sem informação requerida."
                            ),
                        })
                        continue
                    base["Resposta anterior"] = texto(resposta_auditado.get(coluna_real)) or "Vazio"
                    if not valor:
                        pendencias.append({**base, "Motivo da pendência": origem_valor})
                        continue
                    propostas.append({
                        **base,
                        "Resposta afirmada": valor,
                        "Resposta ajustada": valor,
                        "Origem do valor positivo": origem_valor,
                        "Justificativa": (
                            f"Ajuste pós-comentários do gestor: situação {estado_temporal}. "
                            f"{texto(conclusao.get('justificativa'))}"
                        ),
                        "observacao": texto(conclusao.get("justificativa")),
                        "Avaliação do auditor revisor": "",
                        "Justificativa do auditor revisor": "",
                    })
                    if codigo_item.casefold() == "q0103[g]" and texto(resposta_original).casefold() == "sim":
                        codigo_companheiro = "q0103[D]"
                        coluna_companheira = colunas.get(codigo_companheiro.casefold())
                        if coluna_companheira and resposta_auditado is not None:
                            propostas.append({
                                **base,
                                "Código do item avaliado": codigo_companheiro,
                                "Questão-base": base_item(codigo_companheiro),
                                "Resposta anterior": texto(resposta_auditado.get(coluna_companheira)) or "Vazio",
                                "Resposta afirmada": "Sim",
                                "Resposta ajustada": "Sim",
                                "Origem do valor positivo": (
                                    "Ajuste complementar: a reversão de q0103[G] exige reconhecer em q0103[D] "
                                    "as atribuições formais de governança, planejamento ou gestão de TIC."
                                ),
                                "Justificativa": (
                                    f"Ajuste complementar pós-comentários do gestor: situação {estado_temporal}. "
                                    f"{texto(conclusao.get('justificativa'))}"
                                ),
                                "observacao": texto(conclusao.get("justificativa")),
                                "Avaliação do auditor revisor": "",
                                "Justificativa do auditor revisor": "",
                            })
    return propostas, pendencias


def mapear_ajustes_secao2(
    *,
    consolidado: Path,
    ajustes_pos_avaliacao_evidencias: Path,
    questionario: Path,
    saneados_secao1: set[tuple[str, str]],
) -> list[dict[str, Any]]:
    itens = carregar_itens_nao_conformes(ajustes_pos_avaliacao_evidencias, questionario)
    por_item = {(item.auditado, item.codigo): item for grupo in itens.values() for item in grupo}
    propostas: list[dict[str, Any]] = []
    for registro in _ler_jsonl(consolidado):
        if registro.get("status") != "completed":
            continue
        if validar_resultado_no_escopo(registro, registro.get("result")):
            continue
        auditado = texto(registro.get("auditado")).upper()
        for conclusao in (registro.get("result") or {}).get("conclusoes") or []:
            if texto(conclusao.get("estado")) != "conforme":
                continue
            codigo = texto(conclusao.get("item_codigo"))
            if (auditado, codigo) in saneados_secao1:
                continue
            anterior = por_item.get((auditado, codigo))
            if not anterior:
                continue
            valor = anterior.resposta_afirmada or "Vazio"
            propostas.append({
                "Auditado": auditado, "Código do item avaliado": codigo,
                "Questão-base": anterior.base, "Resposta afirmada": valor, "Resposta ajustada": valor,
                "Resposta anterior": "", "Achado": "", "Código da situação": "", "Situação": "",
                "Estado temporal": "", "IDs dos motivos": "", "IDs das ações": "",
                "Fonte da ação": "avaliacao_evidencias_ajustes", "Origem": "secao_2_reavaliacao",
                "Origem do valor positivo": "Resposta afirmada restaurada após reavaliação da nova evidência.",
                "Resultado da avaliação do juiz": conclusao.get("estado", ""),
                "Justificativa do juiz": conclusao.get("justificativa", ""),
                "Avaliação do auditor revisor": "", "Justificativa do auditor revisor": "",
                "Justificativa": "Retificação pós-comentários do gestor: nova evidência considerada suficiente.",
                "observacao": conclusao.get("justificativa", ""),
                "Case ID": registro.get("case_id", ""), "Identidade do parecer": registro.get("identity", ""),
                "Data do parecer": registro.get("finished_at", ""),
            })
    return propostas


def carregar_revisoes_respostas(
    *,
    path: Path,
    consolidado_secao1: Path,
    respostas_base: Path,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Materializa decisões técnicas humanas explicitamente aprovadas e rastreáveis."""
    if not path.is_file():
        return [], []
    payload = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    revisoes = payload.get("revisoes")
    if not isinstance(revisoes, list):
        raise ValueError(f"revisões de respostas sem lista 'revisoes': {path}")
    respostas, colunas = _lookup_respostas_base(respostas_base)
    pareceres = {
        (texto(item.get("auditado")).upper(), texto(item.get("codigo")), texto(item.get("case_id"))): item
        for item in _ler_jsonl(consolidado_secao1)
        if item.get("status") == "completed"
    }
    propostas: list[dict[str, Any]] = []
    pendencias: list[dict[str, Any]] = []
    for revisao in revisoes:
        if not isinstance(revisao, dict) or revisao.get("aprovada") is not True:
            continue
        auditado = texto(revisao.get("auditado")).upper()
        codigo_situacao = texto(revisao.get("codigo_situacao"))
        cid = texto(revisao.get("case_id"))
        parecer = pareceres.get((auditado, codigo_situacao, cid))
        ajustes = revisao.get("ajustes") if isinstance(revisao.get("ajustes"), list) else []
        if not parecer:
            pendencias.append({
                "Auditado": auditado, "Código do item avaliado": "",
                "Código da situação": codigo_situacao, "Case ID": cid,
                "Motivo da pendência": "Revisão humana não corresponde a parecer consolidado vigente.",
            })
            continue
        identidade_revisada = texto(revisao.get("identidade_parecer"))
        if not identidade_revisada or identidade_revisada != texto(parecer.get("identity")):
            pendencias.append({
                "Auditado": auditado, "Código do item avaliado": "",
                "Código da situação": codigo_situacao, "Case ID": cid,
                "Motivo da pendência": (
                    "Revisão humana ainda não revalidada para a identidade do novo parecer consolidado."
                ),
            })
            continue
        for ajuste in ajustes:
            if not isinstance(ajuste, dict):
                continue
            codigo = texto(ajuste.get("codigo_item"))
            valor = texto(ajuste.get("resposta_ajustada"))
            coluna = colunas.get(codigo.casefold())
            linha = respostas.get(auditado)
            if not codigo or not valor or not coluna or linha is None:
                pendencias.append({
                    "Auditado": auditado, "Código do item avaliado": codigo,
                    "Código da situação": codigo_situacao, "Case ID": cid,
                    "Motivo da pendência": "Revisão humana sem item, valor ou coluna válida na base de respostas.",
                })
                continue
            fundamento = texto(ajuste.get("fundamento"))
            referencias = ajuste.get("referencias")
            if not isinstance(referencias, list):
                referencias = revisao.get("referencias") if isinstance(revisao.get("referencias"), list) else []
            conclusoes_parecer = (parecer.get("result") or {}).get("conclusoes") or []
            justificativa_parecer = texto(conclusoes_parecer[0].get("justificativa")) if conclusoes_parecer else ""
            propostas.append({
                "Auditado": auditado,
                "Código do item avaliado": codigo,
                "Questão-base": base_item(codigo),
                "Resposta anterior": texto(linha.get(coluna)) or "Vazio",
                "Resposta afirmada": valor,
                "Resposta ajustada": valor,
                "Achado": revisao.get("achado", ""),
                "Código da situação": codigo_situacao,
                "Situação": texto(revisao.get("situacao")),
                "Estado temporal": "mantida com fundamento parcialmente afastado",
                "Estado do motivo": "decisão técnica revisada",
                "IDs dos motivos": texto(ajuste.get("id_motivo")),
                "IDs das ações": "",
                "Fonte da ação": "revisao_tecnica_documentada",
                "Origem": "revisao_tecnica_documentada",
                "Origem do valor positivo": fundamento,
                "Resultado da avaliação do juiz": "revisão humana aprovada",
                "Justificativa do juiz": justificativa_parecer,
                "Justificativa": fundamento,
                "observacao": fundamento,
                "Referências documentais": json.dumps(referencias, ensure_ascii=False),
                "Case ID": cid,
                "Identidade do parecer": parecer.get("identity", ""),
                "Data do parecer": parecer.get("finished_at", ""),
                "Avaliação do auditor revisor": "Aprovada",
                "Justificativa do auditor revisor": fundamento,
            })
    return propostas, pendencias


def _agregar_propostas_ajuste(
    propostas: list[dict[str, Any]],
    pendencias: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    grupos: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for proposta in propostas:
        grupos[(texto(proposta.get("Auditado")).upper(), texto(proposta.get("Código do item avaliado")))].append(proposta)
    finais: list[dict[str, Any]] = []
    for (auditado, codigo), grupo in grupos.items():
        valores = {texto(p.get("Resposta ajustada")) for p in grupo}
        if len(valores) != 1:
            pendencias.append({
                **grupo[0],
                "Motivo da pendência": f"Valores positivos conflitantes: {', '.join(sorted(valores))}.",
            })
            continue
        base = dict(grupo[0])
        for campo in [
            "Código da situação", "Situação", "Estado temporal", "IDs dos motivos", "IDs das ações",
            "Origem", "Case ID", "Identidade do parecer", "Justificativa do juiz",
        ]:
            valores_campo = []
            for proposta in grupo:
                valor = texto(proposta.get(campo))
                if valor and valor not in valores_campo:
                    valores_campo.append(valor)
            base[campo] = "\n".join(valores_campo)
        base["Auditado"] = auditado
        base["Código do item avaliado"] = codigo
        finais.append(base)
    finais.sort(key=lambda r: (texto(r.get("Auditado")), texto(r.get("Código do item avaliado"))))
    pendencias.sort(key=lambda r: (texto(r.get("Auditado")), texto(r.get("Código do item avaliado"))))
    return finais, pendencias


def gravar_ajustes_combinados_xlsx(
    path: Path,
    *,
    ajustes_secao1: list[dict[str, Any]],
    ajustes_secao2: list[dict[str, Any]],
    pendencias: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    finais, pendencias = _agregar_propostas_ajuste([*ajustes_secao1, *ajustes_secao2], pendencias)
    headers = [
        "Auditado", "Código do item avaliado", "Resposta afirmada", "Resposta anterior", "Resposta ajustada",
        "Achado", "Código da situação", "Situação", "Estado temporal", "IDs dos motivos", "IDs das ações",
        "Fonte da ação", "Origem", "Origem do valor positivo", "Resultado da avaliação do juiz",
        "Justificativa do juiz", "Avaliação do auditor revisor", "Justificativa do auditor revisor",
        "Justificativa", "observacao", "Case ID", "Identidade do parecer", "Data do parecer",
    ]
    wb = Workbook()
    ws = wb.active
    ws.title = "Ajustes"
    ws.append(headers)
    for row in finais:
        ws.append([row.get(h, "") for h in headers])
    pend_ws = wb.create_sheet("Pendências")
    pend_headers = headers + ["Motivo da pendência"]
    pend_ws.append(pend_headers)
    for row in pendencias:
        pend_ws.append([row.get(h, "") for h in pend_headers])
    saneados_ws = wb.create_sheet("Saneados seção 1")
    saneados_ws.append(headers)
    for row in ajustes_secao1:
        saneados_ws.append([row.get(h, "") for h in headers])
    for aba in wb.worksheets:
        _formatar_aba_tabela(aba)
    sanitizar_workbook_para_excel(wb)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return finais, pendencias


def gerar_painel_evidencias_pos_comentarios(
    *,
    painel_origem: Path,
    output: Path,
    ajustes: list[dict[str, Any]],
) -> int:
    """Limpa no painel as não conformidades saneadas pelos comentários do gestor."""
    wb = load_workbook(painel_origem)
    ws = wb.worksheets[0]
    headers = {texto(cell.value): cell.column for cell in ws[1] if texto(cell.value)}
    auditado_col = headers.get("Auditado")
    if auditado_col is None:
        wb.close()
        raise ValueError(f"painel sem coluna Auditado: {painel_origem}")
    linhas_auditados = {
        texto(ws.cell(row=row, column=auditado_col).value).upper(): row
        for row in range(2, ws.max_row + 1)
        if texto(ws.cell(row=row, column=auditado_col).value)
    }
    alteracoes = 0
    for ajuste in ajustes:
        auditado = texto(ajuste.get("Auditado")).upper()
        codigo = texto(ajuste.get("Código do item avaliado"))
        row = linhas_auditados.get(auditado)
        if row is None or codigo not in headers:
            continue
        for coluna in [codigo, f"{codigo}__resposta_afirmada", f"{codigo}__justificativa", f"{codigo}__pratica"]:
            if coluna in headers:
                ws.cell(row=row, column=headers[coluna]).value = None
        alteracoes += 1
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    wb.close()
    return alteracoes


def _registro_avaliacao_utilizavel(registro: dict[str, Any] | None) -> bool:
    if not registro or registro.get("status") != "completed":
        return False
    result = registro.get("result") if isinstance(registro.get("result"), dict) else {}
    estrutura_valida = (
        result.get("status", "completed") == "completed"
        and isinstance(result.get("conclusoes"), list)
        and bool(result.get("conclusoes"))
    )
    if not estrutura_valida:
        return False
    return not validar_resultado_no_escopo(
        registro,
        result,
        validar_temporal=texto(registro.get("secao")) in {"1", "situacoes"},
    )


def _latest_by_case_model(
    paths: list[Path],
    *,
    expected_models: set[str] | None = None,
) -> dict[str, dict[str, dict[str, Any]]]:
    grupos: dict[str, dict[str, dict[str, Any]]] = defaultdict(dict)
    for path in paths:
        for registro in _ler_jsonl(path):
            if not _registro_avaliacao_utilizavel(registro) or not registro.get("case_id"):
                continue
            model_key = texto(registro.get("model_key")) or texto(registro.get("model"))
            if expected_models is not None and model_key not in expected_models:
                continue
            anterior = grupos[registro["case_id"]].get(model_key)
            if anterior is None or texto(registro.get("finished_at")) >= texto(anterior.get("finished_at")):
                grupos[registro["case_id"]][model_key] = registro
    return grupos


def _itens_autoritativos_caso(opiniao: dict[str, Any]) -> list[ItemAfirmado]:
    """Reconstrói o escopo do juiz a partir do caso, nunca da união de opiniões."""
    itens: list[ItemAfirmado] = []
    for item in opiniao.get("itens") or []:
        codigo = texto(item.get("codigo") if isinstance(item, dict) else getattr(item, "codigo", ""))
        if not codigo:
            continue
        itens.append(
            ItemAfirmado(
                codigo=codigo,
                texto=texto(item.get("texto") if isinstance(item, dict) else getattr(item, "texto", "")),
                afirmacao=texto(
                    item.get("afirmacao") if isinstance(item, dict) else getattr(item, "afirmacao", "")
                ),
            )
        )
    if itens:
        return itens
    codigo = texto(opiniao.get("codigo"))
    if texto(opiniao.get("secao")) in {"1", "situacoes"} and codigo:
        return [ItemAfirmado(codigo=codigo, texto="", afirmacao="")]
    return []


def _itens_para_consolidacao(opinioes: list[dict[str, Any]]) -> list[ItemAfirmado]:
    if not opinioes:
        return []
    autoritativos = _itens_autoritativos_caso(opinioes[0])
    if autoritativos:
        return autoritativos
    # Compatibilidade fechada para checkpoints antigos: somente aceita o escopo
    # quando todas as opiniões trazem exatamente o mesmo conjunto de códigos.
    conjuntos: list[list[str]] = []
    primeira_por_codigo: dict[str, dict[str, Any]] = {}
    for opiniao in opinioes:
        conclusoes = (opiniao.get("result") or {}).get("conclusoes") or []
        codigos = [texto(item.get("item_codigo")) for item in conclusoes if isinstance(item, dict)]
        codigos = [codigo for codigo in codigos if codigo]
        if not codigos or len(codigos) != len(set(codigos)):
            return []
        conjuntos.append(codigos)
        if not primeira_por_codigo:
            primeira_por_codigo = {
                texto(item.get("item_codigo")): item for item in conclusoes if isinstance(item, dict)
            }
    if any(set(codigos) != set(conjuntos[0]) for codigos in conjuntos[1:]):
        return []
    return [
        ItemAfirmado(
            codigo=codigo,
            texto=texto(primeira_por_codigo[codigo].get("item_texto")),
            afirmacao=texto(primeira_por_codigo[codigo].get("afirmacao_auditado")),
        )
        for codigo in conjuntos[0]
    ]


def registro_tem_evidencia_documental(registro: dict[str, Any]) -> bool:
    if "evidencia_documental_processada" in registro:
        return bool(registro.get("evidencia_documental_processada"))
    if registro.get("evidence_paths"):
        return True
    return any(
        texto(item.get("status")) == "capturado"
        for item in registro.get("links_manifestacao") or []
        if isinstance(item, dict)
    )


def validar_justificativa_publicavel(
    secao: str,
    result: dict[str, Any],
    *,
    evidencia_documental_disponivel: bool | None = None,
) -> list[str]:
    conclusoes = result.get("conclusoes") or []
    if not conclusoes:
        return ["resultado sem conclusões para validar a justificativa publicável"]
    justificativas = [texto(item.get("justificativa")) for item in conclusoes]
    violacoes: list[str] = []
    proibidas = re.compile(
        r"data[- ]base|período auditado|período da auditoria|correç(?:ão|ao) posterior|"
        r"posterior(?:es)? à data|momento da regularização",
        flags=re.I,
    )
    artefatos_internos = re.compile(
        r"\b(?:manifestacao_gestor\.json|contexto_situacao\.json|"
        r"comentario_gestor\.txt|avaliacao_consolidada_anterior\.json|"
        r"avaliacoes_individuais\.json|avisos_processamento_anexos\.json|"
        r"nao_conformidade_original\.json)\b",
        flags=re.I,
    )
    marcadores_nao_adaptados = re.compile(r"\b(?:o\s*\(\s*a\s*\)|a\s*\(\s*o\s*\))", flags=re.I)
    prefixo_tecnico_upload = re.compile(r"\b\d{5}_\d{2}_[^\s,;]+", flags=re.I)
    for justificativa in justificativas:
        if proibidas.search(justificativa):
            violacoes.append("justificativa destinada ao auditado contém referência temporal proibida")
        if marcadores_nao_adaptados.search(justificativa):
            violacoes.append("justificativa contém marcador de gênero não adaptado")
        if prefixo_tecnico_upload.search(justificativa):
            violacoes.append("justificativa expõe nome técnico de armazenamento do anexo")
    for conclusao in conclusoes:
        conteudo_publicavel = json.dumps(conclusao, ensure_ascii=False, default=str)
        if artefatos_internos.search(conteudo_publicavel):
            violacoes.append("conclusão destinada ao auditado cita artefato interno")
        if evidencia_documental_disponivel is False and conclusao.get("arquivos_referenciados"):
            violacoes.append("conclusão referencia documentos sem anexo ou link comprobatório")
    if evidencia_documental_disponivel is False:
        for justificativa in justificativas:
            normalizada = chave(justificativa)
            if "a organizacao nao apresentou documentacao comprobatoria" not in normalizada:
                violacoes.append(
                    "caso sem anexo ou link deve informar que a organização não apresentou documentação comprobatória"
                )
            if "apresentou como evidencia" in normalizada:
                violacoes.append("caso sem anexo ou link não pode declarar evidência apresentada")

    if secao == "1":
        motivos = conclusoes[0].get("conclusoes_motivos") or []
        estados = [texto(item.get("estado_motivo")) for item in motivos]
        afastados = sum(estado == "afastado" for estado in estados)
        if estados and afastados == len(estados):
            abertura = "a manifestacao foi acolhida"
            encerramento = "a situacao e considerada sanada"
        elif afastados:
            abertura = "a manifestacao foi parcialmente acolhida"
            encerramento = "a inconformidade permanece"
            if not re.search(r"documentacao(?: apresentada)? e insuficiente", chave(justificativas[0])):
                violacoes.append(
                    "acolhimento parcial deve declarar que a documentação é insuficiente para demonstrar integralmente a prática"
                )
        else:
            abertura = "a manifestacao nao foi acolhida"
            encerramento = "a inconformidade permanece"
        normalizada = chave(justificativas[0])
        if not normalizada.startswith(abertura):
            violacoes.append(f"justificativa deve iniciar com '{abertura}'")
        if encerramento not in normalizada:
            violacoes.append(f"justificativa deve concluir que '{encerramento}'")
    else:
        if len(set(justificativas)) != 1:
            violacoes.append("todas as conclusões da seção 2 devem repetir a mesma justificativa consolidada")
        estados = [texto(item.get("estado")) for item in conclusoes]
        conformes = sum(estado == "conforme" for estado in estados)
        if conformes == len(estados):
            abertura = "a reavaliacao foi acolhida"
        elif conformes:
            abertura = "a reavaliacao foi parcialmente acolhida"
            if not re.search(r"documentacao(?: apresentada)? e insuficiente", chave(justificativas[0])):
                violacoes.append(
                    "acolhimento parcial deve declarar que a documentação é insuficiente para demonstrar integralmente a prática"
                )
        else:
            abertura = "a reavaliacao nao foi acolhida"
        normalizada = chave(justificativas[0])
        if not normalizada.startswith(abertura):
            violacoes.append(f"justificativa deve iniciar com '{abertura}'")
        for conclusao in conclusoes:
            codigo = chave(conclusao.get("item_codigo"))
            if codigo and codigo not in normalizada:
                violacoes.append(
                    f"justificativa consolidada da seção 2 não menciona o item {conclusao.get('item_codigo')}"
                )
    return list(dict.fromkeys(violacoes))


def _materializar_consolidado(
    checkpoint: Path,
    clean: Path,
    *,
    deterministicos: list[dict[str, Any]] | None = None,
    expected_case_ids: set[str] | None = None,
    expected_identities: dict[str, str] | None = None,
) -> list[dict[str, Any]]:
    vigentes: dict[str, dict[str, Any]] = {}
    for registro in [*_ler_jsonl(checkpoint), *(deterministicos or [])]:
        if registro.get("status") != "completed" or not registro.get("case_id"):
            continue
        if expected_case_ids is not None and registro["case_id"] not in expected_case_ids:
            continue
        if (
            expected_identities is not None
            and texto(registro.get("origem_decisao")) != "regra_deterministica"
            and registro.get("identity") != expected_identities.get(registro["case_id"])
        ):
            continue
        origem = texto(registro.get("origem_decisao"))
        resultado = registro.get("result")
        possui_resultado = isinstance(resultado, dict) and isinstance(resultado.get("conclusoes"), list)
        if possui_resultado and origem != "regra_deterministica" and validar_resultado_no_escopo(
            registro,
            resultado,
            validar_temporal=texto(registro.get("secao")) in {"1", "situacoes"},
        ):
            continue
        if (
            possui_resultado
            and origem != "regra_deterministica"
            and texto(registro.get("provider")) != "fake"
        ):
            secao_publicacao = "1" if texto(registro.get("secao")) in {"1", "situacoes"} else "2"
            if validar_justificativa_publicavel(
                secao_publicacao,
                resultado,
                evidencia_documental_disponivel=registro_tem_evidencia_documental(registro),
            ):
                continue
        atual = vigentes.get(registro["case_id"])
        if atual is None or texto(registro.get("finished_at")) >= texto(atual.get("finished_at")):
            vigentes[registro["case_id"]] = registro
    registros = sorted(vigentes.values(), key=lambda r: (texto(r.get("auditado")), texto(r.get("codigo"))))
    gravar_jsonl(clean, registros)
    return registros


def _lista_texto(valor: Any) -> str:
    if not isinstance(valor, list):
        return texto(valor)
    partes: list[str] = []
    for item in valor:
        if isinstance(item, dict):
            partes.append(json.dumps(item, ensure_ascii=False, sort_keys=True))
        else:
            partes.append(texto(item))
    return "\n".join(p for p in partes if p)


def _sanitizar_objeto_para_xlsx(valor: Any) -> Any:
    """Copia estruturas removendo controles que o openpyxl rejeita ao anexar."""
    if isinstance(valor, str):
        return ILLEGAL_CHARACTERS_RE.sub("", valor)
    if isinstance(valor, dict):
        return {chave: _sanitizar_objeto_para_xlsx(item) for chave, item in valor.items()}
    if isinstance(valor, list):
        return [_sanitizar_objeto_para_xlsx(item) for item in valor]
    if isinstance(valor, tuple):
        return tuple(_sanitizar_objeto_para_xlsx(item) for item in valor)
    return valor


def _formatar_aba_tabela(ws: Any, *, larguras: dict[str, int] | None = None) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for coluna, largura in (larguras or {}).items():
        ws.column_dimensions[coluna].width = largura


def gravar_avaliacoes_modelos_xlsx(
    path: Path,
    analyses_files: list[Path],
    *,
    expected_models: Iterable[str] = (),
    expected_case_ids: set[str] | None = None,
    itens_excluidos: list[dict[str, Any]] | None = None,
) -> Path:
    """Materializa visão tabular dos checkpoints individuais, inclusive erros."""
    registros: list[dict[str, Any]] = []
    for arquivo in analyses_files:
        registros.extend(_ler_jsonl(arquivo))
    registros = [_sanitizar_objeto_para_xlsx(registro) for registro in registros]
    itens_excluidos = _sanitizar_objeto_para_xlsx(itens_excluidos or [])
    if expected_case_ids is not None:
        registros = [r for r in registros if texto(r.get("case_id")) in expected_case_ids]
    registros.sort(
        key=lambda r: (
            texto(r.get("auditado")), texto(r.get("codigo")),
            texto(r.get("model_key") or r.get("model")), texto(r.get("finished_at")),
        )
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Avaliações"
    ws.append([
        "Case ID", "Seção", "Auditado", "Código", "Provider", "Modelo", "Model key",
        "Status", "Estado", "Estado temporal", "Item", "Afirmação", "Justificativa",
        "Lacunas", "Arquivos referenciados", "Trechos ou elementos", "Páginas ou localização",
        "Providências informadas", "Comentário encaminhamento", "Consequências práticas",
        "Alternativas propostas", "Início", "Fim", "Erro",
    ])
    motivos_ws = wb.create_sheet("Motivos")
    motivos_ws.append([
        "Case ID", "Auditado", "Código", "Provider", "Modelo", "Model key",
        "Origem", "ID motivo", "Estado do motivo", "Justificativa",
    ])
    erros_ws = wb.create_sheet("Erros e ausências")
    erros_ws.append(["Case ID", "Auditado", "Código", "Provider", "Modelo", "Model key", "Status", "Erro"])

    presentes_por_caso: dict[str, set[str]] = defaultdict(set)
    metadados_caso: dict[str, dict[str, Any]] = {}
    for registro in registros:
        cid = texto(registro.get("case_id"))
        model_key = texto(registro.get("model_key") or registro.get("model"))
        metadados_caso[cid] = registro
        conclusoes = (registro.get("result") or {}).get("conclusoes") or []
        if registro.get("status") == "completed" and conclusoes:
            presentes_por_caso[cid].add(model_key)
        else:
            erros_ws.append([
                cid, registro.get("auditado", ""), registro.get("codigo", ""),
                registro.get("provider", ""), registro.get("model", ""), model_key,
                registro.get("status", ""), registro.get("error", ""),
            ])
        for conclusao in conclusoes or [{}]:
            ws.append([
                cid, registro.get("secao", ""), registro.get("auditado", ""), registro.get("codigo", ""),
                registro.get("provider", ""), registro.get("model", ""), model_key,
                registro.get("status", ""), conclusao.get("estado", ""), conclusao.get("estado_temporal", ""),
                conclusao.get("item_texto", ""), conclusao.get("afirmacao_auditado", ""),
                conclusao.get("justificativa", ""), _lista_texto(conclusao.get("lacunas")),
                _lista_texto(conclusao.get("arquivos_referenciados")),
                _lista_texto(conclusao.get("trechos_ou_elementos")),
                _lista_texto(conclusao.get("paginas_ou_localizacao")),
                _lista_texto(conclusao.get("providencias_informadas")),
                conclusao.get("comentarios_encaminhamento", ""),
                _lista_texto(conclusao.get("consequencias_praticas")),
                _lista_texto(conclusao.get("alternativas_propostas")),
                registro.get("started_at", ""), registro.get("finished_at", ""), registro.get("error", ""),
            ])
            for motivo in conclusao.get("conclusoes_motivos") or []:
                if isinstance(motivo, dict):
                    motivos_ws.append([
                        cid, registro.get("auditado", ""), registro.get("codigo", ""),
                        registro.get("provider", ""), registro.get("model", ""), model_key,
                        "modelo", motivo.get("id_motivo", ""), motivo.get("estado_motivo", ""),
                        motivo.get("justificativa", ""),
                    ])

    esperados = set(expected_models)
    if expected_case_ids is not None:
        for cid in sorted(expected_case_ids):
            meta = metadados_caso.get(cid, {})
            for model_key in sorted(esperados - presentes_por_caso.get(cid, set())):
                erros_ws.append([
                    cid, meta.get("auditado", ""), meta.get("codigo", ""), "", "", model_key,
                    "ausente", "Sem avaliação válida no checkpoint vigente.",
                ])

    escopo_ws = wb.create_sheet("Escopo seção 2")
    escopo_ws.append([
        "Auditado", "Código do item", "Questão-base", "Motivo", "Código da situação",
        "Situação", "Estado temporal", "Resposta ajustada", "Case ID seção 1", "Identidade do parecer",
    ])
    for item in itens_excluidos or []:
        escopo_ws.append([
            item.get("Auditado", ""), item.get("Código do item avaliado", ""), item.get("Questão-base", ""),
            item.get("Motivo da exclusão", ""), item.get("Código da situação", ""), item.get("Situação", ""),
            item.get("Estado temporal", ""), item.get("Resposta ajustada", ""), item.get("Case ID", ""),
            item.get("Identidade do parecer", ""),
        ])

    for aba in wb.worksheets:
        _formatar_aba_tabela(aba)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _gravar_pareceres_xlsx(
    path: Path,
    registros: list[dict[str, Any]],
    *,
    grupos: dict[str, dict[str, dict[str, Any]]] | None = None,
) -> None:
    registros = [_sanitizar_objeto_para_xlsx(registro) for registro in registros]
    grupos = _sanitizar_objeto_para_xlsx(grupos or {})
    wb = Workbook()
    ws = wb.active
    ws.title = "Pareceres"
    headers = [
        "Case ID", "Seção", "Auditado", "Código", "Item", "Afirmação", "Estado",
        "Estado temporal", "Justificativa", "Opiniões válidas", "Avaliadores ausentes",
        "Identidade do parecer", "Substitui identidade", "Origem da decisão",
        "Lacunas", "Arquivos referenciados", "Trechos ou elementos", "Páginas ou localização",
        "Providências informadas", "Comentário encaminhamento", "Consequências práticas",
        "Alternativas propostas",
        "Avaliação do auditor revisor", "Justificativa do auditor revisor",
    ]
    ws.append(headers)
    for registro in registros:
        conclusoes = (registro.get("result") or {}).get("conclusoes") or []
        for conclusao in conclusoes:
            ws.append(
                [
                    registro.get("case_id", ""), registro.get("secao", ""), registro.get("auditado", ""),
                    registro.get("codigo", ""), conclusao.get("item_texto", ""), conclusao.get("afirmacao_auditado", ""),
                    conclusao.get("estado", ""), conclusao.get("estado_temporal", ""), conclusao.get("justificativa", ""),
                    registro.get("opinioes_validas", 0), "; ".join(registro.get("avaliadores_ausentes", []) or []),
                    registro.get("identity", ""), registro.get("supersedes_identity", ""),
                    registro.get("origem_decisao", "juiz_ia"),
                    _lista_texto(conclusao.get("lacunas")), _lista_texto(conclusao.get("arquivos_referenciados")),
                    _lista_texto(conclusao.get("trechos_ou_elementos")),
                    _lista_texto(conclusao.get("paginas_ou_localizacao")),
                    _lista_texto(conclusao.get("providencias_informadas")),
                    conclusao.get("comentarios_encaminhamento", ""),
                    _lista_texto(conclusao.get("consequencias_praticas")),
                    _lista_texto(conclusao.get("alternativas_propostas")), "", "",
                ]
            )
    motivos_ws = wb.create_sheet("Motivos")
    motivos_ws.append([
        "Case ID", "Auditado", "Código", "Origem", "Provider", "Modelo", "Model key",
        "ID motivo", "Estado do motivo", "Justificativa",
    ])
    for registro in registros:
        for conclusao in (registro.get("result") or {}).get("conclusoes") or []:
            for motivo in conclusao.get("conclusoes_motivos") or []:
                if isinstance(motivo, dict):
                    motivos_ws.append([
                        registro.get("case_id", ""), registro.get("auditado", ""), registro.get("codigo", ""),
                        "juiz", registro.get("provider", ""), registro.get("model", ""), "",
                        motivo.get("id_motivo", ""), motivo.get("estado_motivo", ""), motivo.get("justificativa", ""),
                    ])
    for cid, por_modelo in (grupos or {}).items():
        for model_key, registro in por_modelo.items():
            for conclusao in (registro.get("result") or {}).get("conclusoes") or []:
                for motivo in conclusao.get("conclusoes_motivos") or []:
                    if isinstance(motivo, dict):
                        motivos_ws.append([
                            cid, registro.get("auditado", ""), registro.get("codigo", ""), "modelo",
                            registro.get("provider", ""), registro.get("model", ""), model_key,
                            motivo.get("id_motivo", ""), motivo.get("estado_motivo", ""), motivo.get("justificativa", ""),
                        ])
    _formatar_aba_tabela(ws)
    _formatar_aba_tabela(motivos_ws)
    sanitizar_workbook_para_excel(wb)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def consolidar_casos(
    *,
    analyses_files: list[Path],
    out_dir: Path,
    secao: str,
    expected_models: list[str],
    judge_provider: str = "gemini",
    judge_model: str = "gemini-3.1-flash-lite",
    min_opinions: int = 2,
    reasoning: str = "high",
    rpm: int = 0,
    tpm: int = 0,
    max_parallel: int = 1,
    pdf2md: bool = False,
    docx2html: bool = False,
    pdf_detail: str = "auto",
    catalog_comentarios: Path = DEFAULT_CATALOG_COMENTARIOS,
    deterministic_path: Path | None = None,
    expected_case_ids: set[str] | None = None,
    selected_case_ids: set[str] | None = None,
    refresh_links: bool = False,
    quiet: bool = False,
    verbose: bool = False,
) -> dict[str, Any]:
    """Consolida opiniões por case_id e refaz o parecer quando o conjunto muda."""
    if max_parallel < 1:
        raise ValueError("max_parallel deve ser inteiro positivo")
    if tpm < 0:
        raise ValueError("tpm deve ser inteiro não negativo")
    modelos_esperados = set(expected_models)
    grupos = _latest_by_case_model(analyses_files, expected_models=modelos_esperados)
    if expected_case_ids is not None:
        grupos = {cid: por_modelo for cid, por_modelo in grupos.items() if cid in expected_case_ids}
    for cid in expected_case_ids or set():
        grupos.setdefault(cid, {})
    prompts = carregar_prompts_comentarios(catalog_comentarios)
    prompt = prompts["juiz_situacoes" if secao == "1" else "juiz_reavaliacao"]
    response_profile = "manager_comments_temporal" if secao == "1" else "evidence"
    checkpoint = out_dir / "consolidated.jsonl"
    clean = out_dir / "consolidated_clean.jsonl"
    anteriores = _ler_jsonl(checkpoint)
    por_identidade = {r.get("identity"): r for r in anteriores if r.get("identity")}
    por_caso_anterior: dict[str, dict[str, Any]] = {}
    for r in anteriores:
        if r.get("status") == "completed" and r.get("case_id"):
            atual = por_caso_anterior.get(r["case_id"])
            if atual is None or texto(r.get("finished_at")) >= texto(atual.get("finished_at")):
                por_caso_anterior[r["case_id"]] = r
    limiter = RequestsPerMinuteLimiter(rpm)
    processados = concluidos = erros = pulados = pendentes = 0
    pendencias: list[dict[str, Any]] = []
    identidades_geracao: dict[str, str] = {}
    tarefas: list[dict[str, Any]] = []

    for index, (cid, por_modelo) in enumerate(sorted(grupos.items()), start=1):
        if selected_case_ids is not None and cid not in selected_case_ids:
            continue
        opinioes = list(por_modelo.values())
        presentes = set(por_modelo)
        ausentes = sorted(modelos_esperados - presentes)
        if len(opinioes) < min_opinions:
            pendentes += 1
            pendencias.append({"case_id": cid, "opinioes_validas": len(opinioes), "avaliadores_ausentes": ausentes})
            continue
        primeira = opinioes[0]
        pacote_contexto = primeira.get("pacote_contexto_consolidacao") or {}
        documentos = list(pacote_contexto.get("documentos") or [])
        links_paths, links_manifestacao, erros_links = capturar_links_manifestacao(
            documentos,
            out_dir / "links-manifestacoes" / cid,
            refresh=refresh_links,
        )
        identity = _hash_json(
            {
                "case_id": cid,
                "opinioes": sorted(r.get("identity", "") for r in opinioes),
                "judge_provider": judge_provider,
                "judge_model": judge_model,
                "prompt_hash": _hash_json(prompt),
                "response_profile": response_profile,
                "reasoning": reasoning,
                "pdf2md": pdf2md,
                "docx2html": docx2html,
                "pdf_detail": pdf_detail,
                "links_manifestacao": [
                    {
                        "url_original": item.get("url_original"),
                        "url_final": item.get("url_final"),
                        "sha256": item.get("sha256"),
                        "status": item.get("status"),
                    }
                    for item in links_manifestacao
                ],
            }
        )
        identidades_geracao[cid] = identity
        # Uma seleção explícita representa reparo do parecer vigente. Nesse caso,
        # não reutilize o checkpoint que motivou o reparo, ainda que sua identidade
        # lógica continue a mesma; as avaliações individuais permanecem reutilizadas.
        if selected_case_ids is None and _registro_avaliacao_utilizavel(por_identidade.get(identity)):
            pulados += 1
            continue
        tarefas.append({
            "index": index,
            "cid": cid,
            "opinioes": opinioes,
            "ausentes": ausentes,
            "primeira": primeira,
            "documentos": documentos,
            "identity": identity,
            "links_paths": links_paths,
            "links_manifestacao": links_manifestacao,
            "erros_links": erros_links,
        })

    key_pool = None
    if judge_provider == "gemini" and tarefas:
        key_pool = ExclusiveApiKeyPool(api_key("gemini"), rpm=rpm, tpm=tpm)
        effective_workers = min(max_parallel, key_pool.key_count, len(tarefas))
    elif judge_provider == "fake":
        effective_workers = min(max_parallel, len(tarefas)) if tarefas else 1
    else:
        # O pool exclusivo solicitado e especifico do Gemini. Os demais
        # providers preservam a execucao sequencial e sua rotacao historica.
        effective_workers = 1

    def executar_chamada(
        *,
        tarefa: dict[str, Any],
        prompt_atual: str,
        pacote: dict[str, Any],
        itens_autoritativos: list[ItemAfirmado],
    ) -> dict[str, Any]:
        primeira = tarefa["primeira"]
        cid = tarefa["cid"]
        if key_pool is None:
            if judge_provider != "fake" and rpm:
                espera = limiter.wait_seconds()
                limiter.wait_and_mark(espera)
            return executar_provider(
                provider=judge_provider,
                model=judge_model,
                api_key=api_key(judge_provider),
                prompt=prompt_atual,
                auditado=texto(primeira.get("auditado")),
                questao_base=texto(primeira.get("codigo")),
                coluna_evidencia=texto(primeira.get("coluna_evidencia")),
                itens_afirmados=itens_autoritativos,
                pacote=pacote,
                reasoning_effort=reasoning,
                response_profile=response_profile,
                pdf_detail=pdf_detail,
            )

        tokens_info = estimar_tokens_payload(
            prompt=prompt_atual,
            auditado=texto(primeira.get("auditado")),
            questao_base=texto(primeira.get("codigo")),
            coluna_evidencia=texto(primeira.get("coluna_evidencia")),
            itens_afirmados=itens_autoritativos,
            pacote=pacote,
            provider=judge_provider,
            response_profile=response_profile,
        )
        while True:
            lease = key_pool.acquire(
                tokens=tokens_info["tokens_total"],
                on_wait=(
                    lambda fields: log_event(
                        "comment_judge_key_wait",
                        "Aguardando chave Gemini livre dentro dos limites por chave.",
                        case_id=cid,
                        index=tarefa["index"],
                        total=len(grupos),
                        secao=secao,
                        rpm=rpm,
                        tpm=tpm,
                        **fields,
                    )
                ) if verbose and not quiet else None,
            )
            if verbose and not quiet:
                log_event(
                    "comment_judge_key_acquired",
                    "Chave Gemini locada exclusivamente para o parecer.",
                    case_id=cid,
                    key=lease.label,
                )
            result: dict[str, Any] | None = None
            cooldown = 0.0
            try:
                result = executar_provider(
                    provider=judge_provider,
                    model=judge_model,
                    api_key=lease.key,
                    prompt=prompt_atual,
                    auditado=texto(primeira.get("auditado")),
                    questao_base=texto(primeira.get("codigo")),
                    coluna_evidencia=texto(primeira.get("coluna_evidencia")),
                    itens_afirmados=itens_autoritativos,
                    pacote=pacote,
                    reasoning_effort=reasoning,
                    response_profile=response_profile,
                    pdf_detail=pdf_detail,
                )
                if resultado_429(result):
                    cooldown = cooldown_429(result)
            finally:
                lease.release(cooldown_seconds=cooldown)
            if not resultado_429(result):
                if verbose and not quiet:
                    log_event(
                        "comment_judge_key_released",
                        "Chave Gemini liberada após a chamada ao juiz.",
                        case_id=cid,
                        key=lease.label,
                    )
                return result or {"status": "error", "error": "resultado vazio do provider"}
            if verbose and not quiet:
                log_event(
                    "comment_judge_key_cooldown",
                    "Chave Gemini recebeu 429 e entrou em cooldown; o caso será repetido.",
                    level="warning",
                    case_id=cid,
                    key=lease.label,
                    cooldown_seconds=cooldown,
                )

    def processar_tarefa(tarefa: dict[str, Any]) -> dict[str, Any]:
        cid = tarefa["cid"]
        opinioes = tarefa["opinioes"]
        primeira = tarefa["primeira"]
        documentos = list(tarefa["documentos"])
        documentos.append(_documento("avaliacoes_individuais.json", "avaliacoes_individuais", [
            {
                "provider": r.get("provider"), "model": r.get("model"),
                "model_key": r.get("model_key"), "result": r.get("result"),
            }
            for r in opinioes
        ]))
        evidence_paths: list[Path] = []
        caminhos_vistos: set[Path] = set()
        for opiniao in opinioes:
            for valor in opiniao.get("evidence_paths") or []:
                caminho = Path(valor)
                if caminho.is_file() and caminho not in caminhos_vistos:
                    caminhos_vistos.add(caminho)
                    evidence_paths.append(caminho)
        for caminho in tarefa["links_paths"]:
            if caminho not in caminhos_vistos:
                caminhos_vistos.add(caminho)
                evidence_paths.append(caminho)

        started = dt.datetime.now(dt.timezone.utc)
        avisos_processamento: list[dict[str, str]] = []
        evidencia_documental_processada = False
        try:
            if tarefa["erros_links"]:
                raise ValueError(
                    "links .gov.br não puderam ser processados: "
                    + "; ".join(tarefa["erros_links"])
                )
            with tempfile.TemporaryDirectory() as tmp:
                pacote = {
                    "documentos": documentos,
                    "inventario": [d.get("nome", "") for d in documentos],
                    "erro": "",
                    "arquivos_upload": [],
                }
                conversion_context = _EVIDENCE_CONVERSION_LOCK if (pdf2md or docx2html) else nullcontext()
                with conversion_context:
                    for caminho in evidence_paths:
                        preparado, arquivos, erro = preparar_evidencia_para_provider(
                            caminho, tmp, pdf2md=pdf2md, docx2html=docx2html, dpi=150
                        )
                        if erro:
                            _registrar_anexo_nao_processado(
                                avisos_processamento, caminho, f"erro ao preparar: {erro}"
                            )
                            continue
                        bloqueante = erro_tecnico_bloqueante_pacote(preparado, arquivos)
                        if bloqueante:
                            _registrar_anexo_nao_processado(
                                avisos_processamento, caminho, bloqueante
                            )
                            continue
                        if preparado.documentos or arquivos:
                            evidencia_documental_processada = True
                        pacote["documentos"].extend(preparado.documentos)
                        pacote["inventario"].extend(preparado.inventario)
                        pacote["arquivos_upload"].extend(arquivos)
                if not pacote["documentos"] and not pacote["arquivos_upload"]:
                    raise ValueError("consolidação sem opiniões, contexto ou evidência processável")
                _adicionar_avisos_anexos_ao_pacote(pacote, avisos_processamento)
                itens_autoritativos = _itens_para_consolidacao(opinioes)
                if not itens_autoritativos:
                    raise ValueError("caso sem itens autoritativos para consolidação")
                prompt_tentativa = prompt
                for tentativa in range(3):
                    result = executar_chamada(
                        tarefa=tarefa,
                        prompt_atual=prompt_tentativa,
                        pacote=pacote,
                        itens_autoritativos=itens_autoritativos,
                    )
                    if secao == "1" and judge_provider == "fake" and result.get("status") == "completed":
                        result = _resultado_fake_temporal(
                            result, (primeira.get("contexto") or {}).get("motivos", [])
                        )
                    if result.get("status") != "completed":
                        break
                    registro_escopo = {
                        "secao": secao,
                        "codigo": primeira.get("codigo"),
                        "itens": [asdict(item) for item in itens_autoritativos],
                        "contexto": primeira.get("contexto", {}),
                    }
                    violacoes_escopo = validar_resultado_no_escopo(
                        registro_escopo,
                        result,
                        validar_temporal=secao == "1",
                    )
                    violacoes_publicacao = (
                        validar_justificativa_publicavel(
                            secao,
                            result,
                            evidencia_documental_disponivel=evidencia_documental_processada,
                        )
                        if judge_provider != "fake" and not violacoes_escopo
                        else []
                    )
                    if not violacoes_escopo and not violacoes_publicacao:
                        break
                    detalhes = []
                    if violacoes_escopo:
                        detalhes.append("escopo lógico: " + formatar_violacoes(violacoes_escopo))
                    if violacoes_publicacao:
                        detalhes.append("texto publicável: " + "; ".join(violacoes_publicacao))
                    erro_validacao = " | ".join(detalhes)
                    if judge_provider == "fake" or tentativa == 2:
                        result = {
                            "status": "error",
                            "error": "parecer inválido após correção: " + erro_validacao,
                        }
                        break
                    prompt_tentativa = (
                        prompt
                        + "\n\nCORREÇÃO OBRIGATÓRIA: a resposta anterior foi rejeitada por: "
                        + erro_validacao
                        + ". Refaça integralmente a consolidação e cumpra exatamente o schema e o padrão textual."
                    )
            status = result.get("status", "error")
            error = result.get("error", "")
        except Exception as exc:
            result = {"status": "error", "error": str(exc)}
            status = "error"
            error = str(exc)
        finished = dt.datetime.now(dt.timezone.utc)
        return {
            "identity": tarefa["identity"], "case_id": cid, "secao": primeira.get("secao"),
            "auditado": primeira.get("auditado"), "codigo": primeira.get("codigo"),
            "questao": primeira.get("codigo"), "coluna_evidencia": primeira.get("coluna_evidencia"),
            "evidencia": primeira.get("evidencia", ""), "provider": judge_provider, "model": judge_model,
            "status": status, "error": error, "result": result,
            "opinioes_validas": len(opinioes), "opinioes_esperadas": len(expected_models),
            "avaliadores_ausentes": tarefa["ausentes"],
            "opinioes": [
                {
                    "identity": r.get("identity"), "provider": r.get("provider"),
                    "model": r.get("model"), "model_key": r.get("model_key"),
                }
                for r in opinioes
            ],
            "contexto": primeira.get("contexto", {}),
            "itens": [asdict(item) for item in _itens_para_consolidacao(opinioes)],
            "evidence_paths": [str(path) for path in evidence_paths],
            "evidencia_documental_processada": evidencia_documental_processada,
            "links_manifestacao": tarefa["links_manifestacao"],
            "avisos_processamento_evidencias": avisos_processamento,
            "started_at": started.isoformat(), "finished_at": finished.isoformat(),
            "duration_seconds": round((finished - started).total_seconds(), 3),
            "supersedes_identity": por_caso_anterior.get(cid, {}).get("identity", ""),
            "_index": tarefa["index"],
        }

    if tarefas:
        log_event(
            "comment_consolidation_parallel_started",
            "Consolidação paralela de comentários iniciada.",
            quiet=quiet,
            secao=secao,
            max_parallel=max_parallel,
            effective_workers=effective_workers,
            gemini_keys=key_pool.key_count if key_pool is not None else 0,
            rpm_per_key=rpm if key_pool is not None else 0,
            tpm_per_key=tpm if key_pool is not None else 0,
        )
        with ThreadPoolExecutor(max_workers=effective_workers) as executor:
            futures = {executor.submit(processar_tarefa, tarefa): tarefa for tarefa in tarefas}
            for future in as_completed(futures):
                registro = future.result()
                index = registro.pop("_index")
                _append_jsonl(checkpoint, registro)
                por_identidade[registro["identity"]] = registro
                if registro["status"] == "completed":
                    por_caso_anterior[registro["case_id"]] = registro
                    concluidos += 1
                else:
                    erros += 1
                processados += 1
                log_event(
                    "comment_consolidation_recorded", "Parecer de comentário consolidado.", quiet=quiet,
                    case_id=registro["case_id"], index=index, total=len(grupos), secao=secao,
                    status=registro["status"], opinioes=registro["opinioes_validas"],
                    avaliadores_ausentes=registro["avaliadores_ausentes"], error=registro["error"],
                )
    deterministicos = _ler_jsonl(deterministic_path) if deterministic_path else []
    ids_vigentes = None
    if expected_case_ids is not None:
        ids_vigentes = set(expected_case_ids)
        ids_vigentes.update(
            registro["case_id"]
            for registro in deterministicos
            if registro.get("case_id")
        )
    vigentes = _materializar_consolidado(
        checkpoint,
        clean,
        deterministicos=deterministicos,
        expected_case_ids=ids_vigentes,
        expected_identities=identidades_geracao if expected_case_ids is not None else None,
    )
    xlsx = out_dir / "pareceres_consolidados.xlsx"
    _gravar_pareceres_xlsx(xlsx, vigentes, grupos=grupos)
    pendencias_path = out_dir / "pendencias_quorum.json"
    pendencias_path.parent.mkdir(parents=True, exist_ok=True)
    pendencias_path.write_text(json.dumps(pendencias, ensure_ascii=False, indent=2), encoding="utf-8")
    return {
        "secao": secao, "grupos": len(grupos), "processados": processados, "concluidos": concluidos,
        "erros": erros, "pulados": pulados, "pendentes_quorum": pendentes,
        "max_parallel": max_parallel, "effective_workers": effective_workers,
        "gemini_keys": key_pool.key_count if key_pool is not None else 0,
        "rpm_per_key": rpm if key_pool is not None else 0,
        "tpm_per_key": tpm if key_pool is not None else 0,
        "checkpoint": str(checkpoint), "clean": str(clean), "xlsx": str(xlsx),
        "vigentes": len(vigentes), "pendencias": str(pendencias_path),
    }
