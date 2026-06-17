#!/usr/bin/env python3
"""Avalia evidencias do questionario iGovTI com prompts da matriz.

O script processa a planilha de respostas do LimeSurvey, localiza as
evidencias fisicas em uma pasta informada como argumento, executa os prompts da
planilha de verificacao de evidencias e gera:

- JSONL incremental de resultados;
- JSON em lista com os mesmos resultados;
- XLSX consolidado por auditado, com colunas como q1001evi e
  q1001evi_justificativa.

Foram reaproveitadas as ideias do pipeline de
`/home/acba/workspace/md2survey/avaliacao_evidencias`: providers isolados,
checkpoint incremental, resolucao de nomes exportados pelo LimeSurvey, retry
para erros transientes e saida JSON validada.
"""

from __future__ import annotations

import argparse
import base64
import dataclasses
import datetime as dt
import difflib
import hashlib
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import unicodedata
import urllib.error
import urllib.request
import zipfile
from dataclasses import dataclass
from email.utils import parsedate_to_datetime
from pathlib import Path
from typing import Any, Callable, Iterable
from urllib.parse import unquote

from openpyxl import Workbook, load_workbook


REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_RESPOSTAS = REPO_ROOT / "02-Execucao/01-Coleta_Dados/20260607-respostas-questionario.xlsx"
DEFAULT_MATRIZ_VERIFICACAO = (
    REPO_ROOT
    / "01-Planejamento/03-Estrategia_e_Plano/04-Matriz_Planejamento/matriz_verificaca_evidencias.xlsx"
)
DEFAULT_MATRIZ_PLANEJAMENTO = (
    REPO_ROOT
    / "01-Planejamento/03-Estrategia_e_Plano/04-Matriz_Planejamento/matriz_planejamento.md"
)
DEFAULT_OUT_DIR = REPO_ROOT / "02-Execucao/02-Testes_Auditoria/avaliacao_evidencias"

RESULTADO_CONFORME = "Conforme"
RESULTADO_NAO_CONFORME = "Não Conforme"
RESULTADO_PENDENTE = "Pendente"
RESULTADOS_VALIDOS = {RESULTADO_CONFORME, RESULTADO_NAO_CONFORME}
REMOTE_PROVIDERS = {"openrouter", "gemini"}
RETRYABLE_PROVIDER_STATUSES = {429, 500, 502, 503, 504}
DEFAULT_TRANSIENT_RETRY_DELAYS = (30.0, 60.0, 120.0)
SOURCE_COLUMNS = ["id", "token", "submitdate", "firstname", "lastname", "email"]
WORD_EXTENSIONS_TO_PDF = {".doc", ".docx"}


@dataclass(frozen=True)
class AvaliacaoPrompt:
    id_avaliacao: str
    questao_auditoria: str
    situacao_encontrada: str
    item_evidencia: str
    descricao_avaliacao: str
    prompt_checklist: str
    referenciada_na_matriz: bool
    prompt_hash: str


@dataclass(frozen=True)
class EvidenciaResolvida:
    caminho: Path | None
    nome_decodificado: str
    upload: dict[str, Any]
    erro: str = ""


@dataclass(frozen=True)
class PacoteEvidencia:
    documentos: list[dict[str, Any]]
    inventario: list[str]
    arquivos: list[dict[str, Any]]
    erros: list[str]


@dataclass
class RequestsPerMinuteLimiter:
    rpm: int
    clock: Callable[[], float] = time.monotonic
    sleeper: Callable[[float], None] = time.sleep
    last_started_at: float | None = None

    def wait_seconds(self) -> float:
        if self.rpm <= 0 or self.last_started_at is None:
            return 0.0
        elapsed = self.clock() - self.last_started_at
        return max(0.0, (60.0 / self.rpm) - elapsed)

    def wait_and_mark(self) -> float:
        wait_seconds = self.wait_seconds()
        if wait_seconds > 0:
            self.sleeper(wait_seconds)
        self.last_started_at = self.clock()
        return wait_seconds


def json_default(value: Any) -> Any:
    if isinstance(value, Path):
        return str(value)
    if dataclasses.is_dataclass(value):
        return dataclasses.asdict(value)
    raise TypeError(f"Objeto nao serializavel: {type(value).__name__}")


def dumps_json(payload: Any, *, indent: int | None = None) -> str:
    return json.dumps(payload, ensure_ascii=False, sort_keys=True, default=json_default, indent=indent)


def log_event(event: str, message: str, *, quiet: bool = False, level: str = "info", **fields: Any) -> None:
    if quiet:
        return
    payload = {
        "ts": dt.datetime.now(dt.timezone.utc).isoformat(),
        "level": level,
        "event": event,
        "message": message,
        **fields,
    }
    print(dumps_json(payload), flush=True)


def validar_rpm(value: str) -> int:
    try:
        rpm = int(value)
    except ValueError as exc:
        raise argparse.ArgumentTypeError("--rpm deve ser um inteiro maior ou igual a zero") from exc
    if rpm < 0:
        raise argparse.ArgumentTypeError("--rpm deve ser maior ou igual a zero")
    return rpm


def sha256_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def rows_from_xlsx(path: Path, *, sheet_name: str | None = None) -> tuple[list[str], list[dict[str, Any]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name] if sheet_name else workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
        data = [dict(zip(headers, values)) for values in rows]
        return headers, data
    finally:
        workbook.close()


def colunas_evidencia(headers: Iterable[str]) -> list[str]:
    return [name for name in headers if "evi" in name and not name.endswith("[filecount]")]


def extrair_targets_matriz(path: Path) -> set[str]:
    text = path.read_text(encoding="utf-8")
    return set(re.findall(r"avaliacao\[([^\]]+)\]", text))


def carregar_avaliacoes(path: Path, matrix_targets: set[str]) -> list[AvaliacaoPrompt]:
    headers, rows = rows_from_xlsx(path, sheet_name="Avaliacoes")
    required = {
        "id_avaliacao",
        "questao_auditoria",
        "situacao_encontrada",
        "item_evidencia",
        "descricao_avaliacao",
        "prompt_checklist",
    }
    missing = required.difference(headers)
    if missing:
        raise ValueError(f"Planilha de verificacao sem colunas obrigatorias: {', '.join(sorted(missing))}")

    avaliacoes: list[AvaliacaoPrompt] = []
    seen: set[str] = set()
    for row_number, row in enumerate(rows, start=2):
        id_avaliacao = str(row.get("id_avaliacao") or "").strip()
        if not id_avaliacao:
            continue
        if id_avaliacao in seen:
            raise ValueError(f"id_avaliacao duplicado na linha {row_number}: {id_avaliacao}")
        seen.add(id_avaliacao)
        prompt = str(row.get("prompt_checklist") or "").strip()
        if not prompt:
            raise ValueError(f"prompt_checklist vazio na linha {row_number}: {id_avaliacao}")
        item_evidencia = str(row.get("item_evidencia") or "").strip()
        if not item_evidencia:
            raise ValueError(f"item_evidencia vazio na linha {row_number}: {id_avaliacao}")
        avaliacoes.append(
            AvaliacaoPrompt(
                id_avaliacao=id_avaliacao,
                questao_auditoria=str(row.get("questao_auditoria") or "").strip(),
                situacao_encontrada=str(row.get("situacao_encontrada") or "").strip(),
                item_evidencia=item_evidencia,
                descricao_avaliacao=str(row.get("descricao_avaliacao") or "").strip(),
                prompt_checklist=prompt,
                referenciada_na_matriz=id_avaliacao in matrix_targets,
                prompt_hash=sha256_text(prompt),
            )
        )
    return avaliacoes


def validar_matriz(avaliacoes: list[AvaliacaoPrompt], matrix_targets: set[str]) -> dict[str, Any]:
    xlsx_ids = {avaliacao.id_avaliacao for avaliacao in avaliacoes}
    return {
        "targets_matriz": sorted(matrix_targets),
        "ids_planilha": sorted(xlsx_ids),
        "ids_planilha_nao_referenciados_na_matriz": sorted(xlsx_ids - matrix_targets),
        "targets_matriz_sem_prompt_na_planilha": sorted(matrix_targets - xlsx_ids),
    }


def parse_uploads(value: Any) -> tuple[list[dict[str, Any]], str]:
    if value in (None, ""):
        return [], ""
    if isinstance(value, str):
        try:
            parsed = json.loads(value)
        except json.JSONDecodeError as exc:
            return [], f"metadados de upload invalidos: {exc.msg}"
    else:
        parsed = value
    if isinstance(parsed, dict):
        parsed = [parsed]
    if not isinstance(parsed, list):
        return [], "metadados de upload devem ser uma lista"
    uploads: list[dict[str, Any]] = []
    for index, item in enumerate(parsed, start=1):
        if not isinstance(item, dict):
            return [], f"upload {index} deve ser um objeto"
        name = item.get("name")
        filename = item.get("filename")
        if not name and not filename:
            return [], f"upload {index} sem name ou filename"
        uploads.append(item)
    return uploads, ""


def _nome_exportado_key(name: str) -> tuple[str, str]:
    path = Path(name)
    stem = unicodedata.normalize("NFKD", path.stem.casefold())
    stem = "".join(char for char in stem if not unicodedata.combining(char))
    stem = re.sub(r"[^a-z0-9]+", "-", stem)
    stem = re.sub(r"-+", "-", stem).strip("-")
    return stem, path.suffix.casefold()


def _remover_prefixo_exportacao_limesurvey(name: str) -> str:
    return re.sub(r"^\d+_\d+_", "", name)


def _prefixo_exportacao_limesurvey(resposta_id: Any, evidence_index: int | None) -> str:
    if evidence_index is None:
        return ""
    try:
        return f"{int(resposta_id):05d}_{int(evidence_index):02d}_"
    except (TypeError, ValueError):
        return ""


def _melhor_match_nome_exportado(nome_decodificado: str, candidates: list[Path]) -> Path | None:
    name_key, suffix = _nome_exportado_key(nome_decodificado)
    matches: list[tuple[float, str, Path]] = []
    for candidate in candidates:
        candidate_without_prefix = _remover_prefixo_exportacao_limesurvey(candidate.name)
        candidate_key, candidate_suffix = _nome_exportado_key(candidate_without_prefix)
        if suffix and candidate_suffix != suffix:
            continue
        score = 1.0 if candidate_key == name_key else difflib.SequenceMatcher(None, name_key, candidate_key).ratio()
        if score >= 0.92:
            matches.append((score, candidate.name, candidate))
    if not matches:
        return None
    matches.sort(key=lambda item: (-item[0], item[1]))
    return matches[0][2]


def _find_in_dir(
    directory: Path,
    names: list[str],
    *,
    resposta_id: Any,
    evidence_index: int | None,
) -> Path | None:
    if not directory.is_dir():
        return None
    for name in names:
        if not name:
            continue
        candidate = directory / name
        if candidate.is_file():
            return candidate

    candidates = [path for path in directory.iterdir() if path.is_file()]
    prefix = _prefixo_exportacao_limesurvey(resposta_id, evidence_index)
    if prefix:
        prefixed = [path for path in candidates if path.name.startswith(prefix)]
        for name in names:
            match = _melhor_match_nome_exportado(name, prefixed)
            if match:
                return match

    for name in names:
        match = _melhor_match_nome_exportado(name, candidates)
        if match:
            return match
    return None


def _find_recursive(root: Path, auditado: str, names: list[str]) -> Path | None:
    normalized_targets = {_nome_exportado_key(name) for name in names if name}
    exact_names = {name for name in names if name}
    matches: list[tuple[int, str, Path]] = []
    for path in root.rglob("*"):
        if not path.is_file():
            continue
        if path.name in exact_names or _nome_exportado_key(_remover_prefixo_exportacao_limesurvey(path.name)) in normalized_targets:
            contains_auditado = auditado and auditado.casefold() in str(path.parent).casefold()
            matches.append((0 if contains_auditado else 1, str(path), path))
    if not matches:
        return None
    matches.sort(key=lambda item: (item[0], item[1]))
    return matches[0][2]


def resolver_upload(
    root: Path,
    auditado: str,
    upload: dict[str, Any],
    *,
    resposta_id: Any,
    evidence_index: int | None,
    recursive: bool,
) -> EvidenciaResolvida:
    raw_name = str(upload.get("name") or "").strip()
    decoded_name = unquote(raw_name) if raw_name else ""
    filename = str(upload.get("filename") or "").strip()
    ext = str(upload.get("ext") or "").strip().lstrip(".")
    names = [decoded_name, raw_name, filename]
    if filename and ext:
        names.append(f"{filename}.{ext}")
    names = [name for name in dict.fromkeys(names) if name]
    if not names:
        return EvidenciaResolvida(caminho=None, nome_decodificado="", upload=upload, erro="upload sem nome localizavel")

    for directory in [root / auditado, root]:
        match = _find_in_dir(directory, names, resposta_id=resposta_id, evidence_index=evidence_index)
        if match:
            return EvidenciaResolvida(caminho=match, nome_decodificado=decoded_name or match.name, upload=upload)

    if recursive:
        match = _find_recursive(root, auditado, names)
        if match:
            return EvidenciaResolvida(caminho=match, nome_decodificado=decoded_name or match.name, upload=upload)

    expected = root / auditado / (decoded_name or filename)
    return EvidenciaResolvida(
        caminho=None,
        nome_decodificado=decoded_name or filename,
        upload=upload,
        erro=f"arquivo de evidencia nao encontrado: {expected}",
    )


def _zip_member_safe(name: str) -> bool:
    path = Path(name)
    return not path.is_absolute() and ".." not in path.parts


def _read_text_file(path: Path) -> tuple[str, str]:
    try:
        return path.read_text(encoding="utf-8"), ""
    except UnicodeDecodeError:
        try:
            return path.read_text(encoding="latin-1"), ""
        except UnicodeDecodeError as exc:
            return "", f"erro ao ler texto: {exc}"


def _truncate(text: str, max_chars: int) -> str:
    if len(text) <= max_chars:
        return text
    return text[:max_chars] + f"\n\n[texto truncado em {max_chars} caracteres]"


def _extrair_texto_pdf(path: Path, *, max_chars: int) -> tuple[list[dict[str, Any]], str]:
    try:
        from pypdf import PdfReader
    except ModuleNotFoundError:
        return [], "pypdf nao instalado para extrair texto de PDF"
    try:
        reader = PdfReader(path)
        documentos = []
        for index, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                documentos.append({"nome": path.name, "pagina": index, "texto": _truncate(text, max_chars)})
        if not documentos:
            return [], "pdf sem texto extraivel"
        return documentos, ""
    except Exception as exc:
        return [], f"erro ao extrair texto de PDF: {exc}"


def _extrair_texto_pdf_bytes(name: str, data: bytes, *, max_chars: int) -> tuple[list[dict[str, Any]], str]:
    try:
        from pypdf import PdfReader
    except ModuleNotFoundError:
        return [], "pypdf nao instalado para extrair texto de PDF"
    try:
        reader = PdfReader(io.BytesIO(data))
        documentos = []
        for index, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                documentos.append({"nome": name, "pagina": index, "texto": _truncate(text, max_chars)})
        if not documentos:
            return [], "pdf sem texto extraivel"
        return documentos, ""
    except Exception as exc:
        return [], f"erro ao extrair texto de PDF: {exc}"


def _extrair_texto_com_markitdown(path: Path) -> tuple[str, str]:
    try:
        from markitdown import MarkItDown
    except ModuleNotFoundError:
        return "", "markitdown nao instalado no ambiente virtual"
    try:
        md = MarkItDown()
        result = md.convert(str(path))
        return result.text_content, ""
    except Exception as exc:
        return "", f"erro ao extrair texto com markitdown: {exc}"


def _extrair_texto_bytes_com_markitdown(nome: str, data: bytes, sufixo: str) -> tuple[str, str]:
    with tempfile.TemporaryDirectory() as tmp_dir:
        temp_file = Path(tmp_dir) / f"temp_{_slug_ascii(Path(nome).stem)}{sufixo}"
        temp_file.write_bytes(data)
        return _extrair_texto_com_markitdown(temp_file)


def normalizar_arquivo(path: Path, *, max_chars: int) -> tuple[list[dict[str, Any]], list[str], list[str]]:
    suffix = path.suffix.casefold()
    if suffix in {".txt", ".md", ".csv"}:
        text, error = _read_text_file(path)
        if error:
            return [], [path.name], [error]
        return [{"nome": path.name, "texto": _truncate(text, max_chars)}], [path.name], []

    if suffix == ".pdf":
        documentos, error = _extrair_texto_pdf(path, max_chars=max_chars)
        return documentos, [path.name], [error] if error else []

    if suffix == ".xlsx":
        text, error = _extrair_texto_com_markitdown(path)
        if error:
            return [], [path.name], [error]
        return [{"nome": path.name, "texto": _truncate(text, max_chars)}], [path.name], []
    if suffix in WORD_EXTENSIONS_TO_PDF:
        return [
            {
                "nome": path.name,
                "tipo_convertido": "pdf",
                "observacao": "Documento Word preservado para avaliacao visual por conversao para PDF no upload.",
            }
        ], [path.name], []

    if suffix == ".zip":
        documentos: list[dict[str, Any]] = []
        inventario: list[str] = []
        erros: list[str] = []
        try:
            with zipfile.ZipFile(path) as archive:
                for name in archive.namelist():
                    inventario.append(name)
                    if name.endswith("/"):
                        continue
                    if not _zip_member_safe(name):
                        erros.append(f"zip contem caminho inseguro: {name}")
                        continue
                    inner_suffix = Path(name).suffix.casefold()
                    data = archive.read(name)
                    if inner_suffix in {".txt", ".md", ".csv"}:
                        try:
                            text = data.decode("utf-8")
                        except UnicodeDecodeError:
                            text = data.decode("latin-1", errors="replace")
                        documentos.append({"nome": name, "texto": _truncate(text, max_chars)})
                    elif inner_suffix == ".pdf":
                        pdf_docs, error = _extrair_texto_pdf_bytes(name, data, max_chars=max_chars)
                        documentos.extend(pdf_docs)
                        if error:
                            erros.append(f"{name}: {error}")
                    elif inner_suffix == ".xlsx":
                        text, error = _extrair_texto_bytes_com_markitdown(name, data, inner_suffix)
                        if error:
                            erros.append(f"{name}: {error}")
                        else:
                            documentos.append({"nome": name, "texto": _truncate(text, max_chars)})
                    elif inner_suffix in WORD_EXTENSIONS_TO_PDF:
                        documentos.append(
                            {
                                "nome": name,
                                "tipo_convertido": "pdf",
                                "observacao": "Documento Word preservado para avaliacao visual por conversao para PDF no upload.",
                            }
                        )
                    else:
                        documentos.append({"nome": name, "tipo_nao_extraido": inner_suffix.lstrip(".") or "desconhecido"})
            return documentos, inventario, erros
        except zipfile.BadZipFile:
            return [], [path.name], ["zip invalido"]

    if suffix == ".zip":
        documentos: list[dict[str, Any]] = []
        inventario: list[str] = []
        erros: list[str] = []
        try:
            with zipfile.ZipFile(path) as archive:
                for name in archive.namelist():
                    inventario.append(name)
                    if name.endswith("/"):
                        continue
                    if not _zip_member_safe(name):
                        erros.append(f"zip contem caminho inseguro: {name}")
                        continue
                    inner_suffix = Path(name).suffix.casefold()
                    data = archive.read(name)
                    if inner_suffix in {".txt", ".md", ".csv"}:
                        try:
                            text = data.decode("utf-8")
                        except UnicodeDecodeError:
                            text = data.decode("latin-1", errors="replace")
                        documentos.append({"nome": name, "texto": _truncate(text, max_chars)})
                    elif inner_suffix == ".pdf":
                        pdf_docs, error = _extrair_texto_pdf_bytes(name, data, max_chars=max_chars)
                        documentos.extend(pdf_docs)
                        if error:
                            erros.append(f"{name}: {error}")
                    else:
                        documentos.append({"nome": name, "tipo_nao_extraido": inner_suffix.lstrip(".") or "desconhecido"})
            return documentos, inventario, erros
        except zipfile.BadZipFile:
            return [], [path.name], ["zip invalido"]

    if suffix in {".png", ".jpg", ".jpeg", ".gif", ".webp"}:
        return (
            [{"nome": path.name, "tipo": "imagem", "observacao": "imagem disponivel para upload ao Gemini; sem texto extraido"}],
            [path.name],
            [],
        )

    return [], [path.name], [f"tipo de evidencia nao suportado para extracao textual: {suffix or path.name}"]


def montar_pacote_evidencia(resolvidas: list[EvidenciaResolvida], *, max_chars: int) -> PacoteEvidencia:
    documentos: list[dict[str, Any]] = []
    inventario: list[str] = []
    arquivos: list[dict[str, Any]] = []
    erros: list[str] = []
    for resolved in resolvidas:
        if resolved.erro or not resolved.caminho:
            erros.append(resolved.erro or "evidencia sem caminho")
            continue
        file_hash = sha256_file(resolved.caminho)
        arquivos.append(
            {
                "nome_original": resolved.nome_decodificado,
                "caminho": str(resolved.caminho),
                "arquivo": resolved.caminho.name,
                "hash_sha256": file_hash,
                "tamanho_bytes": resolved.caminho.stat().st_size,
                "upload": resolved.upload,
            }
        )
        docs, inv, errs = normalizar_arquivo(resolved.caminho, max_chars=max_chars)
        documentos.extend(docs)
        inventario.extend(inv)
        erros.extend(errs)
    return PacoteEvidencia(documentos=documentos, inventario=inventario, arquivos=arquivos, erros=erros)


def respostas_declaradas(row: dict[str, Any]) -> dict[str, Any]:
    declared: dict[str, Any] = {}
    for key, value in row.items():
        if value in (None, ""):
            continue
        if key in SOURCE_COLUMNS:
            continue
        if "evi" in key or key.endswith("[filecount]"):
            continue
        if key.startswith("q"):
            declared[key] = value
    return declared


def montar_payload_prompt(
    avaliacao: AvaliacaoPrompt,
    row: dict[str, Any],
    auditado: str,
    pacote: PacoteEvidencia,
) -> str:
    payload = {
        "prompt_checklist": avaliacao.prompt_checklist,
        "contexto_avaliacao": {
            "auditado": auditado,
            "resposta_id": row.get("id"),
            "id_avaliacao": avaliacao.id_avaliacao,
            "questao_auditoria": avaliacao.questao_auditoria,
            "situacao_encontrada": avaliacao.situacao_encontrada,
            "item_evidencia": avaliacao.item_evidencia,
            "descricao_avaliacao": avaliacao.descricao_avaliacao,
            "referenciada_na_matriz_planejamento": avaliacao.referenciada_na_matriz,
        },
        "respostas_declaradas_do_questionario": respostas_declaradas(row),
        "pacote_evidencia": {
            "arquivos": pacote.arquivos,
            "inventario": pacote.inventario,
            "documentos": pacote.documentos,
            "erros_tecnicos": pacote.erros,
        },
        "contrato_de_saida": {
            "resultado": "Conforme ou Não conforme",
            "justificativa_sintetica": "fundamento objetivo da decisão",
            "elementos_comprovados": [],
            "lacunas": [],
            "inconsistencias": [],
        },
        "instrucao_final": (
            "Responda exclusivamente em JSON valido. Nao inclua texto fora do JSON. "
            "Use somente os dados fornecidos no pacote de evidencia e nas respostas declaradas."
        ),
    }
    return dumps_json(payload, indent=2)


def resultado_nao_conforme(justificativa: str, *, lacunas: list[str] | None = None) -> dict[str, Any]:
    return {
        "resultado": RESULTADO_NAO_CONFORME,
        "justificativa_sintetica": justificativa,
        "elementos_comprovados": [],
        "lacunas": lacunas or [justificativa],
        "inconsistencias": [],
    }


def resultado_fake() -> dict[str, Any]:
    return resultado_nao_conforme(
        "Provider fake nao executa avaliacao substantiva; use openrouter ou gemini para avaliar a evidencia.",
        lacunas=["Analise real de IA nao executada."],
    )


def _extrair_bloco_json(text: str) -> str:
    match = re.search(r"```(?:json)?\s*(.*?)```", text, flags=re.IGNORECASE | re.DOTALL)
    if match:
        return match.group(1).strip()
    return text.strip()


def carregar_json_modelo(text: str) -> dict[str, Any]:
    raw = _extrair_bloco_json(text)
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError:
        try:
            from json_repair import repair_json
        except ModuleNotFoundError as exc:
            raise ValueError("resposta nao e JSON valido e json-repair nao esta instalado") from exc
        repaired = repair_json(raw)
        if isinstance(repaired, str):
            parsed = json.loads(repaired)
        else:
            parsed = repaired
    if not isinstance(parsed, dict):
        raise ValueError("resposta do modelo precisa ser um objeto JSON")
    return parsed


def normalizar_resultado_texto(value: Any) -> str:
    raw = str(value or "").strip().casefold()
    raw_ascii = unicodedata.normalize("NFKD", raw).encode("ascii", "ignore").decode("ascii")
    if raw_ascii in {"conforme", "sim"}:
        return RESULTADO_CONFORME
    if raw_ascii in {"nao conforme", "naoconforme", "nao_conforme", "non compliant", "non-compliant"}:
        return RESULTADO_NAO_CONFORME
    raise ValueError(f"resultado invalido: {value!r}")


def _as_list(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, list):
        return [str(item) for item in value if item not in (None, "")]
    return [str(value)]


def validar_resultado_modelo(payload: dict[str, Any]) -> dict[str, Any]:
    result = normalizar_resultado_texto(payload.get("resultado"))
    justificativa = str(payload.get("justificativa_sintetica") or payload.get("justificativa") or "").strip()
    if not justificativa:
        raise ValueError("resultado sem justificativa_sintetica")
    return {
        "resultado": result,
        "justificativa_sintetica": justificativa,
        "elementos_comprovados": _as_list(payload.get("elementos_comprovados")),
        "lacunas": _as_list(payload.get("lacunas")),
        "inconsistencias": _as_list(payload.get("inconsistencias")),
    }


def resultado_legado_indica_erro_tecnico(resultado: dict[str, Any] | None) -> str:
    if not isinstance(resultado, dict):
        return ""
    termos = [
        "markitdown",
        "erro tecnico",
        "erro técnico",
        "nao possui acesso ao conteudo",
        "não possui acesso ao conteúdo",
        "conteudo do arquivo inacessivel",
        "conteúdo do arquivo inacessível",
    ]
    campos = [
        resultado.get("justificativa_sintetica", ""),
        " ".join(str(valor) for valor in resultado.get("lacunas", []) if valor is not None)
        if isinstance(resultado.get("lacunas"), list)
        else str(resultado.get("lacunas", "")),
        " ".join(str(valor) for valor in resultado.get("inconsistencias", []) if valor is not None)
        if isinstance(resultado.get("inconsistencias"), list)
        else str(resultado.get("inconsistencias", "")),
    ]
    texto = " ".join(campos).casefold()
    if any(termo in texto for termo in termos):
        return "modelo indicou erro tecnico de acesso/processamento da evidencia"
    return ""


def parse_retry_after(value: str | None, *, now: Callable[[], float] = time.time) -> float | None:
    if not value:
        return None
    raw = value.strip()
    if not raw:
        return None
    try:
        return max(0.0, float(raw))
    except ValueError:
        pass
    try:
        retry_at = parsedate_to_datetime(raw)
    except (TypeError, ValueError, IndexError, OverflowError):
        return None
    if retry_at.tzinfo is None:
        retry_at = retry_at.replace(tzinfo=dt.timezone.utc)
    return max(0.0, retry_at.timestamp() - now())


def _header_retry_after(headers: Any) -> str | None:
    getter = getattr(headers, "get", None)
    if callable(getter):
        return getter("Retry-After") or getter("retry-after")
    if isinstance(headers, dict):
        return headers.get("Retry-After") or headers.get("retry-after")
    return None


def _status_from_exception(exc: BaseException) -> Any:
    status = getattr(exc, "code", None) or getattr(exc, "status", None) or getattr(exc, "status_code", None)
    response = getattr(exc, "response", None)
    if status is None and response is not None:
        status = getattr(response, "status_code", None) or getattr(response, "status", None)
    return status


def _retry_after_from_exception(exc: BaseException) -> float | None:
    response = getattr(exc, "response", None)
    retry_after = _header_retry_after(getattr(exc, "headers", None))
    if retry_after is None and response is not None:
        retry_after = _header_retry_after(getattr(response, "headers", None))
    return parse_retry_after(retry_after)


def executar_com_retry_transiente(
    func: Callable[[], Any],
    *,
    max_retries: int = 3,
    sleeper: Callable[[float], None] = time.sleep,
) -> Any:
    attempt = 0
    while True:
        try:
            return func()
        except Exception as exc:
            if _status_from_exception(exc) not in RETRYABLE_PROVIDER_STATUSES or attempt >= max_retries:
                raise
            delay = _retry_after_from_exception(exc)
            if delay is None:
                delay = DEFAULT_TRANSIENT_RETRY_DELAYS[min(attempt, len(DEFAULT_TRANSIENT_RETRY_DELAYS) - 1)]
            sleeper(delay)
            attempt += 1


def json_schema_response_format() -> dict[str, Any]:
    return {
        "type": "json_schema",
        "json_schema": {
            "name": "resultado_avaliacao_evidencia",
            "strict": True,
            "schema": {
                "type": "object",
                "additionalProperties": False,
                "properties": {
                    "resultado": {"type": "string", "enum": ["Conforme", "Não conforme"]},
                    "justificativa_sintetica": {"type": "string"},
                    "elementos_comprovados": {"type": "array", "items": {"type": "string"}},
                    "lacunas": {"type": "array", "items": {"type": "string"}},
                    "inconsistencias": {"type": "array", "items": {"type": "string"}},
                },
                "required": [
                    "resultado",
                    "justificativa_sintetica",
                    "elementos_comprovados",
                    "lacunas",
                    "inconsistencias",
                ],
            },
        },
    }


def modelo_openrouter_suporta_pdf_nativo(model: str) -> bool:
    normalizado = model.strip().casefold()
    return (
        normalizado.startswith("google/")
        or "gemini" in normalizado
        or normalizado.startswith("openai/")
        or "chatgpt" in normalizado
        or "/gpt-" in normalizado
        or normalizado.startswith("gpt-")
        or re.search(r"(^|/|:)o[134](?:-|$)", normalizado) is not None
    )


def arquivos_pdf_openrouter(upload_files: list[str]) -> list[Path]:
    arquivos = []
    for value in upload_files:
        path = Path(value)
        if path.suffix.casefold() == ".pdf" and path.is_file():
            arquivos.append(path)
    return arquivos


def arquivo_pdf_para_openrouter(path: Path) -> dict[str, Any]:
    data_url = "data:application/pdf;base64," + base64.b64encode(path.read_bytes()).decode("ascii")
    return {
        "type": "file",
        "file": {
            "filename": path.name,
            "file_data": data_url,
        },
    }


def chamar_openrouter(
    *,
    api_key: str,
    model: str,
    prompt_payload: str,
    upload_files: list[str] | None = None,
    reasoning_effort: str = "",
) -> tuple[dict[str, Any], str]:
    pdf_files = arquivos_pdf_openrouter(upload_files or [])
    usar_pdf_nativo = bool(pdf_files) and modelo_openrouter_suporta_pdf_nativo(model)
    content: str | list[dict[str, Any]]
    if usar_pdf_nativo:
        content = [{"type": "text", "text": prompt_payload}]
        content.extend(arquivo_pdf_para_openrouter(path) for path in pdf_files)
    else:
        content = prompt_payload
    body = {
        "model": model,
        "messages": [{"role": "user", "content": content}],
        "temperature": 0,
        "response_format": json_schema_response_format(),
    }
    if reasoning_effort:
        body["reasoning"] = {"effort": reasoning_effort}
    if usar_pdf_nativo:
        body["plugins"] = [{"id": "file-parser", "pdf": {"engine": "native"}}]

    def call() -> dict[str, Any]:
        request = urllib.request.Request(
            "https://openrouter.ai/api/v1/chat/completions",
            data=dumps_json(body).encode("utf-8"),
            headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
            method="POST",
        )
        with urllib.request.urlopen(request, timeout=180) as response:
            return json.loads(response.read().decode("utf-8"))

    payload = executar_com_retry_transiente(call)
    content = payload["choices"][0]["message"]["content"]
    return validar_resultado_modelo(carregar_json_modelo(content)), content


def formatar_erro_http(exc: BaseException) -> str:
    if isinstance(exc, urllib.error.HTTPError):
        detail = ""
        try:
            detail = exc.read().decode("utf-8", errors="replace").strip()
        except Exception:
            detail = ""
        if detail:
            return f"HTTP Error {exc.code}: {exc.reason}; body: {detail[:2000]}"
    return str(exc)


def _slug_ascii(value: str, *, fallback: str = "evidencia", max_len: int = 80) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    ascii_text = normalized.encode("ascii", "ignore").decode("ascii")
    slug = re.sub(r"[^A-Za-z0-9._-]+", "-", ascii_text).strip("._-").lower()
    slug = re.sub(r"-{2,}", "-", slug)
    if not slug:
        slug = fallback
    return slug[:max_len].strip("._-") or fallback


def _unique_path(directory: Path, name: str) -> Path:
    candidate = directory / name
    if not candidate.exists():
        return candidate
    stem = candidate.stem
    suffix = candidate.suffix
    counter = 2
    while True:
        next_candidate = directory / f"{stem}-{counter}{suffix}"
        if not next_candidate.exists():
            return next_candidate
        counter += 1


def _copy_for_upload(source: Path, target_dir: Path, name_hint: str | None = None) -> Path:
    target_dir.mkdir(parents=True, exist_ok=True)
    suffix = source.suffix.lower()
    stem = Path(name_hint or source.name).stem or "evidencia"
    digest = hashlib.sha256(str(source).encode("utf-8", errors="ignore")).hexdigest()[:10]
    target = _unique_path(target_dir, f"{_slug_ascii(stem)}-{digest}{suffix}")
    shutil.copyfile(source, target)
    return target


def _write_text_for_upload(target_dir: Path, name_hint: str, text: str) -> Path:
    target_dir.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256(name_hint.encode("utf-8", errors="ignore")).hexdigest()[:10]
    target = _unique_path(target_dir, f"{_slug_ascii(Path(name_hint).stem)}-{digest}.txt")
    target.write_text(text, encoding="utf-8")
    return target


def _convert_word_to_pdf_for_upload(source: Path, target_dir: Path, name_hint: str | None = None) -> tuple[Path | None, str]:
    converter = shutil.which("soffice") or shutil.which("libreoffice")
    if not converter:
        return None, "LibreOffice/soffice nao encontrado para converter documento Word em PDF"
    target_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        try:
            result = subprocess.run(
                [
                    converter,
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(tmp_dir),
                    str(source),
                ],
                check=False,
                capture_output=True,
                text=True,
                timeout=120,
            )
        except subprocess.TimeoutExpired:
            return None, "timeout ao converter documento Word em PDF"
        except OSError as exc:
            return None, f"erro ao executar conversor de documento Word para PDF: {exc}"
        if result.returncode != 0:
            detail = (result.stderr or result.stdout or "").strip()
            return None, f"erro ao converter documento Word em PDF: {detail or result.returncode}"
        converted = tmp_dir / f"{source.stem}.pdf"
        if not converted.is_file():
            candidates = sorted(tmp_dir.glob("*.pdf"))
            if not candidates:
                return None, "conversor de documento Word para PDF nao gerou arquivo PDF"
            converted = candidates[0]
        original = name_hint or source.name
        digest = hashlib.sha256(original.encode("utf-8", errors="ignore")).hexdigest()[:10]
        target = _unique_path(target_dir, f"{_slug_ascii(Path(original).stem)}-{digest}.pdf")
        shutil.copyfile(converted, target)
        return target, ""


def preparar_uploads_gemini(paths: list[Path], target_dir: Path) -> list[str]:
    uploadable = {".pdf", ".txt", ".md", ".csv", ".png", ".jpg", ".jpeg", ".webp"}
    prepared: list[str] = []
    for path in paths:
        suffix = path.suffix.casefold()
        if suffix in uploadable:
            prepared.append(str(_copy_for_upload(path, target_dir)))
        elif suffix in WORD_EXTENSIONS_TO_PDF:
            converted, error = _convert_word_to_pdf_for_upload(path, target_dir)
            if error:
                raise RuntimeError(error)
            if converted:
                prepared.append(str(converted))
        elif suffix == ".xlsx":
            text, error = _extrair_texto_com_markitdown(path)
            if not error:
                prepared.append(str(_write_text_for_upload(target_dir, path.name, text)))
        elif suffix == ".zip":
            with zipfile.ZipFile(path) as archive:
                for name in archive.namelist():
                    if name.endswith("/") or not _zip_member_safe(name):
                        continue
                    member_suffix = Path(name).suffix.casefold()
                    if member_suffix in uploadable:
                        target = _unique_path(target_dir, f"{_slug_ascii(Path(name).stem)}{member_suffix}")
                        target.write_bytes(archive.read(name))
                        prepared.append(str(target))
                    elif member_suffix in WORD_EXTENSIONS_TO_PDF:
                        staging_dir = target_dir / "_word_sources"
                        staging_dir.mkdir(parents=True, exist_ok=True)
                        data = archive.read(name)
                        staging = _unique_path(staging_dir, f"{_slug_ascii(Path(name).stem)}{member_suffix}")
                        staging.write_bytes(data)
                        converted, error = _convert_word_to_pdf_for_upload(staging, target_dir, name)
                        if error:
                            raise RuntimeError(f"{name}: {error}")
                        if converted:
                            prepared.append(str(converted))
                    elif member_suffix == ".xlsx":
                        data = archive.read(name)
                        text, error = _extrair_texto_bytes_com_markitdown(name, data, member_suffix)
                        if not error:
                            prepared.append(str(_write_text_for_upload(target_dir, name, text)))
    return prepared


def chamar_gemini(*, api_key: str, model: str, prompt_payload: str, upload_files: list[str]) -> tuple[dict[str, Any], str]:
    try:
        from google import genai
    except Exception as exc:
        raise RuntimeError(f"google-genai nao disponivel: {exc}") from exc

    client = genai.Client(api_key=api_key)
    uploaded = []
    for file_path in upload_files:
        uploaded.append(executar_com_retry_transiente(lambda file_path=file_path: client.files.upload(file=file_path)))

    response = executar_com_retry_transiente(
        lambda: client.models.generate_content(
            model=model,
            contents=[prompt_payload, *uploaded],
            config={"response_mime_type": "application/json", "temperature": 0},
        )
    )
    raw_text = response.text
    return validar_resultado_modelo(carregar_json_modelo(raw_text)), raw_text


def executar_provider(
    *,
    provider: str,
    model: str,
    api_key: str,
    prompt_payload: str,
    evidence_paths: list[Path],
    reasoning_effort: str = "",
) -> dict[str, Any]:
    if provider == "dry-run":
        return {"status": "dry_run", "result": None, "prompt_payload": prompt_payload}
    if provider == "fake":
        return {"status": "completed", "result": resultado_fake(), "raw_response_excerpt": ""}
    if provider == "openrouter":
        if not api_key:
            return {"status": "error", "error": "OPENROUTER_API_KEY nao configurada"}
        try:
            with tempfile.TemporaryDirectory() as tmp:
                upload_files = preparar_uploads_gemini(evidence_paths, Path(tmp))
                result, raw = chamar_openrouter(
                    api_key=api_key,
                    model=model,
                    prompt_payload=prompt_payload,
                    upload_files=upload_files,
                    reasoning_effort=reasoning_effort,
                )
            return {"status": "completed", "result": result, "raw_response_excerpt": raw[:2000]}
        except Exception as exc:
            return {"status": "error", "error": f"erro ao chamar OpenRouter: {formatar_erro_http(exc)}"}
    if provider == "gemini":
        if not api_key:
            return {"status": "error", "error": "GEMINI_API_KEY nao configurada"}
        try:
            with tempfile.TemporaryDirectory() as tmp:
                upload_files = preparar_uploads_gemini(evidence_paths, Path(tmp))
                result, raw = chamar_gemini(
                    api_key=api_key,
                    model=model,
                    prompt_payload=prompt_payload,
                    upload_files=upload_files,
                )
            return {"status": "completed", "result": result, "raw_response_excerpt": raw[:2000]}
        except Exception as exc:
            return {"status": "error", "error": f"erro ao chamar Gemini: {exc}"}
    return {"status": "error", "error": f"provider nao suportado: {provider}"}


def calcular_identidade(
    *,
    row: dict[str, Any],
    avaliacao: AvaliacaoPrompt,
    upload_value: Any,
    evidence_hashes: list[str],
    provider: str,
    model: str,
    reasoning_effort: str = "",
) -> str:
    payload = {
        "resposta_id": row.get("id"),
        "token": row.get("token"),
        "id_avaliacao": avaliacao.id_avaliacao,
        "item_evidencia": avaliacao.item_evidencia,
        "upload_value": upload_value,
        "evidence_hashes": evidence_hashes,
        "prompt_hash": avaliacao.prompt_hash,
        "provider": provider,
        "model": model,
        "reasoning_effort": reasoning_effort,
    }
    return hashlib.sha256(dumps_json(payload).encode("utf-8")).hexdigest()


def load_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.is_file():
        return []
    records = []
    for line_number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), start=1):
        if not line.strip():
            continue
        try:
            records.append(json.loads(line))
        except json.JSONDecodeError as exc:
            raise ValueError(f"JSONL invalido em {path}:{line_number}: {exc.msg}") from exc
    return records


def records_by_identity(records: Iterable[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    result = {}
    for record in records:
        identity = record.get("identity")
        if isinstance(identity, str) and identity:
            result[identity] = record
    return result


def append_jsonl(path: Path, record: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as file:
        file.write(dumps_json(record))
        file.write("\n")


def escrever_json_lista(records: list[dict[str, Any]], json_path: Path) -> int:
    json_path.parent.mkdir(parents=True, exist_ok=True)
    json_path.write_text(dumps_json(records, indent=2), encoding="utf-8")
    return len(records)


def _resultado_record(record: dict[str, Any]) -> str:
    if record.get("status") == "dry_run":
        return RESULTADO_PENDENTE
    result = record.get("result") if isinstance(record.get("result"), dict) else {}
    if record.get("status") == "completed" and result.get("resultado") in RESULTADOS_VALIDOS:
        return str(result.get("resultado"))
    return RESULTADO_NAO_CONFORME


def _justificativa_record(record: dict[str, Any]) -> str:
    if record.get("status") == "dry_run":
        return "Dry-run: prompt montado, mas provider nao executado."
    result = record.get("result") if isinstance(record.get("result"), dict) else {}
    if result.get("justificativa_sintetica"):
        return str(result.get("justificativa_sintetica"))
    if record.get("error"):
        return str(record.get("error"))
    return ""


def _join(parts: Iterable[str]) -> str:
    return "\n".join(part for part in parts if part)


def consolidar_item(records: list[dict[str, Any]]) -> tuple[str, str]:
    if not records:
        return "", ""
    resultados = [_resultado_record(record) for record in records]
    if RESULTADO_PENDENTE in resultados:
        consolidated = RESULTADO_PENDENTE
    elif all(result == RESULTADO_CONFORME for result in resultados):
        consolidated = RESULTADO_CONFORME
    else:
        consolidated = RESULTADO_NAO_CONFORME
    justification = _join(
        f"[{record.get('id_avaliacao', '')}] {_justificativa_record(record)}"
        for record in sorted(records, key=lambda item: str(item.get("id_avaliacao") or ""))
    )
    return consolidated, justification


def gerar_planilha_resultados(
    *,
    responses: list[dict[str, Any]],
    avaliacoes: list[AvaliacaoPrompt],
    records: list[dict[str, Any]],
    validation: dict[str, Any],
    destino: Path,
    auditado_col: str,
    include_unsubmitted: bool,
) -> None:
    evidence_items = []
    for avaliacao in avaliacoes:
        if avaliacao.item_evidencia not in evidence_items:
            evidence_items.append(avaliacao.item_evidencia)

    by_response_item: dict[tuple[str, str], list[dict[str, Any]]] = {}
    for record in records:
        key = (str(record.get("resposta_id") or ""), str(record.get("item_evidencia") or ""))
        by_response_item.setdefault(key, []).append(record)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Resultados"
    headers = ["id", "token", "auditado", "submitdate", "firstname", "lastname", "email"]
    for item in evidence_items:
        headers.extend([item, f"{item}_justificativa"])
    sheet.append(headers)

    for row in responses:
        if not include_unsubmitted and not row.get("submitdate"):
            continue
        response_id = str(row.get("id") or "")
        auditado = str(row.get(auditado_col) or row.get("firstname") or row.get("token") or response_id).strip()
        values = [
            row.get("id"),
            row.get("token"),
            auditado,
            row.get("submitdate"),
            row.get("firstname"),
            row.get("lastname"),
            row.get("email"),
        ]
        for item in evidence_items:
            result, justification = consolidar_item(by_response_item.get((response_id, item), []))
            values.extend([result, justification])
        sheet.append(values)

    details = workbook.create_sheet("Avaliacoes")
    detail_headers = [
        "resposta_id",
        "auditado",
        "id_avaliacao",
        "questao_auditoria",
        "situacao_encontrada",
        "item_evidencia",
        "resultado",
        "justificativa",
        "status",
        "error",
        "provider",
        "model",
        "referenciada_na_matriz",
        "evidencias",
        "finished_at",
    ]
    details.append(detail_headers)
    for record in sorted(records, key=lambda item: (str(item.get("auditado") or ""), str(item.get("id_avaliacao") or ""))):
        details.append(
            [
                record.get("resposta_id"),
                record.get("auditado"),
                record.get("id_avaliacao"),
                record.get("questao_auditoria"),
                record.get("situacao_encontrada"),
                record.get("item_evidencia"),
                _resultado_record(record),
                _justificativa_record(record),
                record.get("status"),
                record.get("error"),
                record.get("provider"),
                record.get("model"),
                record.get("referenciada_na_matriz"),
                _join(record.get("evidencias") or []),
                record.get("finished_at"),
            ]
        )

    validation_sheet = workbook.create_sheet("Validacao_Matriz")
    validation_sheet.append(["campo", "valor"])
    for key, value in validation.items():
        if isinstance(value, list):
            validation_sheet.append([key, "\n".join(str(item) for item in value)])
        else:
            validation_sheet.append([key, str(value)])

    destino.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destino)


def build_record(
    *,
    identity: str,
    row: dict[str, Any],
    auditado: str,
    avaliacao: AvaliacaoPrompt,
    provider: str,
    model: str,
    status: str,
    result: dict[str, Any] | None,
    error: str = "",
    evidencias: list[str] | None = None,
    prompt_payload: str | None = None,
    raw_response_excerpt: str = "",
) -> dict[str, Any]:
    record = {
        "identity": identity,
        "status": status,
        "resposta_id": row.get("id"),
        "token": row.get("token"),
        "auditado": auditado,
        "id_avaliacao": avaliacao.id_avaliacao,
        "questao_auditoria": avaliacao.questao_auditoria,
        "situacao_encontrada": avaliacao.situacao_encontrada,
        "item_evidencia": avaliacao.item_evidencia,
        "descricao_avaliacao": avaliacao.descricao_avaliacao,
        "referenciada_na_matriz": avaliacao.referenciada_na_matriz,
        "provider": provider,
        "model": model,
        "result": result,
        "error": error,
        "evidencias": evidencias or [],
        "prompt_hash": avaliacao.prompt_hash,
        "raw_response_excerpt": raw_response_excerpt,
        "finished_at": dt.datetime.now(dt.timezone.utc).isoformat(),
    }
    if prompt_payload is not None:
        record["prompt_payload"] = prompt_payload
    return record


def processar(
    *,
    args: argparse.Namespace,
    responses_headers: list[str],
    responses: list[dict[str, Any]],
    avaliacoes: list[AvaliacaoPrompt],
    validation: dict[str, Any],
) -> int:
    out_dir = Path(args.out_dir)
    jsonl_path = out_dir / "resultados_avaliacao_evidencias.jsonl"
    json_path = out_dir / "resultados_avaliacao_evidencias.json"
    xlsx_path = out_dir / "resultados_avaliacao_evidencias.xlsx"
    existing_records = records_by_identity(load_jsonl(jsonl_path))
    evidence_index = {name: index for index, name in enumerate(colunas_evidencia(responses_headers), start=1)}
    limiter = RequestsPerMinuteLimiter(args.rpm)
    api_key = args.api_key
    if not api_key and args.provider == "openrouter":
        api_key = os.environ.get("OPENROUTER_API_KEY", "")
    if not api_key and args.provider == "gemini":
        api_key = os.environ.get("GEMINI_API_KEY", "")

    total = 0
    skipped = 0
    errors = 0
    completed = 0
    dry_runs = 0

    for row_index, row in enumerate(responses, start=2):
        if not args.include_unsubmitted and not row.get("submitdate"):
            continue
        if args.only_auditado:
            row_auditado = str(row.get(args.auditado_col) or row.get("firstname") or "").strip()
            if row_auditado != args.only_auditado:
                continue
        auditado = str(row.get(args.auditado_col) or row.get("firstname") or row.get("token") or row.get("id")).strip()
        for avaliacao in avaliacoes:
            if args.only_id_avaliacao and avaliacao.id_avaliacao != args.only_id_avaliacao:
                continue
            upload_value = row.get(avaliacao.item_evidencia)
            uploads, upload_error = parse_uploads(upload_value)
            resolvidas: list[EvidenciaResolvida] = []
            if uploads:
                for upload in uploads:
                    resolvidas.append(
                        resolver_upload(
                            Path(args.evidencias),
                            auditado,
                            upload,
                            resposta_id=row.get("id"),
                            evidence_index=evidence_index.get(avaliacao.item_evidencia),
                            recursive=not args.no_recursive_evidence_search,
                        )
                    )
            pacote = montar_pacote_evidencia(resolvidas, max_chars=args.max_document_chars)
            evidence_hashes = [arquivo["hash_sha256"] for arquivo in pacote.arquivos if arquivo.get("hash_sha256")]
            identity = calcular_identidade(
                row=row,
                avaliacao=avaliacao,
                upload_value=upload_value,
                evidence_hashes=evidence_hashes,
                provider=args.provider,
                model=args.model,
                reasoning_effort=args.reasoning_effort,
            )
            if identity in existing_records and existing_records[identity].get("status") in {"completed", "dry_run"}:
                skipped += 1
                continue

            log_event(
                "avaliacao_started",
                "Iniciando avaliacao de evidencia.",
                quiet=args.quiet,
                resposta_id=row.get("id"),
                row_index=row_index,
                auditado=auditado,
                id_avaliacao=avaliacao.id_avaliacao,
                item_evidencia=avaliacao.item_evidencia,
                provider=args.provider,
                model=args.model,
            )

            if upload_error:
                result = resultado_nao_conforme(upload_error)
                record = build_record(
                    identity=identity,
                    row=row,
                    auditado=auditado,
                    avaliacao=avaliacao,
                    provider=args.provider,
                    model=args.model,
                    status="completed",
                    result=result,
                    evidencias=[],
                )
            elif not uploads:
                result = resultado_nao_conforme(
                    f"Nenhuma evidencia foi anexada em {avaliacao.item_evidencia}.",
                    lacunas=[f"Coluna {avaliacao.item_evidencia} sem anexo."],
                )
                record = build_record(
                    identity=identity,
                    row=row,
                    auditado=auditado,
                    avaliacao=avaliacao,
                    provider=args.provider,
                    model=args.model,
                    status="completed",
                    result=result,
                    evidencias=[],
                )
            elif any(resolved.erro for resolved in resolvidas):
                missing = [resolved.erro for resolved in resolvidas if resolved.erro]
                result = resultado_nao_conforme("Uma ou mais evidencias anexadas nao foram localizadas.", lacunas=missing)
                record = build_record(
                    identity=identity,
                    row=row,
                    auditado=auditado,
                    avaliacao=avaliacao,
                    provider=args.provider,
                    model=args.model,
                    status="completed",
                    result=result,
                    evidencias=[str(resolved.caminho) for resolved in resolvidas if resolved.caminho],
                )
            elif pacote.erros:
                record = build_record(
                    identity=identity,
                    row=row,
                    auditado=auditado,
                    avaliacao=avaliacao,
                    provider=args.provider,
                    model=args.model,
                    status="error",
                    result=None,
                    error=f"erro tecnico ao processar evidencia: {'; '.join(pacote.erros)}",
                    evidencias=[str(resolved.caminho) for resolved in resolvidas if resolved.caminho],
                )
            else:
                prompt_payload = montar_payload_prompt(avaliacao, row, auditado, pacote)
                evidence_paths = [resolved.caminho for resolved in resolvidas if resolved.caminho]
                if args.provider in REMOTE_PROVIDERS:
                    wait_seconds = limiter.wait_and_mark()
                    if wait_seconds:
                        log_event(
                            "rate_limit_wait",
                            "Aguardando limite de requests por minuto.",
                            quiet=args.quiet,
                            wait_seconds=round(wait_seconds, 3),
                            rpm=args.rpm,
                    )
                provider_result = executar_provider(
                    provider=args.provider,
                    model=args.model,
                    api_key=api_key,
                    prompt_payload=prompt_payload,
                    evidence_paths=evidence_paths,
                    reasoning_effort=args.reasoning_effort,
                )
                erro_tecnico_resultado = resultado_legado_indica_erro_tecnico(provider_result.get("result"))
                if provider_result.get("status") == "completed" and erro_tecnico_resultado:
                    provider_result = {
                        "status": "error",
                        "error": erro_tecnico_resultado,
                        "result": None,
                        "raw_response_excerpt": provider_result.get("raw_response_excerpt", ""),
                    }
                record = build_record(
                    identity=identity,
                    row=row,
                    auditado=auditado,
                    avaliacao=avaliacao,
                    provider=args.provider,
                    model=args.model,
                    status=provider_result.get("status", "error"),
                    result=provider_result.get("result"),
                    error=provider_result.get("error", ""),
                    evidencias=[str(path) for path in evidence_paths],
                    prompt_payload=prompt_payload if args.store_prompts or args.provider == "dry-run" else None,
                    raw_response_excerpt=provider_result.get("raw_response_excerpt", ""),
                )

            append_jsonl(jsonl_path, record)
            existing_records[identity] = record
            total += 1
            if record["status"] == "completed":
                completed += 1
            elif record["status"] == "dry_run":
                dry_runs += 1
            else:
                errors += 1

            log_event(
                "avaliacao_finished",
                "Avaliacao registrada.",
                quiet=args.quiet,
                status=record["status"],
                resultado=_resultado_record(record),
                resposta_id=row.get("id"),
                auditado=auditado,
                id_avaliacao=avaliacao.id_avaliacao,
                item_evidencia=avaliacao.item_evidencia,
                error=record.get("error", ""),
            )

            if args.limit and total >= args.limit:
                break
        if args.limit and total >= args.limit:
            break

    latest_records = list(records_by_identity(load_jsonl(jsonl_path)).values())
    current_records = [
        record
        for record in latest_records
        if record.get("provider") == args.provider and record.get("model") == args.model
    ]
    total_json = escrever_json_lista(current_records, json_path)
    gerar_planilha_resultados(
        responses=responses,
        avaliacoes=avaliacoes,
        records=current_records,
        validation=validation,
        destino=xlsx_path,
        auditado_col=args.auditado_col,
        include_unsubmitted=args.include_unsubmitted,
    )
    log_event(
        "processamento_finished",
        "Processamento finalizado.",
        quiet=args.quiet,
        processados=total,
        pulados_por_checkpoint=skipped,
        concluidos=completed,
        dry_runs=dry_runs,
        erros=errors,
        jsonl=str(jsonl_path),
        json=str(json_path),
        registros_json=total_json,
        xlsx=str(xlsx_path),
    )
    return 0


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Avalia evidencias do questionario iGovTI 2026.")
    parser.add_argument("evidencias", help="Pasta que contem os arquivos fisicos das evidencias.")
    parser.add_argument("--respostas", default=str(DEFAULT_RESPOSTAS), help="Planilha de respostas do questionario.")
    parser.add_argument("--matriz-verificacao", default=str(DEFAULT_MATRIZ_VERIFICACAO), help="Planilha de prompts.")
    parser.add_argument("--matriz-planejamento", default=str(DEFAULT_MATRIZ_PLANEJAMENTO), help="Matriz de planejamento em Markdown.")
    parser.add_argument("--out-dir", default=str(DEFAULT_OUT_DIR), help="Diretorio de saida.")
    parser.add_argument("--provider", choices=["dry-run", "fake", "openrouter", "gemini"], default="dry-run")
    parser.add_argument("--dry-run", action="store_true", help="Monta prompts e saidas sem chamar provider remoto.")
    parser.add_argument("--model", default="", help="Modelo do provider. Padrao depende do provider.")
    parser.add_argument(
        "--reasoning",
        "--reasoning-effort",
        dest="reasoning_effort",
        choices=["low", "medium", "high"],
        default="",
        help="Nivel de reasoning a solicitar ao provider quando suportado: low, medium ou high.",
    )
    parser.add_argument("--api-key", default="", help="Chave de API. Se omitida, usa OPENROUTER_API_KEY ou GEMINI_API_KEY.")
    parser.add_argument("--rpm", type=validar_rpm, default=12, help="Limite de requests por minuto; use 0 para desativar.")
    parser.add_argument("--include-unsubmitted", action="store_true", help="Inclui respostas sem submitdate.")
    parser.add_argument("--auditado-col", default="firstname", help="Coluna que identifica o auditado e subpasta de evidencias.")
    parser.add_argument("--only-auditado", default="", help="Processa somente o auditado informado.")
    parser.add_argument("--only-id-avaliacao", default="", help="Processa somente um id_avaliacao.")
    parser.add_argument("--limit", type=int, default=0, help="Limita a quantidade de avaliacoes processadas nesta execucao.")
    parser.add_argument("--max-document-chars", type=int, default=60_000, help="Limite de caracteres por documento enviado no prompt.")
    parser.add_argument("--store-prompts", action="store_true", help="Armazena prompt_payload tambem em execucoes reais.")
    parser.add_argument("--no-recursive-evidence-search", action="store_true", help="Nao procura evidencias recursivamente na pasta.")
    parser.add_argument("--strict-matrix", action="store_true", help="Falha se IDs da planilha de prompts divergirem dos targets da matriz.")
    parser.add_argument("--quiet", action="store_true", help="Nao imprime logs estruturados.")
    args = parser.parse_args(argv)
    if args.dry_run:
        args.provider = "dry-run"
    if not args.model:
        if args.provider == "openrouter":
            args.model = "google/gemini-2.5-flash"
        elif args.provider == "gemini":
            args.model = "gemini-2.5-flash"
        elif args.provider == "fake":
            args.model = "fake"
        else:
            args.model = "dry-run"
    if args.limit < 0:
        parser.error("--limit deve ser maior ou igual a zero")
    if args.max_document_chars <= 0:
        parser.error("--max-document-chars deve ser maior que zero")
    return args


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    respostas_path = Path(args.respostas)
    evidencias_path = Path(args.evidencias)
    matriz_verificacao_path = Path(args.matriz_verificacao)
    matriz_planejamento_path = Path(args.matriz_planejamento)
    if not respostas_path.is_file():
        raise FileNotFoundError(f"Planilha de respostas nao encontrada: {respostas_path}")
    if not evidencias_path.is_dir():
        raise FileNotFoundError(f"Pasta de evidencias nao encontrada: {evidencias_path}")
    if not matriz_verificacao_path.is_file():
        raise FileNotFoundError(f"Planilha de verificacao nao encontrada: {matriz_verificacao_path}")
    if not matriz_planejamento_path.is_file():
        raise FileNotFoundError(f"Matriz de planejamento nao encontrada: {matriz_planejamento_path}")

    response_headers, responses = rows_from_xlsx(respostas_path)
    matrix_targets = extrair_targets_matriz(matriz_planejamento_path)
    avaliacoes = carregar_avaliacoes(matriz_verificacao_path, matrix_targets)
    validation = validar_matriz(avaliacoes, matrix_targets)

    if args.strict_matrix and (
        validation["ids_planilha_nao_referenciados_na_matriz"] or validation["targets_matriz_sem_prompt_na_planilha"]
    ):
        raise ValueError(
            "IDs da planilha de prompts divergem dos targets avaliacao[...] da matriz. "
            "Execute sem --strict-matrix para processar todas as linhas da planilha."
        )

    log_event(
        "validation_completed",
        "Validacao cruzada entre matriz e planilha de prompts concluida.",
        quiet=args.quiet,
        total_respostas=len(responses),
        total_avaliacoes=len(avaliacoes),
        ids_planilha_nao_referenciados=len(validation["ids_planilha_nao_referenciados_na_matriz"]),
        targets_matriz_sem_prompt=len(validation["targets_matriz_sem_prompt_na_planilha"]),
    )
    return processar(
        args=args,
        responses_headers=response_headers,
        responses=responses,
        avaliacoes=avaliacoes,
        validation=validation,
    )


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        print("Interrompido pelo usuario.", file=sys.stderr)
        raise SystemExit(130)
