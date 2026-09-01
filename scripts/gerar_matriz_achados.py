#!/usr/bin/env python3
"""Gera a matriz de achados iGovTI 2026 preservando o OOXML do modelo."""

from __future__ import annotations

import argparse
import copy
import importlib.util
import json
import re
import sys
import zipfile
from pathlib import Path

from lxml import etree
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent / "resources"))
from aplicabilidade_juridica import PerfilAuditado
from matriz_aplicabilidade import carregar_catalogo_matriz

try:
    from scripts.gerar_matriz_planejamento import (
        AUDITED_ENTITIES_HEADER,
        AUDIT_OBJECTIVE_HEADER,
        display_identifier,
        parse_matrix as parse_planning_matrix,
    )
except ImportError:
    from gerar_matriz_planejamento import (  # type: ignore
        AUDITED_ENTITIES_HEADER,
        AUDIT_OBJECTIVE_HEADER,
        display_identifier,
        parse_matrix as parse_planning_matrix,
    )


NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "xml": "http://www.w3.org/XML/1998/namespace",
}
W = f"{{{NS['w']}}}"


EFEITOS_POR_SITUACAO = {
    "S1.1": ("Fragilidade na responsabilização", "Ausência de unidade ou função institucionalmente reconhecida para coordenar o uso da tecnologia da informação e responder por seus resultados."),
    "S1.2": ("Atuação reativa e fragmentada", "Falta de clareza sobre as responsabilidades de planejamento, coordenação, gestão, execução, monitoramento e controle da TIC."),
    "S1.3": ("Baixa influência institucional", "Participação insuficiente da TIC em decisões estratégicas, orçamentárias, contratuais e de gestão de riscos."),
    "S2.1": ("Direcionamento insuficiente da TIC", "Baixa clareza sobre objetivos, indicadores, metas e acompanhamento do desempenho da TIC."),
    "S2.2": ("Priorização deficiente", "Ausência de instância colegiada para deliberar sobre prioridades, projetos, riscos, serviços, orçamento e contratações de TIC."),
    "S2.3": ("Governança meramente formal", "Inexistência de reuniões, deliberações e acompanhamento efetivo das decisões relacionadas à TIC."),
    "S3.1": ("Atuação reativa", "Planejamento de TIC sem processo formal e sem participação adequada das áreas demandantes."),
    "S3.2": ("Baixa legitimidade institucional", "Plano de TIC sem aprovação formal suficiente para orientar projetos, orçamento e contratações."),
    "S3.4": ("Desalinhamento institucional", "Execução de ações de TIC com baixo valor para os objetivos e resultados da organização."),
    "S3.5": ("Ineficiência na alocação de recursos", "Aquisições reativas, não priorizadas ou desconectadas do orçamento e do plano de contratações."),
    "S3.6": ("Desatualização das prioridades", "Manutenção de metas e iniciativas incompatíveis com mudanças institucionais, orçamentárias ou tecnológicas."),
    "S4.1": ("Execução insuficiente de funções críticas", "Risco de incapacidade para planejar, gerir, proteger, contratar, fiscalizar e sustentar serviços e ativos de TIC."),
    "S4.2": ("Subdimensionamento da força de trabalho", "Ausência de base estruturada para estimar, alocar e acompanhar a capacidade necessária de pessoal de TIC e segurança da informação."),
    "S4.3": ("Baixa clareza de responsabilidades", "Ausência de cargos ou funções formalmente atribuídos à TIC ou à segurança da informação."),
    "S4.6": ("Dependência excessiva de terceiros", "Perda de conhecimento, redução da governabilidade e risco de descontinuidade dos serviços quando não houver capacidade interna mínima de coordenação e fiscalização."),
    "S5.1": ("Prestação reativa de serviços", "Falta de definição clara, padronizada e transparente dos serviços de TIC oferecidos aos usuários e às áreas demandantes."),
    "S5.2": ("Ausência de parâmetros de desempenho", "Impossibilidade de avaliar objetivamente a qualidade, a disponibilidade e o atendimento dos principais serviços de TIC."),
    "S5.3": ("Controle inadequado dos ativos", "Informações insuficientes sobre dispositivos e softwares utilizados pela organização."),
    "S5.4": ("Controle inadequado da configuração", "Informações insuficientes sobre itens de configuração e seus relacionamentos com sistemas, infraestrutura e serviços."),
    "S5.5": ("Tratamento deficiente de incidentes", "Resposta não padronizada, intempestiva ou sem rastreabilidade, com maior impacto sobre a continuidade e a qualidade dos serviços."),
    "S6.1": ("Instrução processual frágil", "Falta de clareza sobre etapas, responsabilidades, instâncias decisórias, modelos e critérios de aprovação das contratações de TIC."),
    "S6.2": ("Soluções incompatíveis", "Contratação de soluções incompatíveis com padrões tecnológicos ou requisitos institucionais."),
    "S6.3": ("Contratações desalinhadas", "Contratações de TIC desconectadas dos instrumentos de planejamento aplicáveis."),
    "S6.4": ("Participação técnica insuficiente", "Planejamento da contratação sem equipe formalmente designada ou sem participação técnica da área de TIC."),
}


def carregar_modulo_matriz(repo: Path):
    path = repo / ".agents/skills/preencher-matriz-procedimentos-auditoria/scripts/gerar_matriz_procedimentos.py"
    spec = importlib.util.spec_from_file_location("gerar_matriz_procedimentos", path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Não foi possível carregar {path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def carregar_achados(repo: Path):
    module = carregar_modulo_matriz(repo)
    matriz = repo / "01-Planejamento/03-Estrategia_e_Plano/04-Matriz_Planejamento/matriz_planejamento-pos-comentarios-gestor.md"
    planning_matrix = parse_planning_matrix(matriz.read_text(encoding="utf-8"))
    questions = {question.id: question for question in planning_matrix.questions}
    return [
        (
            finding,
            {risk.id: risk.text for risk in questions[finding.questao_codigo].riscos},
            questions[finding.questao_codigo],
        )
        for finding in module.parse_matrix(matriz)
    ]


def carregar_perfis_auditados(path: Path) -> dict[str, PerfilAuditado]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook["auditados"]
        headers = {
            str(worksheet.cell(1, column).value or "").strip(): column
            for column in range(1, worksheet.max_column + 1)
        }
        obrigatorias = {
            "sigla",
            "segmento_institucional",
            "natureza_administrativa",
            "tags_aplicabilidade",
        }
        faltantes = sorted(obrigatorias - set(headers))
        if faltantes:
            raise ValueError(
                "Cadastro de auditados sem colunas de aplicabilidade: "
                + ", ".join(faltantes)
                + "."
            )
        perfis = {}
        for row in range(2, worksheet.max_row + 1):
            dados = {
                nome: worksheet.cell(row, column).value
                for nome, column in headers.items()
            }
            perfil = PerfilAuditado.from_mapping(dados)
            if perfil.sigla:
                perfis[perfil.sigla] = perfil
        return perfis
    finally:
        workbook.close()


def carregar_ocorrencias(
    resultado_path: Path,
    catalogo,
    perfis: dict[str, PerfilAuditado],
) -> dict[str, dict]:
    """Cruza ocorrências finais com as variantes explicitamente declaradas na matriz."""
    dados = json.loads(resultado_path.read_text(encoding="utf-8"))
    resolvedor = catalogo.criar_resolvedor()
    ocorrencias: dict[str, dict] = {}

    for registro in dados.values():
        if not isinstance(registro, dict):
            continue
        sigla = str(registro.get("sigla") or "").strip().upper()
        if not sigla:
            continue
        perfil = perfis.get(sigla)
        if perfil is None:
            raise ValueError(f"{sigla}: auditado do resultado final ausente do cadastro institucional.")
        for procedimento in registro.get("procedimentos_executados", []):
            achado = procedimento.get("achado") or {}
            for situacao in achado.get("situacoes_detalhadas", []) or []:
                id_situacao = str(situacao.get("id_situacao") or "").strip()
                if not id_situacao:
                    raise ValueError(
                        f"{sigla}: situação final sem identificador estável em resultado_auditoria.json."
                    )
                resolvida = resolvedor.resolver(perfil, id_situacao)
                registro_situacao = ocorrencias.setdefault(
                    id_situacao,
                    {"auditados": set(), "variantes": {}},
                )
                registro_situacao["auditados"].add(sigla)
                registro_variante = registro_situacao["variantes"].setdefault(
                    resolvida.variante.id,
                    {
                        "variante": resolvida.variante,
                        "criterios": resolvida.criterios,
                        "auditados": set(),
                    },
                )
                registro_variante["auditados"].add(sigla)

    return ocorrencias


def carregar_acoes(repo: Path):
    path = repo / "02-Execucao/03-Execucao_Procedimentos/01-Insumos/mapa-verificacao-achados-pos-comentarios-gestor.xlsx"
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        ws_acoes = workbook["Ações de Verificação"]
        headers = {ws_acoes.cell(3, col).value: col for col in range(1, ws_acoes.max_column + 1)}
        acoes = {}
        for row in range(4, ws_acoes.max_row + 1):
            action_id = ws_acoes.cell(row, headers["id"]).value
            if not action_id:
                continue
            acoes[str(action_id)] = {
                name: ws_acoes.cell(row, col).value for name, col in headers.items()
            }

        ws_proc = workbook["Procedimentos de Auditoria"]
        proc_headers = {ws_proc.cell(3, col).value: col for col in range(1, ws_proc.max_column + 1)}
        por_achado = {}
        for row in range(4, ws_proc.max_row + 1):
            numero = ws_proc.cell(row, proc_headers["numero_achado"]).value
            logica = str(ws_proc.cell(row, proc_headers["logica_achado"]).value or "")
            if numero:
                por_achado[f"A{int(numero)}"] = re.findall(r"\bAV\d+\b", logica)
        return acoes, por_achado
    finally:
        workbook.close()


def contar_auditados(repo: Path) -> int:
    respostas_dir = repo / "02-Execucao/01-Questionario/03-Respostas_Processadas"
    candidatos = [
        respostas_dir / "20260716-respostas-questionario-pos-comentarios-gestor.xlsx",
        respostas_dir / "20260621-respostas-questionario-02-pos-avaliacao-evidencias.xlsx",
        respostas_dir / "20260621-respostas-questionario-01-pos-ajuste-inicial.xlsx",
    ]
    path = next((candidato for candidato in candidatos if candidato.exists()), None)
    if path is None:
        raise FileNotFoundError(
            "Nenhuma base processada do questionário foi encontrada para contar os auditados."
        )
    workbook = load_workbook(path, read_only=False, data_only=True)
    try:
        ws = workbook.active
        headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
        column = headers.index("firstname") + 1
        return sum(1 for row in range(2, ws.max_row + 1) if ws.cell(row, column).value)
    finally:
        workbook.close()


def formatar_condicao(value) -> str:
    value = str(value or "").strip()
    fixed = {
        "(Não | N/A)": 'resposta "Não" ou "N/A"',
        "Não": 'resposta "Não"',
        "Sim": 'resposta "Sim"',
        "0": "valor igual a zero",
        "True": "resultado verdadeiro",
    }
    if value in fixed:
        return fixed[value]
    if value.startswith("~(") and value.endswith(")"):
        marker = re.match(r"([a-z])\)", value[2:], flags=re.IGNORECASE)
        return f"resposta distinta da alternativa {marker.group(1).lower()})" if marker else "resposta distinta da alternativa indicada"
    if value.startswith("(") and value.endswith(")") and " | " in value:
        options = [item.strip() for item in value[1:-1].split(" | ")]
        markers = [re.match(r"([a-z])\)", item, flags=re.IGNORECASE) for item in options]
        if all(markers):
            return "uma das alternativas " + ", ".join(f"{match.group(1).lower()})" for match in markers)
        return "um dos valores " + ", ".join(f'"{item}"' for item in options)
    marker = re.match(r"([a-z])\)", value, flags=re.IGNORECASE)
    if marker:
        return f"alternativa {marker.group(1).lower()})"
    if value.startswith((">", "<")):
        return f"condição {value}"
    return f'resposta "{value}"'


def listar_campos(fields) -> str:
    fields = list(dict.fromkeys(str(field) for field in fields))
    if len(fields) == 1:
        return fields[0]
    return ", ".join(fields[:-1]) + " e " + fields[-1]


def evidencias_do_achado(achado, action_ids, acoes, ocorrencias):
    por_descricao = {}
    for action_id in action_ids:
        action = acoes[action_id]
        por_descricao.setdefault(action["descricao_situacao_inconforme"], []).append(action)

    evidencias = []
    for situacao in achado.situacoes:
        actions = por_descricao.get(situacao.descricao, [])
        conditions = {}
        for action in actions:
            field = action["informacao_requerida"]
            condition = formatar_condicao(action["situacao_inconforme"])
            conditions.setdefault(condition, []).append(field)
        items = [
            f"{condition} em {listar_campos(fields)}"
            for condition, fields in conditions.items()
        ] or ["Condição prevista na Matriz de Procedimentos de Auditoria"]
        evidencias.append({
            "situacao": f"{situacao.codigo} - {situacao.descricao.rstrip('.')}",
            "quantidade": len(ocorrencias[situacao.codigo]["auditados"]),
            "itens": items,
        })
    return evidencias


def _id_local_criterio(criterion_id: str) -> str:
    return criterion_id.rsplit(".", 1)[-1]


def criterios_do_achado(question, achado, variantes_especificas):
    """Materializa a visão sequencial dos critérios usados pelo achado."""
    usados = {
        reference
        for situacao in achado.situacoes
        for reference in situacao.criterios
    }
    usados.update(
        _id_local_criterio(reference)
        for variante in variantes_especificas
        for reference in variante.criterios
    )
    display_mapping = {
        criterion.id: f"C{index}"
        for index, criterion in enumerate(
            (criterion for criterion in question.criterios if criterion.id in usados),
            start=1,
        )
    }

    gerais = []
    especificos = {}
    for criterion in question.criterios:
        if criterion.id not in usados:
            continue
        item = {
            "id_exibicao": display_identifier(criterion.id, display_mapping),
            "descricao": criterion.descricao,
        }
        if criterion.especifico:
            especificos.setdefault(criterion.publico or "Público específico", []).append(item)
        else:
            gerais.append(item)
    return gerais, especificos, display_mapping


def _juntar_publicos(publicos: list[str]) -> str:
    if len(publicos) <= 1:
        return publicos[0] if publicos else ""
    return ", ".join(publicos[:-1]) + " e " + publicos[-1]


def _rotulo_destinatarios(publico: str, auditados: set[str]) -> str:
    quantidade = len(auditados)
    complemento = "organização" if quantidade == 1 else "organizações"
    if publico:
        return f"{publico} — {quantidade} {complemento} com ocorrência"
    return f"{quantidade} {complemento} com ocorrência"


def encaminhamentos_do_achado(achado, ocorrencias, display_mapping):
    """Exibe somente encaminhamentos destinados a ocorrências do resultado final."""
    resultado = []
    for situacao in achado.situacoes:
        registro_situacao = ocorrencias[situacao.codigo]
        variantes_ocorridas = sorted(
            registro_situacao["variantes"].values(),
            key=lambda item: (
                not item["variante"].geral,
                item["variante"].rotulo_publico,
            ),
        )
        tem_especifica = any(not item["variante"].geral for item in variantes_ocorridas)
        membros = []
        for item in variantes_ocorridas:
            variante = item["variante"]
            if variante.geral:
                publico = "Demais jurisdicionados" if tem_especifica else "Público geral"
            else:
                publico = variante.rotulo_publico
            membros.append({
                "geral": variante.geral,
                "publico": publico,
                "tipo": variante.tipo_encaminhamento,
                "encaminhamento": variante.encaminhamento,
                "criterios": [_id_local_criterio(value) for value in variante.criterios],
                "auditados": set(item["auditados"]),
            })

        grupos = []
        por_conteudo = {}
        for membro in membros:
            key = (
                membro["tipo"].strip().casefold(),
                membro["encaminhamento"].strip().rstrip(".").casefold(),
            )
            if key not in por_conteudo:
                por_conteudo[key] = len(grupos)
                grupos.append({"membros": [], "tipo": membro["tipo"],
                               "encaminhamento": membro["encaminhamento"]})
            grupos[por_conteudo[key]]["membros"].append(membro)

        total_membros = len(membros)
        for grupo in grupos:
            membros_grupo = grupo["membros"]
            contem_geral = any(membro["geral"] for membro in membros_grupo)
            auditados = set().union(*(membro["auditados"] for membro in membros_grupo))
            if len(membros_grupo) == total_membros and contem_geral and tem_especifica:
                publico = "Todos os grupos alcançados"
            else:
                publico = _juntar_publicos([
                    membro["publico"] for membro in membros_grupo if membro["publico"]
                ])

            criterios = []
            for membro in membros_grupo:
                for criterion_id in membro["criterios"]:
                    shown = display_identifier(criterion_id, display_mapping)
                    if shown not in criterios:
                        criterios.append(shown)
            grupo["publico"] = publico
            grupo["destinatarios"] = _rotulo_destinatarios(publico, auditados)
            grupo["quantidade"] = len(auditados)
            grupo["criterios"] = criterios
            del grupo["membros"]

        resultado.append({
            "situacao": f"{situacao.codigo} - {situacao.descricao.rstrip('.')}",
            "grupos": grupos,
        })
    return resultado


def dividir_dispositivo_criterio(descricao: str) -> tuple[str, str]:
    """Separa o dispositivo de sua explicação para aplicar negrito no DOCX."""
    candidates = []
    for separator in (" — ", " – ", " - ", ": "):
        position = descricao.find(separator)
        if position >= 0:
            candidates.append((position, separator))
    if not candidates:
        return descricao, ""
    position, separator = min(candidates, key=lambda item: item[0])
    return descricao[:position], descricao[position:]


def clone_properties(element, child_name):
    child = element.find(f"w:{child_name}", NS)
    return copy.deepcopy(child) if child is not None else None


def run_properties(paragraph, *, bold=None, underline=None):
    for run in paragraph.findall(".//w:r", NS):
        rpr = run.find("w:rPr", NS)
        has_bold = rpr is not None and rpr.find("w:b", NS) is not None
        has_underline = rpr is not None and rpr.find("w:u", NS) is not None
        if (bold is None or has_bold == bold) and (underline is None or has_underline == underline):
            return copy.deepcopy(rpr)
    return None


def make_run(text: str, properties=None):
    run = etree.Element(W + "r")
    if properties is not None:
        run.append(copy.deepcopy(properties))
    node = etree.SubElement(run, W + "t")
    if text.startswith(" ") or text.endswith(" "):
        node.set(f"{{{NS['xml']}}}space", "preserve")
    node.text = text
    return run


def make_paragraph(template, parts, *, left_indent=None, hanging=None):
    paragraph = etree.Element(W + "p")
    ppr = clone_properties(template, "pPr")
    if ppr is not None:
        paragraph.append(ppr)
    if left_indent is not None or hanging is not None:
        if ppr is None:
            ppr = etree.Element(W + "pPr")
            paragraph.insert(0, ppr)
        indent = ppr.find("w:ind", NS)
        if indent is None:
            indent = etree.SubElement(ppr, W + "ind")
        if left_indent is not None:
            indent.set(W + "left", str(left_indent))
        if hanging is not None:
            indent.set(W + "hanging", str(hanging))
    for text, properties in parts:
        paragraph.append(make_run(text, properties))
    return paragraph


def blank_paragraph(template):
    paragraph = etree.Element(W + "p")
    ppr = clone_properties(template, "pPr")
    if ppr is not None:
        paragraph.append(ppr)
    return paragraph


def replace_cell(cell, paragraphs):
    for child in list(cell):
        if child.tag != W + "tcPr":
            cell.remove(child)
    for paragraph in paragraphs:
        cell.append(paragraph)


def atualizar_documento(document_xml: bytes, dados):
    root = etree.fromstring(document_xml)
    body = root.find("w:body", NS)
    tables = body.findall("w:tbl", NS)
    if not dados:
        raise ValueError("O resultado final não contém achados para compor a matriz.")
    if len(tables) < len(dados):
        raise ValueError(
            f"O modelo contém {len(tables)} tabelas, mas o resultado exige {len(dados)}."
        )
    used_tables = tables[:len(dados)]
    final_section_properties = copy.deepcopy(body.find("w:sectPr", NS))
    if len(tables) > len(used_tables):
        children = list(body)
        start = children.index(used_tables[-1]) + 1
        stop = children.index(tables[len(used_tables)])
        for child in children[start:stop]:
            section_properties = child.find("w:pPr/w:sectPr", NS)
            if section_properties is not None:
                final_section_properties = copy.deepcopy(section_properties)
                break

    sample_cells = tables[0].findall("w:tr", NS)[1].findall("w:tc", NS)
    templates = {
        "finding_bold": sample_cells[0].findall("w:p", NS)[0],
        "finding_blank": sample_cells[0].findall("w:p", NS)[1],
        "finding_normal": sample_cells[0].findall("w:p", NS)[2],
        "criterion": sample_cells[1].findall("w:p", NS)[0],
        "criterion_blank": sample_cells[1].findall("w:p", NS)[1],
        "evidence_intro": sample_cells[2].findall("w:p", NS)[0],
        "evidence_item": sample_cells[2].findall("w:p", NS)[1],
        "cause": sample_cells[3].findall("w:p", NS)[0],
        "effect": sample_cells[4].findall("w:p", NS)[0],
        "effect_blank": sample_cells[4].findall("w:p", NS)[1],
        "referral": sample_cells[5].findall("w:p", NS)[0],
        "referral_blank": sample_cells[5].findall("w:p", NS)[1],
    }
    props = {
        "finding_bold": run_properties(templates["finding_bold"], bold=True),
        "normal": run_properties(templates["finding_normal"], bold=False),
        "bold": run_properties(templates["criterion"], bold=True),
        "criterion_normal": run_properties(templates["criterion"], bold=False),
        "evidence_bold": run_properties(templates["evidence_intro"], bold=True),
        "evidence_normal": run_properties(templates["evidence_intro"], bold=False),
        "cause_bold": run_properties(templates["cause"], bold=True),
        "effect_bold": run_properties(templates["effect"], bold=True),
        "effect_normal": run_properties(templates["effect"], bold=False),
        # A matriz oficial também é o modelo da próxima geração. As
        # propriedades procuradas podem estar em parágrafos posteriores ao
        # título da situação, por isso a busca abrange toda a célula.
        "referral_label": run_properties(sample_cells[5], bold=True, underline=True),
        "referral_normal": run_properties(sample_cells[5], bold=False),
    }

    for index, item in enumerate(dados):
        achado = item["achado"]
        cells = used_tables[index].findall("w:tr", NS)[1].findall("w:tc", NS)

        finding_paragraphs = [
            make_paragraph(templates["finding_bold"], [(f"ACHADO {int(achado.codigo[1:]):02d}", props["finding_bold"])]),
            blank_paragraph(templates["finding_blank"]),
            make_paragraph(templates["finding_normal"], [(achado.nome.rstrip("."), props["normal"])]),
            blank_paragraph(templates["finding_blank"]),
            make_paragraph(templates["finding_normal"], [
                ("Achado composto pelas seguintes situações efetivamente identificadas:", props["normal"]),
            ]),
        ]
        for situacao in achado.situacoes:
            quantidade = item["ocorrencias"][situacao.codigo]["auditados"]
            finding_paragraphs.append(make_paragraph(
                templates["finding_normal"],
                [(
                    f"• {situacao.codigo} - {situacao.descricao.rstrip('.')} "
                    f"({len(quantidade)} ocorrência{'s' if len(quantidade) != 1 else ''})",
                    props["normal"],
                )],
                left_indent=300,
                hanging=180,
            ))
        replace_cell(cells[0], finding_paragraphs)

        paragraphs = []
        for criterion in item["criterios"]:
            dispositivo, explicacao = dividir_dispositivo_criterio(criterion["descricao"])
            paragraphs.append(make_paragraph(templates["criterion"], [
                (f'{criterion["id_exibicao"]}: {dispositivo}', props["bold"]),
                (explicacao, props["criterion_normal"]),
            ]))
            paragraphs.append(blank_paragraph(templates["criterion_blank"]))
        for publico, criterios in item.get("criterios_especificos", {}).items():
            paragraphs.append(make_paragraph(templates["criterion"], [(publico, props["bold"])]))
            for criterio in criterios:
                dispositivo, explicacao = dividir_dispositivo_criterio(criterio["descricao"])
                paragraphs.append(make_paragraph(templates["criterion"], [
                    (f'{criterio["id_exibicao"]}: {dispositivo}', props["bold"]),
                    (explicacao, props["criterion_normal"]),
                ]))
                paragraphs.append(blank_paragraph(templates["criterion_blank"]))
        if paragraphs:
            paragraphs.pop()
        replace_cell(cells[1], paragraphs)

        paragraphs = [make_paragraph(templates["evidence_intro"], [
            ("Respostas ao Questionário iGovTI 2026", props["evidence_bold"]),
            (" - Situações caracterizadas pelas respostas declaradas às questões indicadas:", props["evidence_normal"]),
        ])]
        for evidence_index, evidence in enumerate(item["evidencias"]):
            if evidence_index:
                paragraphs.append(blank_paragraph(templates["criterion_blank"]))
            paragraphs.append(make_paragraph(templates["evidence_item"], [
                (evidence["situacao"], props["bold"]),
            ]))
            paragraphs.append(make_paragraph(templates["evidence_item"], [
                (
                    f"Ocorrência no resultado final: {evidence['quantidade']} "
                    f"{'organizações' if evidence['quantidade'] != 1 else 'organização'}.",
                    props["evidence_normal"],
                ),
            ]))
            for evidence_item in evidence["itens"]:
                paragraphs.append(make_paragraph(templates["evidence_item"], [
                    (f"• {evidence_item.rstrip('.')}.", props["evidence_normal"]),
                ]))
        replace_cell(cells[2], paragraphs)

        replace_cell(cells[3], [
            make_paragraph(templates["cause"], [("Não investigada.", props["cause_bold"])])
        ])

        paragraphs = []
        for effect_index, (label, text) in enumerate(item["efeitos"]):
            paragraphs.append(make_paragraph(templates["effect"], [
                (label + ":", props["effect_bold"]),
                (" " + text, props["effect_normal"]),
            ]))
            if effect_index < len(item["efeitos"]) - 1:
                paragraphs.append(blank_paragraph(templates["effect_blank"]))
        replace_cell(cells[4], paragraphs)

        paragraphs = []
        for referral_index, situacao in enumerate(item["encaminhamentos"]):
            if referral_index:
                paragraphs.append(blank_paragraph(templates["referral_blank"]))
            paragraphs.append(make_paragraph(templates["referral"], [
                (situacao["situacao"].split(" - ", 1)[-1], props["bold"]),
            ]))
            for grupo in situacao["grupos"]:
                paragraphs.append(make_paragraph(templates["referral"], [
                    (grupo["destinatarios"], props["bold"]),
                ]))
                codigo_situacao = situacao["situacao"].split(" - ", 1)[0]
                referencias = ", ".join([codigo_situacao, *grupo["criterios"]])
                paragraphs.append(make_paragraph(templates["referral"], [
                    ("• ", props["referral_normal"]),
                    (f'Comunicação com {grupo["tipo"]}', props["referral_label"]),
                    (f' para que {grupo["encaminhamento"].rstrip(".")} [{referencias}].',
                     props["referral_normal"]),
                ], left_indent=360, hanging=180))
        replace_cell(cells[5], paragraphs)

    # A propriedade pageBreakBefore dentro da primeira célula de uma tabela
    # não é interpretada de modo consistente pelo Word. Remove-se a marcação
    # legada e insere-se uma quebra estrutural imediatamente antes da tabela.
    for table in used_tables[1:]:
        first_paragraph = table.find("w:tr/w:tc/w:p", NS)
        ppr = first_paragraph.find("w:pPr", NS)
        if ppr is not None:
            page_break_before = ppr.find("w:pageBreakBefore", NS)
            if page_break_before is not None:
                ppr.remove(page_break_before)

    for child in list(body):
        if child.tag == W + "p" and child.find(".//w:br[@w:type='page']", NS) is not None:
            body.remove(child)

    for table in tables[len(used_tables):]:
        body.remove(table)

    for table in used_tables[1:]:
        paragraph = etree.Element(W + "p")
        run = etree.SubElement(paragraph, W + "r")
        page_break = etree.SubElement(run, W + "br")
        page_break.set(W + "type", "page")
        body.insert(body.index(table), paragraph)

    last_table = used_tables[-1]
    remove = False
    for child in list(body):
        if child is last_table:
            remove = True
            continue
        if remove and child.tag != W + "sectPr":
            body.remove(child)
    sect_pr = body.find("w:sectPr", NS)
    if final_section_properties is not None:
        body.replace(sect_pr, final_section_properties)
        sect_pr = final_section_properties
    final_paragraph = etree.Element(W + "p")
    body.insert(body.index(sect_pr), final_paragraph)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")


def atualizar_cabecalho(header_xml: bytes):
    root = etree.fromstring(header_xml)
    cell = root.find(".//w:tbl/w:tr/w:tc[3]", NS)
    paragraphs = cell.findall("w:p", NS)
    if len(paragraphs) < 3:
        raise ValueError("Cabeçalho do modelo incompatível")

    templates = list(paragraphs[:3])
    bold = [run_properties(paragraph, bold=True) for paragraph in templates]
    normal = [run_properties(paragraph, bold=False) for paragraph in templates]
    new_paragraphs = [
        make_paragraph(templates[0], [("FISCALIZAÇÃO", bold[0]), (": 18/2026", normal[0])]),
        make_paragraph(templates[1], [
            ("JURISDICIONADOS", bold[1]),
            (f": {AUDITED_ENTITIES_HEADER}", normal[1]),
        ]),
        make_paragraph(templates[2], [
            ("OBJETIVO DA AUDITORIA", bold[2]),
            (f": {AUDIT_OBJECTIVE_HEADER}", normal[2]),
        ]),
    ]
    replace_cell(cell, new_paragraphs)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")


def montar_dados(
    repo: Path,
    resultado_path: Path | None = None,
    auditados_path: Path | None = None,
):
    achados = carregar_achados(repo)
    acoes, por_achado = carregar_acoes(repo)
    matriz_path = repo / "01-Planejamento/03-Estrategia_e_Plano/04-Matriz_Planejamento/matriz_planejamento-pos-comentarios-gestor.md"
    catalogo = carregar_catalogo_matriz(matriz_path)
    resultado_path = resultado_path or (
        repo
        / "02-Execucao/03-Execucao_Procedimentos/02-Resultados_Auditoria/"
        "03-pos-comentarios-gestor/resultado_auditoria.json"
    )
    auditados_path = auditados_path or (
        repo / "02-Execucao/03-Execucao_Procedimentos/01-Insumos/bd_auditados.xlsx"
    )
    perfis = carregar_perfis_auditados(auditados_path)
    ocorrencias = carregar_ocorrencias(resultado_path, catalogo, perfis)
    dados = []
    for achado, riscos, question in achados:
        achado = copy.deepcopy(achado)
        achado.situacoes = [
            situacao
            for situacao in achado.situacoes
            if situacao.codigo in ocorrencias
        ]
        if not achado.situacoes:
            continue
        ids_situacoes = {situacao.codigo for situacao in achado.situacoes}
        ocorrencias_achado = {
            id_situacao: ocorrencias[id_situacao]
            for id_situacao in ids_situacoes
        }
        variantes_especificas = [
            registro["variante"]
            for ocorrencia in ocorrencias_achado.values()
            for registro in ocorrencia["variantes"].values()
            if not registro["variante"].geral
        ]
        criterios, criterios_especificos, display_mapping = criterios_do_achado(
            question, achado, variantes_especificas
        )
        dados.append({
            "achado": achado,
            "criterios": criterios,
            "evidencias": evidencias_do_achado(
                achado,
                por_achado[achado.codigo],
                acoes,
                ocorrencias_achado,
            ),
            "riscos": list(riscos.values()),
            "efeitos": [
                EFEITOS_POR_SITUACAO[situacao.codigo]
                for situacao in achado.situacoes
            ],
            "criterios_especificos": criterios_especificos,
            "ocorrencias": ocorrencias_achado,
            "encaminhamentos": encaminhamentos_do_achado(
                achado, ocorrencias_achado, display_mapping
            ),
        })
    return dados


def construir_painel_ocorrencias(dados):
    situacoes = [
        (situacao.codigo, situacao.descricao.rstrip("."))
        for item in dados
        for situacao in item["achado"].situacoes
    ]
    valores = {}
    auditados = set()
    for item in dados:
        for id_situacao, ocorrencia in item["ocorrencias"].items():
            for registro in ocorrencia["variantes"].values():
                tipo = registro["variante"].tipo_encaminhamento.strip().casefold()
                marca = "D" if tipo == "determinação" else "R"
                for sigla in registro["auditados"]:
                    chave = (sigla, id_situacao)
                    anterior = valores.get(chave)
                    if anterior and anterior != marca:
                        raise ValueError(
                            f"{sigla}/{id_situacao}: tipos de encaminhamento conflitantes."
                        )
                    valores[chave] = marca
                    auditados.add(sigla)
    return situacoes, sorted(auditados), valores


def _definir_margens_celula(cell, valor: int = 36):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for lado in ("top", "left", "bottom", "right"):
        margem = tc_mar.find(qn(f"w:{lado}"))
        if margem is None:
            margem = OxmlElement(f"w:{lado}")
            tc_mar.append(margem)
        margem.set(qn("w:w"), str(valor))
        margem.set(qn("w:type"), "dxa")


def _impedir_quebra_linha_tabela(row):
    from docx.oxml import OxmlElement

    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def _repetir_cabecalho_tabela(row):
    from docx.oxml import OxmlElement

    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:tblHeader"))


def _sombrear_celula(cell, preenchimento: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), preenchimento)


def anexar_painel_ocorrencias(saida: Path, dados) -> None:
    from docx import Document
    from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    situacoes, auditados, valores = construir_painel_ocorrencias(dados)
    if not situacoes or not auditados:
        return

    document = Document(str(saida))
    document.add_page_break()

    titulo = document.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = titulo.add_run("QUADRO CONSOLIDADO DE SITUAÇÕES ENCONTRADAS POR AUDITADO")
    run.bold = True
    run.font.size = Pt(10)

    legenda = document.add_paragraph(
        "D = determinação; R = recomendação; célula vazia = situação não identificada. "
        "As colunas correspondem às situações descritas nos achados anteriores."
    )
    legenda.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in legenda.runs:
        run.font.size = Pt(7)

    table = document.add_table(rows=1, cols=len(situacoes) + 1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    header = table.rows[0]
    _repetir_cabecalho_tabela(header)
    header.cells[0].text = "Auditado"
    for index, (id_situacao, _) in enumerate(situacoes, start=1):
        header.cells[index].text = id_situacao

    largura_auditado = Cm(3.2)
    largura_situacao = Cm(0.99)
    for index, cell in enumerate(header.cells):
        cell.width = largura_auditado if index == 0 else largura_situacao
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        _definir_margens_celula(cell)
        _sombrear_celula(cell, "D9E2F3")
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(6.5)

    for sigla in auditados:
        row = table.add_row()
        _impedir_quebra_linha_tabela(row)
        for index, cell in enumerate(row.cells):
            cell.width = largura_auditado if index == 0 else largura_situacao
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _definir_margens_celula(cell)
        row.cells[0].text = sigla
        for index, (id_situacao, _) in enumerate(situacoes, start=1):
            row.cells[index].text = valores.get((sigla, id_situacao), "")
        for index, cell in enumerate(row.cells):
            for paragraph in cell.paragraphs:
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
                )
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(6.5)
                    run.bold = index > 0 and bool(run.text)

    document.save(str(saida))


def gerar(
    modelo: Path,
    saida: Path,
    repo: Path,
    resultado_path: Path,
    auditados_path: Path,
):
    dados = montar_dados(repo, resultado_path, auditados_path)
    auditados_avaliados = contar_auditados(repo)
    with zipfile.ZipFile(modelo, "r") as source:
        archive = [(copy.copy(info), source.read(info.filename)) for info in source.infolist()]
    with zipfile.ZipFile(saida, "w") as target:
        for info, original_content in archive:
            content = original_content
            if info.filename == "word/document.xml":
                content = atualizar_documento(content, dados)
            elif info.filename == "word/header1.xml":
                content = atualizar_cabecalho(content)
            target.writestr(info, content)
    anexar_painel_ocorrencias(saida, dados)
    print(
        f"OK: {len(dados)} achados e {auditados_avaliados} organizações avaliadas -> {saida}"
    )


def parse_args():
    repo_root = Path(__file__).resolve().parent.parent
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--modelo",
        type=Path,
        default=repo_root / "02-Execucao/04-Matriz_Achados/AN06 – Matriz de achados.docx",
    )
    parser.add_argument(
        "--saida",
        type=Path,
        default=repo_root / "02-Execucao/04-Matriz_Achados/AN06 – Matriz de achados.docx",
    )
    parser.add_argument(
        "--resultado-auditoria",
        type=Path,
        default=(
            repo_root
            / "02-Execucao/03-Execucao_Procedimentos/02-Resultados_Auditoria/"
            "03-pos-comentarios-gestor/resultado_auditoria.json"
        ),
        help="Resultado final que define quais situações e destinatários integram a matriz.",
    )
    parser.add_argument(
        "--auditados",
        type=Path,
        default=(
            repo_root
            / "02-Execucao/03-Execucao_Procedimentos/01-Insumos/bd_auditados.xlsx"
        ),
        help="Cadastro institucional usado para resolver as variantes explicitamente declaradas.",
    )
    parser.add_argument("--repo", type=Path, default=repo_root)
    return parser.parse_args()


def main():
    args = parse_args()
    args.saida.parent.mkdir(parents=True, exist_ok=True)
    gerar(
        args.modelo.resolve(),
        args.saida.resolve(),
        args.repo.resolve(),
        args.resultado_auditoria.resolve(),
        args.auditados.resolve(),
    )


if __name__ == "__main__":
    main()
