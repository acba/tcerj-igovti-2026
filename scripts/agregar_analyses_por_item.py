"""Agrega conclusoes dos arquivos JSONL de avaliacao por questao avaliada."""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ESTADOS = ("conforme", "nao_conforme", "inconclusivo")
PRIORIDADE_EXEMPLOS = {"nao_conforme": 0, "inconclusivo": 1, "conforme": 2}


def questao_base(registro: dict[str, Any], conclusao: dict[str, Any]) -> str:
    candidatos = (
        registro.get("questao"),
        conclusao.get("item_codigo"),
        registro.get("coluna_evidencia"),
    )
    for candidato in candidatos:
        match = re.search(r"q\d{4}", str(candidato or ""), flags=re.IGNORECASE)
        if match:
            return match.group(0).lower()
    return "sem_item"


def ler_registros(
    caminhos: Iterable[Path],
    referencias: set[str] | None = None,
) -> tuple[list[dict[str, Any]], list[str]]:
    conclusoes: list[dict[str, Any]] = []
    erros: list[str] = []
    for caminho in caminhos:
        with caminho.open("r", encoding="utf-8-sig") as arquivo:
            for numero_linha, linha in enumerate(arquivo, start=1):
                if not linha.strip():
                    continue
                try:
                    registro = json.loads(linha)
                except json.JSONDecodeError as exc:
                    erros.append(f"{caminho}:{numero_linha}: JSON invalido: {exc}")
                    continue

                modelo = str(registro.get("model") or "").strip()
                if referencias and modelo.casefold() not in referencias:
                    continue

                resultado = registro.get("result") or {}
                itens = resultado.get("conclusoes") or []
                if not isinstance(itens, list):
                    erros.append(f"{caminho}:{numero_linha}: conclusoes nao e uma lista")
                    continue

                for conclusao in itens:
                    if not isinstance(conclusao, dict):
                        continue
                    estado = str(conclusao.get("estado") or "").strip().lower()
                    conclusoes.append(
                        {
                            "item": questao_base(registro, conclusao),
                            "estado": estado,
                            "auditado": str(registro.get("auditado") or ""),
                            "provider": str(registro.get("provider") or ""),
                            "modelo": modelo,
                            "evidencia": str(registro.get("evidencia") or ""),
                            "coluna_evidencia": str(registro.get("coluna_evidencia") or ""),
                            "item_codigo": str(conclusao.get("item_codigo") or ""),
                            "afirmacao": str(conclusao.get("afirmacao_auditado") or ""),
                            "justificativa": str(conclusao.get("justificativa") or ""),
                            "lacunas": " | ".join(map(str, conclusao.get("lacunas") or [])),
                            "arquivo_origem": str(caminho),
                        }
                    )
    return conclusoes, erros


def selecionar_exemplos(registros: list[dict[str, Any]], limite: int = 2) -> list[dict[str, Any]]:
    ordenados = sorted(
        registros,
        key=lambda item: (
            PRIORIDADE_EXEMPLOS.get(item["estado"], 9),
            -len(item["justificativa"]),
            item["auditado"],
            item["modelo"],
        ),
    )
    escolhidos: list[dict[str, Any]] = []
    estados_usados: set[str] = set()
    auditados_usados: set[str] = set()

    for item in ordenados:
        if item["estado"] in estados_usados:
            continue
        escolhidos.append(item)
        estados_usados.add(item["estado"])
        auditados_usados.add(item["auditado"])
        if len(escolhidos) == limite:
            return escolhidos

    for item in ordenados:
        if item in escolhidos or item["auditado"] in auditados_usados:
            continue
        escolhidos.append(item)
        auditados_usados.add(item["auditado"])
        if len(escolhidos) == limite:
            return escolhidos

    for item in ordenados:
        if item not in escolhidos:
            escolhidos.append(item)
        if len(escolhidos) == limite:
            break
    return escolhidos


def formatar_planilha(ws, larguras: dict[int, float]) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    for celula in ws[1]:
        celula.fill = fill
        celula.font = Font(color="FFFFFF", bold=True)
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for linha in ws.iter_rows(min_row=2):
        for celula in linha:
            celula.alignment = Alignment(vertical="top", wrap_text=True)
    for indice, largura in larguras.items():
        ws.column_dimensions[get_column_letter(indice)].width = largura


def gerar_xlsx(
    conclusoes: list[dict[str, Any]],
    fontes: list[Path],
    saida: Path,
    erros: list[str],
    referencias: list[str],
) -> None:
    agrupados: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for conclusao in conclusoes:
        agrupados[conclusao["item"]].append(conclusao)

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo por item"
    cabecalho = [
        "Item",
        "Total de avaliações",
        "Conforme",
        "Não conforme",
        "Inconclusivo",
        "Erros",
        "% conforme",
        "% não conforme",
        "% inconclusivo",
        "Auditados distintos",
        "Modelos distintos",
        "Exemplo 1 - auditado",
        "Exemplo 1 - modelo",
        "Exemplo 1 - estado",
        "Exemplo 1 - evidência",
        "Exemplo 1 - afirmação avaliada",
        "Exemplo 1 - justificativa",
        "Exemplo 2 - auditado",
        "Exemplo 2 - modelo",
        "Exemplo 2 - estado",
        "Exemplo 2 - evidência",
        "Exemplo 2 - afirmação avaliada",
        "Exemplo 2 - justificativa",
    ]
    ws.append(cabecalho)

    for item in sorted(agrupados):
        registros = agrupados[item]
        contagem = Counter(registro["estado"] for registro in registros)
        total = len(registros)
        erros_estado = contagem["erro"]
        exemplos = selecionar_exemplos(registros)
        exemplos += [{}] * (2 - len(exemplos))

        def percentual(estado: str) -> float:
            return contagem[estado] / total if total else 0

        linha: list[Any] = [
            item,
            total,
            contagem["conforme"],
            contagem["nao_conforme"],
            contagem["inconclusivo"],
            erros_estado,
            percentual("conforme"),
            percentual("nao_conforme"),
            percentual("inconclusivo"),
            len({registro["auditado"] for registro in registros if registro["auditado"]}),
            len({registro["modelo"] for registro in registros if registro["modelo"]}),
        ]
        for exemplo in exemplos:
            linha.extend(
                [
                    exemplo.get("auditado", ""),
                    exemplo.get("modelo", ""),
                    exemplo.get("estado", ""),
                    exemplo.get("evidencia", ""),
                    exemplo.get("afirmacao", ""),
                    exemplo.get("justificativa", ""),
                ]
            )
        ws.append(linha)

    for coluna in (7, 8, 9):
        for celula in ws.iter_cols(min_col=coluna, max_col=coluna, min_row=2):
            for item in celula:
                item.number_format = "0.0%"
    formatar_planilha(
        ws,
        {
            1: 12, 2: 18, 3: 12, 4: 15, 5: 15, 6: 14,
            7: 13, 8: 16, 9: 15, 10: 18, 11: 16,
            12: 15, 13: 22, 14: 15, 15: 32, 16: 48, 17: 70,
            18: 15, 19: 22, 20: 15, 21: 32, 22: 48, 23: 70,
        },
    )
    if ws.max_row > 1:
        tabela = Table(displayName="ResumoAvaliacoes", ref=ws.dimensions)
        tabela.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tabela)

    detalhes = wb.create_sheet("Rastreabilidade")
    detalhes.append(
        [
            "Item", "Auditado", "Provider", "Modelo", "Estado", "Coluna de evidência",
            "Evidência", "Item original", "Afirmação avaliada", "Justificativa", "Lacunas",
            "Arquivo JSONL",
        ]
    )
    for registro in sorted(conclusoes, key=lambda x: (x["item"], x["auditado"], x["modelo"], x["estado"])):
        detalhes.append(
            [
                registro["item"], registro["auditado"], registro["provider"], registro["modelo"],
                registro["estado"], registro["coluna_evidencia"], registro["evidencia"],
                registro["item_codigo"], registro["afirmacao"], registro["justificativa"],
                registro["lacunas"], registro["arquivo_origem"],
            ]
        )
    formatar_planilha(
        detalhes,
        {1: 12, 2: 15, 3: 14, 4: 24, 5: 15, 6: 20, 7: 40, 8: 45, 9: 55, 10: 75, 11: 55, 12: 65},
    )

    metadados = wb.create_sheet("Metadados")
    metadados.append(["Campo", "Valor"])
    metadados.append(["Arquivos processados", len(fontes)])
    metadados.append(["Conclusões processadas", len(conclusoes)])
    metadados.append(["Itens consolidados", len(agrupados)])
    metadados.append(["Linhas JSON inválidas", len(erros)])
    metadados.append(["Referência de modelo", ", ".join(referencias) if referencias else "Todos os modelos"])
    for fonte in fontes:
        metadados.append(["Fonte", str(fonte)])
    for erro in erros:
        metadados.append(["Erro de leitura", erro])
    formatar_planilha(metadados, {1: 24, 2: 110})

    saida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(saida)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("analyses", nargs="+", type=Path, help="Arquivos JSONL de avaliacao de entrada.")
    parser.add_argument("--output", required=True, type=Path, help="Planilha XLSX de saída.")
    parser.add_argument(
        "--referencia",
        action="append",
        default=[],
        help="Modelo a considerar, conforme o campo 'model' do JSONL. Pode ser repetido.",
    )
    args = parser.parse_args()
    fontes = [caminho.resolve() for caminho in args.analyses]
    referencias = [referencia.strip() for referencia in args.referencia if referencia.strip()]
    filtro_referencias = {referencia.casefold() for referencia in referencias} or None
    conclusoes, erros = ler_registros(fontes, filtro_referencias)
    gerar_xlsx(conclusoes, fontes, args.output.resolve(), erros, referencias)
    print(f"Arquivos processados: {len(fontes)}")
    print(f"Referencia: {', '.join(referencias) if referencias else 'todos os modelos'}")
    print(f"Conclusoes processadas: {len(conclusoes)}")
    print(f"Planilha gerada: {args.output.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
